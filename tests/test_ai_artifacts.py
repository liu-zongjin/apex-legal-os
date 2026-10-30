import json
import io
import sqlite3
import tempfile
from pathlib import Path
import unittest
from unittest.mock import AsyncMock, patch
from fastapi import HTTPException
from backend import main
from backend.ai_artifacts import extract_artifacts
from docx import Document

SVG = '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="200"><rect width="200" height="200" fill="#1E6FD9"/></svg>'
HTML = '<!DOCTYPE html><html lang="zh-CN"><meta charset="UTF-8"><title>蓝色正方形</title><div style="width:200px;height:200px;background:#1E6FD9"></div></html>'


class ArtifactTests(unittest.TestCase):
    def test_historical_and_named_artifacts(self):
        answer = f'🙂已生成。\n```svg filename="蓝色正方形.svg"\n{SVG}\n```\n\n```html\n{HTML}\n```'
        items = extract_artifacts(answer)
        self.assertEqual([item['format'] for item in items], ['svg', 'html'])
        self.assertEqual(items[0]['filename'], '蓝色正方形.svg')
        self.assertEqual(items[1]['index'], 1)
        self.assertNotIn('```', items[0]['content'])
        self.assertIn('```svg', items[0]['source'])

    def test_invalid_incomplete_oversized_and_unsafe_names(self):
        for answer in ['```svg\n<svg>\n```', '```json\n{bad}\n```', '```html\n<div>partial</div>\n```',
                       '```txt\n' + '大' * 100000 + '\n```', f'```svg\n{SVG}']:
            self.assertEqual(extract_artifacts(answer), [])
        items = extract_artifacts('```txt filename="../../坏名字.html"\n你好\n```')
        self.assertEqual(items[0]['filename'], '_.._坏名字.txt')
        self.assertEqual(len(extract_artifacts('```txt\nhi\n```\n' * 12)), 8)

    def test_windows_line_endings_still_produce_an_artifact(self):
        raw = '已生成。\r\n```md filename="报告.md"\r\n# 标题\r\n正文\r\n```\r\n'
        artifacts = extract_artifacts(raw)
        self.assertEqual(len(artifacts), 1)
        self.assertIn(artifacts[0]['source'], raw)
        self.assertEqual(artifacts[0]['filename'], '报告.md')

    def test_download_exact_file_not_instructions(self):
        data = {'question': '导出html', 'answer': f'已生成\n```html filename="正方形.html"\n{HTML}\n```\n使用方法'}
        with patch.object(main, '_load_ai_ask_export', return_value=data):
            response = main.export_ai_ask_message('anonymous', 'html', artifact_index=0)
            self.assertEqual(response.body.decode(), HTML + '\n')
            self.assertIn('text/html', response.headers['content-type'])
            self.assertIn('filename*=', response.headers['content-disposition'])
            with self.assertRaises(HTTPException) as mismatch:
                main.export_ai_ask_message('anonymous', 'svg', artifact_index=0)
            self.assertEqual(mismatch.exception.status_code, 400)
            with self.assertRaises(HTTPException) as missing:
                main.export_ai_ask_message('anonymous', 'html', artifact_index=1)
            self.assertEqual(missing.exception.status_code, 404)

    def test_formats_and_markdown_download_compatibility(self):
        for fmt, content in [('json', '{"标题":"文件"}'), ('csv', '名称,金额\n测试,2'), ('txt', '你好'), ('md', '# 文件')]:
            data = {'question': '生成文件', 'answer': f'```{fmt}\n{content}\n```'}
            with patch.object(main, '_load_ai_ask_export', return_value=data):
                response = main.export_ai_ask_message('anonymous', fmt, artifact_index=0)
                self.assertEqual(response.body.decode(), content + '\n')

    def test_word_conversion_uses_artifact_body(self):
        data = {'question':'导出Word', 'answer':'```md filename="说明.md"\n# 说明\n\n这是文件正文。\n```'}
        with patch.object(main, '_load_ai_ask_export', return_value=data):
            response = main.export_ai_ask_message('anonymous', 'docx', layout='legal_standard', artifact_index=0)
        text = '\n'.join(p.text for p in Document(io.BytesIO(response.body)).paragraphs)
        self.assertIn('这是文件正文。', text)
        self.assertNotIn('filename=', text)
        self.assertNotIn('导出Word', text)
        self.assertNotIn('法律操作系统 AI增强', text)
        self.assertNotIn('知识范围', text)

    def test_empty_or_bare_file_claim_is_not_a_success(self):
        from backend.ai_artifacts import require_deliverable_answer
        for value in ('', '已生成文件。', '成功创建附件', '已生成文件。请点击下载'):
            with self.subTest(value=value), self.assertRaises(HTTPException) as error:
                require_deliverable_answer(value)
            self.assertEqual(error.exception.status_code, 502)
        require_deliverable_answer('已生成文件。\n```md filename="报告.md"\n# 报告\n正文\n```')
        require_deliverable_answer('未找到案例，请调整检索条件。')

    def test_word_request_without_fence_builds_real_download_from_answer(self):
        from backend.ai_artifacts import ensure_requested_document
        body = '# 匿名文稿\n\n' + '这是经过整理的真实正文，保留用户提供的全部内容。' * 3
        answer = ensure_requested_document('请生成 Word 文档', body)
        artifacts = extract_artifacts(answer)
        self.assertEqual(len(artifacts), 1)
        self.assertEqual(artifacts[0]['content'].strip(), body)
        with patch.object(main, '_load_ai_ask_export', return_value={'question': '请生成 Word 文档', 'answer': answer}):
            response = main.export_ai_ask_message('anonymous', 'docx', layout='legal_standard', artifact_index=0)
        doc = Document(io.BytesIO(response.body))
        self.assertIn('这是经过整理的真实正文', '\n'.join(p.text for p in doc.paragraphs))

    def test_document_fallback_does_not_invent_or_convert_instructions(self):
        from backend.ai_artifacts import ensure_requested_document
        body = '这是正常说明。' * 12
        self.assertEqual(ensure_requested_document('请解释如何生成Word文件', body), body)
        self.assertEqual(ensure_requested_document('检索支持预约合同观点的裁判案例', body), body)
        self.assertEqual(len(extract_artifacts(ensure_requested_document('生成Word文件', '你好'))), 1)
        self.assertEqual(ensure_requested_document('生成Word文件', '已生成文件。请点击下载'), '已生成文件。请点击下载')
        refusal = '我无法生成 Word 文档，因为未能读取任何原始材料。' * 3
        self.assertEqual(ensure_requested_document('生成Word文件', refusal), refusal)
        malformed = '正文说明\n```md\n没有闭合的片段'
        self.assertEqual(ensure_requested_document('生成Word文件', malformed), malformed)

    def test_saved_history_restores_legacy_artifact_and_download(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'anonymous.db'
            def db():
                conn = sqlite3.connect(path)
                conn.row_factory = sqlite3.Row
                return conn
            with db() as conn:
                conn.executescript('''
                    CREATE TABLE ai_ask_conversations(id INTEGER PRIMARY KEY,history_key TEXT,context_json TEXT,created_at TEXT,updated_at TEXT);
                    CREATE TABLE ai_ask_messages(id TEXT PRIMARY KEY,conversation_id INTEGER,role TEXT,content TEXT,sources_json TEXT,cited_sources_json TEXT,reasoning TEXT,is_reasoning_model INTEGER,model TEXT,case_id INTEGER,research_question_id INTEGER,strategy_version_id INTEGER,strategy_version TEXT,strategy_decision_json TEXT,created_at_ms INTEGER);
                    CREATE TABLE case_manager_cases(id INTEGER PRIMARY KEY,title TEXT);
                    CREATE TABLE research_project_questions(id INTEGER PRIMARY KEY,project_id INTEGER);
                    CREATE TABLE research_projects(id INTEGER PRIMARY KEY);
                ''')
            with db() as conn:
                main.user_memory.add_schema(conn)
                main.ai_context.add_schema(conn)
            records = [dict(id='q',role='user',content='导出html',ts=1),
                       dict(id='a',role='ai',content=f'```html\n{HTML}\n```',ts=2)]
            with patch.object(main, 'get_db', db):
                main.save_ai_ask_history(main.AIAskHistoryPayload(history_key='ask_history_direct_ai',context=None,records=records))
                loaded = main.get_ai_ask_history()['conversations'][0]['records'][1]
                self.assertEqual(loaded['artifacts'][0]['format'], 'html')
                response = main.export_ai_ask_message('a', 'html', artifact_index=0)
                self.assertEqual(response.body.decode(), HTML + '\n')


class ArtifactDirectTests(unittest.IsolatedAsyncioTestCase):
    async def test_draw_then_export_html_with_same_context(self):
        call = AsyncMock(side_effect=[
            ({'choices': [{'message': {'content': f'已绘制。\n```svg filename="正方形.svg"\n{SVG}\n```'}}]}, 'deepseek-flash'),
            ({'choices': [{'message': {'content': f'已生成。\n```html filename="正方形.html"\n{HTML}\n```'}}]}, 'deepseek-flash'),
        ])
        with patch.object(main, '_call_ai_with_fallback', call), patch.object(main, '_ask_skill_instruction', return_value=''):
            first = await main._ask_direct_ai('画一个蓝色的正方形给我', 'deepseek-flash', None, None, None)
            history = json.dumps([{'role':'user', 'content':'画一个蓝色的正方形给我'}, {'role':'ai', 'content':first['answer']}])
            second = await main._ask_direct_ai('导出html', 'deepseek-flash', history, None, None)
        self.assertEqual(first['artifacts'][0]['format'], 'svg')
        self.assertEqual(second['artifacts'][0]['format'], 'html')
        self.assertTrue(any(SVG in item['content'] for item in call.call_args.args[1]))
        self.assertTrue(any('可下载' in item['content'] for item in call.call_args.args[1]))
