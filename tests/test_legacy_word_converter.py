"""Run the real legacy-document loader without importing its database-owning module."""
import ast
import asyncio
import hashlib
import os
from pathlib import Path
import subprocess
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch
from docx import Document


class LegacyWordConverterTests(unittest.TestCase):
    def exercise(self, platform):
        source = Path(__file__).parents[1] / 'backend/main.py'
        tree = ast.parse(source.read_text(encoding='utf-8'))
        function = next(n for n in tree.body if isinstance(n, ast.AsyncFunctionDef) and n.name == '_load_official_law_document')
        module = ast.Module(body=[ast.ImportFrom(module='__future__', names=[ast.alias(name='annotations')], level=0), function], type_ignores=[])
        ast.fix_missing_locations(module)
        calls = []
        content = '匿名官方法规转换测试，验证使用随包程序并保留用户资料。'
        def convert(command, **kwargs):
            calls.append(command)
            if platform == 'win32':
                self.assertEqual(command[0], r'C:\Stella\converter\program\soffice.com')
                self.assertTrue(any(arg.startswith('-env:UserInstallation=file:') for arg in command))
                output = Path(command[command.index('--outdir') + 1]) / 'official.docx'
            else:
                self.assertEqual(command[0], 'textutil')
                output = Path(command[command.index('-output') + 1])
            document = Document()
            document.add_paragraph(content)
            document.save(output)
            return subprocess.CompletedProcess(command, 0, stdout='', stderr='')
        def parse(path, preserve_source_format=False):
            self.assertTrue(preserve_source_format)
            return {'content': ''.join(p.text for p in Document(path).paragraphs)}
        scope = {'asyncio': asyncio, 'tempfile': tempfile, 'Path': Path, 'os': os,
                 'sys': SimpleNamespace(platform=platform), 'subprocess': SimpleNamespace(run=convert, SubprocessError=subprocess.SubprocessError),
                 'download_national_law_docx': AsyncMock(return_value=bytes.fromhex('d0cf11e0a1b11ae1') + b'anonymous legacy file'),
                 'parse_docx': parse, 'NationalLawError': ValueError,
                 'compute_hash': lambda text: hashlib.sha256(text.encode()).hexdigest()}
        exec(compile(module, str(source), 'exec'), scope)
        with patch.dict(os.environ, {'LEGAL_KB_SOFFICE': r'C:\Stella\converter\program\soffice.com'}):
            detail, result, digest = asyncio.run(scope['_load_official_law_document']('anonymous', {'title': '匿名测试'}))
        self.assertEqual(result, content)
        self.assertEqual(detail['title'], '匿名测试')
        self.assertEqual(len(calls), 1)

    def test_windows_uses_bundled_converter_without_global_office(self):
        self.exercise('win32')

    def test_mac_keeps_builtin_textutil_first(self):
        self.exercise('darwin')


if __name__ == '__main__':
    unittest.main()
