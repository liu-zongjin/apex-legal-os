import sqlite3
import tempfile
import unittest
import io
from pathlib import Path

import openpyxl
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.series_manager import register_series_manager_routes


BASE_SCHEMA = """
CREATE TABLE case_manager_cases (
    id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL UNIQUE,
    case_type TEXT DEFAULT '', current_stage TEXT DEFAULT '', court_contact TEXT DEFAULT '',
    party_contact TEXT DEFAULT '', memo TEXT DEFAULT '', source_status TEXT DEFAULT 'active',
    deleted_at TEXT DEFAULT '', archived_at TEXT DEFAULT '', pinned INTEGER DEFAULT 0,
    created_at TEXT NOT NULL, updated_at TEXT NOT NULL
);
CREATE TABLE case_manager_parties (
    id INTEGER PRIMARY KEY AUTOINCREMENT, case_id INTEGER NOT NULL, party_group TEXT DEFAULT 'client',
    role TEXT DEFAULT '', name TEXT DEFAULT '', credit_code TEXT DEFAULT '', client_id INTEGER,
    opponent_id INTEGER, sort_order INTEGER DEFAULT 0, created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
    FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
);
CREATE TABLE case_manager_case_numbers (
    id INTEGER PRIMARY KEY AUTOINCREMENT, case_id INTEGER NOT NULL, number_type TEXT DEFAULT '',
    case_no TEXT DEFAULT '', court TEXT DEFAULT '', handler_contact TEXT DEFAULT '', sort_order INTEGER DEFAULT 0,
    first_filing_at TEXT DEFAULT '', first_judgment_served_at TEXT DEFAULT '', second_filing_at TEXT DEFAULT '',
    second_judgment_effective_at TEXT DEFAULT '', created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
    FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
);
CREATE TABLE case_manager_clients (
    id INTEGER PRIMARY KEY AUTOINCREMENT, client_type TEXT DEFAULT '单位', name TEXT UNIQUE,
    identity_code TEXT DEFAULT '', address TEXT DEFAULT '', contact_person TEXT DEFAULT '', phone TEXT DEFAULT '',
    invoice_info TEXT DEFAULT '', created_at TEXT NOT NULL, updated_at TEXT NOT NULL
);
CREATE TABLE case_manager_opponents (
    id INTEGER PRIMARY KEY AUTOINCREMENT, opponent_type TEXT DEFAULT '单位', name TEXT UNIQUE,
    identity_code TEXT DEFAULT '', address TEXT DEFAULT '', contact_person TEXT DEFAULT '', phone TEXT DEFAULT '',
    notes TEXT DEFAULT '', created_at TEXT NOT NULL, updated_at TEXT NOT NULL
);
CREATE TABLE case_manager_todos (
    id INTEGER PRIMARY KEY AUTOINCREMENT, case_id INTEGER NOT NULL, title TEXT NOT NULL,
    todo_title TEXT DEFAULT '', remind_at TEXT, notes TEXT DEFAULT '', status TEXT DEFAULT 'pending',
    created_at TEXT NOT NULL, updated_at TEXT NOT NULL, source_key TEXT UNIQUE,
    FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
);
CREATE TABLE case_manager_hearings (
    id INTEGER PRIMARY KEY AUTOINCREMENT, case_id INTEGER NOT NULL, title TEXT NOT NULL,
    reminder_type TEXT DEFAULT '开庭', hearing_at TEXT, location TEXT DEFAULT '', case_no TEXT DEFAULT '',
    court TEXT DEFAULT '', court_contact TEXT DEFAULT '', party_contact TEXT DEFAULT '', notes TEXT DEFAULT '',
    created_at TEXT NOT NULL, updated_at TEXT NOT NULL, source_key TEXT UNIQUE,
    FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
);
CREATE TABLE case_manager_preservations (
    id INTEGER PRIMARY KEY AUTOINCREMENT, case_id INTEGER NOT NULL, title TEXT NOT NULL,
    preservation_type TEXT DEFAULT '', asset_content TEXT DEFAULT '', start_at TEXT, deadline_at TEXT,
    period_years REAL, notes TEXT DEFAULT '', created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
    source_key TEXT UNIQUE, FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
);
CREATE TABLE case_manager_appeals (
    id INTEGER PRIMARY KEY AUTOINCREMENT, case_id INTEGER NOT NULL, title TEXT NOT NULL,
    appeal_type TEXT DEFAULT '', start_at TEXT, period_days INTEGER, deadline_at TEXT, case_no TEXT DEFAULT '',
    court TEXT DEFAULT '', court_contact TEXT DEFAULT '', party_contact TEXT DEFAULT '', notes TEXT DEFAULT '',
    created_at TEXT NOT NULL, updated_at TEXT NOT NULL, source_key TEXT UNIQUE,
    FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
);
CREATE TABLE case_manager_audit_logs (
    id INTEGER PRIMARY KEY AUTOINCREMENT, case_id INTEGER NOT NULL, entity_type TEXT NOT NULL,
    entity_id INTEGER DEFAULT 0, action TEXT NOT NULL, summary TEXT DEFAULT '', changes_json TEXT DEFAULT '[]',
    actor TEXT DEFAULT '', created_at TEXT NOT NULL,
    FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
);
"""


class SeriesManagerTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp_dir.name) / "series.db"
        conn = sqlite3.connect(self.db_path)
        conn.executescript(BASE_SCHEMA)
        conn.execute(
            "INSERT INTO case_manager_cases(title,created_at,updated_at) VALUES ('既有普通案件','2026-01-01','2026-01-01')"
        )
        conn.commit()
        conn.close()

        def get_db():
            database = sqlite3.connect(self.db_path)
            database.row_factory = sqlite3.Row
            database.execute("PRAGMA foreign_keys=ON")
            return database

        self.get_db = get_db
        app = FastAPI()
        register_series_manager_routes(app, get_db)
        self.client = TestClient(app)

    def tearDown(self):
        self.temp_dir.cleanup()

    def create_series(self):
        response = self.client.post(
            "/api/case-manager/series",
            json={
                "title": "商品房系列纠纷",
                "case_type": "商品房预售合同纠纷",
                "profile": {"common_facts": "共同案情", "common_contract_terms": "共同条款"},
                "common_claims": [
                    {"claim_code": "R1", "claim_text": "返还购房款", "amount_text": "100.10", "currency": "CNY"}
                ],
                "children": [
                    {"series_code": "001", "display_name": "诉示一某案", "current_stage": "一审"},
                    {"series_code": "002", "display_name": "诉丙公司案", "current_stage": "执行"},
                ],
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_bulk_hearing_numbers_sync_only_to_selected_child(self):
        series = self.create_series()
        children = self.client.get(f"/api/case-manager/series/{series['id']}/children").json()['items']
        child_id = children[0]['id']
        response = self.client.post(f"/api/case-manager/series/{series['id']}/children/bulk", json={
            'operation': 'add_hearing', 'child_ids': [child_id],
            'payload': {'case_no': '（2026）测试民终123号', 'court': '匿名中级法院', 'hearing_at': '2026-11-03 15:00'}
        })
        self.assertEqual(response.status_code, 200, response.text)
        with self.get_db() as conn:
            rows = conn.execute('SELECT case_id,number_type,case_no,court FROM case_manager_case_numbers WHERE case_no=?', ('（2026）测试民终123号',)).fetchall()
        self.assertEqual([tuple(row) for row in rows], [(child_id, '二审', '（2026）测试民终123号', '匿名中级法院')])

    def test_migration_preserves_existing_cases_as_ordinary(self):
        with self.get_db() as conn:
            row = conn.execute("SELECT case_kind,parent_series_id FROM case_manager_cases WHERE title='既有普通案件'").fetchone()
        self.assertEqual(row["case_kind"], "ordinary")
        self.assertIsNone(row["parent_series_id"])

    def test_series_name_conflict_ignores_deleted_records_but_blocks_active_records(self):
        series = self.create_series()
        duplicate = self.client.post("/api/case-manager/series", json={"title": series["title"]})
        self.assertEqual(duplicate.status_code, 409, duplicate.text)
        self.assertIn("在办案件中已有同名系列案件", duplicate.json()["detail"])

        with self.get_db() as conn:
            conn.execute(
                "UPDATE case_manager_cases SET source_status='deleted',deleted_at='2026-07-27' WHERE id=?",
                (series["id"],),
            )
            conn.commit()

        recreated = self.client.post("/api/case-manager/series", json={"title": series["title"]})
        self.assertEqual(recreated.status_code, 200, recreated.text)
        self.assertNotEqual(recreated.json()["id"], series["id"])
        with self.get_db() as conn:
            deleted = conn.execute(
                "SELECT title FROM case_manager_cases WHERE id=?",
                (series["id"],),
            ).fetchone()
        self.assertEqual(deleted["title"], f"{series['title']} · 已删除 #{series['id']}")

    def test_series_children_and_dynamic_claim_inheritance(self):
        series = self.create_series()
        self.assertEqual(series["stats"]["total"], 2)
        self.assertEqual(series["children"][0]["title"], "商品房系列纠纷·诉示一某案")
        common_id = series["common_claims"][0]["id"]
        child = series["children"][0]
        response = self.client.put(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}",
            json={
                "series_code": "001",
                "display_name": "诉示一某案",
                "current_stage": "一审",
                "facts_mode": "supplement",
                "facts_difference": "示一某个别事实",
                "contract_mode": "inherit",
                "claims": [
                    {"parent_claim_id": common_id, "claim_mode": "supplement", "claim_text": "另加资金占用费"}
                ],
            },
        )
        self.assertEqual(response.status_code, 200, response.text)

        response = self.client.put(
            f"/api/case-manager/series/{series['id']}",
            json={
                "title": series["title"],
                "case_type": series["case_type"],
                "profile": {"common_facts": "更新后的共同案情", "common_contract_terms": "共同条款"},
                "common_claims": [
                    {"claim_code": "R1", "claim_text": "返还全部购房款", "amount_text": "200.20", "currency": "CNY"}
                ],
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["common_claims"][0]["id"], common_id)

        detail = self.client.get(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}"
        ).json()
        self.assertIn("更新后的共同案情", detail["effective_facts"])
        self.assertIn("示一某个别事实", detail["effective_facts"])
        self.assertIn("返还全部购房款", detail["effective_claims"][0]["claim_text"])
        self.assertIn("另加资金占用费", detail["effective_claims"][0]["claim_text"])

        renamed = self.client.put(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}",
            json={
                "display_name": "18-01",
                "parties": [
                    {"party_group": "client", "role": "委托人", "name": "示丆某"},
                    {"party_group": "other", "role": "申请人", "name": "钟晓勇"},
                ],
            },
        )
        self.assertEqual(renamed.status_code, 200, renamed.text)
        self.assertEqual(renamed.json()["title"], "商品房系列纠纷·钟晓勇案")

    def test_series_public_custom_fields_are_persisted(self):
        series = self.create_series()
        response = self.client.put(
            f"/api/case-manager/series/{series['id']}",
            json={
                "title": series["title"],
                "case_type": series["case_type"],
                "current_stage": series["current_stage"],
                "profile": {
                    "common_facts": series["profile"]["common_facts"],
                    "common_contract_terms": series["profile"]["common_contract_terms"],
                    "series_memo": "共同备忘",
                    "custom_fields": [
                        {"key": "custom_focus", "label": "共同争议焦点", "value": "房屋交付与违约责任"}
                    ],
                },
                "parties": series["parties"],
                "common_claims": series["common_claims"],
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        detail = response.json()
        self.assertEqual(detail["profile"]["custom_fields"], [
            {"key": "custom_focus", "label": "共同争议焦点", "value": "房屋交付与违约责任"}
        ])
        refreshed = self.client.get(f"/api/case-manager/series/{series['id']}")
        self.assertEqual(refreshed.status_code, 200, refreshed.text)
        self.assertEqual(refreshed.json()["profile"]["custom_fields"][0]["value"], "房屋交付与违约责任")

    def test_series_public_contacts_and_memo_use_one_shared_source(self):
        court_contact = (
            "【联调/诉讼/仲裁阶段】\n"
            "【联系人组】\n组ID：series-contact-a\n身份：秘书\n姓名：王静\n电话：13800000200\n邮箱：wang@example.com\n备注：共同联系人"
        )
        party_contact = (
            "【委托方】\n"
            "【联系人组】\n组ID：series-contact-b\n身份：联系人\n姓名：李明\n电话：13900000600\n邮箱：li@example.com\n备注：委托方联系人"
        )
        created = self.client.post(
            "/api/case-manager/series",
            json={
                "title": "联系人一致性系列案",
                "profile": {"series_memo": "创建时的共同备忘"},
                "court_contact": court_contact,
                "party_contact": party_contact,
            },
        )
        self.assertEqual(created.status_code, 200, created.text)
        series = created.json()
        self.assertEqual(series["memo"], "创建时的共同备忘")
        self.assertEqual(series["profile"]["series_memo"], "创建时的共同备忘")
        self.assertEqual(series["court_contact"], court_contact)
        self.assertEqual(series["party_contact"], party_contact)

        updated = self.client.put(
            f"/api/case-manager/series/{series['id']}",
            json={
                "title": series["title"],
                "court_contact": court_contact,
                "party_contact": party_contact,
                "profile": {
                    "common_facts": "",
                    "common_contract_terms": "",
                    "series_memo": "编辑后的共同备忘",
                },
            },
        )
        self.assertEqual(updated.status_code, 200, updated.text)
        detail = updated.json()
        self.assertEqual(detail["memo"], "编辑后的共同备忘")
        self.assertEqual(detail["profile"]["series_memo"], "编辑后的共同备忘")
        self.assertEqual(detail["court_contact"], court_contact)
        self.assertEqual(detail["party_contact"], party_contact)

    def test_series_detail_field_mappings_are_persisted_and_validate_sources(self):
        series = self.create_series()
        columns = [
            *series["profile"]["table_columns"],
            {"key": "custom_applicant", "label": "申请人", "field_type": "text", "source": "custom"},
            {"key": "custom_secretary", "label": "秘书", "field_type": "text", "source": "custom"},
        ]
        self.assertEqual(
            self.client.put(f"/api/case-manager/series/{series['id']}/columns", json={"columns": columns}).status_code,
            200,
        )
        response = self.client.put(
            f"/api/case-manager/series/{series['id']}/detail-field-mappings",
            json={"mappings": [
                {"source_key": "custom_applicant", "target_key": "parties.other"},
                {"source_key": "custom_secretary", "target_key": "contacts"},
                {"source_key": "missing_column", "target_key": "memo"},
                {"source_key": "custom_secretary", "target_key": "not_a_detail_field"},
            ]},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["mappings"], [
            {"source_key": "custom_applicant", "target_key": "parties.other"},
            {"source_key": "custom_secretary", "target_key": "contacts"},
        ])
        refreshed = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        self.assertEqual(refreshed["profile"]["detail_field_mappings"], response.json()["mappings"])

    def test_series_child_contact_display_selection_is_persisted(self):
        series = self.create_series()
        child = series["children"][0]
        columns = [
            *series["profile"]["table_columns"],
            {"key": "custom_contact", "label": "主表联系人", "field_type": "text", "source": "custom"},
        ]
        self.assertEqual(
            self.client.put(f"/api/case-manager/series/{series['id']}/columns", json={"columns": columns}).status_code,
            200,
        )
        contact_text = (
            "【联调/诉讼/仲裁阶段】\n"
            "【联系人组】\n组ID：contact-a\n身份：联系人\n姓名：甲\n电话：13800000200\n邮箱：\n备注：第一组\n\n"
            "【联系人组】\n组ID：contact-b\n身份：联系人\n姓名：乙\n电话：13900000600\n邮箱：\n备注：第二组"
        )
        response = self.client.put(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}",
            json={
                "series_code": child["series_code"],
                "display_name": child["display_name"],
                "court_contact": contact_text,
                "detail_contact_selection": {"custom_contact": ["contact-b"]},
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["detail_contact_selection"], {"custom_contact": ["contact-b"]})
        refreshed = self.client.get(f"/api/case-manager/series/{series['id']}/children/{child['id']}").json()
        self.assertEqual(refreshed["detail_contact_selection"], {"custom_contact": ["contact-b"]})
        with self.get_db() as conn:
            raw = conn.execute("SELECT detail_contact_selection_json FROM case_manager_cases WHERE id=?", (child["id"],)).fetchone()[0]
        self.assertEqual(raw, '{"custom_contact": ["contact-b"]}')

    def test_series_child_rows_can_be_hidden_and_restored(self):
        series = self.create_series()
        ids = [item["id"] for item in series["children"]]
        hidden = self.client.post(
            f"/api/case-manager/series/{series['id']}/children/bulk",
            json={"child_ids": [ids[0]], "operation": "hide", "payload": {}},
        )
        self.assertEqual(hidden.status_code, 200, hidden.text)
        detail = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        children = {item["id"]: item for item in detail["children"]}
        self.assertEqual(children[ids[0]]["series_row_hidden"], 1)
        self.assertEqual(children[ids[1]]["series_row_hidden"], 0)

        restored = self.client.post(
            f"/api/case-manager/series/{series['id']}/children/bulk",
            json={"child_ids": [ids[0]], "operation": "unhide", "payload": {}},
        )
        self.assertEqual(restored.status_code, 200, restored.text)
        refreshed = self.client.get(f"/api/case-manager/series/{series['id']}/children/{ids[0]}").json()
        self.assertEqual(refreshed["series_row_hidden"], 0)

    def test_bulk_progress_reminders_and_lifecycle(self):
        series = self.create_series()
        ids = [item["id"] for item in series["children"]]
        response = self.client.post(
            f"/api/case-manager/series/{series['id']}/children/bulk",
            json={
                "child_ids": ids,
                "operation": "add_milestone",
                "payload": {
                    "milestone": {"milestone_type": "开庭", "title": "第一次开庭", "milestone_status": "planned", "happened_at": "2026-08-01T09:00"},
                    "sync_stage": True,
                    "current_stage": "一审",
                },
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        response = self.client.post(
            f"/api/case-manager/series/{series['id']}/children/bulk",
            json={"child_ids": ids, "operation": "add_todo", "payload": {"todo_title": "提交证据", "remind_at": "2026-08-01T09:00"}},
        )
        self.assertEqual(response.status_code, 200, response.text)
        groups = self.client.get(f"/api/case-manager/series/{series['id']}/daily-reminders").json()["groups"]
        self.assertEqual(groups[0]["count"], 2)

        one_archived = self.client.post(
            f"/api/case-manager/series/{series['id']}/children/bulk",
            json={"child_ids": [ids[0]], "operation": "archive", "payload": {}},
        )
        self.assertEqual(one_archived.status_code, 200, one_archived.text)
        detail = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        self.assertEqual(detail["stats"]["active"], 1)
        self.assertEqual(detail["stats"]["archived"], 1)
        one_restored = self.client.post(
            f"/api/case-manager/series/{series['id']}/children/bulk",
            json={"child_ids": [ids[0]], "operation": "restore", "payload": {}},
        )
        self.assertEqual(one_restored.status_code, 200, one_restored.text)

        blocked = self.client.post(f"/api/case-manager/series/{series['id']}/archive", json={"force": False})
        self.assertEqual(blocked.status_code, 409)
        archived = self.client.post(f"/api/case-manager/series/{series['id']}/archive", json={"force": True})
        self.assertEqual(archived.status_code, 200, archived.text)
        restored = self.client.post(f"/api/case-manager/series/{series['id']}/restore")
        self.assertEqual(restored.status_code, 200, restored.text)
        detail = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        self.assertEqual(detail["source_status"], "active")
        self.assertTrue(all(child["source_status"] == "archived" for child in detail["children"]))

    def test_excel_export_preview_and_roundtrip(self):
        series = self.create_series()
        exported = self.client.get(f"/api/case-manager/series/{series['id']}/export")
        self.assertEqual(exported.status_code, 200)
        self.assertGreater(len(exported.content), 1000)
        workbook = openpyxl.load_workbook(io.BytesIO(exported.content), data_only=False)
        self.assertEqual(workbook.active.title, "系列案件总表")
        self.assertEqual(
            [sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"],
            ["系列案件总表"],
        )
        self.assertTrue(all(sheet.sheet_state == "veryHidden" for sheet in workbook.worksheets[1:]))
        main_sheet = workbook["系列案件总表"]
        self.assertIn("系列案件总表", main_sheet["A1"].value)
        self.assertEqual(main_sheet["A15"].value, "编号")
        self.assertEqual(main_sheet["B15"].value, "子案名称")
        self.assertEqual(main_sheet.auto_filter.ref.split(":")[0], "A15")
        self.assertEqual(main_sheet.freeze_panes, "C16")
        preview = self.client.post(
            f"/api/case-manager/series/{series['id']}/import-preview",
            files={"file": ("series.xlsx", exported.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(preview.status_code, 200, preview.text)
        self.assertEqual(preview.json()["errors"], [])
        self.assertEqual(preview.json()["mode"], "smart_table")
        self.assertEqual(preview.json()["counts"]["children"], 2)
        committed = self.client.post(
            f"/api/case-manager/series/{series['id']}/import-commit",
            json={"preview_id": preview.json()["preview_id"]},
        )
        self.assertEqual(committed.status_code, 200, committed.text)
        detail = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        self.assertEqual(detail["stats"]["total"], 2)
        self.assertEqual(detail["common_claims"][0]["claim_code"], "R1")

    def test_excel_import_previews_add_update_delete_and_restores_history(self):
        series = self.create_series()
        exported = self.client.get(f"/api/case-manager/series/{series['id']}/export")
        workbook = openpyxl.load_workbook(io.BytesIO(exported.content), data_only=False)
        sheet = workbook["系列案件总表"]
        headers = [cell.value for cell in sheet[15]]
        code_column = headers.index("编号") + 1
        name_column = headers.index("子案名称") + 1
        sheet.cell(row=16, column=name_column).value = "诉示一某案（Excel修改）"
        sheet.delete_rows(17, 1)
        sheet.append(["003", "Excel新增子案", *([""] * (len(headers) - 2))])
        output = io.BytesIO()
        workbook.save(output)

        preview = self.client.post(
            f"/api/case-manager/series/{series['id']}/import-preview",
            files={"file": ("edited.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(preview.status_code, 200, preview.text)
        payload = preview.json()
        self.assertTrue(payload["has_stable_key"])
        self.assertEqual(payload["missing_stable_key_rows"], [])
        self.assertEqual(payload["diff"]["summary"]["added"], 1)
        self.assertEqual(payload["diff"]["summary"]["updated"], 1)
        self.assertEqual(payload["diff"]["summary"]["deleted"], 1)
        self.assertTrue(any(item["type"] == "deleted" and item["code"] == "002" for item in payload["diff"]["changes"]))

        committed = self.client.post(
            f"/api/case-manager/series/{series['id']}/import-commit",
            json={"preview_id": payload["preview_id"]},
        )
        self.assertEqual(committed.status_code, 200, committed.text)
        self.assertTrue(committed.json()["history_id"])
        detail = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        by_code = {item["series_code"]: item for item in detail["children"]}
        self.assertEqual(by_code["001"]["display_name"], "诉示一某案（Excel修改）")
        self.assertEqual(by_code["002"]["source_status"], "deleted")
        self.assertEqual(by_code["003"]["display_name"], "Excel新增子案")

        history = self.client.get(f"/api/case-manager/series/{series['id']}/history")
        self.assertEqual(history.status_code, 200, history.text)
        self.assertEqual(len(history.json()["items"]), 1)
        original_history_id = history.json()["items"][0]["history_id"]

        current_export = self.client.get(f"/api/case-manager/series/{series['id']}/export")
        second_workbook = openpyxl.load_workbook(io.BytesIO(current_export.content), data_only=False)
        second_sheet = second_workbook["系列案件总表"]
        second_headers = [cell.value for cell in second_sheet[15]]
        second_name_column = second_headers.index("子案名称") + 1
        second_sheet.cell(row=16, column=second_name_column).value = "再次修改"
        second_output = io.BytesIO()
        second_workbook.save(second_output)
        second_preview = self.client.post(
            f"/api/case-manager/series/{series['id']}/import-preview",
            files={"file": ("edited-again.xlsx", second_output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(second_preview.status_code, 200, second_preview.text)
        second_commit = self.client.post(
            f"/api/case-manager/series/{series['id']}/import-commit",
            json={"preview_id": second_preview.json()["preview_id"]},
        )
        self.assertEqual(second_commit.status_code, 200, second_commit.text)

        restored = self.client.post(
            f"/api/case-manager/series/{series['id']}/history/{original_history_id}/restore"
        )
        self.assertEqual(restored.status_code, 200, restored.text)
        restored_detail = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        restored_by_code = {item["series_code"]: item for item in restored_detail["children"]}
        self.assertEqual(restored_by_code["001"]["display_name"], "诉示一某案")
        self.assertEqual(restored_by_code["002"]["source_status"], "active")
        self.assertEqual(restored_by_code["003"]["source_status"], "deleted")

    def test_dynamic_columns_inline_edit_and_database_options(self):
        series = self.create_series()
        child = series["children"][0]
        columns = [
            {"key": "series_code", "label": "编号", "field_type": "text", "source": "builtin"},
            {"key": "display_name", "label": "子案名称", "field_type": "text", "source": "builtin"},
            {"key": "custom_house", "label": "房号", "field_type": "text", "source": "custom", "width": 236, "manual_width": True, "hidden": True},
        ]
        response = self.client.put(f"/api/case-manager/series/{series['id']}/columns", json={"columns": columns})
        self.assertEqual(response.status_code, 200, response.text)
        response = self.client.patch(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}/cell",
            json={"field": "custom_house", "value": "A-1802"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["custom_values"]["custom_house"], "A-1802")
        formatted = self.client.patch(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}/cell-format",
            json={"field": "custom_house", "format": {"cell_background": "yellow"}},
        )
        self.assertEqual(formatted.status_code, 200, formatted.text)
        audit_summaries = [item["summary"] for item in self.client.get(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}"
        ).json()["audit_logs"]]
        self.assertIn("表格内编辑：房号", audit_summaries)
        self.assertIn("表格单元格格式：房号", audit_summaries)
        self.assertNotIn("custom_house", "\n".join(audit_summaries))
        response = self.client.patch(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}/cell",
            json={"field": "clients", "value": "示丆某、示一某"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["clients"], "示丆某、示一某")
        detail_response = self.client.get(f"/api/case-manager/series/{series['id']}/children/{child['id']}")
        self.assertEqual(detail_response.status_code, 200, detail_response.text)
        self.assertEqual(detail_response.json()["clients"], "示丆某、示一某")
        response = self.client.put(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}",
            json={
                "display_name": child["display_name"],
                "custom_values": {"custom_house": "B-1903"},
                "court_contact": "书记员：020-12345678",
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["custom_values"]["custom_house"], "B-1903")
        self.assertEqual(response.json()["court_contact"], "书记员：020-12345678")
        refreshed_child = next(
            item for item in self.client.get(f"/api/case-manager/series/{series['id']}").json()["children"]
            if item["id"] == child["id"]
        )
        self.assertEqual(refreshed_child["custom_values"]["custom_house"], "B-1903")
        options = self.client.get("/api/case-manager/series-options", params={"kind": "clients", "q": "示丆某"})
        self.assertEqual(options.status_code, 200, options.text)
        self.assertEqual(options.json()["items"][0]["value"], "示丆某")
        settings_response = self.client.put(
            f"/api/case-manager/series/{series['id']}/columns",
            json={"columns": columns, "settings": {"freeze_columns": 1, "freeze_rows": 2}},
        )
        self.assertEqual(settings_response.status_code, 200, settings_response.text)
        refreshed = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        self.assertEqual(refreshed["profile"]["table_settings"], {"freeze_columns": 1, "freeze_rows": 2})
        house_column = next(item for item in refreshed["profile"]["table_columns"] if item["key"] == "custom_house")
        self.assertEqual(house_column["width"], 236)
        self.assertTrue(house_column["manual_width"])
        self.assertTrue(house_column["hidden"])

        response = self.client.put(
            f"/api/case-manager/series/{series['id']}/columns",
            json={"columns": [{**item, "hidden": False if item["key"] == "custom_house" else item.get("hidden", False)} for item in refreshed["profile"]["table_columns"]]},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertFalse(next(item for item in response.json()["items"] if item["key"] == "custom_house")["hidden"])

    def test_cell_text_formats_are_persisted_and_cleared_when_text_changes(self):
        series = self.create_series()
        child = series["children"][0]
        formatted = self.client.patch(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}/cell-format",
            json={
                "field": "display_name",
                "format": {"runs": [{"start": 0, "end": 2, "bold": True, "highlight": "yellow"}]},
            },
        )
        self.assertEqual(formatted.status_code, 200, formatted.text)
        self.assertEqual(formatted.json()["format"]["runs"][0]["highlight"], "yellow")
        detail = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        refreshed = next(item for item in detail["children"] if item["id"] == child["id"])
        self.assertTrue(refreshed["cell_formats"]["display_name"]["runs"][0]["bold"])

        changed = self.client.patch(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}/cell",
            json={"field": "display_name", "value": "更新后的子案名称"},
        )
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertNotIn("display_name", changed.json().get("cell_formats", {}))

    def test_cell_background_format_is_persisted_and_survives_text_changes(self):
        series = self.create_series()
        child = series["children"][0]
        url = f"/api/case-manager/series/{series['id']}/children/{child['id']}/cell-format"

        formatted = self.client.patch(
            url,
            json={
                "field": "display_name",
                "format": {"runs": [{"start": 0, "end": 2, "bold": True}], "cell_background": "yellow"},
            },
        )
        self.assertEqual(formatted.status_code, 200, formatted.text)
        self.assertEqual(formatted.json()["format"]["cell_background"], "yellow")

        changed = self.client.patch(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}/cell",
            json={"field": "display_name", "value": "更新后的高亮子案名称"},
        )
        self.assertEqual(changed.status_code, 200, changed.text)
        persisted = changed.json()["cell_formats"]["display_name"]
        self.assertEqual(persisted["cell_background"], "yellow")
        self.assertNotIn("runs", persisted)

        cleared = self.client.patch(url, json={"field": "display_name", "format": {}})
        self.assertEqual(cleared.status_code, 200, cleared.text)
        self.assertEqual(cleared.json()["format"], {})

    def test_table_number_column_replaces_all_lines_without_duplication(self):
        series = self.create_series()
        child_id = series["children"][0]["id"]
        base = f"/api/case-manager/series/{series['id']}/children/{child_id}"
        numbers = [
            {"case_no": "（2026）示例初字1号", "court": "示例一院", "number_type": "一审", "first_filing_at": "2026-01-02"},
            {"case_no": "（2026）示例终字2号", "court": "示例二院", "number_type": "二审"},
        ]
        response = self.client.put(base, json={"case_numbers": numbers})
        self.assertEqual(response.status_code, 200, response.text)
        value = "（2026）示例初字3号\n（2026）示例终字2号"
        for _ in range(3):
            response = self.client.patch(base + "/cell", json={"field": "case_no", "value": value})
            self.assertEqual(response.status_code, 200, response.text)
            result = self.client.get(base).json()
            self.assertEqual(result["case_no"], value)
            self.assertEqual(len(result["case_numbers"]), 2)
            self.assertEqual(result["case_numbers"][0]["court"], "示例一院")
            self.assertEqual(result["case_numbers"][0]["first_filing_at"], "2026-01-02")
        response = self.client.patch(base + "/cell", json={"field": "court", "value": "新示例一院\n新示例二院"})
        self.assertEqual(response.json()["court"], "新示例一院\n新示例二院")
        response = self.client.patch(base + "/cell", json={"field": "case_no", "value": "单一案号"})
        self.assertEqual(response.json()["case_no"], "单一案号")
        self.assertEqual(self.client.get(base).json()["case_numbers"][1]["court"], "新示例二院")
        response = self.client.patch(base + "/cell", json={"field": "case_no", "value": ""})
        self.assertEqual(response.json()["case_no"], "")
        self.assertEqual(response.json()["court"], "新示例一院\n新示例二院")

    def test_cell_alignment_and_multiline_value_are_persisted(self):
        series = self.create_series()
        child = series["children"][0]
        url = f"/api/case-manager/series/{series['id']}/children/{child['id']}/cell-format"

        aligned = self.client.patch(url, json={"field": "display_name", "format": {"align": "center"}})
        self.assertEqual(aligned.status_code, 200, aligned.text)
        self.assertEqual(aligned.json()["format"]["align"], "center")

        combined = self.client.patch(
            url,
            json={"field": "display_name", "format": {"runs": [{"start": 0, "end": 1, "bold": True}], "align": "right"}},
        )
        self.assertEqual(combined.status_code, 200, combined.text)
        self.assertEqual(combined.json()["format"]["align"], "right")
        self.assertTrue(combined.json()["format"]["runs"][0]["bold"])

        detail = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        refreshed = next(item for item in detail["children"] if item["id"] == child["id"])
        self.assertEqual(refreshed["cell_formats"]["display_name"]["align"], "right")

        cleared = self.client.patch(url, json={"field": "display_name", "format": {"runs": [{"start": 0, "end": 1, "bold": True}]}})
        self.assertEqual(cleared.status_code, 200, cleared.text)
        self.assertTrue(cleared.json()["format"]["runs"][0]["bold"])
        self.assertNotIn("align", cleared.json()["format"])
        rejected = self.client.patch(url, json={"field": "display_name", "format": {"align": "middle"}})
        self.assertEqual(rejected.status_code, 200, rejected.text)
        self.assertNotIn("align", rejected.json()["format"])

        multiline = self.client.patch(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}/cell",
            json={"field": "display_name", "value": "钟晓勇\n【已调解】"},
        )
        self.assertEqual(multiline.status_code, 200, multiline.text)
        self.assertEqual(multiline.json()["display_name"], "钟晓勇\n【已调解】")
        listing = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        listed = next(item for item in listing["children"] if item["id"] == child["id"])
        self.assertEqual(listed["display_name"], "钟晓勇\n【已调解】")

    def test_cell_vertical_alignment_is_persisted_and_survives_text_changes(self):
        series = self.create_series()
        child = series["children"][0]
        url = f"/api/case-manager/series/{series['id']}/children/{child['id']}/cell-format"

        aligned = self.client.patch(url, json={"field": "display_name", "format": {"vertical_align": "bottom"}})
        self.assertEqual(aligned.status_code, 200, aligned.text)
        self.assertEqual(aligned.json()["format"]["vertical_align"], "bottom")

        changed = self.client.patch(
            f"/api/case-manager/series/{series['id']}/children/{child['id']}/cell",
            json={"field": "display_name", "value": "换行后的子案名称\n第二行"},
        )
        self.assertEqual(changed.status_code, 200, changed.text)
        self.assertEqual(changed.json()["cell_formats"]["display_name"]["vertical_align"], "bottom")

        rejected = self.client.patch(url, json={"field": "display_name", "format": {"vertical_align": "middle"}})
        self.assertEqual(rejected.status_code, 200, rejected.text)
        self.assertNotIn("vertical_align", rejected.json()["format"])

    def test_arbitrary_spreadsheet_adjusts_columns_and_imports_rows(self):
        series = self.create_series()
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "业主明细"
        sheet.append(["序号", "业主姓名", "案由", "阶段", "委托人", "对方当事人", "法院", "房号", "合同金额"])
        sheet.append(["101", "示七某案", "商品房预售合同纠纷", "一审", "示七某", "开发商甲", "广州市中级人民法院", "18-02", "1200000.00"])
        output = io.BytesIO()
        workbook.save(output)
        preview = self.client.post(
            f"/api/case-manager/series/{series['id']}/import-preview",
            files={"file": ("owners.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(preview.status_code, 200, preview.text)
        payload = preview.json()
        self.assertEqual(payload["mode"], "smart_table")
        self.assertEqual(payload["errors"], [])
        self.assertTrue(any(item["label"] == "房号" and item["source"] == "custom" for item in payload["columns"]))
        amount_column = next(item for item in payload["columns"] if item["label"] == "合同金额")
        self.assertEqual(amount_column["semantic_type"], "amount")
        self.assertEqual(amount_column["recognized_as"], "金额列（自动汇总）")
        committed = self.client.post(
            f"/api/case-manager/series/{series['id']}/import-commit",
            json={"preview_id": payload["preview_id"]},
        )
        self.assertEqual(committed.status_code, 200, committed.text)
        detail = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        imported = next(item for item in detail["children"] if item["series_code"] == "101")
        self.assertEqual(imported["display_name"], "示七某案")
        self.assertEqual(imported["clients"], "示七某")
        self.assertEqual(imported["opponents"], "开发商甲")
        self.assertIn("18-02", imported["custom_values"].values())
        self.assertEqual(detail["stats"]["amount_source"], "table_columns")
        self.assertEqual(detail["stats"]["claim_totals"]["CNY"], "1200000.00")
        self.assertIn("合同金额", detail["stats"]["amount_columns"])
        exported = self.client.get(f"/api/case-manager/series/{series['id']}/export")
        workbook = openpyxl.load_workbook(io.BytesIO(exported.content), data_only=False)
        main_sheet = workbook["系列案件总表"]
        headers = [cell.value for cell in main_sheet[15]]
        amount_column = headers.index("合同金额") + 1
        code_column = headers.index("序号") + 1
        imported_row = next(
            row for row in range(16, main_sheet.max_row + 1)
            if str(main_sheet.cell(row=row, column=code_column).value) == "101"
        )
        amount_cell = main_sheet.cell(row=imported_row, column=amount_column)
        self.assertIsInstance(amount_cell.value, (int, float))
        self.assertEqual(amount_cell.number_format, "#,##0.00")

    def test_create_series_with_initial_spreadsheet_preview(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "首批子案"
        sheet.append(["序号", "案件名称", "委托人", "对方当事人", "标的金额"])
        sheet.append(["001", "示七某诉示丆某", "示七某", "示丆某", "￥1,250,000.50"])
        sheet.append(["002", "李四诉示丆某", "李四", "示丆某", "800000"])
        output = io.BytesIO()
        workbook.save(output)
        preview = self.client.post(
            "/api/case-manager/series-create/import-preview",
            files={"file": ("initial.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        self.assertEqual(preview.status_code, 200, preview.text)
        preview_data = preview.json()
        self.assertEqual(preview_data["counts"]["children"], 2)
        self.assertTrue(any(item["semantic_type"] == "amount" for item in preview_data["columns"]))
        created = self.client.post(
            "/api/case-manager/series",
            json={
                "title": "创建时导入系列",
                "case_type": "合同纠纷",
                "parties": [{"party_group": "client", "role": "委托人", "name": "共同委托人"}],
                "initial_import_preview_id": preview_data["preview_id"],
            },
        )
        self.assertEqual(created.status_code, 200, created.text)
        data = created.json()
        self.assertEqual(data["stats"]["total"], 2)
        self.assertEqual(data["stats"]["claim_totals"]["CNY"], "2050000.50")
        self.assertEqual(data["profile"]["table_settings"]["freeze_columns"], 2)

    def test_insert_multiple_children_at_position_and_reorder(self):
        series = self.create_series()
        original_ids = [item["id"] for item in series["children"]]
        inserted = self.client.post(
            f"/api/case-manager/series/{series['id']}/children/insert",
            json={"position": 1, "count": 2},
        )
        self.assertEqual(inserted.status_code, 200, inserted.text)
        inserted_ids = [item["id"] for item in inserted.json()["items"]]
        detail = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        visible_ids = [item["id"] for item in detail["children"] if item["source_status"] != "deleted"]
        self.assertEqual(visible_ids, [original_ids[0], *inserted_ids, original_ids[1]])

        requested = [original_ids[1], original_ids[0], *inserted_ids]
        reordered = self.client.put(
            f"/api/case-manager/series/{series['id']}/children/order",
            json={"child_ids": requested},
        )
        self.assertEqual(reordered.status_code, 200, reordered.text)
        detail = self.client.get(f"/api/case-manager/series/{series['id']}").json()
        visible_ids = [item["id"] for item in detail["children"] if item["source_status"] != "deleted"]
        self.assertEqual(visible_ids, requested)

        invalid = self.client.put(
            f"/api/case-manager/series/{series['id']}/children/order",
            json={"child_ids": requested[:-1]},
        )
        self.assertEqual(invalid.status_code, 400, invalid.text)


if __name__ == "__main__":
    unittest.main()
