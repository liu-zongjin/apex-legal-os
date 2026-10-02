from io import BytesIO
import zipfile

from backend.main import _article_editor_docx_bytes
from docx import Document
from PIL import Image


def test_editor_docx_keeps_a4_page_field_and_native_comments():
    content = (
        '<h1 style="font-family:黑体;color:#8b1a1a">一级标题</h1>'
        '<p style="text-indent:2em;line-height:1;margin-bottom:.5em">'
        '这是<strong>带批注的正文</strong>。</p>'
    )
    exported = _article_editor_docx_bytes(
        "测试文章",
        content,
        [{"content": "侧边栏批注", "selected_text": "带批注的正文", "created_at": "2026-08-10T12:00:00"}],
    )

    with zipfile.ZipFile(BytesIO(exported)) as package:
        assert "word/comments.xml" in package.namelist()
        assert b"commentRangeStart" in package.read("word/document.xml")
        assert "侧边栏批注".encode() in package.read("word/comments.xml")
        assert b"PAGE" in package.read("word/footer1.xml")


def test_editor_docx_without_notes_does_not_add_comments_part():
    exported = _article_editor_docx_bytes("无批注", "<p>正文</p>", [])
    with zipfile.ZipFile(BytesIO(exported)) as package:
        assert "word/comments.xml" not in package.namelist()


def test_editor_docx_uses_requested_page_margins():
    exported = _article_editor_docx_bytes(
        "自定义页边距", "<p>正文</p>", [],
        {"top": 2.2, "bottom": 2.8, "left": 2.6, "right": 3.1},
    )
    section = Document(BytesIO(exported)).sections[0]
    assert round(section.top_margin.cm, 1) == 2.2
    assert round(section.bottom_margin.cm, 1) == 2.8
    assert round(section.left_margin.cm, 1) == 2.6
    assert round(section.right_margin.cm, 1) == 3.1


def test_editor_docx_keeps_title_table_nested_lists_and_local_image(tmp_path, monkeypatch):
    from backend import main
    image_dir = tmp_path / "images"
    image_dir.mkdir()
    filename = "article-1-" + "a" * 32 + ".png"
    Image.new("RGB", (20, 20), "red").save(image_dir / filename)
    monkeypatch.setattr(main, "ARTICLE_IMAGE_DIR", image_dir)
    content = (
        '<p>正文<img src="/api/article-images/' + filename + '"></p>'
        '<table><tr><td>甲</td><td>乙</td></tr></table>'
        '<ol><li>第一项<ol><li>子项</li></ol></li></ol>'
    )
    doc = Document(BytesIO(_article_editor_docx_bytes("文档标题", content, [])))
    assert doc.paragraphs[0].text == "文档标题"
    assert len(doc.inline_shapes) == 1
    assert [[cell.text for cell in row.cells] for row in doc.tables[0].rows] == [["甲", "乙"]]
    assert "第一项" in [p.text for p in doc.paragraphs]
    assert "子项" in [p.text for p in doc.paragraphs]


def test_editor_docx_embeds_fetched_external_image():
    picture = BytesIO()
    Image.new("RGB", (20, 20), "blue").save(picture, format="PNG")
    url = "https://example.com/picture.png"
    raw = _article_editor_docx_bytes("标题", f'<p><img src="{url}"></p>', [], external_images={url: picture.getvalue()})
    assert len(Document(BytesIO(raw)).inline_shapes) == 1


def test_editor_docx_preserves_visual_page_breaks():
    body = '<p>第一页<span class="body-editor-export-page-break"></span>第二页</p>'
    raw = _article_editor_docx_bytes("标题", body, [])
    with zipfile.ZipFile(BytesIO(raw)) as package:
        assert b'w:type="page"' in package.read("word/document.xml")
