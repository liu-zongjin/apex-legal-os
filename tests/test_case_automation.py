import asyncio
import json
import sqlite3
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException

from backend.case_automation import (
    CONTRACT_SUBDIR,
    COURT_DOCUMENT_SUBDIR,
    ENTRUST_TYPE_DEFINITIONS,
    _set_court_document_default_tag,
    ensure_court_document_directory,
    _chinese_year_month,
    _contract_profile_fingerprint,
    analyze_docx,
    build_fee_clause,
    ensure_case_automation_schema,
    fill_contract_docx,
    get_folder_binding,
    register_case_automation_routes,
    safe_subfolder_name,
    series_child_folder_name,
    store_case_file,
    validate_case_folder,
)


class CaseAutomationTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.database = self.root / "case-manager.db"

        def get_db():
            connection = sqlite3.connect(self.database)
            connection.row_factory = sqlite3.Row
            connection.execute("PRAGMA foreign_keys=ON")
            return connection

        self.get_db = get_db
        with self.get_db() as connection:
            connection.execute(
                "CREATE TABLE case_manager_cases(id INTEGER PRIMARY KEY,title TEXT)"
            )
            connection.execute(
                "INSERT INTO case_manager_cases(id,title) VALUES (1,'测试案件')"
            )
            connection.commit()
        ensure_case_automation_schema(self.get_db)
        self.case_folder = self.root / "具体案件"
        self.case_folder.mkdir()
        stat = self.case_folder.stat()
        with self.get_db() as connection:
            connection.execute(
                """INSERT INTO case_manager_folder_bindings
                   (case_id,folder_path,real_path,device_id,inode,status,subdirs_json,
                    last_verified_at,created_at,updated_at)
                   VALUES (1,?,?,?,?,?,'{}','2026-07-28','2026-07-28','2026-07-28')""",
                (
                    str(self.case_folder),
                    str(self.case_folder.resolve()),
                    stat.st_dev,
                    stat.st_ino,
                    "active",
                ),
            )
            connection.commit()

    def tearDown(self):
        self.temporary.cleanup()

    def test_subdirectories_are_created_only_when_used_and_never_overwrite(self):
        self.assertFalse((self.case_folder / COURT_DOCUMENT_SUBDIR).exists())
        self.assertFalse((self.case_folder / CONTRACT_SUBDIR).exists())

        first = store_case_file(
            self.get_db, 1, "court_document", "裁定书.pdf", b"%PDF-test"
        )
        second = store_case_file(
            self.get_db, 1, "court_document", "裁定书.pdf", b"%PDF-second"
        )

        self.assertTrue((self.case_folder / COURT_DOCUMENT_SUBDIR).is_dir())
        self.assertFalse((self.case_folder / CONTRACT_SUBDIR).exists())
        self.assertEqual(Path(first["absolute_path"]).name, "裁定书.pdf")
        self.assertEqual(Path(second["absolute_path"]).name, "裁定书_2.pdf")
        self.assertEqual(Path(first["absolute_path"]).read_bytes(), b"%PDF-test")

    def test_legacy_judicial_delivery_folder_is_renamed_and_index_is_updated(self):
        legacy = self.case_folder / "#法院文书"
        legacy.mkdir()
        archived = legacy / "传票.pdf"
        archived.write_bytes(b"%PDF-test")
        with self.get_db() as connection:
            connection.execute(
                """INSERT INTO case_manager_files
                   (case_id,category,relative_path,absolute_path,sha256,size_bytes,created_at)
                   VALUES (1,'court_document',?,?,?,?,?)""",
                ("#法院文书/传票.pdf", str(archived), "digest", 9, "2026-08-21"),
            )
            connection.commit()

        binding = get_folder_binding(self.get_db, 1, required=True)
        target = self.case_folder / COURT_DOCUMENT_SUBDIR

        self.assertFalse(legacy.exists())
        self.assertEqual((target / "传票.pdf").read_bytes(), b"%PDF-test")
        self.assertEqual(binding["subdirs"]["court_document"], COURT_DOCUMENT_SUBDIR)
        with self.get_db() as connection:
            indexed = connection.execute(
                "SELECT relative_path,absolute_path FROM case_manager_files WHERE case_id=1"
            ).fetchone()
        self.assertEqual(indexed["relative_path"], f"{COURT_DOCUMENT_SUBDIR}/传票.pdf")
        self.assertEqual(indexed["absolute_path"], str((target / "传票.pdf").resolve()))

    def test_apex_archive_merges_collisions_and_updates_each_index(self):
        target = self.case_folder / COURT_DOCUMENT_SUBDIR
        target.mkdir()
        (target / "传票.pdf").write_bytes(b"existing")
        for legacy_name in ("#APEX电子送达归档", "#法院文书"):
            legacy = self.case_folder / legacy_name
            legacy.mkdir()
            archived = legacy / "传票.pdf"
            archived.write_bytes(legacy_name.encode())
            with self.get_db() as connection:
                connection.execute(
                    "INSERT INTO case_manager_files "
                    "(case_id,category,relative_path,absolute_path,sha256,size_bytes,created_at) "
                    "VALUES (1,'court_document',?,?,?,?,?)",
                    (f"{legacy_name}/传票.pdf", str(archived), "digest", 9, "2026-10-13"),
                )
        with patch("backend.case_automation._set_court_document_default_tag") as tag:
            get_folder_binding(self.get_db, 1, required=True)
            tag.assert_called_once_with(target.resolve())
        self.assertEqual((target / "传票.pdf").read_bytes(), b"existing")
        with self.get_db() as connection:
            rows = connection.execute("SELECT * FROM case_manager_files ORDER BY id").fetchall()
        for index, row in enumerate(rows, 2):
            self.assertEqual(row["relative_path"], f"{COURT_DOCUMENT_SUBDIR}/传票_{index}.pdf")
            self.assertTrue(Path(row["absolute_path"]).is_file())
        self.assertFalse((self.case_folder / "#APEX电子送达归档").exists())
        self.assertFalse((self.case_folder / "#法院文书").exists())

    def test_archive_creation_tags_only_folder_and_preserves_later_user_edits(self):
        with patch("backend.case_automation._set_court_document_default_tag") as tag:
            first = store_case_file(self.get_db, 1, "court_document", "传票.pdf", b"first")
            store_case_file(self.get_db, 1, "court_document", "传票.pdf", b"second")
            tag.assert_called_once_with((self.case_folder / COURT_DOCUMENT_SUBDIR).resolve())
        self.assertEqual(COURT_DOCUMENT_SUBDIR, "#Stella电子送达归档")
        self.assertEqual(Path(first["absolute_path"]).parent.name, COURT_DOCUMENT_SUBDIR)

    def test_archive_default_tag_is_important_red_and_retains_other_tags(self):
        folder = self.case_folder / COURT_DOCUMENT_SUBDIR
        folder.mkdir()
        with patch("backend.case_automation.sys.platform", "darwin"), \
             patch("backend.case_files._finder_tags_for_path", return_value=[
                 {"name": "待核对", "color_index": 3},
                 {"name": "重要", "color_index": 5},
             ]), \
             patch("backend.case_files._write_finder_tag_values") as write, \
             patch("backend.case_files._set_finder_info_color") as color:
            _set_court_document_default_tag(folder)
        write.assert_called_once_with(folder, ["待核对\n3", "重要\n6"])
        color.assert_called_once_with(folder, 6)

    def test_archive_tag_failure_does_not_block_writes(self):
        with patch("backend.case_automation.sys.platform", "darwin"), \
             patch("backend.case_files._write_finder_tag_values", side_effect=OSError("unsupported")):
            result = store_case_file(self.get_db, 1, "court_document", "传票.pdf", b"safe")
        self.assertEqual(Path(result["absolute_path"]).read_bytes(), b"safe")

    def test_archive_directory_never_follows_symlink(self):
        outside = self.root / "outside"
        outside.mkdir()
        (self.case_folder / COURT_DOCUMENT_SUBDIR).symlink_to(outside, target_is_directory=True)
        with self.assertRaises(HTTPException):
            ensure_court_document_directory(self.case_folder)
        self.assertEqual(list(outside.iterdir()), [])

    def test_broad_directories_are_rejected(self):
        for path in ("/", str(Path.home())):
            with self.assertRaises(HTTPException):
                validate_case_folder(path)

    def test_binding_survives_device_id_change_and_repairs_stored_identity(self):
        current = self.case_folder.stat()
        with self.get_db() as connection:
            connection.execute(
                "UPDATE case_manager_folder_bindings SET device_id=? WHERE case_id=1",
                (int(current.st_dev) + 1,),
            )
            connection.commit()

        binding = get_folder_binding(self.get_db, 1, required=True)

        self.assertTrue(binding["available"])
        self.assertEqual(binding["status"], "active")
        self.assertEqual(binding["device_id"], current.st_dev)
        with self.get_db() as connection:
            stored = connection.execute(
                "SELECT device_id FROM case_manager_folder_bindings WHERE case_id=1"
            ).fetchone()
        self.assertEqual(stored["device_id"], current.st_dev)

    def test_binding_still_rejects_a_replaced_directory(self):
        original = self.root / "原目录"
        self.case_folder.rename(original)
        self.case_folder.mkdir()

        binding = get_folder_binding(self.get_db, 1)

        self.assertFalse(binding["available"])
        self.assertEqual(binding["status"], "replaced")
        with self.assertRaises(HTTPException):
            get_folder_binding(self.get_db, 1, required=True)

    def _seed_series_cases(self):
        with self.get_db() as connection:
            connection.execute("ALTER TABLE case_manager_cases ADD COLUMN case_kind TEXT")
            connection.execute("ALTER TABLE case_manager_cases ADD COLUMN parent_series_id INTEGER")
            connection.execute(
                "UPDATE case_manager_cases SET case_kind='series', title='测试系列案' WHERE id=1"
            )
            connection.execute(
                """INSERT INTO case_manager_cases(id,title,case_kind,parent_series_id)
                   VALUES (2,'测试系列案·范国华案','series_child',1)"""
            )
            connection.execute(
                """INSERT INTO case_manager_cases(id,title,case_kind,parent_series_id)
                   VALUES (3,'测试系列案·林远合案','series_child',1)"""
            )
            connection.execute(
                """INSERT INTO case_manager_cases(id,title,case_kind,parent_series_id)
                   VALUES (4,'未绑定系列案','series',NULL)"""
            )
            connection.execute(
                """INSERT INTO case_manager_cases(id,title,case_kind,parent_series_id)
                   VALUES (5,'未绑定系列案·赵太宇案','series_child',4)"""
            )
            connection.commit()

    @staticmethod
    def _route_app():
        class RouteApp:
            def __init__(self):
                self.routes = {}

            def _decorator(self, method, path):
                def register(fn):
                    self.routes[(method, path)] = fn
                    return fn
                return register

            def get(self, path): return self._decorator("GET", path)
            def put(self, path): return self._decorator("PUT", path)
            def post(self, path): return self._decorator("POST", path)
            def delete(self, path): return self._decorator("DELETE", path)

        return RouteApp()

    @staticmethod
    def _json_request(payload):
        class JsonRequest:
            async def json(self):
                return payload

        return JsonRequest()

    def test_series_child_folder_is_created_and_bound_inside_series_folder(self):
        self._seed_series_cases()
        app = self._route_app()
        register_case_automation_routes(app, self.get_db)

        result = asyncio.run(
            app.routes[("POST", "/api/case-manager/cases/{case_id}/folder-binding/create-in-series")](
                2, self._json_request({})
            )
        )

        expected = self.case_folder / "范国华案"
        self.assertTrue(expected.is_dir())
        self.assertEqual(Path(result["real_path"]).resolve(), expected.resolve())
        binding = get_folder_binding(self.get_db, 2)
        self.assertTrue(binding["bound"])
        self.assertTrue(binding["available"])
        stored = store_case_file(self.get_db, 2, "court_document", "送达回证.pdf", b"%PDF-test")
        self.assertEqual(
            Path(stored["absolute_path"]).resolve().parent,
            expected.resolve() / COURT_DOCUMENT_SUBDIR,
        )

    def test_series_child_folder_name_conflict_gets_a_suffix(self):
        self._seed_series_cases()
        (self.case_folder / "范国华案").mkdir()
        app = self._route_app()
        register_case_automation_routes(app, self.get_db)

        result = asyncio.run(
            app.routes[("POST", "/api/case-manager/cases/{case_id}/folder-binding/create-in-series")](
                2, self._json_request({})
            )
        )

        self.assertTrue((self.case_folder / "范国华案_2").is_dir())
        self.assertEqual(Path(result["real_path"]).name, "范国华案_2")

    def test_series_child_folder_accepts_sanitized_custom_name(self):
        self._seed_series_cases()
        app = self._route_app()
        register_case_automation_routes(app, self.get_db)

        result = asyncio.run(
            app.routes[("POST", "/api/case-manager/cases/{case_id}/folder-binding/create-in-series")](
                3, self._json_request({"name": "林远合/二审:庭"})
            )
        )

        expected = self.case_folder / "林远合_二审_庭"
        self.assertTrue(expected.is_dir())
        self.assertEqual(Path(result["real_path"]).name, "林远合_二审_庭")

    def test_series_child_folder_rejects_non_series_child_and_unbound_series(self):
        self._seed_series_cases()
        app = self._route_app()
        register_case_automation_routes(app, self.get_db)
        create = app.routes[("POST", "/api/case-manager/cases/{case_id}/folder-binding/create-in-series")]

        with self.assertRaises(HTTPException) as not_child:
            asyncio.run(create(1, self._json_request({})))
        self.assertEqual(not_child.exception.status_code, 400)

        with self.assertRaises(HTTPException) as unbound:
            asyncio.run(create(5, self._json_request({})))
        self.assertEqual(unbound.exception.status_code, 409)
        self.assertIn("系列案件", unbound.exception.detail)

    def test_series_child_folder_name_helpers(self):
        self.assertEqual(series_child_folder_name("保利系列案", "保利系列案·范国华案"), "范国华案")
        self.assertEqual(series_child_folder_name("保利系列案", "自定义子案"), "自定义子案")
        self.assertEqual(safe_subfolder_name('a/b<c>:d"e|f?g*h'), "a_b_c__d_e_f_g_h")
        self.assertEqual(safe_subfolder_name(" .. "), "")

    def test_series_child_binds_existing_folder_inside_series_folder(self):
        self._seed_series_cases()
        existing = self.case_folder / "现有子文件夹"
        existing.mkdir()
        app = self._route_app()
        register_case_automation_routes(app, self.get_db)
        bind = app.routes[("POST", "/api/case-manager/cases/{case_id}/folder-binding/bind-in-series")]

        result = asyncio.run(bind(2, self._json_request({"name": "现有子文件夹"})))

        self.assertEqual(Path(result["real_path"]).resolve(), existing.resolve())
        binding = get_folder_binding(self.get_db, 2)
        self.assertTrue(binding["bound"])
        self.assertTrue(binding["available"])

        with self.assertRaises(HTTPException) as missing:
            asyncio.run(bind(3, self._json_request({"name": "不存在"})))
        self.assertEqual(missing.exception.status_code, 404)

    def test_series_folder_options_list_existing_folders(self):
        self._seed_series_cases()
        (self.case_folder / "范国华案旧档").mkdir()
        (self.case_folder / COURT_DOCUMENT_SUBDIR).mkdir()
        (self.case_folder / CONTRACT_SUBDIR).mkdir()
        (self.case_folder / ".隐藏目录").mkdir()
        (self.case_folder / "普通文件.txt").write_bytes(b"text")
        app = self._route_app()
        register_case_automation_routes(app, self.get_db)
        options = app.routes[("GET", "/api/case-manager/cases/{case_id}/folder-binding/series-options")]

        bound = options(2)
        self.assertTrue(bound["bound"])
        self.assertEqual([item["name"] for item in bound["folders"]], ["范国华案旧档"])

        unbound = options(5)
        self.assertFalse(unbound["bound"])
        self.assertEqual(unbound["folders"], [])


    def test_split_run_placeholder_and_table_placeholder_are_filled(self):
        from docx import Document

        template = self.root / "template.docx"
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("案件：{{案件")
        paragraph.add_run("名称}}")
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "委托人：{{委托人}}"
        document.save(template)

        _, mappings, warnings = analyze_docx(template)
        self.assertFalse(warnings)
        for mapping in mappings:
            mapping["confirmed"] = True
        content, missing = fill_contract_docx(
            template,
            mappings,
            {"case.title": "示丆某诉示一某案", "client.name": "示丆某"},
        )
        output = self.root / "filled.docx"
        output.write_bytes(content)
        rendered = Document(output)
        self.assertIn("示丆某诉示一某案", rendered.paragraphs[0].text)
        self.assertIn("示丆某", rendered.tables[0].cell(0, 0).text)
        self.assertEqual(missing, [])

    def test_unselected_checkbox_does_not_loop_during_contract_fill(self):
        from docx import Document

        template = self.root / "checkbox-template.docx"
        document = Document()
        document.add_paragraph("☐ 代为承认、放弃、变更诉讼请求")
        document.save(template)

        content, missing = fill_contract_docx(
            template,
            [],
            {
                "engagement.authorization_ids": "",
                "engagement.fee_clause": "",
            },
        )
        output = self.root / "checkbox-filled.docx"
        output.write_bytes(content)
        rendered = Document(output)
        self.assertEqual(rendered.paragraphs[0].text, "☐ 代为承认、放弃、变更诉讼请求")
        self.assertEqual(missing, [])

    def test_fee_presets_generate_selected_calculation_method(self):
        recovery = build_fee_clause(
            {
                "fee_plan": "base_risk",
                "fee_config": {
                    "base_amount": "30000",
                    "risk_method": "recovery_rate",
                    "rate": "15",
                    "payment_terms": "签约后三日内支付基础费",
                },
            }
        )
        reduction = build_fee_clause(
            {
                "fee_plan": "pure_risk",
                "fee_config": {"risk_method": "reduction_rate", "rate": "20"},
            }
        )
        self.assertIn("基础律师服务费为30000元", recovery)
        self.assertIn("实际收回的款项", recovery)
        self.assertIn("15%", recovery)
        self.assertIn("不预先收取基础律师服务费", reduction)
        self.assertIn("最终承担金额之差额", reduction)

    def test_known_template_fee_section_and_authorizations_are_applied(self):
        from docx import Document

        template = self.root / "known-template.docx"
        document = Document()
        document.add_paragraph("□调查取证；")
        document.add_paragraph("□提起诉讼；")
        document.add_paragraph("□申请执行；")
        document.add_paragraph("（一）收费方案")
        document.add_paragraph("方案A. 固定收费人民币【】元（大写【】元整）")
        document.add_paragraph("方案B. 分阶段固定收费")
        document.add_paragraph("一审代理费人民币【】元（大写人民币【】万元整）")
        document.add_paragraph("二审代理费人民币【】元（大写人民币【】万元整）")
        document.add_paragraph("方案C. 基础+风险收费")
        document.add_paragraph("基础代理费人民币【】元（大写人民币【】万元整）")
        document.add_paragraph("2.风险代理费")
        document.add_paragraph("差额部分的【】%作为风险代理费")
        document.add_paragraph("无需付款时支付【】万元")
        document.add_paragraph("方案D. 分段回款风险收费")
        document.add_paragraph("（二）收款账号")
        document.save(template)

        content, _ = fill_contract_docx(
            template,
            [],
            {
                "engagement.authorization_ids": "investigate,enforce",
                "engagement.fee_plan": "fixed",
                "engagement.fee_config_json": json.dumps({"base_amount": "30000"}),
            },
        )
        output = self.root / "known-filled.docx"
        output.write_bytes(content)
        paragraphs = [item.text for item in Document(output).paragraphs]
        self.assertIn("☑调查取证；", paragraphs)
        self.assertIn("☐提起诉讼；", paragraphs)
        self.assertIn("☑申请执行；", paragraphs)
        self.assertIn("固定收费人民币30000元（大写叁万元整）", paragraphs)
        self.assertFalse(
            any("方案B" in text or "方案C" in text or "方案D" in text for text in paragraphs)
        )

    def test_contract_authority_text_exactly_matches_editor_selection(self):
        from docx import Document

        template = self.root / "authority-lines-template.docx"
        document = Document()
        document.add_paragraph("第二条 委托代理权限（划√为确认，划×为否认）")
        document.add_paragraph("特别授权，包括（选择项以划√为准）：")
        document.add_paragraph("□调查取证；")
        document.add_paragraph("□提起诉讼；")
        document.add_paragraph("□应诉答辩；")
        document.add_paragraph("第三条 乙方的义务")
        document.save(template)

        selected = [
            "代为申请强制执行、参与执行程序、代为协助法院执行；",
            "代为领取案款；",
            "代为签收相关法律文件；",
        ]
        content, _ = fill_contract_docx(
            template,
            [],
            {
                "engagement.authority_lines": "\n".join(selected),
                "engagement.authorization_ids": "investigate",
                "engagement.fee_clause": "",
            },
        )
        output = self.root / "authority-lines-filled.docx"
        output.write_bytes(content)
        paragraphs = [item.text for item in Document(output).paragraphs]
        self.assertEqual(
            paragraphs[2:5],
            [f"☑{line}" for line in selected],
        )
        self.assertNotIn("□提起诉讼；", paragraphs)

    def test_contract_matter_and_signature_party_follow_current_case(self):
        from docx import Document

        template = self.root / "matter-template.docx"
        document = Document()
        document.add_paragraph(
            "乙方接受甲方委托，委派律师在下列□事项□案件中担任甲方的委托代理人："
        )
        document.add_paragraph("甲方：旧案件委托人        乙方：示例律师事务所（虚构）")
        document.save(template)

        content, _ = fill_contract_docx(
            template,
            [],
            {
                "client.name": "当前案件委托人",
                "engagement.scope": "担任再审审查程序申请人的委托代理人",
                "engagement.matter_kind": "matter",
                "engagement.authorization_ids": "",
                "engagement.fee_clause": "",
            },
        )
        output = self.root / "matter-filled.docx"
        output.write_bytes(content)
        paragraphs = [item.text for item in Document(output).paragraphs]
        self.assertEqual(
            paragraphs[0],
            "乙方接受甲方委托，委派律师在下列事项中担任甲方的委托代理人：",
        )
        self.assertIn("甲方：当前案件委托人", paragraphs[1])
        self.assertNotIn("旧案件委托人", paragraphs[1])

    def test_contract_number_year_and_cover_month_follow_generation_date(self):
        from docx import Document

        template = self.root / "date-template.docx"
        document = Document()
        document.add_paragraph("（2025）粤明思律民字第        号")
        document.add_paragraph("二〇二五年一月")
        document.save(template)

        today = date.today()
        content, _ = fill_contract_docx(
            template,
            [],
            {
                "contract.current_year": str(today.year),
                "contract.cover_month": _chinese_year_month(today),
                "engagement.authorization_ids": "",
                "engagement.fee_clause": "",
            },
        )
        output = self.root / "date-filled.docx"
        output.write_bytes(content)
        paragraphs = [item.text for item in Document(output).paragraphs]
        self.assertEqual(
            paragraphs[0],
            f"（{today.year}）粤明思律民字第        号",
        )
        self.assertEqual(paragraphs[1], _chinese_year_month(today))

    def test_dynamic_stage_and_single_risk_tier_remove_template_markers(self):
        from docx import Document

        template = self.root / "fee-options.docx"
        document = Document()
        for text in (
            "（一）收费方案",
            "方案A.固定收费人民币【】元（大写【】元整）",
            "方案B.分阶段固定收费",
            "一审代理费人民币【】元",
            "二审代理费人民币【】元",
            "方案C.基础风险收费",
            "基础代理费人民币【】元",
            "风险代理费",
            "差额部分【】%",
            "无需付款支付【】万元",
            "方案D.分段回款风险收费",
            "基础代理费人民币【】元",
            "全额回款达到【】万元支付【】万元，未达到【】万元",
            "第1档【】万元以下按【】%",
            "第2档【】至【】万元按【】%",
            "第3档【】至【】万元按【】%",
            "第4档【】万元以上按【】%",
            "（二）收款账号",
        ):
            document.add_paragraph(text)
        document.save(template)

        staged, _ = fill_contract_docx(
            template, [], {
                "engagement.authorization_ids": "",
                "engagement.fee_plan": "staged_fixed",
                "engagement.fee_config_json": json.dumps({
                    "stage_fees": [
                        {"stage": "再审审查", "amount": "40000"},
                        {"stage": "强制执行", "amount": "20000"},
                    ]
                }, ensure_ascii=False),
            }
        )
        staged_path = self.root / "staged.docx"
        staged_path.write_bytes(staged)
        staged_text = "\n".join(p.text for p in Document(staged_path).paragraphs)
        self.assertIn("再审审查代理费", staged_text)
        self.assertIn("强制执行代理费", staged_text)
        self.assertNotIn("一审代理费", staged_text)
        self.assertNotRegex(staged_text, r"方案[ABCD]|【|】")

        tiered, _ = fill_contract_docx(
            template, [], {
                "engagement.authorization_ids": "",
                "engagement.fee_plan": "recovery_tiered",
                "engagement.fee_config_json": json.dumps({
                    "base_amount": "10000",
                    "tiers": [{"lower": "", "upper": "", "rate": "12"}],
                }, ensure_ascii=False),
            }
        )
        tiered_path = self.root / "tiered.docx"
        tiered_path.write_bytes(tiered)
        tiered_text = "\n".join(p.text for p in Document(tiered_path).paragraphs)
        self.assertIn("全部回款金额，按 12% 计收风险代理费", tiered_text)
        self.assertNotIn("全额回款达到", tiered_text)
        self.assertNotRegex(tiered_text, r"方案[ABCD]|【|】")

    def test_footer_uses_dynamic_total_pages_field(self):
        from docx import Document

        template = self.root / "footer.docx"
        document = Document()
        document.add_paragraph("正文")
        document.sections[0].footer.paragraphs[0].text = "第 1 页 共 11 页"
        document.save(template)
        content, _ = fill_contract_docx(
            template, [], {
                "engagement.authorization_ids": "",
                "engagement.fee_clause": "",
            }
        )
        output = self.root / "footer-filled.docx"
        output.write_bytes(content)
        with zipfile.ZipFile(output) as archive:
            footer = archive.read("word/footer1.xml").decode("utf-8")
            settings = archive.read("word/settings.xml").decode("utf-8")
        self.assertIn("NUMPAGES", footer)
        self.assertNotIn("共 11 页", footer)
        self.assertIn("updateFields", settings)

    def test_contract_entrust_types_share_authorization_generator_wording(self):
        arbitration = next(
            item for item in ENTRUST_TYPE_DEFINITIONS if item["id"] == "arbApplicant"
        )
        execution = next(
            item for item in ENTRUST_TYPE_DEFINITIONS if item["id"] == "execApplicant"
        )
        self.assertEqual(arbitration["stage"], "仲裁")
        self.assertEqual(
            arbitration["matter"],
            "在本案仲裁程序中担任申请人的委托代理人",
        )
        self.assertIn("代为选择仲裁员；", arbitration["authority_lines"])
        self.assertIn("代为领取案款；", execution["authority_lines"])

    def test_preview_fingerprint_rejects_changed_form_but_ignores_timestamps(self):
        profile = {
            "template_id": 2,
            "lawyers": [{"name": "示丁某"}],
            "representation_stage": "再审审查",
            "fee_config": {"base_amount": "50000"},
            "updated_at": "2026-07-28 23:00:00",
        }
        first = _contract_profile_fingerprint(74, 2, profile)
        profile["updated_at"] = "2026-07-28 23:10:00"
        self.assertEqual(first, _contract_profile_fingerprint(74, 2, profile))
        profile["fee_config"]["base_amount"] = "60000"
        self.assertNotEqual(first, _contract_profile_fingerprint(74, 2, profile))


if __name__ == "__main__":
    unittest.main()
