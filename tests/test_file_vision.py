"""Capability/transport/file integration, without OCR or paid model requests."""
import asyncio
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch
from fastapi import FastAPI
from fastapi.testclient import TestClient
import httpx

from backend import ai_capabilities, ai_custom, ai_input, ai_context, ai_responses, case_files, cc2api_bridge, main


class FileVisionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.db_path = self.root / 'test.db'
        self.db_patch = patch.object(main, 'DB_PATH', self.db_path); self.db_patch.start()
        main.init_db()
        with main.get_db() as conn:
            conn.execute("INSERT INTO case_manager_cases(id,title,created_at,updated_at) VALUES(901,'匿名案件','2026-10-28','2026-10-28')")
        import fitz
        doc = fitz.open()
        from PIL import Image
        import io
        buffer = io.BytesIO(); Image.new('RGB', (500, 500), 'white').save(buffer, format='PNG')
        for index in range(3):
            doc.new_page().insert_image(fitz.Rect(0, 0, 500, 500), stream=buffer.getvalue())
        self.pdf = self.root / 'scan.pdf'; doc.save(self.pdf); doc.close()
        self.captured = []
        async def call(model, messages):
            ai_context.guard_messages(messages, 5000, model=model)
            self.captured.append(messages)
            return '已逐页读取，关键条款见 [FILE:1] 和 [FILE:3]。', '', model
        app = FastAPI()
        case_files.register_case_file_routes(app, main.get_db, self.root,
            parse_file=lambda *args: (_ for _ in ()).throw(AssertionError('visual route must bypass OCR/text parser')),
            call_ai=call)
        self.client = TestClient(app)
        self.attach_patch = patch.object(case_files, '_AI_ATTACHMENT_ROOT', self.root / 'sessions'); self.attach_patch.start()
        self.bound_patch = patch.object(case_files, '_bound_root', return_value=self.root); self.bound_patch.start()

    def tearDown(self):
        self.bound_patch.stop(); self.attach_patch.stop(); self.db_patch.stop(); self.temp.cleanup()

    def assert_visual_request(self, result):
        self.assertEqual(result.status_code, 200, result.text)
        self.assertTrue(result.json()['file_reading']['complete'])
        images = [block for message in self.captured[-1] if isinstance(message['content'], list)
                  for block in message['content'] if block['type'] == 'image_url']
        self.assertEqual(len(images), 3)
        self.assertTrue(all(block['image_url']['url'].startswith('data:image/jpeg;base64,') for block in images))
        self.assertNotIn('base64', result.text)
        self.assertTrue(any('第 3 页' in source['locator'] for source in result.json()['cited_sources']))

    def test_upload_uses_selected_model_for_all_visual_pages_without_ocr(self):
        prepared = self.client.post('/api/ai-ask/local-files/prepare', data={'model': 'deepseek-flash'},
            files={'files': ('scan.pdf', self.pdf.read_bytes(), 'application/pdf')})
        self.assertEqual(prepared.status_code, 200, prepared.text)
        answer = self.client.post('/api/ai-ask/local-files/ask', data={
            'session_id': prepared.json()['session_id'], 'model': 'deepseek-flash', 'question': '完整分析这份扫描件'})
        self.assert_visual_request(answer)

    def test_case_files_use_same_visual_pipeline(self):
        prepared = self.client.post('/api/case-manager/cases/901/local-files/ai/prepare', json={
            'model': 'deepseek-flash', 'paths': ['scan.pdf']})
        self.assertEqual(prepared.status_code, 200, prepared.text)
        self.assertEqual(prepared.json()['consent_required'], [])
        self.assertTrue(prepared.json()['items'][0]['vision'])
        answer = self.client.post('/api/case-manager/cases/901/local-files/ask', data={
            'model': 'deepseek-flash', 'paths': '["scan.pdf"]', 'question': '完整分析扫描件'})
        self.assert_visual_request(answer)

    def test_native_pdf_uses_complete_text_and_zero_image_requests(self):
        import fitz
        doc = fitz.open()
        for index in range(3):
            doc.new_page().insert_text((50, 50), f'PAGE_{index + 1}_COMPLETE_NATIVE_TEXT')
        native = self.root / 'native.pdf'; doc.save(native); doc.close()
        prepared = self.client.post('/api/ai-ask/local-files/prepare', data={'model': 'deepseek-flash'},
            files={'files': ('native.pdf', native.read_bytes(), 'application/pdf')})
        answer = self.client.post('/api/ai-ask/local-files/ask', data={
            'session_id': prepared.json()['session_id'], 'model': 'deepseek-flash', 'question': '总结文件内容'})
        self.assertEqual(answer.status_code, 200, answer.text)
        self.assertEqual(len(self.captured), 1)
        payload = json.dumps(self.captured[0])
        self.assertNotIn('image_url', payload)
        self.assertTrue(all(f'PAGE_{index}_COMPLETE_NATIVE_TEXT' in payload for index in (1, 2, 3)))

    def test_capability_is_not_tied_to_deepseek_and_custom_declaration_is_validated(self):
        for model in ('deepseek-flash', 'doubao:Doubao-Seed-2.0-pro',
                      'Pro/moonshotai/Kimi-K2.6', 'cc2api:claude-sonnet-4-6'):
            self.assertTrue(ai_capabilities.image_input(model), model)
        self.assertFalse(ai_capabilities.image_input('deepseek-v4-pro'))
        with patch.object(ai_custom, 'PATH', self.root / 'custom.json'):
            item = ai_custom.save({'name': '匿名 API', 'api_url': 'https://anonymous.invalid/v1',
                                   'api_key': '', 'protocol': 'chat', 'models': ['new-visual-model'],
                                   'vision_models': ['new-visual-model']})
            self.assertTrue(ai_capabilities.image_input(item['id'] + ':new-visual-model'))
            with self.assertRaises(ValueError):
                ai_custom.save({'vision_models': ['not-configured']}, item['id'])

    def test_visual_batches_send_native_material_only_once_and_keep_every_page(self):
        from backend import file_vision
        rows=[{'relative_path':'text','display_name':'匿名文字材料','chunk_index':0,'locator':'全文','chunk_text':'UNIQUE_NATIVE_ORIGINAL\n'+'原文否定条件。'*300}]
        sent=[]
        async def call(model,messages):
            ai_context.guard_messages(messages,5000,model=model)
            sent.append(json.dumps(messages,ensure_ascii=False))
            return '本批已阅读，文字原文约定未经书面同意不得转让。','',model
        async def run():
            token=ai_input.SCOPE.set(ai_input.new_scope('完整阅读'))
            try:
                return await file_vision.answer('完整阅读并比较',rows,[{'path':self.pdf,'name':'scan.pdf','visual_pages':list(range(8))}],'fixture-small','',call)
            finally:ai_input.SCOPE.reset(token)
        with patch.object(file_vision,'render_page',return_value='data:image/jpeg;base64,fixture'):
            result=asyncio.run(run())
        self.assertGreater(result[-1],1)
        self.assertEqual(sum('UNIQUE_NATIVE_ORIGINAL' in text for text in sent),1)
        # Eight original pages reach exactly one visual request apiece; synthesis has no images.
        self.assertEqual(sum(text.count('image_url')//2 for text in sent),8)
        self.assertEqual(len(sent),result[-1]+1)

    def test_unfinishable_visual_task_stops_before_paid_reads_or_rendering(self):
        from backend import file_vision
        from fastapi import HTTPException
        async def run():
            token=ai_input.SCOPE.set({'usage':{'model_requests':0,'limit':4}})
            try:
                return await file_vision.answer('完整分析',[],[{'path':self.pdf,'name':'scan.pdf','visual_pages':list(range(100))}],'fixture-small','',AsyncMock())
            finally:ai_input.SCOPE.reset(token)
        with patch.object(file_vision,'render_page') as render:
            with self.assertRaises(HTTPException) as error:asyncio.run(run())
            self.assertEqual(error.exception.status_code,413);render.assert_not_called()

    def test_old_word_upload_is_repaired_and_explicit_research_really_runs(self):
        from docx import Document
        path = self.root / 'report.docx'
        doc = Document(); doc.add_paragraph('第一类：合同继续履行。')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '第二类：解除权行使期间。'
        doc.add_paragraph('第三类：文末合同目的已经实现。'); doc.save(path)
        questions = []
        async def call(model, messages):
            payload = json.dumps(messages, ensure_ascii=False)
            self.assertIn('第二类：解除权行使期间', payload)
            self.assertIn('第三类：文末合同目的已经实现', payload)
            self.captured.append(messages)
            if '为实际检索提炼任务说明' in payload:
                return '检索第一类继续履行、第二类解除权期间、第三类合同目的的更高层级案例。', '', model
            self.assertIn('已实际检索到的匿名案例', payload)
            return '已检索并对比三类争点。[FILE:1]', '', model
        async def external(question, model):
            questions.append(question)
            return {'answer': '已实际检索到的匿名案例。',
                    'external_sources': [{'title': '匿名权威来源', 'url': 'https://example.org/judgment'}]}
        app = FastAPI()
        case_files.register_case_file_routes(app, main.get_db, self.root,
            parse_file=main._parse_case_local_file, call_ai=call, authoritative_answer=external)
        client = TestClient(app)
        prepared = client.post('/api/ai-ask/local-files/prepare', data={'model': 'deepseek-flash'},
            files={'files': ('report.docx', path.read_bytes(), 'application/octet-stream')})
        self.assertEqual(prepared.status_code, 200, prepared.text)
        session = self.root / 'sessions' / prepared.json()['session_id']
        manifest = json.loads((session / 'manifest.json').read_text())
        manifest[0]['parser_version'] = 'case-files-v2'
        manifest[0]['chunks'] = [{'chunk_index': 0, 'chunk_text': '仅剩旧版首段', 'locator': '旧索引'}]
        (session / 'manifest.json').write_text(json.dumps(manifest))
        with patch.object(case_files, '_general_web_context', AsyncMock(return_value=('公开网页原文。', [], ''))):
            answer = client.post('/api/ai-ask/local-files/ask', data={
                'model': 'deepseek-flash', 'session_id': prepared.json()['session_id'],
                'question': '全网检索该报告中三类问题更贴切、层级更高的裁判案例'})
        self.assertEqual(answer.status_code, 200, answer.text)
        self.assertEqual(len(questions), 3)
        self.assertIn('解除权行使期间', questions[1])
        self.assertIn('合同目的已经实现', questions[2])
        self.assertEqual(len(self.captured), 1)  # Local outline avoids a paid planning read.
        self.assertTrue(answer.json()['external_sources'])
        self.assertTrue(answer.json()['file_reading']['complete'])
        self.assertNotIn('仅剩旧版首段', json.dumps(self.captured[-1], ensure_ascii=False))


class VisionTransportTests(unittest.IsolatedAsyncioTestCase):
    async def test_request_budget_is_shared_by_nested_research_and_reserves_final_calls(self):
        from backend.ai_transport import FINAL
        token = ai_input.SCOPE.set(ai_input.new_scope('匿名工作'))
        try:
            ai_input.SCOPE.get()['usage']['limit'] = 4
            ai_input.before_request(); ai_input.before_request()
            child = ai_input.new_scope('内部补充检索')
            self.assertIs(child['usage'], ai_input.SCOPE.get()['usage'])
            with self.assertRaises(Exception): ai_input.before_request()
            final = FINAL.set(True)
            try:
                ai_input.before_request(); ai_input.before_request()
                with self.assertRaises(Exception): ai_input.before_request()
            finally: FINAL.reset(final)
            self.assertEqual(child['usage']['model_requests'], 4)
        finally:
            ai_input.SCOPE.reset(token)
    async def test_provider_metadata_discovers_unknown_visual_model_without_sending_a_picture(self):
        seen = []
        async def handler(request):
            seen.append(request)
            return httpx.Response(200, json={'data': [{'id': 'anonymous-vision', 'architecture': {'input_modalities': ['text', 'image']}}]})
        from contextlib import asynccontextmanager
        @asynccontextmanager
        async def client(**kwargs):
            async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as value:
                yield value
        with patch('backend.ai_client.api_config', return_value=('https://anonymous.invalid/v1/chat/completions', 'fixture-key', 'anonymous-vision')), \
             patch('backend.ai_transport.ai_http_client', client):
            self.assertTrue(await ai_capabilities.resolve_image_input('anonymous-vision-fixture'))
            self.assertTrue(await ai_capabilities.resolve_image_input('anonymous-vision-fixture'))
        self.assertEqual(len(seen), 1)
        self.assertEqual(seen[0].method, 'GET')
        self.assertEqual(str(seen[0].url), 'https://anonymous.invalid/v1/models')

    async def test_responses_translates_multimodal_parts_and_does_not_store_remote_history(self):
        seen = []
        def handler(request):
            seen.append(json.loads(request.content))
            return httpx.Response(200, json={'status': 'completed', 'output': [{'type': 'message', 'content': [{'type': 'output_text', 'text': 'answer'}]}]})
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            result = await ai_responses.call_responses(client, 'https://anonymous.invalid/responses', 'anonymous',
                [{'role': 'user', 'content': [{'type': 'text', 'text': 'read'},
                 {'type': 'image_url', 'image_url': {'url': 'data:image/jpeg;base64,YQ==', 'detail': 'high'}}]}], 1000, {})
        self.assertEqual(result.status_code, 200)
        self.assertEqual(seen[0]['input'][0]['content'][1]['type'], 'input_image')
        self.assertFalse(seen[0]['store'])

    async def test_claude_bridge_sends_images_via_structured_stdin_with_tools_disabled(self):
        seen = {}
        class Process:
            returncode = 0
            async def communicate(self, data=None):
                seen['input'] = json.loads(data)
                return (json.dumps({'type': 'result', 'subtype': 'success', 'result': '视觉回答。'}).encode() + b'\n', b'')
        async def create(*args, **kwargs):
            seen['args'] = args
            return Process()
        with patch.object(cc2api_bridge, 'client_path', return_value='/anonymous/claude'), \
             patch.object(cc2api_bridge.asyncio, 'create_subprocess_exec', create):
            result = await cc2api_bridge.call_cc2api('https://anonymous.invalid/v1/messages', 'fixture-key',
                'claude-sonnet-4-6', [{'role': 'user', 'content': [{'type': 'text', 'text': '第1页'},
                 {'type': 'image_url', 'image_url': {'url': 'data:image/jpeg;base64,YQ=='}}]}])
        self.assertEqual(result.status_code, 200)
        self.assertIn('--input-format', seen['args'])
        self.assertIn('--no-session-persistence', seen['args'])
        self.assertEqual(seen['args'][seen['args'].index('--tools') + 1], '')
        self.assertEqual(seen['input']['message']['content'][1]['type'], 'image')
        self.assertNotIn('fixture-key', ' '.join(seen['args']))

    async def test_image_data_is_not_counted_as_base64_text_tokens(self):
        messages = [{'role': 'user', 'content': [{'type': 'text', 'text': '第1页'},
                    {'type': 'image_url', 'image_url': {'url': 'data:image/jpeg;base64,' + 'A' * 1000000}}]}]
        self.assertLess(ai_input.size(messages), 5000)
