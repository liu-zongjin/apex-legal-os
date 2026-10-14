"""Isolated anonymous fixtures; never import the live application's database."""
import asyncio
import json
import sqlite3
import tempfile
import time
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

import fitz
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from docx import Document

from backend.case_workflows import register_case_workflow_routes, validate_classification, claims_from_draft


def pdf(path, pages=1, rotation=0):
    with fitz.open() as doc:
        for i in range(pages):
            page = doc.new_page()
            page.insert_text((60,80), f'Anonymous evidence page {i+1}')
            if rotation:
                page.set_rotation(rotation)
        doc.save(path)


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='stella-workflow-tests-')
        self.folder = Path(self.temp.name)
        self.root = self.folder / 'case'
        self.root.mkdir()
        db = self.folder / 'test.db'
        conn = sqlite3.connect(db)
        conn.execute('CREATE TABLE case_manager_cases(id INTEGER PRIMARY KEY,title TEXT,case_type TEXT,source_status TEXT)')
        conn.executemany('INSERT INTO case_manager_cases VALUES(?,?,?,?)', [(1,'匿名案件','民事','active'),(2,'另案','民事','active')])
        conn.execute('CREATE TABLE case_manager_folder_bindings(case_id INTEGER PRIMARY KEY,real_path TEXT,inode INTEGER,device_id INTEGER)')
        conn.execute('INSERT INTO case_manager_folder_bindings VALUES(1,?,?,?)',(str(self.root),self.root.stat().st_ino,self.root.stat().st_dev))
        conn.execute('INSERT INTO case_manager_folder_bindings VALUES(2,?,?,?)',(str(self.root),self.root.stat().st_ino,self.root.stat().st_dev))
        conn.commit();conn.close()
        @contextmanager
        def get_db():
            conn = sqlite3.connect(db);conn.row_factory=sqlite3.Row
            try:yield conn
            finally:conn.close()
        def parse(path, consent):
            return {'content':path.read_text()}
        async def ai(model, messages):
            data=json.loads(messages[-1]['content'])
            if 'claims' not in data and 'document_items' not in data:
                return json.dumps({'events':[{'source_id':source['id'],'date_text':'','actors':'','description':'材料记载事件','topic':'事项','quote':source['text'].strip()[:800]} for source in data['sources']]}), '', 'fixture-model'
            if 'document_items' in data:
                source=data['sources'][0]
                return json.dumps({'title':'匿名证据','summary':'材料摘要','purpose':'证明合同事实','quotes':[{'source_id':source['id'],'quote':source['text'].strip()}]}), '', 'fixture-model'
            results=[]
            for claim in data['claims']:
                quote=data['sources'][0]['text'].strip()
                results.append({'id':claim['id'],'status':'supported','reason':'匿名测试结果','evidence':[{'source_id':data['sources'][0]['id'],'quote':quote}]})
            return json.dumps({'results':results}), '', 'fixture-model'
        self.app=FastAPI()
        self.service=register_case_workflow_routes(self.app,get_db,Path(__file__).resolve().parents[1]/'frontend',parse_file=parse,call_ai=ai,require_action=lambda request:None)
        self.client=TestClient(self.app).__enter__()
        (self.root/'合同.txt').write_text('合同约定价款10000元。付款日为2026年1月1日。')

    def tearDown(self):
        self.client.__exit__(None,None,None)
        self.temp.cleanup()

    def create(self,body):
        return self.client.post('/api/case-workflows/cases/1/tasks',json=body)

    def done(self,job):
        for _ in range(200):
            result=self.client.get(f'/api/case-workflows/cases/1/tasks/{job["id"]}').json()
            if result['status'] not in {'queued','running'}:return result
            time.sleep(.01)
        self.fail('task did not finish')

    def fact(self):
        response=self.create({'kind':'fact_check','paths':['合同.txt'],'draft':'价款10000元。付款日为2026年1月1日。'})
        self.assertEqual(response.status_code,200,response.text)
        return self.done(response.json())

    def plan(self):
        pdf(self.root/'证据A.pdf',3)
        pdf(self.root/'证据B.pdf',2)
        response=self.create({'kind':'evidence_plan','paths':['证据A.pdf','证据B.pdf'],'draft':'合同约定价款。\n请求返还价款。'})
        self.assertEqual(response.status_code,200,response.text)
        job=self.done(response.json())
        self.assertEqual(job['status'],'completed',job)
        return job

    def test_document_plan_save_group_select_pages_compile_and_reference(self):
        job=self.plan()
        self.assertEqual(len(job['plan_entries']),2)
        self.assertTrue(job['plan_entries'][0]['summary_refs'])
        self.assertEqual(job['claims'][1]['document_locator'],'第2段 · 事项2')
        prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
        self.assertEqual(self.client.get(prefix+'/material-info',params={'path':'证据A.pdf'}).json()['pages'],3)
        self.assertEqual(self.client.get(prefix+'/material-preview',params={'path':'证据A.pdf'}).status_code,200)
        pixels=self.client.get(prefix+'/material-page',params={'path':'证据A.pdf','page':3})
        self.assertEqual(pixels.status_code,200);self.assertTrue(pixels.content.startswith(b'\x89PNG'))
        self.assertEqual(self.client.get(prefix+'/material-page',params={'path':'证据A.pdf','page':4}).status_code,422)
        self.assertEqual(self.client.get(prefix+'/material-info',params={'path':'合同.txt'}).status_code,404)
        self.assertEqual(self.client.get(prefix.replace('/cases/1/','/cases/2/')+'/material-info',params={'path':'证据A.pdf'}).status_code,404)
        entry={'title':'合同及付款','purpose':'证明履行情况','summary':'各材料的客观摘要','origin':'当事人提供','group':'履行情况',
            'claim_ids':['1'], 'materials':[{'path':'证据A.pdf','pages':'1,3'},{'path':'证据B.pdf','pages':'2'}]}
        saved=self.client.put(prefix+'/plan',json={'revision':1,'entries':[entry]})
        self.assertEqual(saved.status_code,200,saved.text)
        self.assertEqual(saved.json()['revision'],2)
        self.assertEqual(self.client.put(prefix+'/plan',json={'revision':1,'entries':[entry]}).status_code,409)
        body={'kind':'evidence_book','plan_id':job['id'],'plan_revision':2,'paths':['证据A.pdf','证据B.pdf'],'entries':[entry],'confirmed':True}
        self.assertEqual(self.create(body).status_code,422)  # Unassociated matter must be acknowledged.
        body['gaps_acknowledged']=True
        response=self.create(body);self.assertEqual(response.status_code,200,response.text)
        book=self.done(response.json());self.assertEqual(book['status'],'completed',book)
        mapping=book['mapping'][0]
        self.assertEqual((mapping['start'],mapping['end']),(1,3))
        self.assertEqual(mapping['source_mappings'][0]['selected_pages'],[1,3])
        self.assertEqual(mapping['source_mappings'][1]['start'],3)
        directory=self.root/'Stella办案成果'/book['id']
        with fitz.open(directory/next(a['path'] for a in book['artifacts'] if a['name']=='证据册.pdf')) as compiled:
            self.assertIn('page 1',compiled[0].get_text());self.assertIn('page 3',compiled[1].get_text());self.assertIn('page 2',compiled[2].get_text())
        workbook=load_workbook(directory/next(a['path'] for a in book['artifacts'] if a['name']=='证据目录.xlsx'))
        self.assertEqual(workbook['证据目录'].cell(2,9).value,'各材料的客观摘要')
        self.assertEqual(workbook['证据目录'].cell(2,10).value,'当事人提供')
        self.assertEqual(workbook['文书对应（内部）'].cell(2,6).value,'1-3')
        workbook.close()
        self.assertEqual(book['plan_revision'],2)

    def test_pages_invalid_fail_without_publishing_partial_book(self):
        from backend.evidence_planning import page_selection
        for value in ['0','3-1','1,1','2,1','4','1-2,2','abc']:
            with self.assertRaises(HTTPException):page_selection(value,3)
        job=self.plan()
        entry=job['plan_entries'][0]
        entry['materials'][0]['pages']='99'
        entry['purpose']='证明事实'
        response=self.create({'kind':'evidence_book','plan_id':job['id'],'plan_revision':1,'paths':['证据A.pdf','证据B.pdf'],
            'entries':[entry],'confirmed':True,'gaps_acknowledged':True})
        self.assertEqual(response.status_code,200,response.text)
        failed=self.done(response.json());self.assertEqual(failed['status'],'failed');self.assertFalse(failed['artifacts'])
        self.assertIn('选页',failed['error'])

    def test_legal_claim_cannot_be_certified_as_fact_and_omitted_page_is_gap(self):
        from backend.evidence_planning import gaps
        source={'id':'1','text':'合同约定价款10000元。','path':'a.pdf','sha256':'h','locator':'第 3 页','chunk_index':0}
        row=validate_classification({'id':'1','claim':'应返还价款。'},{'status':'supported','statement_type':'legal','evidence':[{'source_id':'1','quote':source['text']}]},[source])
        self.assertEqual(row['status'],'inference');self.assertEqual(gaps([row],[]),[])
        row['statement_type']='fact';row['status']='supported'
        entry={'claim_ids':['1'],'materials':[{'path':'a.pdf','pages':'1-2'}]}
        self.assertIn('选页',gaps([row],[entry])[0]['gap'])
        entry['materials'][0]['pages']='1,3';self.assertEqual(gaps([row],[entry]),[])

    def test_fact_review_carries_into_plan_without_another_classification(self):
        pdf(self.root/'证据A.pdf')
        fact=self.done(self.create({'kind':'fact_check','paths':['证据A.pdf'],'draft':'合同事实。'}).json())
        prefix=f'/api/case-workflows/cases/1/tasks/{fact["id"]}'
        self.client.put(prefix+'/review',json={'revision':1,'reviews':[{'id':'1','status':'rejected','note':'证明范围不足'}]})
        plan=self.done(self.create({'kind':'evidence_plan','paths':['证据A.pdf'],'draft':'合同事实。','reference_task_id':fact['id']}).json())
        self.assertEqual(plan['status'],'completed',plan)
        self.assertEqual(plan['results'][0]['review_status'],'rejected')
        self.assertEqual(plan['results'][0]['review_note'],'证明范围不足')
        self.assertEqual(plan['reference_revision'],2)
        response=self.create({'kind':'evidence_plan','paths':['证据A.pdf'],'draft':'新事实。','reference_task_id':fact['id']})
        self.assertEqual(response.status_code,409)

    def test_plan_summary_failure_resumes_without_reclassifying_or_duplicating_entries(self):
        pdf(self.root/'a.pdf');pdf(self.root/'b.pdf')
        original=self.service.call_ai
        calls={'summary':0,'classification':0}
        async def provider(model,messages):
            data=json.loads(messages[-1]['content'])
            if 'claims' not in data and 'document_items' not in data:
                return json.dumps({'events':[{'source_id':source['id'],'date_text':'','actors':'','description':'材料记载事件','topic':'事项','quote':source['text'].strip()[:800]} for source in data['sources']]}), '', 'fixture-model'
            if 'document_items' in data:
                calls['summary']+=1
                if calls['summary']==2:return 'invalid-json','','fixture'
            else:calls['classification']+=1
            return await original(model,messages)
        self.service.call_ai=provider
        job=self.done(self.create({'kind':'evidence_plan','paths':['a.pdf','b.pdf'],'draft':'合同事实。'}).json())
        self.assertEqual(job['status'],'failed');self.assertEqual(len(job['plan_entries']),1);self.assertFalse(job['artifacts'])
        prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
        resumed=self.client.post(prefix+'/resume')
        self.assertEqual(resumed.status_code,200)
        finished=self.done(resumed.json());self.assertEqual(finished['status'],'completed',finished)
        self.assertEqual(len(finished['plan_entries']),2)
        self.assertEqual(calls,{'summary':3,'classification':1})

    def test_plan_rejects_forged_associations_and_changed_document(self):
        job=self.plan();prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
        entries=job['plan_entries'];entries[0]['claim_ids']=['999']
        self.assertEqual(self.client.put(prefix+'/plan',json={'revision':1,'entries':entries}).status_code,422)
        entries[0]['claim_ids']=['1'];entries[0]['materials'][0]['path']='../outside.pdf'
        self.assertEqual(self.client.put(prefix+'/plan',json={'revision':1,'entries':entries}).status_code,422)
        # An actual document file remains a provenance input even after its text is parsed.
        (self.root/'定稿.txt').write_text('合同事实。')
        response=self.create({'kind':'evidence_plan','paths':['证据A.pdf'],'draft_path':'定稿.txt'})
        plan=self.done(response.json());self.assertEqual(plan['status'],'completed',plan)
        (self.root/'定稿.txt').write_text('改动后的定稿。')
        body={'kind':'evidence_book','plan_id':plan['id'],'plan_revision':1,'paths':['证据A.pdf'],
            'entries':[{'title':'证据','purpose':'证明合同事实','claim_ids':['1'],'materials':[{'path':'证据A.pdf','pages':''}]}],
            'confirmed':True,'gaps_acknowledged':True}
        self.assertEqual(self.create(body).status_code,409)

    def test_document_paragraph_numbers_do_not_shift_between_sentences(self):
        rows=claims_from_draft('第一句。第二句。\n\n第三句。')
        self.assertEqual([row['document_locator'] for row in rows],['第1段 · 事项1','第1段 · 事项2','第2段 · 事项3'])

    def test_legacy_stopped_book_resumes_after_document_plan_upgrade(self):
        pdf(self.root/'a.pdf')
        root,job=self.service.create(1,{'kind':'evidence_book','paths':['a.pdf'],'confirmed':True,
            'entries':[{'path':'a.pdf','title':'旧证据','purpose':'证明事实'}]})
        job['entries']=[{'path':'a.pdf','title':'旧证据','purpose':'证明事实'}]
        job['status']='cancelled'
        self.service.save(root,job)
        response=self.client.post(f'/api/case-workflows/cases/1/tasks/{job["id"]}/resume')
        self.assertEqual(response.status_code,200,response.text)
        finished=self.done(response.json())
        self.assertEqual(finished['status'],'completed',finished)
        self.assertEqual(finished['mapping'][0]['source_mappings'][0]['selected_pages'],[1])

    def test_converted_preview_tamper_is_rejected(self):
        from PIL import Image
        Image.new('RGB',(120,160),'white').save(self.root/'scan.png')
        root,job=self.service.create(1,{'kind':'evidence_book','paths':['scan.png'],'confirmed':True,
            'entries':[{'path':'scan.png','title':'图片','purpose':'证明事实'}]})
        job['status']='cancelled';self.service.save(root,job)
        prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
        self.assertEqual(self.client.get(prefix+'/material-info',params={'path':'scan.png'}).status_code,200)
        preview=root/'Stella办案成果'/job['id']/('preview-'+job['sources'][0]['sha256'][:20])/'material.pdf'
        preview.write_bytes(b'not the original preview')
        self.assertEqual(self.client.get(prefix+'/material-info',params={'path':'scan.png'}).status_code,409)

    def test_new_task_is_not_marked_interrupted_before_launch(self):
        root,job=self.service.create(1,{'kind':'fact_check','paths':['合同.txt'],'draft':'合同事实。'})
        self.assertEqual(self.service.load(1,job['id'])[1]['status'],'queued')
        self.service.pending[job['id']]=0  # An abandoned creation can still recover.
        self.assertEqual(self.service.load(1,job['id'])[1]['status'],'interrupted')

    def test_partial_reference_coverage_requires_review(self):
        from backend.evidence_planning import gaps
        row={'id':'1','claim':'合同与履行事实','statement_type':'fact','status':'supported','review_status':'pending',
            'evidence':[{'path':'a.pdf','locator':'第1页'},{'path':'b.pdf','locator':'第2页'}]}
        entry={'claim_ids':['1'],'materials':[{'path':'a.pdf','pages':'1'}]}
        self.assertTrue(gaps([row],[entry]))
        entry['materials'].append({'path':'b.pdf','pages':'2'})
        self.assertEqual(gaps([row],[entry]),[])

    def test_timeline_dates_keep_precision_and_reject_guessing(self):
        from backend.case_timeline import date_bounds,grounded_event,issues
        self.assertEqual(date_bounds('2024年2月')['sort_end'],'2024-02-29')
        self.assertEqual(date_bounds('2026年')['date_precision'],'year')
        self.assertEqual(date_bounds('2026年1月1日至2026年2月1日')['date_precision'],'range')
        for text in ['次日','2026年1月初','2026年2月30日','2月1日','不明']:
            self.assertEqual(date_bounds(text)['sort_start'],'')
        source={'id':'1','path':'a.txt','sha256':'h','locator':'第1段','chunk_index':0,'text':'甲于2026年1月1日签约，同日付款。'}
        row=grounded_event({'source_id':'1','quote':source['text'],'date_text':'2026年1月2日','actors':'甲','description':'签约','topic':'签约'},[source])
        self.assertEqual(row['date_text'],'');self.assertIn('未采用',row['date_warning'])
        proposal={'source_id':'1','quote':source['text'],'date_text':'2026年1月1日','actors':'甲','description':'签约','topic':'签约'}
        first=grounded_event(proposal,[source]);second=grounded_event({**proposal,'description':'付款'},[source])
        self.assertNotEqual(first['id'],second['id'])  # Distinct actions in the same quote cannot be lost.
        self.assertIsNone(grounded_event({**proposal,'quote':'原文没有这一句'},[source]))

    def test_timeline_material_only_review_manual_add_versions_and_source_changes(self):
        (self.root/'事件.txt').write_text('甲于2026年1月1日签订合同。\n甲于2026年2月支付价款。')
        job=self.done(self.create({'kind':'timeline','paths':['事件.txt']}).json())
        self.assertEqual(job['status'],'completed',job)
        self.assertTrue(job['timeline_events']);self.assertEqual(job['claims'],[])
        self.assertTrue(any(row['name']=='事实时间轴.pdf' for row in job['artifacts']))
        prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
        event=job['timeline_events'][0]
        self.assertEqual(self.client.get(prefix+'/timeline-source',params={'event_id':event['id']}).status_code,200)
        self.assertEqual(self.client.get(prefix.replace('/cases/1/','/cases/2/')+'/timeline-fragments').status_code,404)
        update={'id':event['id'],'date_text':'2026年1月1日','actors':'甲','description':event['description'],'status':'confirmed','note':''}
        self.assertEqual(self.client.put(prefix+'/timeline',json={'revision':1,'events':[update]}).status_code,422)
        update['status']='amended';update['note']='人工核对合同日期'
        fragments=self.client.get(prefix+'/timeline-fragments').json()['items']
        manual={'source_id':fragments[0]['id'],'quote':'甲于2026年2月支付价款。','date_text':'2026年2月','actors':'甲','description':'支付价款','topic':'付款'}
        response=self.client.put(prefix+'/timeline',json={'revision':1,'events':[update],'manual_events':[manual]})
        self.assertEqual(response.status_code,200,response.text)
        reviewed=response.json();self.assertEqual(reviewed['revision'],2);self.assertEqual(len(reviewed['timeline_events']),2)
        self.assertEqual(reviewed['timeline_events'][1]['date_precision'],'month')
        self.assertEqual(self.client.get(prefix+'/artifacts/0?revision=1').status_code,200)
        self.assertEqual(self.client.put(prefix+'/timeline',json={'revision':1,'events':[update]}).status_code,409)
        (self.root/'事件.txt').write_text('改动后的材料')
        self.assertEqual(self.client.get(prefix+'/timeline-source',params={'event_id':event['id']}).status_code,409)

    def test_timeline_reference_document_and_changed_source_during_export(self):
        from unittest.mock import patch
        from backend.case_timeline import artifacts
        job=self.done(self.create({'kind':'timeline','paths':['合同.txt'],'draft':'价款10000元。'}).json())
        self.assertEqual(job['status'],'completed')
        self.assertTrue(job['results'])
        self.assertTrue(any(row['name']=='事实核对报告.docx' for row in job['artifacts']))
        prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
        updates=[{'id':row['id'],'status':'confirmed','note':''} for row in job['timeline_events']]
        def altered(job,folder):
            artifacts(job,folder)
            (self.root/'合同.txt').write_text('导出中发生变化')
        with patch('backend.case_timeline.artifacts',side_effect=altered):
            response=self.client.put(prefix+'/timeline',json={'revision':1,'events':updates})
        self.assertEqual(response.status_code,409)
        self.assertEqual(self.client.get(prefix).json()['revision'],1)

    def test_timeline_duplicates_date_differences_and_exclusion(self):
        from backend.case_timeline import grounded_event,issues,review_events
        sources=[{'id':str(i),'path':f'{i}.txt','sha256':'h','locator':'第1段','chunk_index':0,'text':f'甲于2026年1月{i}日签约。'} for i in [1,2]]
        events=[grounded_event({'source_id':source['id'],'quote':source['text'],'date_text':f'2026年1月{source["id"]}日','actors':'甲','description':'签约','topic':'签约'},sources) for source in sources]
        self.assertEqual(issues(events)[0]['type'],'date_difference')
        updates=[{'id':row['id'],'status':'confirmed','note':''} for row in events]
        updates[1].update(status='excluded',note='同名事项实际为另案，复核后排除')
        reviewed=review_events(events,updates);self.assertEqual(issues(reviewed),[])
        with self.assertRaises(HTTPException):review_events(events,[{'id':'forged','status':'confirmed'}]*2)

    def test_timeline_failure_resumes_exact_extraction_checkpoints(self):
        original=self.service.call_ai;calls={'n':0}
        (self.root/'长材料.txt').write_text(('甲于2026年1月1日签约。'+'说明文字'*300+'\n')*9)
        async def provider(model,messages):
            calls['n']+=1
            if calls['n']==2:return 'broken','','fixture'
            return await original(model,messages)
        self.service.call_ai=provider
        job=self.done(self.create({'kind':'timeline','paths':['长材料.txt']}).json())
        self.assertEqual(job['status'],'failed');self.assertFalse(job['artifacts']);self.assertEqual(job['timeline_batches'],[0])
        self.client.post(f'/api/case-workflows/cases/1/tasks/{job["id"]}/resume')
        finished=self.done(job);self.assertEqual(finished['status'],'completed',finished)
        self.assertEqual(len(finished['timeline_batches']),len(set(finished['timeline_batches'])))

    def test_input_snapshot_is_independent_and_tampering_fails(self):
        root, job = self.service.create(1, {'kind':'fact_check','paths':['合同.txt'],'draft':'价款10000元。'})
        item = job['sources'][0]
        original = (root / item['path']).read_bytes()
        (root / item['path']).write_text('原件后来变化')
        copied = self.service.input_path(root, job, item)
        self.assertEqual(copied.read_bytes(), original)
        copied.write_text('篡改副本')
        with self.assertRaises(HTTPException) as failure:
            self.service.input_path(root, job, item)
        self.assertEqual(failure.exception.status_code, 409)

    def test_pdf_visible_annotations_and_widgets_survive_without_mutating_source(self):
        from backend.pdf_assembly import numbered_copy, append_document
        with fitz.open() as source:
            page = source.new_page(width=300, height=300)
            page.insert_text((20,40), 'Evidence')
            annotation = page.add_rect_annot(fitz.Rect(30,70,210,130))
            annotation.set_colors(stroke=(1,0,0), fill=(1,1,0))
            annotation.update()
            widget = fitz.Widget()
            widget.field_name = 'amount'
            widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
            widget.rect = fitz.Rect(30,170,200,210)
            widget.field_value = '10000'
            page.add_widget(widget)
            before = page.get_pixmap(alpha=False).samples
            with fitz.open() as merged:
                append_document(merged, source, Path('source.pdf'))
                output = numbered_copy(merged)
            with output:
                actual = output[0].get_pixmap(clip=fitz.Rect(0,0,300,300),alpha=False).samples
                self.assertEqual(actual, before)
                self.assertEqual(output[0].rect.height, 324)
            self.assertIsNotNone(source[0].first_annot)
            self.assertIsNotNone(source[0].first_widget)

    def test_fact_delivers_quotes_review_versions_and_cross_case_denial(self):
        job=self.fact();self.assertEqual(job['status'],'completed',job)
        self.assertEqual(len(job['results']),2)
        self.assertTrue(all(row['review_status']=='pending' for row in job['results']))
        prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
        self.assertEqual(self.client.get(prefix+'/source?source_id=1').status_code,200)
        self.assertEqual(self.client.get(f'/api/case-workflows/cases/2/tasks/{job["id"]}').status_code,404)
        response=self.client.put(prefix+'/review',json={'revision':1,'reviews':[{'id':row['id'],'status':'confirmed','note':'已核'} for row in job['results']]})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['revision'],2)
        self.assertEqual(self.client.get(prefix+'/artifacts/0?revision=1').status_code,200)
        self.assertEqual(self.client.get(prefix+'/artifacts/0?revision=99').status_code,404)
        self.assertEqual(self.client.put(prefix+'/review',json={'revision':1,'reviews':[]}).status_code,409)
        for artifact in response.json()['artifacts']:
            path=self.root/'Stella办案成果'/job['id']/artifact['path']
            if path.suffix=='.xlsx':
                wb=load_workbook(path);self.assertEqual(wb.active.max_row,3);wb.close()
            if path.suffix=='.docx':self.assertIn('事实核对',Document(path).paragraphs[0].text)

    def test_hallucinated_quote_is_never_support(self):
        row=validate_classification({'id':'1','claim':'主张'}, {'status':'supported','evidence':[{'source_id':'1','quote':'原文中没有这一句'}]}, [{'id':'1','text':'合同价款10000元。','path':'a','sha256':'h','locator':'第1页','chunk_index':0}])
        self.assertEqual(row['status'],'not_found');self.assertEqual(row['evidence'],[])

    def test_review_export_does_not_block_async_service_requests(self):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Event
        from unittest.mock import patch
        @self.app.get('/audit-ping')
        async def ping():return {'ok':True}
        job=self.fact()
        prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
        entered,released,responded=Event(),Event(),Event()
        original=self.service.fact_artifacts
        def slow(job,folder):
            entered.set();released.wait(3);original(job,folder)
        def request_ping():
            response=self.client.get('/audit-ping');responded.set();return response
        with patch.object(self.service,'fact_artifacts',side_effect=slow),ThreadPoolExecutor(max_workers=2) as workers:
            saving=workers.submit(self.client.put,prefix+'/review',json={'revision':1,'reviews':[{'id':row['id'],'status':'confirmed','note':''} for row in job['results']]})
            try:
                self.assertTrue(entered.wait(2))
                health=workers.submit(request_ping)
                self.assertTrue(responded.wait(1),'导出阻塞了事件循环')
            finally:released.set()
            self.assertEqual(saving.result(timeout=3).status_code,200)
            self.assertEqual(health.result(timeout=3).status_code,200)

    def test_review_rejects_structured_status_without_publishing(self):
        job=self.done(self.create({'kind':'fact_check','paths':['合同.txt'],'draft':'价款10000元。'}).json())
        prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
        response=self.client.put(prefix+'/review',json={'revision':1,'reviews':[{'id':row['id'],'status':{},'note':''} for row in job['results']]})
        self.assertEqual(response.status_code,422)
        self.assertEqual(self.client.get(prefix).json()['revision'],1)

    def test_fact_and_plan_export_rechecks_originals_before_publication(self):
        from unittest.mock import patch
        for kind,route,writer in [('fact_check','review','fact_artifacts'),('evidence_plan','plan','plan_artifacts')]:
            with self.subTest(kind=kind):
                pdf(self.root/'审查.pdf')
                job=self.done(self.create({'kind':kind,'paths':['审查.pdf'],'draft':'价款10000元。'}).json())
                self.assertEqual(job['status'],'completed')
                prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
                body={'revision':1}
                if kind=='fact_check':body['reviews']=[{'id':row['id'],'status':'confirmed','note':''} for row in job['results']]
                else:body['entries']=job['plan_entries']
                original=getattr(self.service,writer)
                def changed(job,folder):
                    original(job,folder)
                    (self.root/'审查.pdf').write_text('导出中原件改变')
                with patch.object(self.service,writer,side_effect=changed):
                    response=self.client.put(prefix+'/'+route,json=body)
                self.assertEqual(response.status_code,409,response.text)
                saved=self.client.get(prefix).json()
                self.assertEqual(saved['revision'],1)
                self.assertEqual(len(saved['versions']),1)
                self.assertFalse(any(p.name.startswith('v2-') for p in self.service.directory(self.root,job['id']).iterdir()))

    def test_material_change_blocks_reference_review_and_resume(self):
        job=self.fact();(self.root/'合同.txt').write_text('新的材料')
        prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
        self.assertEqual(self.client.get(prefix+'/source?source_id=1').status_code,409)
        response=self.client.put(prefix+'/review',json={'revision':1,'reviews':[{'id':row['id'],'status':'confirmed','note':''} for row in job['results']]})
        self.assertEqual(response.status_code,409)

    def test_evidence_order_page_mapping_sources_and_unique_outputs(self):
        pdf(self.root/'a.pdf',2);pdf(self.root/'b.pdf',1,90)
        before={name:(self.root/name).read_bytes() for name in ['a.pdf','b.pdf']}
        body={'kind':'evidence_book','paths':['b.pdf','a.pdf'],'confirmed':True,'entries':[{'path':'b.pdf','title':'第二份先编','purpose':'证明事实B'},{'path':'a.pdf','title':'第一份后编','purpose':'证明事实A'}]}
        jobs=[]
        for _ in range(2):
            response=self.create(body);self.assertEqual(response.status_code,200,response.text);jobs.append(self.done(response.json()))
        for job in jobs:
            self.assertEqual(job['status'],'completed',job)
            self.assertEqual([(m['start'],m['end']) for m in job['mapping']],[(1,1),(2,3)])
            directory=self.root/'Stella办案成果'/job['id']
            artifact=next(a for a in job['artifacts'] if a['name']=='证据册.pdf')
            with fitz.open(directory/artifact['path']) as doc:
                self.assertEqual(len(doc),3);self.assertEqual(len(doc.get_toc()),2)
                self.assertIn('Stella',doc[0].get_text());self.assertIn('Anonymous',doc[0].get_text())
            workbook=load_workbook(directory/next(a for a in job['artifacts'] if a['name']=='证据目录.xlsx')['path'])
            self.assertEqual(workbook.active.cell(2,5).value,1);self.assertEqual(workbook.active.cell(3,6).value,3);workbook.close()
        self.assertNotEqual(jobs[0]['id'],jobs[1]['id'])
        for name,value in before.items():self.assertEqual((self.root/name).read_bytes(),value)

    def test_validation_traversal_symlinks_empty_and_unconfirmed(self):
        for path in ['../outside.txt','/tmp/file','Stella办案成果/a']:
            self.assertGreaterEqual(self.create({'kind':'fact_check','paths':[path],'draft':'事实。'}).status_code,400)
        (self.root/'alias.txt').symlink_to(self.root/'合同.txt')
        self.assertEqual(self.create({'kind':'fact_check','paths':['alias.txt'],'draft':'事实。'}).status_code,409)
        pdf(self.root/'a.pdf')
        self.assertEqual(self.create({'kind':'evidence_book','paths':['a.pdf'],'entries':[]}).status_code,422)
        self.assertEqual(self.create({'kind':'fact_check','paths':['合同.txt'],'draft':''}).status_code,422)

    def test_model_missing_claim_fails_without_publishing(self):
        async def broken(*args):return '{"results":[]}', '', 'fixture'
        self.service.call_ai=broken
        job=self.fact();self.assertEqual(job['status'],'failed');self.assertFalse(job['artifacts'])

    def test_stop_and_resume_checkpoint_and_restart(self):
        started=thread_event=__import__('threading').Event()
        async def slow(model,messages):
            started.set();await asyncio.sleep(100)
        self.service.call_ai=slow
        response=self.create({'kind':'fact_check','paths':['合同.txt'],'draft':'价款10000元。'})
        job=response.json();self.assertTrue(started.wait(2))
        prefix=f'/api/case-workflows/cases/1/tasks/{job["id"]}'
        result=self.client.post(prefix+'/cancel');self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()['status'],'cancelled')
        async def good(model,messages):
            data=json.loads(messages[-1]['content']);return json.dumps({'results':[{'id':c['id'],'status':'not_found','reason':'未找到','evidence':[]} for c in data['claims']]}),'','fixture'
        self.service.call_ai=good
        self.assertEqual(self.client.post(prefix+'/resume').status_code,200)
        self.assertEqual(self.done(job)['status'],'completed')
        root,record=self.service.load(1,job['id']);record['status']='running';self.service.save(root,record)
        self.assertEqual(self.client.get(prefix).json()['status'],'interrupted')

    def test_formula_text_is_not_executable_in_workbook(self):
        job=self.done(self.create({'kind':'fact_check','paths':['合同.txt'],'draft':'=HYPERLINK("bad","x")'}).json())
        artifact=next(a for a in job['artifacts'] if a['name'].endswith('.xlsx'))
        wb=load_workbook(self.root/'Stella办案成果'/job['id']/artifact['path']);self.assertEqual(wb.active.cell(2,2).data_type,'s');wb.close()

    def test_draft_file_is_parsed(self):
        (self.root/'草稿.txt').write_text('价款10000元。')
        response=self.create({'kind':'fact_check','paths':['合同.txt'],'draft_path':'草稿.txt'})
        self.assertEqual(response.status_code,200,response.text)
        job=self.done(response.json());self.assertEqual(job['status'],'completed');self.assertEqual(len(job['claims']),1)

    def test_evidence_conversion_failure_never_publishes(self):
        (self.root/'invalid.docx').write_bytes(b'not-a-word-file')
        body={'kind':'evidence_book','paths':['invalid.docx'],'confirmed':True,'entries':[{'path':'invalid.docx','title':'匿名文书','purpose':'核验转换失败'}]}
        with patch('backend.case_workflows._convert_word',side_effect=RuntimeError('failed')):
            job=self.done(self.create(body).json())
        self.assertEqual(job['status'],'failed');self.assertEqual(job['artifacts'],[])
        self.assertEqual((self.root/'invalid.docx').read_bytes(),b'not-a-word-file')

    def test_entry_cannot_override_source_identity_and_review_requires_note(self):
        pdf(self.root/'a.pdf')
        body={'kind':'evidence_book','paths':['a.pdf'],'confirmed':True,'entries':[{'path':'a.pdf','title':'=Danger','purpose':'目的','sha256':'forged','start':999}]}
        job=self.done(self.create(body).json())
        self.assertNotEqual(job['mapping'][0]['sha256'],'forged');self.assertEqual(job['mapping'][0]['start'],1)
        fact=self.fact();prefix=f'/api/case-workflows/cases/1/tasks/{fact["id"]}'
        response=self.client.put(prefix+'/review',json={'revision':1,'reviews':[{'id':r['id'],'status':'rejected','note':''} for r in fact['results']]})
        self.assertEqual(response.status_code,422)

    def test_rotation_and_crop_preserve_visible_source_pixels(self):
        from PIL import Image, ImageChops
        from backend.pdf_assembly import numbered_copy
        for rotation in [0,90,180,270]:
            for cropped in [False,True]:
                with fitz.open() as source:
                    page=source.new_page();page.insert_text((60,80),'Anonymous orientation')
                    if cropped:page.set_cropbox(fitz.Rect(20,20,550,780))
                    page.set_rotation(rotation)
                    result=numbered_copy(source)
                    try:
                        a=page.get_pixmap();b=result[0].get_pixmap()
                        before=Image.frombytes('RGB',[a.width,a.height],a.samples)
                        after=Image.frombytes('RGB',[b.width,b.height],b.samples).crop((0,0,a.width,a.height))
                        self.assertIsNone(ImageChops.difference(before,after).getbbox())
                        self.assertEqual(page.rotation,rotation)
                    finally:result.close()

    def test_spreadsheet_coordinates_and_missing_formula_cache(self):
        from openpyxl import Workbook
        path=self.root/'表.xlsx';wb=Workbook();ws=wb.active
        ws['A1']='价款';ws['C1']=10000;ws['D1']='=C1*2';wb.save(path)
        parsed=asyncio.run(self.service.parse(path,False))
        self.assertIn('C1: 10000',parsed['content']);self.assertIn('无缓存，未验算',parsed['content'])
        self.assertNotIn('B1:',parsed['content'])

    def test_mixed_pdf_requires_explicit_ocr(self):
        with fitz.open() as document:
            page=document.new_page();page.insert_text((50,80),'Enough readable material in the first page of this anonymous document.')
            document.new_page();document.save(self.root/'mixed.pdf')
        job=self.done(self.create({'kind':'fact_check','paths':['mixed.pdf'],'draft':'事实。'}).json())
        self.assertEqual(job['status'],'failed');self.assertIn('OCR',job['error']);self.assertFalse(job['artifacts'])

if __name__=='__main__':unittest.main()
