import base64
import io
import json
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import httpx
from fastapi import FastAPI
from starlette.requests import Request
from PIL import Image
from backend import appearance, ai_custom, ai_config, ai_client, main


class AppearanceTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.directory=tempfile.TemporaryDirectory();self.path=Path(self.directory.name)/'fixture.db'
        def get_db():
            conn=sqlite3.connect(self.path);conn.row_factory=sqlite3.Row;return conn
        self.get_db=get_db
        with get_db() as conn:conn.execute('CREATE TABLE app_preferences(preference_key TEXT PRIMARY KEY,value_json TEXT,updated_at TEXT)')
        self.app=FastAPI();appearance.register_appearance_routes(self.app,get_db)

    def tearDown(self):self.directory.cleanup()

    async def test_profile_defaults_partial_updates_logo_and_validation(self):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app),base_url='http://test') as client:
            self.assertEqual((await client.get('/api/appearance')).json(),appearance.DEFAULTS)
            self.assertEqual((await client.get('/api/appearance/logo')).status_code,404)
            image=io.BytesIO();Image.new('RGBA',(80,30),(100,0,0,255)).save(image,format='PNG')
            logo='data:image/png;base64,'+base64.b64encode(image.getvalue()).decode()
            result=await client.put('/api/appearance',json={'user_name':'测试律师','firm_name':'测试律师事务所','firm_logo':logo})
            self.assertEqual(result.status_code,200)
            result=await client.put('/api/appearance',json={'user_name':'另一名称'})
            self.assertEqual(result.json()['firm_name'],'测试律师事务所')
            self.assertEqual((await client.put('/api/appearance',json={'firm_code':'测试函字第'})).json()['firm_code'],'测试')
            self.assertEqual((await client.put('/api/appearance',headers={'Origin':'https://wrong.test'},json={'user_name':'bad'})).status_code,403)
            self.assertEqual((await client.get('/api/appearance/logo')).headers['Content-Type'],'image/png')
            for value in [{'user_name':'x\nscript'},{'firm_logo':'data:image/svg+xml;base64,PHN2Zz4='},{'firm_name':23},{'unknown':'x'}]:
                self.assertEqual((await client.put('/api/appearance',json=value)).status_code,400)
            self.assertEqual((await client.put('/api/appearance',json={'firm_logo':'','user_name':''})).json()['user_name'],'新用户')
            self.assertEqual((await client.get('/api/appearance/logo')).status_code,404)

    def test_document_exports_reuse_profile_and_preserve_page_table_formatting(self):
        from docx import Document
        profile={**appearance.DEFAULTS,'firm_name':'测试律所','firm_code':'测试','firm_address':'测试地址','firm_phone':'12345'}
        with self.get_db() as conn:conn.execute('INSERT INTO app_preferences VALUES(?,?,?)',(appearance.KEY,json.dumps(profile),'fixture'))
        args=main.SuohanRequest(court='测试法院',clients=['示丆某'],lawyers=['测试律师'],parties=['示一某'],caseType='合同纠纷',caseNo='案号',stages=['一审'],date='2026-10-20',lawyerContacts=[])
        with patch.object(main,'get_db',self.get_db):
            document=Document(io.BytesIO(main._generate_suohan_docx(args)))
            text='\n'.join(p.text for p in document.paragraphs)
            self.assertIn('测试律所函',text);self.assertIn('测试地址',text);self.assertNotIn('明思',text)
            approval=Document(io.BytesIO(main._generate_business_approval_docx(main.BusinessApprovalRequest(clientName='示丆某'))))
            self.assertEqual(len(approval.tables[0].rows),9);self.assertIn('测试律所',''.join(p.text for p in approval.paragraphs))
            initial_page_width=approval.sections[0].page_width
            with self.get_db() as conn:conn.execute('DELETE FROM app_preferences')
            empty=Document(io.BytesIO(main._generate_business_approval_docx(main.BusinessApprovalRequest(clientName='示丆某'))))
            self.assertEqual(empty.sections[0].page_width,initial_page_width)
            self.assertEqual(len(empty.tables[0].rows),9);self.assertNotIn('明思',''.join(p.text for p in empty.paragraphs))


class CustomAIProviderTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.directory=tempfile.TemporaryDirectory()
        self.patches=[patch.object(ai_custom,'PATH',Path(self.directory.name)/'private.json'),patch.object(ai_config,'AI_PREFS_PATH',Path(self.directory.name)/'prefs.json'),patch.object(main,'_reload_ai_runtime',lambda:None)]
        for p in self.patches:p.start()
        self.request=Request({'type':'http','method':'POST','path':'/','headers':[],'client':('127.0.0.1',1),'scheme':'http','server':('test',80)})
    def tearDown(self):
        for p in reversed(self.patches):p.stop()
        ai_config.refresh_custom_models();self.directory.cleanup()

    async def test_custom_api_private_storage_catalog_routing_purposes_and_deletion(self):
        args=main.AIProviderUpdate(name='OpenAI 测试',api_url='https://api.example.test/v1/responses',api_key='fixture-credential-only',protocol='responses',models=['gpt-test','vendor/model'])
        result=await main.add_custom_ai_provider(args,self.request)
        public=result['provider'];identifier=public['id'];model=identifier+':gpt-test'
        self.assertNotIn('api_key',public);self.assertNotIn('fixture-credential-only',json.dumps(public))
        self.assertEqual(os.stat(ai_custom.PATH).st_mode&0o777,0o600)
        self.assertIn(model,ai_config.ALLOWED_MODELS);self.assertIn(model,main.ALLOWED_MODELS)
        self.assertTrue(ai_config.uses_responses(model));self.assertEqual(main._get_api_config(model),ai_client.api_config(model))
        await main.update_ai_purpose_models(main.AIPurposeModelsUpdate(purpose_models={'judicial_messages':model,'hearing_notice':model}),self.request)
        self.assertEqual(ai_config.effective_purpose_model('hearing_notice'),model)
        await main.update_ai_provider(identifier,main.AIProviderUpdate(api_key='',name='改名测试'),self.request)
        self.assertEqual(ai_custom.resolve(model)[1],'fixture-credential-only')
        await main.update_ai_provider(identifier,main.AIProviderUpdate(clear_key=True),self.request)
        self.assertFalse(ai_client.model_available(model))
        await main.delete_custom_ai_provider(identifier,self.request)
        self.assertNotIn(model,main.ALLOWED_MODELS);self.assertNotEqual(ai_config.effective_purpose_model('hearing_notice'),model)
        with self.assertRaises(ValueError):ai_client.api_config(model)

    async def test_hearing_recognition_uses_selected_purpose_and_reports_it(self):
        import ast
        from backend import case_manager as cm
        item=ai_custom.save({'name':'fixture','api_url':'https://example.test/v1','api_key':'fixture','protocol':'chat','models':['wire']})
        ai_config.refresh_custom_models();model=item['id']+':wire'
        ai_config.save_ai_prefs({'purpose_models':{'hearing_notice':model}})
        calls=[]
        async def call(selected,messages,**options):
            calls.append(selected);return {'choices':[{'message':{'content':'{"title":"匿名通知","confidence":90}'}}]},selected
        tree=ast.parse((Path(__file__).parents[1]/'backend/case_manager.py').read_text())
        register=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='register_case_manager_routes')
        functions=[n for n in register.body if isinstance(n,ast.AsyncFunctionDef) and n.name in {'enhance_hearing_notice_with_ai','recognize_hearing_notice'}]
        for node in functions:node.decorator_list=[]
        scope={**vars(cm),'ai_call':call,'norm':lambda value:str(value or '').strip(),
            'clean_notice_text':lambda value:value,'parse_hearing_notice_text':lambda value:{'case_numbers':[]},
            'normalize_case_no':lambda value:value,'match_case_by_case_numbers':lambda value:None,
            'merge_matched_case_into_hearing':lambda value,matched:value}
        exec(compile(ast.Module(body=functions,type_ignores=[]),'hearing-purpose-regression','exec'),scope)
        result=await scope['recognize_hearing_notice'](text='匿名开庭通知',files=[])
        self.assertEqual(calls,[model]);self.assertEqual(result['ai']['requested_model'],model);self.assertEqual(result['ai']['model'],model)

    async def test_sms_analysis_uses_central_purpose_instead_of_legacy_db_model(self):
        import ast
        from backend import judicial_messages as jm
        item=ai_custom.save({'name':'fixture','api_url':'https://example.test/v1','api_key':'fixture','protocol':'chat','models':['wire']})
        ai_config.refresh_custom_models();model=item['id']+':wire'
        ai_config.save_ai_prefs({'purpose_models':{'judicial_messages':model}})
        database=Path(self.directory.name)/'sms-fixture.db'
        def get_db():
            conn=sqlite3.connect(database);conn.row_factory=sqlite3.Row;return conn
        with get_db() as conn:conn.execute('CREATE TABLE case_manager_cases(id INTEGER PRIMARY KEY,title TEXT)')
        jm._ensure_schema(get_db)
        with get_db() as conn:
            conn.execute("UPDATE case_manager_message_settings SET ai_model='deepseek-chat',ai_enabled=1 WHERE id=1")
            conn.execute("INSERT INTO case_manager_judicial_messages(id,fingerprint,received_at,raw_body,created_at,updated_at) VALUES(1,'fixture','2026-10-20','匿名通知','2026-10-20','2026-10-20')")
        calls=[]
        async def analyze(body,extracted,candidates,selected):calls.append(selected);return {},selected
        tree=ast.parse((Path(__file__).parents[1]/'backend/judicial_messages.py').read_text())
        register=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='register_judicial_message_routes')
        function=next(n for n in register.body if isinstance(n,ast.AsyncFunctionDef) and n.name=='analyze_saved')
        scope={**vars(jm),'get_db':get_db,'discover_documents':lambda *args:None,'case_candidates':lambda *args:(None,0,[],'unmatched'),'ai_analysis':analyze}
        exec(compile(ast.Module(body=[function],type_ignores=[]),'sms-purpose-regression','exec'),scope)
        result=await scope['analyze_saved'](1,False)
        self.assertEqual(calls,[model]);self.assertEqual(result['ai_model'],model)

    def test_custom_validation_rejects_bad_urls_models_and_protocols(self):
        base={'name':'测试','api_url':'https://example.test/v1','api_key':'fixture-credential','protocol':'chat','models':['model']}
        for change in [{'api_url':'http://example.test/v1'},{'api_url':'https://user:pass@example.test/v1'},{'api_url':'https://example.test/v1?key=secret'},{'models':['<script>']},{'models':[]},{'protocol':'unknown'}]:
            with self.assertRaises(ValueError):ai_custom.save({**base,**change})
        provider=ai_custom.save(base);self.assertEqual(provider['api_url'],'https://example.test/v1/chat/completions')
        changed=ai_custom.save({'protocol':'responses'},provider['id'])
        self.assertEqual(changed['api_url'],'https://example.test/v1/responses')
        self.assertEqual(ai_custom.chat_options(provider['id']+':kimi-k3',1024,0)['temperature'],1.0)

    async def test_custom_chat_and_responses_send_the_configured_wire_model(self):
        from contextlib import asynccontextmanager
        sent=[]
        class Client:
            async def post(self,url,headers,json):
                sent.append((url,headers,json))
                if url.endswith('/responses'):
                    return httpx.Response(200,json={'status':'completed','output':[{'type':'message','role':'assistant','content':[{'type':'output_text','text':'OK'}]}]})
                return httpx.Response(200,json={'choices':[{'message':{'content':'OK'}}]})
        @asynccontextmanager
        async def factory(**kwargs):yield Client()
        for protocol in ['chat','responses']:
            item=ai_custom.save({'name':protocol,'api_url':'https://example.test/v1','api_key':'fixture','protocol':protocol,'models':['wire-model']})
            ai_config.refresh_custom_models();model=item['id']+':wire-model'
            with patch.object(ai_client,'ai_http_client',factory):
                result,actual=await ai_client.call_ai_with_fallback(model,[{'role':'user','content':'测试'}])
            self.assertEqual(actual,model);self.assertEqual(result['choices'][0]['message']['content'],'OK')
            self.assertEqual(sent[-1][2]['model'],'wire-model')
            self.assertEqual(sent[-1][1]['Authorization'],'Bearer fixture')
            if protocol=='responses':self.assertFalse(sent[-1][2]['store']);self.assertNotIn('messages',sent[-1][2])
