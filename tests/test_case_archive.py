"""Archive workbench uses anonymous, disposable cases and source files."""

from __future__ import annotations

import hashlib
import os
import sqlite3
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient

import backend.case_archive as archive
from backend.case_archive import _classify, register_case_archive_routes
from backend.migrations import _add_case_archive_workbench


def _pdf(path: Path, pages: int) -> None:
    doc = fitz.open()
    for number in range(pages):
        page = doc.new_page()
        page.insert_text((60, 100), f"Anonymous material {number + 1}")
    doc.save(path)
    doc.close()


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _client(tmp_path: Path, category: str = "民事") -> tuple[TestClient, Path]:
    root = tmp_path / "case"
    root.mkdir()
    db_path = tmp_path / "archive.sqlite"
    conn = sqlite3.connect(db_path)
    conn.execute("CREATE TABLE case_manager_cases(id INTEGER PRIMARY KEY,title TEXT,case_type TEXT,source_status TEXT)")
    conn.execute("INSERT INTO case_manager_cases VALUES(1,'匿名案件',?,'active')", (category,))
    conn.execute("""CREATE TABLE case_manager_folder_bindings
        (case_id INTEGER PRIMARY KEY,real_path TEXT,inode INTEGER,device_id INTEGER)""")
    conn.execute("INSERT INTO case_manager_folder_bindings VALUES(1,?,?,?)",
                 (str(root), root.stat().st_ino, root.stat().st_dev))
    _add_case_archive_workbench(conn)
    conn.commit()
    conn.close()

    @contextmanager
    def get_db():
        db = sqlite3.connect(db_path)
        db.row_factory = sqlite3.Row
        try:
            yield db
        finally:
            db.close()

    app = FastAPI()
    register_case_archive_routes(app, get_db, tmp_path, lambda _request: None)
    return TestClient(app), root


def _check_litigation_catalog_matches_compiled_pdf_and_preserves_sources(tmp_path: Path) -> None:
    client, root = _client(tmp_path)
    contract = root / "1.委托代理协议.pdf"
    mandate = root / "2.授权委托书.pdf"
    _pdf(contract, 2)
    _pdf(mandate, 1)
    _pdf(root / "#合并打印.pdf", 3)
    before = {path.name: _sha(path) for path in root.glob("*.pdf")}

    result = client.post("/api/case-archive/cases/1/scan")
    assert result.status_code == 200, result.text
    workspace = result.json()
    assert len(workspace["checklist"]) == 16
    aggregate = next(item for item in workspace["materials"] if item["name"] == "#合并打印.pdf")
    assert aggregate["aggregate"] and aggregate["state"] == "review"
    for path, code in ((contract, "li_1"), (mandate, "li_2")):
        result = client.put("/api/case-archive/cases/1/materials", json={
            "path": path.name, "item_code": code, "state": "confirmed"})
        assert result.status_code == 200, result.text
    result = client.post("/api/case-archive/cases/1/exports")
    assert result.status_code == 200, result.text
    assert result.json()["pages"] == 3
    export = root / result.json()["directory"]
    pdf = fitz.open(export / "结案材料.pdf")
    assert len(pdf) == 3
    assert "第3 页" in pdf[-1].get_text()
    pdf.close()
    catalog = Document(export / "卷内目录.docx")
    assert catalog.tables[0].cell(1, 4).text == "1-2"
    assert catalog.tables[0].cell(2, 4).text == "3"
    assert any("共 3 页" in paragraph.text for paragraph in catalog.paragraphs)
    assert {path.name: _sha(path) for path in root.glob("*.pdf")} == before
    assert client.post("/api/case-archive/cases/1/exports").status_code == 200
    assert len(list((root / "APEX案卷归档").glob("*/结案材料.pdf"))) == 2


def _check_non_litigation_uses_eight_items_and_flags_changed_or_duplicate(tmp_path: Path) -> None:
    client, root = _client(tmp_path, "非诉专项")
    source = root / "1.委托代理合同.pdf"
    _pdf(source, 2)
    (root / "9.委托代理合同副本.pdf").write_bytes(source.read_bytes())
    _pdf(root / "结案材料.pdf", 14)
    _pdf(root / "#合并打印.pdf", 18)
    external = tmp_path / "outside.pdf"
    _pdf(external, 1)
    (root / "outside-link.pdf").symlink_to(external)

    data = client.post("/api/case-archive/cases/1/scan").json()
    assert len(data["checklist"]) == 8
    assert all(item["name"] != "outside-link.pdf" for item in data["materials"])
    assert any(item["duplicate_of"] for item in data["materials"])
    assert all(item["state"] == "review" for item in data["materials"] if item["aggregate"])
    assert client.put("/api/case-archive/cases/1/materials", json={
        "path": source.name, "item_code": "no_1", "state": "confirmed"}).status_code == 200
    _pdf(root / "new.pdf", 1)
    replacement = root / "replacement.pdf"
    _pdf(replacement, 1)
    os.replace(replacement, source)
    current = client.get("/api/case-archive/cases/1").json()
    changed = next(item for item in current["materials"] if item["name"] == source.name)
    assert changed["state"] == "changed"
    assert client.post("/api/case-archive/cases/1/exports").status_code == 409


def _check_criminal_template_and_checklist_switch(tmp_path: Path) -> None:
    client, root = _client(tmp_path, "刑事辩护")
    _pdf(root / "授权委托书.pdf", 1)
    data = client.post("/api/case-archive/cases/1/scan").json()
    assert data["category"] == "criminal" and len(data["checklist"]) == 18
    template = client.get("/api/case-archive/templates/litigation").json()["items"]
    result = client.put("/api/case-archive/cases/1/checklist", json={"category":"litigation","items":template})
    assert result.status_code == 200 and len(result.json()["checklist"]) == 16


def _check_conflicted_learned_rule_never_autoclassifies() -> None:
    code, state = _classify("授权委托书.pdf", "litigation", {"授权委托书": ("li_1", 2, 1)})
    assert code == "" and state == "ambiguous"


def _check_non_litigation_choose_fourteen_page_version_over_eighteen(tmp_path: Path) -> None:
    client, root = _client(tmp_path, "非诉专项")
    _pdf(root / "结案材料.pdf", 14)
    _pdf(root / "#合并打印.pdf", 18)
    assert client.post("/api/case-archive/cases/1/scan").status_code == 200
    assert client.put("/api/case-archive/cases/1/materials", json={
        "path":"结案材料.pdf","item_code":"no_6","state":"confirmed"}).status_code == 200
    assert client.put("/api/case-archive/cases/1/materials", json={
        "path":"#合并打印.pdf","state":"excluded"}).status_code == 200
    result = client.post("/api/case-archive/cases/1/exports")
    assert result.status_code == 200, result.text
    assert result.json()["pages"] == 14
    catalog = Document(root / result.json()["directory"] / "卷内目录.docx")
    assert catalog.tables[0].cell(6, 4).text == "1-14"
    assert any("共 14 页" in paragraph.text for paragraph in catalog.paragraphs)


def _check_sample_shaped_litigation_volumes() -> None:
    # Anonymous page structures mirror the two archived litigation examples.
    examples = (
        ((2, 3, 4, 1, 1, 6, 2, 50), (1, 1, 1, 2, 2, 7, 9, 14), 69),
        ((16, 1, 1, 1, 1, 4, 4, 2), (1, 2, 2, 3, 3, 7, 7, 14), 30),
    )
    for page_counts, categories, expected in examples:
        with tempfile.TemporaryDirectory(prefix="apex-archive-example-") as directory:
            client, root = _client(Path(directory))
            sources = []
            for index, pages in enumerate(page_counts, 1):
                source = root / f"{index:02d}.匿名材料.pdf"
                _pdf(source, pages)
                sources.append(source)
            assert client.post("/api/case-archive/cases/1/scan").status_code == 200
            for source, category in zip(sources, categories):
                result = client.put("/api/case-archive/cases/1/materials", json={
                    "path": source.name, "item_code": f"li_{category}", "state": "confirmed"})
                assert result.status_code == 200, result.text
            export = client.post("/api/case-archive/cases/1/exports")
            assert export.status_code == 200, export.text
            assert export.json()["pages"] == expected
            directory_path = root / export.json()["directory"]
            with fitz.open(directory_path / "结案材料.pdf") as pdf:
                assert len(pdf) == expected
            catalog = Document(directory_path / "卷内目录.docx")
            assert any(f"共 {expected} 页" in paragraph.text for paragraph in catalog.paragraphs)


def _check_word_conversion_failure_does_not_publish_partial_export(tmp_path: Path) -> None:
    client, root = _client(tmp_path)
    invalid = root / "1.委托代理协议.docx"
    invalid.write_bytes(b"not a Word document")
    original = _sha(invalid)
    assert client.post("/api/case-archive/cases/1/scan").status_code == 200
    assert client.put("/api/case-archive/cases/1/materials", json={
        "path": invalid.name, "item_code": "li_1", "state": "confirmed"}).status_code == 200
    def fail_conversion(_source, _output):
        raise RuntimeError("转换组件失败")
    with patch.object(archive, "_convert_word", fail_conversion):
        result = client.post("/api/case-archive/cases/1/exports")
    assert result.status_code == 409
    assert not list((root / "APEX案卷归档").glob("*/结案材料.pdf"))
    assert _sha(invalid) == original


class CaseArchiveTests(unittest.TestCase):
    def _temp(self, check):
        with tempfile.TemporaryDirectory(prefix="apex-archive-test-") as directory:
            check(Path(directory))

    def test_litigation_catalog_matches_compiled_pdf_and_preserves_sources(self):
        self._temp(_check_litigation_catalog_matches_compiled_pdf_and_preserves_sources)

    def test_non_litigation_uses_eight_items_and_flags_changed_or_duplicate(self):
        self._temp(_check_non_litigation_uses_eight_items_and_flags_changed_or_duplicate)

    def test_criminal_template_and_checklist_switch(self):
        self._temp(_check_criminal_template_and_checklist_switch)

    def test_conflicted_learned_rule_never_autoclassifies(self):
        _check_conflicted_learned_rule_never_autoclassifies()

    def test_non_litigation_choose_fourteen_page_version_over_eighteen(self):
        self._temp(_check_non_litigation_choose_fourteen_page_version_over_eighteen)

    def test_sample_shaped_litigation_volumes(self):
        _check_sample_shaped_litigation_volumes()

    def test_word_conversion_failure_does_not_publish_partial_export(self):
        self._temp(_check_word_conversion_failure_does_not_publish_partial_export)
