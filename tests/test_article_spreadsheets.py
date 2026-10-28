import json
import sqlite3
from io import BytesIO
from unittest.mock import Mock

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from backend.article_spreadsheets import register_article_spreadsheet_routes, parse_spreadsheet


@pytest.fixture
def env(tmp_path):
    uploads = tmp_path / 'uploads'
    uploads.mkdir()
    source = uploads / 'anonymous.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = '清单'
    ws['A1'] = '项目'
    ws['C1'] = '金额'
    ws['C2'] = '=SUM(C3:C4)'
    ws['C3'] = 12
    ws['C3'].font = Font(bold=True)
    ws.merge_cells('A4:B4')
    wb.create_sheet('第二表')['A1'] = '二'
    wb.save(source)
    db = tmp_path / 'test.db'
    def connection():
        conn = sqlite3.connect(db, check_same_thread=False)
        conn.row_factory = sqlite3.Row
        return conn
    with connection() as conn:
        conn.executescript('''CREATE TABLE articles(id INTEGER PRIMARY KEY,title TEXT,content TEXT,summary TEXT,source_filename TEXT,source_localpath TEXT);
        CREATE TABLE app_preferences(preference_key TEXT PRIMARY KEY,value_json TEXT,updated_at TEXT);
        INSERT INTO articles VALUES(1,'匿名表格','旧文本','旧摘要','anonymous.xlsx',NULL);
        INSERT INTO articles VALUES(2,'普通文章','正文','','text.txt',NULL);''')
    app = FastAPI()
    rebuild, invalidate = Mock(), Mock()
    register_article_spreadsheet_routes(app, connection, uploads, tmp_path, rebuild, invalidate)
    return TestClient(app, raise_server_exceptions=False), connection, source, rebuild, invalidate


def save(client, operations, revision=None):
    if revision is None:
        revision = client.get('/api/articles/1/spreadsheet').json()['revision']
    return client.put('/api/articles/1/spreadsheet', json={'revision': revision, 'operations': operations})


def test_read_structure_and_parse_retains_empty_columns_and_formulas(env):
    client, _, path, _, _ = env
    parsed = parse_spreadsheet(path)
    assert '项目\t\t金额' in parsed['content']
    assert '=SUM(C3:C4)' in parsed['content']
    data = client.get('/api/articles/1/spreadsheet').json()
    assert data['sheets'] == ['清单', '第二表']
    assert data['cells'][1][2]['formula']
    assert data['cells'][2][2]['bold']
    assert data['cells'][3][1]['merged']
    assert data['merges'] == ['A4:B4']
    assert client.get('/api/articles/999/spreadsheet').status_code == 404
    assert client.get('/api/articles/2/spreadsheet').status_code == 422
    assert client.get('/api/articles/1/spreadsheet?sheet=missing').status_code == 422


def test_save_format_add_sheet_export_preserves_original_and_indexes(env):
    client, connection, source, rebuild, invalidate = env
    original = source.read_bytes()
    response = save(client, [
        {'kind':'cell','sheet':'清单','row':2,'col':1,'value':'已编辑'},
        {'kind':'row_height','sheet':'清单','row':2,'value':40},
        {'kind':'column_width','sheet':'清单','col':1,'value':25},
        {'kind':'format','sheet':'清单','range':'A1:C3','border':'dashed','wrap':True},
        {'kind':'add_sheet','name':'新表'},
        {'kind':'cell','sheet':'新表','row':1,'col':1,'value':123},
    ])
    assert response.status_code == 200, response.text
    assert source.read_bytes() == original
    assert rebuild.call_count == invalidate.call_count == 1
    exported = client.get('/api/articles/1/spreadsheet/export')
    wb = load_workbook(BytesIO(exported.content))
    assert wb['清单']['A2'].value == '已编辑'
    assert wb['清单']['C2'].value == '=SUM(C3:C4)'
    assert wb['清单']['C3'].font.bold
    assert wb['清单'].row_dimensions[2].height == 40
    assert wb['清单'].column_dimensions['A'].width == 25
    assert wb['清单']['A1'].border.left.style == 'dashed'
    assert wb['清单']['A1'].alignment.wrap_text
    assert wb['新表']['A1'].value == 123
    assert 'A4:B4' in [str(m) for m in wb['清单'].merged_cells.ranges]
    with connection() as conn:
        row=conn.execute('SELECT content,summary FROM articles WHERE id=1').fetchone()
        assert '已编辑' in row['content'] and row['summary'] == ''


def test_conflict_and_failed_index_update_leave_saved_workbook_unchanged(env):
    client, connection, source, rebuild, _ = env
    revision = client.get('/api/articles/1/spreadsheet').json()['revision']
    op = {'kind':'cell','sheet':'清单','row':2,'col':1,'value':'第一次'}
    assert save(client, [op], revision).status_code == 200
    before = client.get('/api/articles/1/spreadsheet/export').content
    assert save(client, [{**op,'value':'过期写入'}], revision).status_code == 409
    rebuild.side_effect = RuntimeError('index failed')
    assert save(client, [{**op,'value':'不应保存'}]).status_code == 500
    assert client.get('/api/articles/1/spreadsheet/export').content == before
    with connection() as conn:
        assert '第一次' in conn.execute('SELECT content FROM articles WHERE id=1').fetchone()[0]


@pytest.mark.parametrize('op',[
    {'kind':'cell','row':4,'col':2,'value':'merged'},
    {'kind':'cell','row':0,'col':1,'value':'bad'},
    {'kind':'cell','row':1,'col':1,'value':'\x00'},
    {'kind':'row_height','row':1,'value':301},
    {'kind':'column_width','col':1,'value':0},
    {'kind':'format','range':'A1:XFD1048576','border':'thin'},
    {'kind':'format','range':'A:A','border':'thin'},
    {'kind':'format','range':'C3:A1','border':'thin'},
    {'kind':'format','range':'A1','border':'url(bad)'},
    {'kind':'add_sheet','name':'清单'},
    {'kind':'add_sheet','name':'bad/name'},
])
def test_invalid_edits_are_atomic(env,op):
    client, _, _, _, _ = env
    before=client.get('/api/articles/1/spreadsheet/export').content
    response=save(client,[{'kind':'cell','sheet':'清单','row':1,'col':1,'value':'partial'}, {'sheet':'清单',**op}])
    assert response.status_code in (413,422), response.text
    assert client.get('/api/articles/1/spreadsheet/export').content == before


def test_text_conflict_and_protected_sheet(env):
    client, connection, source, _, _ = env
    revision=client.get('/api/articles/1/spreadsheet').json()['revision']
    with connection() as conn:
        conn.execute("UPDATE articles SET content='其他编辑' WHERE id=1")
    op={'kind':'cell','sheet':'清单','row':1,'col':1,'value':'overwrite'}
    assert save(client,[op],revision).status_code == 409
    wb=load_workbook(source);wb.active.protection.sheet=True;wb.save(source)
    assert save(client,[op]).status_code == 422


def test_literal_formula_and_outer_border(env):
    client, _, _, _, _ = env
    assert save(client,[{'kind':'cell','sheet':'清单','row':1,'col':2,'value':'=not_a_formula','literal':True},
                        {'kind':'format','sheet':'清单','range':'A1:C3','border':'outer'}]).status_code == 200
    wb=load_workbook(BytesIO(client.get('/api/articles/1/spreadsheet/export').content))
    assert wb.active['B1'].data_type=='s'
    assert wb.active['A1'].border.left.style=='thin'
    assert wb.active['B2'].border.left.style is None


def test_upload_short_workbooks_are_valid_and_same_names_do_not_overwrite(tmp_path, monkeypatch):
    from backend import main
    monkeypatch.setattr(main, 'UPLOAD_DIR', tmp_path)
    monkeypatch.setattr(main, 'save_article', Mock(side_effect=[1, 2]))
    monkeypatch.setattr(main, '_generate_and_store_ai_summary', Mock())
    wb=Workbook();wb.active['A1']='短';raw=BytesIO();wb.save(raw)
    client=TestClient(main.app)
    for _ in range(2):
        response=client.post('/api/submit/file',files={'file':('same.xlsx',raw.getvalue(),'application/octet-stream')})
        assert response.status_code==200,response.text
        assert response.json()['title']=='same'
    paths=list(tmp_path.glob('*.xlsx'))
    assert len(paths)==2 and paths[0].name!=paths[1].name
    assert all(p.read_bytes()==raw.getvalue() for p in paths)


def test_oversized_merged_range_rejected_before_expansion(tmp_path):
    import zipfile
    wb=Workbook();wb.active['A1']='test';raw=BytesIO();wb.save(raw)
    path=tmp_path/'huge.xlsx'
    with zipfile.ZipFile(BytesIO(raw.getvalue())) as source, zipfile.ZipFile(path,'w') as dest:
        for name in source.namelist():
            data=source.read(name)
            if name=='xl/worksheets/sheet1.xml':
                data=data.replace(b'</worksheet>',b'<mergeCells count="1"><mergeCell ref="A1:XFD1048576"/></mergeCells></worksheet>')
            dest.writestr(name,data)
    with pytest.raises(Exception) as error:
        parse_spreadsheet(path)
    assert error.value.status_code==413


def test_legacy_xls_values_dates_booleans_and_literal_text(tmp_path, monkeypatch):
    import xlrd
    from types import SimpleNamespace
    sheet=SimpleNamespace(name='旧表', nrows=1, ncols=3, get_rows=lambda: iter([[SimpleNamespace(ctype=1,value='=literal'),SimpleNamespace(ctype=4,value=1),SimpleNamespace(ctype=3,value=45000)]]))
    legacy=SimpleNamespace(datemode=0,sheets=lambda:[sheet],release_resources=Mock())
    monkeypatch.setattr(xlrd,'open_workbook',lambda *args,**kwargs:legacy)
    path=tmp_path/'legacy.xls';path.write_bytes(b'legacy-test')
    from backend.article_spreadsheets import read_workbook
    wb=read_workbook(path)
    assert wb.active['A1'].value=='=literal' and wb.active['A1'].data_type=='s'
    assert wb.active['B1'].value is True
    assert wb.active['C1'].value==xlrd.xldate_as_datetime(45000,0)
    legacy.release_resources.assert_called_once()


def test_real_legacy_xls_file_is_readable_and_editable(env):
    from pathlib import Path
    from shutil import copyfile
    client, connection, source, _, _ = env
    legacy = source.with_suffix('.xls')
    copyfile(Path(__file__).parent/'fixtures/article-workbook.xls',legacy)
    with connection() as conn:
        conn.execute('UPDATE articles SET source_filename=? WHERE id=1',(legacy.name,))
    data=client.get('/api/articles/1/spreadsheet').json()
    assert data['legacy'] and data['cells'][1][2]['value']==123.5
    assert save(client,[{'kind':'cell','sheet':'匿名旧表','row':2,'col':1,'value':'转换后'}]).status_code==200
    result=client.get('/api/articles/1/spreadsheet/export')
    wb=load_workbook(BytesIO(result.content))
    assert wb.active['A2'].value=='转换后'
    assert not client.get('/api/articles/1/spreadsheet').json()['legacy']


def test_view_uses_the_same_snapshot_as_its_revision(env, monkeypatch):
    from backend import article_spreadsheets as module
    client, _, source, _, _ = env
    original = module.read_workbook
    def mutate_after_resolve(path, payload=None):
        changed=Workbook();changed.active['A1']='外部改变';changed.save(source)
        return original(path,payload)
    monkeypatch.setattr(module,'read_workbook',mutate_after_resolve)
    data=client.get('/api/articles/1/spreadsheet').json()
    assert data['cells'][0][0]['value']=='项目'
    assert load_workbook(source).active['A1'].value=='外部改变'
