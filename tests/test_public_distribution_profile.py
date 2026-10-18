import base64
import importlib.util
import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image, PngImagePlugin
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('public_edition',ROOT/'scripts/distribution/public_edition.py')
public=importlib.util.module_from_spec(spec);spec.loader.exec_module(public)

class PublicDistributionProfileTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.env=patch.dict(os.environ,{'LEGAL_KB_RUNTIME_DIR':self.tmp.name});self.env.start();self.addCleanup(self.env.stop)
        raw=io.BytesIO();meta=PngImagePlugin.PngInfo();meta.add_text('PrivateNote','anonymous-test-metadata')
        Image.new('RGBA',(900,180),(18,80,99,255)).save(raw,format='PNG',pnginfo=meta)
        self.logo='data:image/png;base64,'+base64.b64encode(raw.getvalue()).decode()

    def test_logo_metadata_removed_and_failed_upload_preserves_profile(self):
        saved=public.save_profile({'firm_name':'演示律师事务所','logo':self.logo})
        self.assertEqual(public.public_firm_name(),'演示律师事务所')
        raw=public.document_logo_bytes();self.assertNotIn(b'anonymous-test-metadata',raw)
        for bad in ('data:image/svg+xml;base64,AAAA','bad',42):
            with self.assertRaises(ValueError):public.save_profile({'firm_name':'x','logo':bad})
            self.assertEqual(public.profile(),saved)
        self.assertTrue(public.profile_path().is_relative_to(Path(self.tmp.name)) if hasattr(Path,'is_relative_to') else str(public.profile_path()).startswith(self.tmp.name))

    def test_docx_brand_bounded_and_standard_titles_preserved(self):
        public.save_profile({'firm_name':'演示律师事务所','logo':self.logo})
        for title in ('授权委托书','法定代表人证明书','业务审批表'):
            doc=Document();doc.add_heading(title,0);output=io.BytesIO();doc.save(output)
            branded=Document(io.BytesIO(public.public_docx_logo(output.getvalue())))
            self.assertIn(title,[p.text for p in branded.paragraphs])
            shape=branded.inline_shapes[0]
            self.assertLessEqual(shape.width,4.8*914400);self.assertLessEqual(shape.height,.75*914400+1)
        public.save_profile({'firm_name':'演示律师事务所','logo':''});self.assertIsNone(public.document_logo_bytes())

    def test_customization_boundaries_and_logo_api(self):
        app=FastAPI();public.register_public_edition(app,ROOT/'scripts/distribution');client=TestClient(app)
        self.assertEqual(client.get('/scanner-inbox').status_code,200)
        self.assertIn('详情请联系作者',client.get('/public-customization').text)
        for path in ('/api/mingsi-chapters/submit-and-approve','/api/mingsi-meetings/credentials','/api/scanner-inbox/settings'):
            response=client.post(path,json={'account':'anonymous-fixture'})
            self.assertEqual(response.status_code,501);self.assertEqual(response.json()['code'],'customization_required')
        response=client.put('/api/public-edition/document-brand',json={'firm_name':'演示律师事务所','logo':self.logo})
        self.assertEqual(response.status_code,200)
        self.assertEqual(client.get('/public-firm-logo.png').headers['content-type'],'image/png')
        self.assertEqual(client.put('/api/public-edition/document-brand',json={'firm_name':'x','logo':'bad'}).status_code,400)

if __name__=='__main__':unittest.main()
