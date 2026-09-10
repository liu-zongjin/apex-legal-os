import asyncio
import io
import unittest
from unittest.mock import AsyncMock, patch

import httpx
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from backend import main


SAMPLE_EXPORT = {
    "question": "预约合同能否强制履行？",
    "answer": "## 结论\n\n原则上应结合主要条款是否完备判断。\n\n- 案例一：支持赔偿损失。",
    "reasoning": "（第 1 步）识别争点\n（第 2 步）核验案例\n\n分析草稿：我需要继续组织答案。",
    "model": "deepseek-flash",
    "context": None,
    "sources": [{"title": "元典·案例库 MCP", "url": "https://example.test/case/1"}],
    "case_title": "测试案件",
    "created_at_ms": 1_786_000_000_000,
}


class AiAskExportTests(unittest.TestCase):
    def test_markdown_export_contains_question_answer_sources_and_reasoning(self):
        content = main._build_ai_ask_export_markdown(SAMPLE_EXPORT)
        self.assertIn("# 法律操作系统 AI增强", content)
        self.assertIn("## 问题", content)
        self.assertIn("预约合同能否强制履行？", content)
        self.assertIn("## 回答来源", content)
        self.assertIn("https://example.test/case/1", content)
        self.assertIn("## 检索与核验记录", content)
        self.assertNotIn("分析草稿", content)

    def test_markdown_download_bytes_are_utf8_with_bom(self):
        content = main._utf8_bom_markdown_bytes("# 中文标题\n")
        self.assertTrue(content.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(content.decode("utf-8-sig"), "# 中文标题\n")

    def test_docx_export_is_readable_and_contains_complete_sections(self):
        content = main._generate_ai_ask_export_docx(SAMPLE_EXPORT)
        document = Document(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("法律操作系统 AI增强", text)
        self.assertIn("预约合同能否强制履行？", text)
        self.assertIn("原则上应结合主要条款是否完备判断", text)
        self.assertIn("检索与核验记录", text)

    def test_default_docx_layout_matches_requested_legal_format(self):
        content = main._generate_ai_ask_export_docx(SAMPLE_EXPORT)
        document = Document(io.BytesIO(content))
        section = document.sections[0]
        self.assertAlmostEqual(section.page_width.cm, 21.0, places=1)
        self.assertAlmostEqual(section.page_height.cm, 29.7, places=1)
        self.assertAlmostEqual(section.top_margin.cm, 3.0, places=1)
        self.assertAlmostEqual(section.bottom_margin.cm, 2.5, places=1)
        self.assertAlmostEqual(section.left_margin.cm, 2.5, places=1)
        self.assertAlmostEqual(section.right_margin.cm, 2.5, places=1)

        normal = document.styles["Normal"]
        self.assertEqual(normal._element.rPr.rFonts.get(qn("w:eastAsia")), "仿宋_GB2312")
        self.assertAlmostEqual(normal.font.size.pt, 12.0, places=1)
        self.assertAlmostEqual(normal.paragraph_format.first_line_indent.pt, 24.0, places=1)
        self.assertAlmostEqual(normal.paragraph_format.line_spacing.pt, 24.0, places=1)
        self.assertAlmostEqual(normal.paragraph_format.space_before.pt, 0.0, places=1)
        self.assertAlmostEqual(normal.paragraph_format.space_after.pt, 12.0, places=1)

        heading1 = document.styles["Heading 1"]
        self.assertEqual(heading1._element.rPr.rFonts.get(qn("w:eastAsia")), "黑体")
        self.assertTrue(heading1.font.bold)
        self.assertAlmostEqual(heading1.paragraph_format.space_before.pt, 24.0, places=1)
        self.assertAlmostEqual(heading1.paragraph_format.space_after.pt, 12.0, places=1)

        heading2 = document.styles["Heading 2"]
        self.assertEqual(heading2._element.rPr.rFonts.get(qn("w:eastAsia")), "仿宋_GB2312")
        self.assertAlmostEqual(heading2.paragraph_format.space_before.pt, 0.0, places=1)
        self.assertAlmostEqual(heading2.paragraph_format.space_after.pt, 12.0, places=1)

    def test_each_word_layout_generates_a_readable_docx(self):
        for layout in main._AI_ASK_WORD_LAYOUTS:
            with self.subTest(layout=layout):
                content = main._generate_ai_ask_export_docx(SAMPLE_EXPORT, layout=layout)
                document = Document(io.BytesIO(content))
                self.assertTrue(any("预约合同" in paragraph.text for paragraph in document.paragraphs))

    def test_saved_custom_layout_controls_all_requested_docx_parameters(self):
        custom = main._validated_custom_word_layout({
            "name": "测试格式", "description": "测试",
            "page_size_cm": [20, 28], "margins_cm": [3.1, 2.6, 2.7, 2.8],
            "body_font": "宋体", "body_size": 11,
            "heading1_font": "黑体", "heading1_size": 15, "heading1_bold": True,
            "heading2_font": "楷体_GB2312", "heading2_size": 13, "heading2_bold": False,
            "title_font": "黑体", "title_size": 21,
            "line_spacing": 25, "first_line_chars": 2,
            "body_before": 1, "body_after": 7,
            "heading1_before": 20, "heading1_after": 9,
            "heading2_before": 3, "heading2_after": 6,
            "accent": [0, 0, 0],
        }, layout_id="custom_test")
        with patch.object(main, "_load_custom_word_layouts", return_value=[custom]):
            document = Document(io.BytesIO(main._generate_ai_ask_export_docx(SAMPLE_EXPORT, "custom_test")))
        section = document.sections[0]
        self.assertAlmostEqual(section.page_width.cm, 20, places=1)
        self.assertAlmostEqual(section.page_height.cm, 28, places=1)
        self.assertAlmostEqual(section.top_margin.cm, 3.1, places=1)
        self.assertAlmostEqual(section.right_margin.cm, 2.8, places=1)
        normal = document.styles["Normal"]
        self.assertAlmostEqual(normal.paragraph_format.space_before.pt, 1, places=1)
        self.assertAlmostEqual(normal.paragraph_format.space_after.pt, 7, places=1)
        self.assertAlmostEqual(normal.paragraph_format.line_spacing.pt, 25, places=1)
        heading2 = document.styles["Heading 2"]
        self.assertFalse(heading2.font.bold)
        self.assertEqual(heading2._element.rPr.rFonts.get(qn("w:eastAsia")), "楷体_GB2312")


class ResearchQuestionExportTests(unittest.TestCase):
    def setUp(self):
        self.project = {"name": "测试研究项目", "case_title": "示丆某诉示一某合同纠纷案"}
        self.records = [{
            "question": "合同是否有效？",
            "answer": "## 结论\n\n合同原则上有效。",
            "reasoning": "（第 1 步）核验案件资料",
            "model": "deepseek-flash",
            "created_at": "2026-08-06T10:00:00",
            "sources": [{"title": "案件材料一", "quote": "双方签署合同。"}],
        }]

    def test_markdown_contains_case_record_and_source(self):
        content = main._build_research_questions_export_markdown(self.project, self.records)
        self.assertIn("# 案件资料问答", content)
        self.assertIn("示丆某诉示一某合同纠纷案", content)
        self.assertIn("合同是否有效？", content)
        self.assertIn("案件材料一", content)

    def test_docx_contains_case_and_record(self):
        content = main._research_questions_docx(self.project, self.records)
        document = Document(io.BytesIO(content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("案件资料问答", text)
        self.assertIn("示丆某诉示一某合同纠纷案", text)
        self.assertIn("合同是否有效？", text)
        self.assertIn("合同原则上有效", text)


class AiAskBlankRecoveryTests(unittest.IsolatedAsyncioTestCase):
    async def test_reasoning_only_response_is_repaired_into_final_answer(self):
        responses = [
            ({"choices": [{"message": {"content": "", "reasoning_content": "分析草稿"}, "finish_reason": "length"}]}, "deepseek-flash"),
            ({"choices": [{"message": {"content": "", "reasoning_content": "继续分析"}, "finish_reason": "stop"}]}, "deepseek-flash"),
            ({"choices": [{"message": {"content": "完整最终答案。", "reasoning_content": ""}, "finish_reason": "stop"}]}, "deepseek-chat"),
        ]
        with patch.object(main, "_call_ai_with_fallback", new=AsyncMock(side_effect=responses)):
            answer, reasoning, model = await main._call_ai_answer_with_continuation(
                "deepseek-flash",
                [{"role": "system", "content": "仅依据检索资料"}, {"role": "user", "content": "测试问题"}],
                max_tokens=100,
            )
        self.assertEqual(answer, "完整最终答案。")
        self.assertEqual(reasoning, "分析草稿")
        self.assertEqual(model, "deepseek-chat")


class AiProviderFallbackTests(unittest.IsolatedAsyncioTestCase):
    def test_long_report_models_have_extended_timeout(self):
        self.assertEqual(main._get_timeout("deepseek-flash"), 180.0)
        self.assertEqual(main._get_timeout("deepseek-v4-pro"), 180.0)

    async def test_v4_pro_timeout_falls_back_to_v4_flash(self):
        requested_models = []

        class FakeClient:
            def __init__(self, **_kwargs):
                pass

            async def __aenter__(self):
                return self

            async def __aexit__(self, *_args):
                return False

            async def post(self, url, json, headers):
                requested_models.append(json["model"])
                request = httpx.Request("POST", url)
                if json["model"] == "deepseek-v4-pro":
                    raise httpx.ReadTimeout("", request=request)
                return httpx.Response(
                    200,
                    request=request,
                    json={"choices": [{"message": {"content": "备用模型回答。"}}]},
                )

        with patch.object(
            main,
            "_get_api_config",
            side_effect=lambda model: ("https://example.test/chat", "test-key", model),
        ), patch.object(httpx, "AsyncClient", FakeClient):
            payload, actual_model = await main._call_ai_with_fallback(
                "deepseek-v4-pro", [{"role": "user", "content": "测试"}]
            )

        self.assertEqual(requested_models, ["deepseek-v4-pro", "deepseek-flash"])
        self.assertEqual(actual_model, "deepseek-flash")
        self.assertEqual(payload["choices"][0]["message"]["content"], "备用模型回答。")

    async def test_v4_flash_timeout_uses_same_account_v4_pro_before_external_provider(self):
        requested_models = []

        class FakeClient:
            def __init__(self, **_kwargs):
                pass

            async def __aenter__(self):
                return self

            async def __aexit__(self, *_args):
                return False

            async def post(self, url, json, headers):
                requested_models.append(json["model"])
                request = httpx.Request("POST", url)
                if json["model"] == "deepseek-flash":
                    raise httpx.ReadTimeout("", request=request)
                return httpx.Response(
                    200,
                    request=request,
                    json={"choices": [{"message": {"content": "同账户模型回答。"}}]},
                )

        with patch.object(
            main,
            "_get_api_config",
            side_effect=lambda model: ("https://example.test/chat", "test-key", model),
        ), patch.object(httpx, "AsyncClient", FakeClient):
            payload, actual_model = await main._call_ai_with_fallback(
                "deepseek-flash", [{"role": "user", "content": "测试"}]
            )

        self.assertEqual(requested_models, ["deepseek-flash", "deepseek-v4-pro"])
        self.assertEqual(actual_model, "deepseek-v4-pro")
        self.assertEqual(payload["choices"][0]["message"]["content"], "同账户模型回答。")

    async def test_external_provider_402_is_not_reported_as_deepseek_balance(self):
        class FakeClient:
            def __init__(self, **_kwargs):
                pass

            async def __aenter__(self):
                return self

            async def __aexit__(self, *_args):
                return False

            async def post(self, url, json, headers):
                request = httpx.Request("POST", url)
                return httpx.Response(402, request=request, text="payment required")

        with patch.object(
            main,
            "_get_api_config",
            side_effect=lambda model: ("https://example.test/chat", "test-key", model),
        ), patch.object(httpx, "AsyncClient", FakeClient):
            with self.assertRaises(main.HTTPException) as caught:
                await main._call_ai_with_fallback(
                    "deepseek-flash", [{"role": "user", "content": "测试"}]
                )

        self.assertNotIn("DeepSeek 账户余额不足", caught.exception.detail)
        self.assertIn("opencode:deepseek-flash", caught.exception.detail)
        self.assertIn("计费或模型授权状态异常", caught.exception.detail)

    def test_blank_timeout_has_readable_message(self):
        request = httpx.Request("POST", "https://example.test/chat")
        self.assertEqual(
            main._readable_exception_message(httpx.ReadTimeout("", request=request)),
            "AI 模型响应超时",
        )

    async def test_continuation_fragment_is_rebuilt_as_self_contained_answer(self):
        responses = [
            ({"choices": [{"message": {"content": "（接前文）第三个案例尚未写完", "reasoning_content": "长推理"}, "finish_reason": "stop"}]}, "deepseek-flash"),
            ({"choices": [{"message": {"content": "完整结论。\n\n直接案例一：法院驳回强制缔约请求。", "reasoning_content": ""}, "finish_reason": "stop"}]}, "deepseek-chat"),
        ]
        with patch.object(main, "_call_ai_with_fallback", new=AsyncMock(side_effect=responses)):
            answer, _, model = await main._call_ai_answer_with_continuation(
                "deepseek-flash",
                [{"role": "system", "content": "仅依据检索资料"}, {"role": "user", "content": "测试问题"}],
                max_tokens=100,
            )
        self.assertTrue(answer.startswith("完整结论"))
        self.assertNotIn("接前文", answer)
        self.assertEqual(model, "deepseek-chat")


class ExternalCaseCitationAuditTests(unittest.TestCase):
    def test_rejects_case_number_not_present_in_evidence_package(self):
        audit = main._audit_external_case_citations(
            "可参考（2021）黔23民终810号及（2099）虚01民终1号。",
            "证据：案号（2021）黔23民终810号。",
        )
        self.assertFalse(audit["passed"])
        self.assertEqual(audit["unsupported_case_numbers"], ["（2099）虚01民终1号"])

    def test_accepts_normalised_parentheses(self):
        audit = main._audit_external_case_citations(
            "案例：(2021)黔23民终810号。",
            "案号：（2021）黔23民终810号",
        )
        self.assertTrue(audit["passed"])

    def test_accepts_case_number_from_qcc_evidence_package(self):
        audit = main._audit_external_case_citations(
            "可参考（2022）浙06民终31号民事判决。",
            "【企查查·司法案例 MCP 检索结果】案号：（2022）浙06民终31号",
        )
        self.assertTrue(audit["passed"])
        self.assertEqual(audit["unsupported_case_numbers"], [])


class CaseSearchReportModeTests(unittest.TestCase):
    def test_enabled_mode_requires_a_complete_case_search_report(self):
        instruction = main._case_search_report_instruction(True)
        self.assertIn("案例检索报告", instruction)
        self.assertIn("权威案例", instruction)
        self.assertIn("直接相关普通案例", instruction)
        self.assertIn("限制性或相反案例", instruction)
        self.assertIn("检索局限", instruction)

    def test_disabled_mode_adds_no_report_constraint(self):
        self.assertEqual(main._case_search_report_instruction(False), "")


if __name__ == "__main__":
    unittest.main()
