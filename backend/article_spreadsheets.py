"""Bounded workbook editing; immutable managed copies, never overwrite source files."""
from copy import copy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import datetime
import json
import math
import uuid
import zipfile
from xml.etree.ElementTree import iterparse

from fastapi import HTTPException, Query
from fastapi.responses import FileResponse, Response
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter, range_boundaries

MAX_CELLS = 200_000
MAX_BYTES = 20 * 1024 * 1024


def _source_bytes(path):
    if not path.is_file():
        raise HTTPException(404, '原始表格文件不存在，请重新导入')
    with path.open('rb') as source:
        payload = source.read(MAX_BYTES + 1)
    if len(payload) > MAX_BYTES:
        raise HTTPException(413, '表格超过 20MB，请拆分后导入')
    return payload


def read_workbook(path, payload=None):
    path = Path(path)
    if payload is None:
        payload = _source_bytes(path)
    try:
        if path.suffix.lower() == '.xls':
            import xlrd
            legacy = xlrd.open_workbook(file_contents=payload, on_demand=True)
            try:
                wb = Workbook()
                wb.remove(wb.active)
                count = 0
                for sheet in legacy.sheets():
                    count += sheet.nrows * sheet.ncols
                    if count > MAX_CELLS:
                        raise HTTPException(413, '表格超过 20 万单元格，请拆分后导入')
                    ws = wb.create_sheet(sheet.name)
                    for row in sheet.get_rows():
                        values = []
                        for cell in row:
                            value = cell.value
                            if cell.ctype == xlrd.XL_CELL_DATE:
                                value = xlrd.xldate_as_datetime(value, legacy.datemode)
                            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                                value = bool(value)
                            values.append(value)
                        ws.append(values)
                        for cell in ws[ws.max_row]:
                            if isinstance(cell.value, str):
                                cell.data_type = 's'
                return wb
            finally:
                legacy.release_resources()
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
                raise HTTPException(413, '表格解压后过大，请拆分后导入')
            # Reject pathological sparse/merged ranges before openpyxl expands them.
            count = 0
            for entry in archive.infolist():
                if not entry.filename.startswith('xl/worksheets/') or not entry.filename.endswith('.xml'):
                    continue
                with archive.open(entry) as stream:
                    for _, element in iterparse(stream, events=('end',)):
                        tag = element.tag.rsplit('}', 1)[-1]
                        if tag == 'c':
                            count += 1
                        elif tag == 'mergeCell':
                            c1, r1, c2, r2 = range_boundaries(element.attrib['ref'])
                            count += (r2-r1+1)*(c2-c1+1)
                        if count > MAX_CELLS:
                            raise HTTPException(413, '表格超过 20 万单元格，请拆分后导入')
                        element.clear()
        wb = load_workbook(BytesIO(payload), keep_links=True)
        if sum(len(ws._cells) for ws in wb) > MAX_CELLS:
            wb.close()
            raise HTTPException(413, '表格超过 20 万单元格，请拆分后导入')
        return wb
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, '无法读取表格，请确认文件未加密且格式正确') from exc


def workbook_text(wb):
    blocks = []
    size = 0
    for ws in wb:
        rows = {}
        for (r, c), cell in sorted(ws._cells.items()):
            if cell.value is not None:
                rows.setdefault(r, {})[c] = str(cell.value)
        lines = []
        for row_number, cells in rows.items():
            # Sparse far-right cells must not allocate thousands of empty columns.
            if max(cells) > 200:
                line = '  '.join(f'{get_column_letter(c)}{row_number}: {value}' for c, value in cells.items())
            else:
                line = '\t'.join(cells.get(c, '') for c in range(1, max(cells) + 1))
            size += len(line)
            if size > 4_000_000:
                raise HTTPException(413, '表格文本超过索引上限，请拆分后导入')
            lines.append(line)
        blocks.append(f'【{ws.title}】\n' + '\n'.join(lines))
    return '\n\n'.join(blocks)


def parse_spreadsheet(path):
    wb = read_workbook(path)
    try:
        return {'title': Path(path).stem, 'content': workbook_text(wb)}
    finally:
        wb.close()


def _rgb(color):
    if color is not None and color.type == 'rgb' and isinstance(color.rgb, str):
        return '#' + color.rgb[-6:]
    return None


def _integer(value, low, high):
    if type(value) is not int or not low <= value <= high:
        raise HTTPException(422, '表格坐标或数量超出范围')
    return value


def _number(value, low, high):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not low <= value <= high:
        raise HTTPException(422, '尺寸超出允许范围')
    return value


def apply_operations(wb, operations):
    if not isinstance(operations, list) or not 1 <= len(operations) <= 2000:
        raise HTTPException(422, '每次保存需包含 1 至 2000 项修改')
    budget = 0
    for op in operations:
        if not isinstance(op, dict):
            raise HTTPException(422, '修改格式无效')
        kind = op.get('kind')
        if kind == 'add_sheet':
            if wb.security and wb.security.lockStructure:
                raise HTTPException(422, '工作簿结构受保护，不能添加工作表')
            name = op.get('name')
            if not isinstance(name, str) or not name.strip() or len(name) > 31 or any(c in name for c in '[]:*?/\\') or name.lower() in [s.lower() for s in wb.sheetnames] or len(wb.worksheets) >= 50:
                raise HTTPException(422, '工作表名称无效、重复或工作表已达 50 个')
            wb.create_sheet(name)
            continue
        if op.get('sheet') not in wb.sheetnames:
            raise HTTPException(422, '工作表不存在')
        ws = wb[op['sheet']]
        if ws.protection.sheet:
            raise HTTPException(422, '该工作表受保护，请先在 Excel 中解除保护')
        if kind == 'cell':
            r = _integer(op.get('row'), 1, 1_048_576)
            c = _integer(op.get('col'), 1, 16_384)
            value = op.get('value')
            if value is not None and (type(value) not in (str, int, float, bool) or isinstance(value, str) and len(value) > 32767 or isinstance(value, (int, float)) and not isinstance(value, bool) and (abs(value) > 1e308 or not math.isfinite(value))):
                raise HTTPException(422, '单元格值无效或超过 32767 字符')
            cell = ws.cell(r, c)
            if isinstance(cell, MergedCell):
                raise HTTPException(422, '请编辑合并区域左上角的单元格')
            if isinstance(value, str) and any(ord(ch) < 32 and ch not in '\t\n\r' for ch in value):
                raise HTTPException(422, '单元格包含不支持的控制字符')
            cell.value = value
            if op.get('literal') and isinstance(value, str):
                cell.data_type = 's'
            budget += 1
        elif kind == 'row_height':
            index = _integer(op.get('row'), 1, 1_048_576)
            ws.row_dimensions[index].height = _number(op.get('value'), 8, 300)
        elif kind == 'column_width':
            index = _integer(op.get('col'), 1, 16_384)
            ws.column_dimensions[get_column_letter(index)].width = _number(op.get('value'), 3, 100)
        elif kind == 'format':
            area = op.get('range')
            try:
                c1, r1, c2, r2 = range_boundaries(area)
                for n, limit in ((r1, 1_048_576), (r2, 1_048_576), (c1, 16_384), (c2, 16_384)):
                    _integer(n, 1, limit)
                if r1 > r2 or c1 > c2:
                    raise ValueError()
            except (ValueError, TypeError):
                raise HTTPException(422, '请选择有效的单元格区域，例如 A1:C5')
            budget += (r2-r1+1) * (c2-c1+1)
            if budget > 10_000:
                raise HTTPException(413, '一次最多调整 10000 个单元格')
            border = op.get('border')
            if border is not None and border not in ('none', 'thin', 'medium', 'dashed', 'outer'):
                raise HTTPException(422, '边框样式无效')
            for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
                for cell in row:
                    if border is not None:
                        edge = Side(style=None if border == 'none' else 'thin' if border == 'outer' else border, color='808080')
                        cell.border = Border(left=edge if border != 'outer' or cell.column == c1 else Side(), right=edge if border != 'outer' or cell.column == c2 else Side(), top=edge if border != 'outer' or cell.row == r1 else Side(), bottom=edge if border != 'outer' or cell.row == r2 else Side())
                    if 'bold' in op:
                        if type(op['bold']) is not bool:
                            raise HTTPException(422, '加粗格式无效')
                        font = copy(cell.font)
                        font.bold = op['bold']
                        cell.font = font
                    if 'wrap' in op:
                        if type(op['wrap']) is not bool:
                            raise HTTPException(422, '换行格式无效')
                        align = copy(cell.alignment)
                        align.wrap_text = op['wrap']
                        cell.alignment = align
        else:
            raise HTTPException(422, '不支持的表格操作')
    if sum(len(ws._cells) for ws in wb) > MAX_CELLS:
        raise HTTPException(413, '表格超过 20 万单元格')


def register_article_spreadsheet_routes(app, get_db, upload_dir, frontend_dir, rebuild, invalidate):
    root = Path(upload_dir) / 'article-workbooks'

    def resolve(conn, article_id):
        article = conn.execute('SELECT * FROM articles WHERE id=?', (article_id,)).fetchone()
        if not article:
            raise HTTPException(404, '文章不存在')
        article = dict(article)
        if not str(article.get('source_filename') or article.get('source_localpath') or '').lower().endswith(('.xlsx', '.xls')):
            raise HTTPException(422, '该文章没有 Excel 源文件，请在正文编辑器中添加表格')
        key = f'article_workbook:{article_id}'
        saved = conn.execute('SELECT value_json FROM app_preferences WHERE preference_key=?', (key,)).fetchone()
        if saved:
            meta = json.loads(saved['value_json'])
            path = root / str(article_id) / (meta['revision'] + '.xlsx')
        else:
            filename = article.get('source_filename')
            path = Path(article['source_localpath']) if article.get('source_localpath') else Path(upload_dir) / Path(filename).name
        payload = _source_bytes(path)
        revision = sha256(payload).hexdigest()
        # Also fence a concurrent edit to article text (not just the binary).
        revision = revision + ':' + sha256((article.get('content') or '').encode()).hexdigest()
        return article, path, revision, key, payload

    @app.get('/article-spreadsheet')
    def page():
        return FileResponse(Path(frontend_dir) / 'article-spreadsheet.html')

    @app.get('/api/articles/{article_id}/spreadsheet')
    def read(article_id: int, sheet: str = '', row: int = Query(1, ge=1, le=1_048_576), col: int = Query(1, ge=1, le=16_384)):
        conn = get_db()
        try:
            article, path, revision, _, payload = resolve(conn, article_id)
        finally:
            conn.close()
        wb = read_workbook(path, payload)
        try:
            if sheet and sheet not in wb.sheetnames:
                raise HTTPException(422, '工作表不存在')
            ws = wb[sheet] if sheet else wb.active
            end_r, end_c = min(row + 49, 1_048_576), min(col + 19, 16_384)
            max_row, max_col = ws.max_row, ws.max_column
            cells = []
            for r in range(row, end_r + 1):
                line = []
                for c in range(col, end_c + 1):
                    cell = ws.cell(r, c)
                    value = cell.value
                    display = str(value) if value is not None else ''
                    line.append({'value': value if type(value) in (str, int, float, bool) or value is None else display,
                                 'text': display, 'formula': cell.data_type == 'f',
                                 'merged': isinstance(cell, MergedCell), 'bold': bool(cell.font.bold),
                                 'wrap': bool(cell.alignment.wrap_text), 'align': cell.alignment.horizontal,
                                 'font_size': cell.font.sz,
                                 'font_color': _rgb(cell.font.color),
                                 'fill': _rgb(cell.fill.fgColor) if cell.fill.patternType == 'solid' else None,
                                 'borders': {side: getattr(cell.border, side).style if getattr(cell.border, side) else None for side in ('left', 'right', 'top', 'bottom')}})
                cells.append(line)
            return {'article_id': article_id, 'title': article['title'], 'revision': revision,
                    'legacy': path.suffix.lower() == '.xls', 'protected': bool(ws.protection.sheet),
                    'sheets': wb.sheetnames, 'sheet': ws.title, 'row': row, 'col': col,
                    'max_row': max_row, 'max_col': max_col, 'cells': cells,
                    'heights': [ws.row_dimensions[r].height or 22 for r in range(row, end_r + 1)],
                    'widths': [ws.column_dimensions[get_column_letter(c)].width or 13 for c in range(col, end_c + 1)],
                    'merges': [str(area) for area in ws.merged_cells.ranges if area.min_row <= end_r and area.max_row >= row and area.min_col <= end_c and area.max_col >= col]}
        finally:
            wb.close()

    @app.put('/api/articles/{article_id}/spreadsheet')
    def save(article_id: int, body: dict):
        conn = get_db()
        created = None
        wb = None
        try:
            conn.execute('BEGIN IMMEDIATE')
            article, path, revision, key, payload = resolve(conn, article_id)
            if body.get('revision') != revision:
                raise HTTPException(409, '表格或正文已在其他位置修改，请保留草稿并重新载入核对')
            wb = read_workbook(path, payload)
            apply_operations(wb, body.get('operations'))
            # Do not fabricate formula results; Excel recalculates when opening the exported copy.
            content = workbook_text(wb)
            output = BytesIO()
            wb.save(output)
            raw = output.getvalue()
            if len(raw) > MAX_BYTES:
                raise HTTPException(413, '保存后的表格超过 20MB')
            digest = sha256(raw).hexdigest()
            folder = root / str(article_id)
            folder.mkdir(parents=True, exist_ok=True)
            target = folder / (digest + '.xlsx')
            if not target.exists():
                temporary = folder / (uuid.uuid4().hex + '.tmp')
                try:
                    temporary.write_bytes(raw)
                    temporary.replace(target)
                    created = target
                finally:
                    temporary.unlink(missing_ok=True)
            if content != (article.get('content') or ''):
                invalidate(conn, article_id, '原电子表格已修改，需重新核验')
                conn.execute('UPDATE articles SET content=?,summary=? WHERE id=?', (content, '', article_id))
                rebuild(conn, article_id, content)
            conn.execute('''INSERT INTO app_preferences(preference_key,value_json,updated_at) VALUES(?,?,?)
                ON CONFLICT(preference_key) DO UPDATE SET value_json=excluded.value_json,updated_at=excluded.updated_at''',
                         (key, json.dumps({'revision': digest}), datetime.datetime.now().isoformat()))
            conn.commit()
            return {'success': True, 'revision': digest + ':' + sha256(content.encode()).hexdigest(), 'content': content}
        except Exception:
            conn.rollback()
            if created:
                created.unlink(missing_ok=True)
            raise
        finally:
            if wb:
                wb.close()
            conn.close()

    @app.get('/api/articles/{article_id}/spreadsheet/export')
    def export(article_id: int):
        conn = get_db()
        try:
            _, path, _, _, payload = resolve(conn, article_id)
        finally:
            conn.close()
        # Export exactly the resolved bytes even if the external original changes.
        return Response(payload, headers={'Content-Disposition': 'attachment; filename=article-workbook' + path.suffix.lower()}, media_type='application/vnd.ms-excel' if path.suffix.lower() == '.xls' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
