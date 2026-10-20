from __future__ import annotations

import asyncio
import base64
from contextvars import ContextVar
import datetime as dt
import hashlib
import hmac
import ipaddress
import json
import mimetypes
import os
import re
import secrets
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
import traceback
import time
import urllib.parse
import zipfile
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Sequence, Tuple

import httpx
from fastapi import File, HTTPException, Request, UploadFile

from .ai_client import call_ai_with_fallback, model_available
from .ai_config import ALLOWED_MODELS, DEFAULT_AI_MODEL, effective_purpose_model


MESSAGE_DB = Path.home() / "Library" / "Messages" / "chat.db"
MESSAGE_ATTACHMENT_DIR = Path.home() / "Library" / "Messages" / "Attachments"
JUDICIAL_STAGING_DIR = Path(
    os.environ.get("LEGAL_KB_DATA_DIR", str(Path(__file__).resolve().parent))
) / "judicial_downloads"
# Documents are still buffered in memory during validation and archiving.
# Allow large scanned evidence bundles while keeping a finite resource bound.
MAX_DOCUMENT_MB = 500
MAX_DOCUMENT_BYTES = MAX_DOCUMENT_MB * 1024 * 1024
PROMPT_VERSION = "judicial-sms-v1"
IRRELEVANT_DISPOSITION = "irrelevant-v4"
REPLAYABLE_IRRELEVANT_DISPOSITIONS = ("irrelevant", "irrelevant-v2", "irrelevant-v3")
CATEGORIES = [
    "开庭/庭询/听证",
    "材料或电子送达",
    "立案/受理/平台状态",
    "举证/答辩/补正期限",
    "裁判文书送达",
    "上诉/异议/复议期限",
    "调解/谈话/询问",
    "缴费/退费",
    "保全/执行",
    "变更/取消",
    "其他重要通知",
]
STATUS_LABELS = {
    "pending": "待分类",
    "suggested": "建议确认",
    "confirmed": "已归类/关联",
    "auto_written": "自动写入",
    "ignored": "已忽略",
}


def _browser_executable() -> str:
    """Prefer an installed system browser over Playwright's bundled Chromium."""
    configured = os.environ.get("LEGAL_KB_BROWSER_EXECUTABLE", "").strip()
    candidates = [
        configured,
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
        str(Path.home() / "Applications/Google Chrome.app/Contents/MacOS/Google Chrome"),
        "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge",
        str(Path.home() / "Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge"),
    ]
    return next((item for item in candidates if item and Path(item).is_file()), "")


def _friendly_browser_error(exc: BaseException) -> str:
    text = str(exc or "")
    lowered = text.lower()
    if isinstance(exc, (asyncio.TimeoutError, TimeoutError)) or "timeout" in lowered:
        return "等待下载超时。请重试，完成登录后在法院页面点击下载。"
    if "target page, context or browser has been closed" in lowered or "browser has been closed" in lowered:
        return "接管窗口已关闭，未取得文书。请重试并在窗口内完成登录和下载。"
    if "executable doesn't exist" in lowered or "failed to launch" in lowered:
        return "浏览器未能启动。请确认 Google Chrome 可正常打开后重试。"
    if "download" in lowered and ("canceled" in lowered or "cancelled" in lowered):
        return "文书下载已取消，请重试或使用“手动导入”。"
    return "浏览器接管未取得文书。请重试，或下载后使用“手动导入”。"

INSTITUTION_TERMS = (
    "人民法院", "法院", "仲裁委员会", "仲裁院", "人民法庭", "12368",
    "人民法院在线服务", "中国审判流程信息公开网", "电子送达平台", "案件管理平台",
)
EVENT_TERMS = (
    "开庭", "庭询", "听证", "询问", "谈话", "送达", "传票", "立案", "受理",
    "举证", "答辩", "补正", "判决书", "裁定书", "调解书", "上诉", "异议",
    "复议", "缴费", "诉讼费", "仲裁费", "保全", "执行", "案件", "案号",
    "债权人会议",
)
CATEGORY_RULES: Sequence[Tuple[str, Sequence[str]]] = (
    ("变更/取消", ("取消开庭", "开庭取消", "取消", "延期", "改期", "变更", "撤销通知")),
    ("开庭/庭询/听证", ("开庭", "庭询", "听证", "传唤")),
    ("上诉/异议/复议期限", ("上诉期", "上诉期限", "异议期限", "复议期限", "申请复议")),
    ("举证/答辩/补正期限", ("举证期限", "答辩期限", "补正", "提交材料", "举证", "答辩")),
    ("裁判文书送达", ("判决书", "裁定书", "调解书", "决定书", "法律文书")),
    ("材料或电子送达", ("电子送达", "送达平台", "送达", "传票", "查收文书")),
    ("立案/受理/平台状态", ("立案", "受理", "分案", "案件状态", "审判流程")),
    ("调解/谈话/询问", ("调解", "谈话", "询问")),
    ("缴费/退费", ("缴费", "诉讼费", "仲裁费", "退费")),
    ("保全/执行", ("财产保全", "保全", "执行", "查封", "冻结", "扣押")),
)


def _now() -> str:
    return dt.datetime.now().replace(microsecond=0).isoformat(sep=" ")


def extract_document_urls(text: str) -> List[str]:
    """Extract and normalize HTTP(S) links without accepting local/file URLs."""
    found = re.findall(r"https?://[^\s<>\]）】\"'，。；：！？]+", text or "", flags=re.I)
    result: List[str] = []
    for raw in found:
        url = raw.rstrip("。；;，,!！?？)")
        if url not in result:
            result.append(url)
    return result


def _court_delivery_params(url: str) -> Dict[str, str]:
    """Return the public SMS-delivery parameters embedded in the SPA hash URL."""
    try:
        parsed = urllib.parse.urlparse(url)
        if parsed.hostname != "zxfw.court.gov.cn":
            return {}
        fragment_path, _, fragment_query = parsed.fragment.partition("?")
        if "/pagesAjkj/app/wssd/index" not in f"/{fragment_path.lstrip('/')}":
            return {}
        values = urllib.parse.parse_qs(fragment_query, keep_blank_values=True)
        result = {
            key: str((values.get(key) or [""])[0])
            for key in ("qdbh", "sdbh", "sdsin", "mm")
        }
        return result if result["qdbh"] and result["sdbh"] else {}
    except (TypeError, ValueError):
        return {}


def _is_court_platform_homepage(url: str) -> bool:
    """The optional login/homepage URL in court SMS messages is not a document."""
    try:
        parsed = urllib.parse.urlparse(url)
        return (
            parsed.hostname == "zxfw.court.gov.cn"
            and parsed.path.rstrip("/") in {"", "/zxfw"}
            and not _court_delivery_params(url)
        )
    except (TypeError, ValueError):
        return False


GZAC_SEARCH_HOSTS = {"new.gzac.org", "new.gziac.cn"}
GZAC_API_ORIGIN = "https://new.gzac.org"


def _is_gzac_search_url(url: str) -> bool:
    """Whether ``url`` is Guangzhou Arbitration Commission's SMS lookup page."""
    try:
        parsed = urllib.parse.urlparse(url)
        return (
            parsed.scheme in {"http", "https"}
            and (parsed.hostname or "").lower().rstrip(".") in GZAC_SEARCH_HOSTS
            and parsed.path.rstrip("/") == "/search"
        )
    except (TypeError, ValueError):
        return False


def _gzac_delivery_params(text: str) -> Dict[str, str]:
    """Extract the three public lookup credentials from a GZAC SMS.

    GZAC SMS messages do not contain a document URL.  Their public ``/search``
    page verifies the case year, serial number and six-character lookup code,
    then exposes the delivery attachments.  Keep this intentionally strict so
    unrelated messages containing a short alphanumeric token never become a
    download task.
    """
    message = text or ""
    urls = extract_document_urls(message)
    search_url = next((url for url in urls if _is_gzac_search_url(url)), "")
    if "广州仲裁委员会" not in message and not search_url:
        return {}
    case_match = re.search(
        r"[（(]\s*(20\d{2})\s*[）)]\s*穗仲案字第?\s*(\d{1,6})\s*号",
        message,
        flags=re.I,
    )
    code_match = re.search(
        r"(?:联系方式\s*(?:查询)?码|联系(?:方式)?查询码|查询码)\s*(?:为|是|[:：])?\s*([A-Za-z0-9]{4,6})(?![A-Za-z0-9])",
        message,
        flags=re.I,
    )
    if not case_match or not code_match:
        return {}
    return {
        "case_year": case_match.group(1),
        "case_no": case_match.group(2),
        "query_code": code_match.group(1).lower(),
        "source_url": search_url or f"{GZAC_API_ORIGIN}/search/",
    }


def _gzac_document_params(document: Dict[str, Any]) -> Dict[str, str]:
    """Read the normalized GZAC credentials stored on a document task."""
    values = _loads(document.get("source_path"), {})
    if not isinstance(values, dict):
        return {}
    year = _norm(values.get("case_year"))
    case_no = _norm(values.get("case_no"))
    code = _norm(values.get("query_code")).lower()
    if not re.fullmatch(r"20\d{2}", year) or not re.fullmatch(r"\d{1,6}", case_no):
        return {}
    if not re.fullmatch(r"[a-z0-9]{4,6}", code):
        return {}
    return {"case_year": year, "case_no": case_no, "query_code": code}


def _gzac_delivery_item_key(document: Dict[str, Any]) -> str:
    values = _loads(document.get("source_path"), {})
    return _norm(values.get("key"))[:500] if isinstance(values, dict) else ""


def _gzac_sms_ref_id(url: str) -> str:
    """Return the optional delivery reference embedded in a GZAC SMS link."""
    try:
        values = urllib.parse.parse_qs(urllib.parse.urlparse(url).query)
        return _norm((values.get("mc") or [""])[0])[:160]
    except (TypeError, ValueError):
        return ""


def _gzac_sms_search_id(url: str) -> str:
    """Return the ``s`` delivery id embedded in a GZAC SMS link.

    The portal's own result page passes this value as ``deliveryId`` when it
    loads the delivery materials, so the backend must use the same value.
    """
    try:
        values = urllib.parse.parse_qs(urllib.parse.urlparse(url).query)
        return _norm((values.get("s") or [""])[0])[:80]
    except (TypeError, ValueError):
        return ""


def _gdems_delivery_id(url: str) -> str:
    """Return the opaque id from a Guangdong EMS electronic-delivery URL."""
    try:
        parsed = urllib.parse.urlparse(url)
        if parsed.hostname != "sd.gdems.com":
            return ""
        match = re.fullmatch(r"/v3/dzsd/([A-Za-z0-9_-]{4,80})/?", parsed.path)
        return match.group(1) if match else ""
    except (TypeError, ValueError):
        return ""


def _html_download_candidates(html: str, base_url: str) -> List[str]:
    """Extract likely document/download URLs from an authenticated delivery page."""
    values: List[str] = []
    patterns = (
        r"(?:href|src|data-url|data-href|action)\s*=\s*['\"]([^'\"]+)['\"]",
        r"['\"]((?:https?:)?//[^'\"\s<>]+)['\"]",
        r"['\"](/[^'\"\s<>]+)['\"]",
    )
    for pattern in patterns:
        values.extend(re.findall(pattern, html or "", flags=re.I))
    blocked_suffixes = (".js", ".css", ".woff", ".woff2", ".ttf", ".svg", ".ico")
    candidates: List[tuple[int, str]] = []
    for raw in values:
        value = raw.replace("&amp;", "&").strip()
        if not value or value.startswith(("javascript:", "data:", "#")):
            continue
        url = urllib.parse.urljoin(base_url, value)
        parsed = urllib.parse.urlparse(url)
        if parsed.scheme not in {"http", "https"} or not parsed.hostname:
            continue
        lowered = f"{parsed.path}?{parsed.query}".lower()
        if lowered.endswith(blocked_suffixes) or any(
            marker in lowered for marker in ("validatecodeservlet", "/static/", "getcode", "validatecode")
        ):
            continue
        score = 0
        if re.search(r"\.(?:pdf|docx?|jpe?g|png|tiff?)(?:$|[?#])", lowered):
            score += 5
        if any(marker in lowered for marker in ("download", "file", "wj", "cl", "preview", "view")):
            score += 3
        if score and url not in {item[1] for item in candidates}:
            candidates.append((score, url))
    return [url for _, url in sorted(candidates, key=lambda item: -item[0])[:100]]


def _gdems_document_manifest(html: str, base_url: str) -> List[Dict[str, str]]:
    """Parse document ids and names from the verified Guangdong EMS page.

    The page does not expose ordinary links. Its ``downloadWs`` JavaScript
    combines the SMS link id, a per-delivery token, and each document id.
    """
    function_match = re.search(
        r"function\s+downloadWs\s*\([^)]*\)\s*\{(.*?)\n\s*\}",
        html or "",
        flags=re.I | re.S,
    )
    if not function_match:
        return []
    string_parts = re.findall(r"['\"]([^'\"]+)['\"]", function_match.group(1))
    prefix_text = "".join(string_parts)
    prefix_match = re.search(r"(/v3/dzsd/d/[^\s]+/)", prefix_text)
    if not prefix_match:
        return []
    prefix = urllib.parse.urljoin(base_url, prefix_match.group(1))
    entries = re.findall(
        r"onclick\s*=\s*['\"]downloadWs\(['\"]([^'\"]+)['\"]\)['\"][^>]*>"
        r"(.*?)</div>",
        html or "",
        flags=re.I | re.S,
    )
    result: List[Dict[str, str]] = []
    for index, (document_id, block) in enumerate(entries[:100]):
        label_match = re.search(r"<label[^>]*>(.*?)</label>", block, flags=re.I | re.S)
        label = re.sub(r"<[^>]+>", " ", label_match.group(1) if label_match else "")
        label = re.sub(r"\s+", " ", label).strip()
        label = re.sub(r"\s+\.(pdf|docx?|jpe?g|png|tiff?)$", r".\1", label, flags=re.I)
        filename = re.sub(r'[<>:"/\\|?*]', "_", label).strip(" ._")
        if not filename:
            filename = f"法院文书-{index + 1}.pdf"
        result.append({
            "key": document_id,
            "url": urllib.parse.urljoin(prefix, document_id),
            "filename": filename,
        })
    return result


def _validate_public_url(url: str) -> None:
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise HTTPException(status_code=400, detail="仅允许下载公开的 HTTP/HTTPS 链接")
    host = parsed.hostname.lower().rstrip(".")
    if host in {"localhost", "localhost.localdomain"}:
        raise HTTPException(status_code=400, detail="不允许访问本机或内网链接")
    try:
        addresses = {
            item[4][0] for item in socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80))
        }
    except socket.gaierror as exc:
        raise HTTPException(status_code=502, detail=f"无法解析下载地址：{host}") from exc
    for address in addresses:
        ip = ipaddress.ip_address(address.split("%", 1)[0])
        if (
            ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_multicast
            or ip.is_reserved or ip.is_unspecified
        ):
            raise HTTPException(status_code=400, detail="不允许访问本机或内网链接")


def _safe_attachment_path(path_text: str) -> Path:
    path = Path(path_text.replace("~", str(Path.home()), 1)).expanduser().resolve()
    try:
        path.relative_to(MESSAGE_ATTACHMENT_DIR.resolve())
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="短信附件路径不在信息附件目录内") from exc
    if not path.is_file():
        raise HTTPException(status_code=404, detail="短信附件尚未同步到本机")
    return path


def _valid_document(content: bytes, filename: str, content_type: str = "") -> tuple[bool, str]:
    suffix = Path(filename).suffix.lower()
    if content.startswith(b"%PDF-"):
        return True, ".pdf"
    if content.startswith(b"PK"):
        try:
            from io import BytesIO
            with zipfile.ZipFile(BytesIO(content)) as archive:
                if "word/document.xml" in archive.namelist():
                    return True, ".docx"
        except zipfile.BadZipFile:
            pass
    if "text/html" in (content_type or "").lower() or content.lstrip().lower().startswith(
        (b"<!doctype html", b"<html")
    ):
        # HTML is not itself a document. Distinguish an SMS recipient check
        # from an account login so public electronic-delivery links are not
        # incorrectly reported as requiring a court-platform account.
        html = content[:512_000].decode("utf-8", errors="ignore").lower()
        if (
            "validatecode" in html
            and ("getcode" in html or "短信验证码" in html)
            and ("手机号" in html or "手机号码" in html)
        ):
            return False, "verification"
        if any(marker in html for marker in ("登录", "login", "sign in", "账号", "密码")):
            return False, "login"
        return False, "interaction"
    if suffix in {".jpg", ".jpeg", ".png", ".tif", ".tiff"}:
        return True, suffix
    return False, ""


def _document_case_numbers(content: bytes, suffix: str) -> List[str]:
    """Best-effort text extraction used only to stop conflicting auto-archives."""
    try:
        from io import BytesIO
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader
            text = "\n".join((page.extract_text() or "") for page in PdfReader(BytesIO(content)).pages[:8])
        elif suffix == ".docx":
            from docx import Document
            document = Document(BytesIO(content))
            text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
        else:
            return []
        return extract_case_numbers(text)
    except Exception:
        return []


def _infer_document_type(content: bytes, suffix: str, original: str = "") -> str:
    candidates = (
        "民事判决书", "刑事判决书", "行政判决书", "判决书", "民事裁定书", "刑事裁定书",
        "行政裁定书", "裁定书", "民事调解书", "调解书", "开庭传票", "传票",
        "应诉通知书", "举证通知书", "受理案件通知书", "执行通知书", "送达回证", "通知书",
    )
    original_stem = Path(original).stem
    for name in candidates:
        if name in original_stem:
            return name
    try:
        from io import BytesIO
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except ImportError:
                from PyPDF2 import PdfReader
            text = "\n".join((page.extract_text() or "") for page in PdfReader(BytesIO(content)).pages[:3])
        elif suffix == ".docx":
            from docx import Document
            document = Document(BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs[:80])
        else:
            text = ""
        for name in candidates:
            if name in text:
                return name
    except Exception:
        pass
    return ""


def _document_filename(message: Dict[str, Any], original: str, suffix: str, document_type: str = "") -> str:
    received = str(message.get("received_at") or "")[:10].replace("-", "") or dt.date.today().strftime("%Y%m%d")
    institution = re.sub(r'[<>:"/\\|?*]', "_", str(message.get("institution") or "法院"))
    category = re.sub(r'[<>:"/\\|?*]', "_", str(document_type or message.get("category") or ""))
    if not document_type and category in {"", "材料或电子送达", "裁判文书送达", "其他重要通知"}:
        original_name = re.sub(r'[<>:"/\\|?*]', "_", Path(original).stem).strip(" ._")
        return f"{original_name or '法院文书'}{suffix}"
    category = category or "司法文书"
    case_numbers = _loads(message.get("case_numbers_json"), [])
    case_no = re.sub(r'[<>:"/\\|?*]', "_", str(case_numbers[0] if case_numbers else "待确认案号"))
    return f"{received}_{institution}_{category}_{case_no}{suffix}"


def _message_attachments_by_rowid(path: Path, source_rowids: Sequence[int]) -> Dict[int, List[Dict[str, str]]]:
    if not source_rowids:
        return {}
    conn = sqlite3.connect(f"file:{path}?mode=ro", uri=True, timeout=5)
    conn.row_factory = sqlite3.Row
    try:
        tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if not {"attachment", "message_attachment_join"}.issubset(tables):
            return {}
        placeholders = ",".join("?" for _ in source_rowids)
        rows = conn.execute(
            f"""SELECT j.message_id,a.filename,a.transfer_name,a.mime_type
                FROM message_attachment_join j JOIN attachment a ON a.ROWID=j.attachment_id
                WHERE j.message_id IN ({placeholders})""",
            [int(item) for item in source_rowids],
        ).fetchall()
        result: Dict[int, List[Dict[str, str]]] = {}
        for row in rows:
            filename = str(row["filename"] or "")
            if not filename:
                continue
            result.setdefault(int(row["message_id"]), []).append({
                "path": filename,
                "name": str(row["transfer_name"] or Path(filename).name),
                "mime_type": str(row["mime_type"] or ""),
            })
        return result
    finally:
        conn.close()


def full_disk_access_target() -> str:
    """Return the exact executable macOS TCC attributes as responsible."""
    return os.environ.get("LEGAL_KB_FDA_TARGET", os.path.realpath(sys.executable))


def _norm(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _loads(value: Any, fallback: Any) -> Any:
    try:
        return json.loads(value or "")
    except (TypeError, ValueError, json.JSONDecodeError):
        return fallback


def normalize_case_no(value: Any) -> str:
    text = _norm(value).translate(str.maketrans({"(": "（", ")": "）", "﹙": "（", "﹚": "）"}))
    return re.sub(r"[\s|｜]+", "", text)


def extract_case_numbers(text: str) -> List[str]:
    found: List[str] = []
    patterns = (
        r"(?:案\s*号|编号|原审案号)\s*[：:]?\s*([（(]\s*\d{4}\s*[）)]\s*[\u4e00-\u9fffA-Za-z0-9第#\-\s]{1,70}?号)",
        r"([（(]\s*\d{4}\s*[）)]\s*[\u4e00-\u9fffA-Za-z0-9第#\-\s]{1,70}?号)",
    )
    for pattern in patterns:
        for match in re.finditer(pattern, text or ""):
            value = normalize_case_no(match.group(1).replace("#", "粤"))
            if value and value not in found:
                found.append(value)
    return found


def classify_text(text: str) -> str:
    compact = re.sub(r"\s+", "", text or "")
    if is_bankruptcy_meeting_notice(compact):
        return "其他重要通知"
    for category, words in CATEGORY_RULES:
        if any(word in compact for word in words):
            return category
    return "其他重要通知" if any(word in compact for word in EVENT_TERMS) else ""


def is_bankruptcy_meeting_notice(text: str) -> bool:
    compact = re.sub(r"\s+", "", text or "")
    return "债权人会议" in compact and any(term in compact for term in ("管理人", "投票表决", "破产", "重整"))


def is_verification_code_text(text: str) -> bool:
    """Verification codes are authentication traffic, not case reminders."""
    compact = re.sub(r"\s+", "", text or "")
    if "验证码" not in compact or extract_case_numbers(compact):
        return False
    return bool(re.search(r"验证码(?:是|为)?[:：]?([A-Za-z0-9]{4,8})(?![A-Za-z0-9])", compact))


def extract_institution(text: str) -> str:
    candidates = re.findall(r"([\u4e00-\u9fff]{2,35}(?:人民法院|法院|仲裁委员会|仲裁院|人民法庭))", text or "")
    if candidates:
        return min(candidates, key=len)
    # Bankruptcy notices often end with the debtor group's administrator's signature.
    for line in reversed((text or "").splitlines()):
        match = re.search(r"([\u4e00-\u9fff\d]{2,60}(?:公司|企业)(?:等\d+家关联企业)?管理人)", line.strip())
        if match:
            return match.group(1)
    for term in ("人民法院在线服务", "中国审判流程信息公开网", "电子送达平台", "12368"):
        if term in (text or ""):
            return term
    return ""


def _datetime_candidates(text: str, received_at: str = "", *, allow_inferred_year: bool = True) -> List[str]:
    base = dt.datetime.now()
    try:
        base = dt.datetime.fromisoformat((received_at or "").replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        pass
    values: List[str] = []
    patterns = (
        r"(20\d{2})\s*[年\-/]\s*(\d{1,2})\s*[月\-/]\s*(\d{1,2})\s*日?(?:星期[一二三四五六日天])?[，,\s]*(上午|下午|晚上|中午)?\s*(\d{1,2})(?:\s*[:：时]\s*(\d{1,2}))?",
        r"(?<!\d)(\d{1,2})\s*月\s*(\d{1,2})\s*日(?:星期[一二三四五六日天])?[，,\s]*(上午|下午|晚上|中午)?\s*(\d{1,2})(?:\s*[:：时]\s*(\d{1,2}))?",
    )
    for index, pattern in enumerate(patterns):
        if index == 1 and not allow_inferred_year:
            continue
        for match in re.finditer(pattern, text or ""):
            groups = match.groups()
            if index == 0:
                year, month, day = map(int, groups[:3]); meridiem = groups[3] or ""; hour = int(groups[4]); minute = int(groups[5] or 0)
            else:
                year = base.year; month, day = map(int, groups[:2]); meridiem = groups[2] or ""; hour = int(groups[3]); minute = int(groups[4] or 0)
            if meridiem in {"下午", "晚上"} and hour < 12:
                hour += 12
            if meridiem == "中午" and hour < 11:
                hour += 12
            try:
                parsed = dt.datetime(year, month, day, hour, minute)
                if index == 1 and parsed < base - dt.timedelta(days=180):
                    parsed = parsed.replace(year=year + 1)
                value = parsed.isoformat(sep=" ", timespec="minutes")
                if value not in values:
                    values.append(value)
            except ValueError:
                continue
    return values


def extract_fields(text: str, received_at: str = "", *, allow_inferred_year: bool = True) -> Dict[str, Any]:
    category = classify_text(text)
    datetimes = _datetime_candidates(text, received_at, allow_inferred_year=allow_inferred_year)
    location = ""
    match = re.search(r"(?:开庭地点|应到处所|地点|地址)\s*[：:]?\s*([^。；;\n]{2,120})", text or "")
    if match:
        location = match.group(1).strip()
    event_type = ""
    for token in ("二审庭询", "二审开庭", "再审开庭", "执行异议听证", "开庭", "庭询", "听证", "询问", "谈话", "调解"):
        if token in (text or ""):
            event_type = token
            break
    return {
        "category": category,
        "institution": extract_institution(text),
        "case_numbers": extract_case_numbers(text),
        "event_at": datetimes[0] if datetimes else "",
        "datetimes": datetimes,
        "location": location,
        "event_type": event_type,
        "summary": re.sub(r"\s+", " ", text or "").strip()[:180],
    }


def _apple_date(value: Any) -> str:
    try:
        number = float(value or 0)
    except (TypeError, ValueError):
        return ""
    if not number:
        return ""
    if number > 1e15:
        number /= 1e9
    elif number > 1e12:
        number /= 1e6
    elif number > 1e10:
        number /= 1e3
    moment = dt.datetime(2001, 1, 1, tzinfo=dt.timezone.utc) + dt.timedelta(seconds=number)
    return moment.astimezone().replace(tzinfo=None, microsecond=0).isoformat(sep=" ")


def _attributed_text(value: Any) -> str:
    if not value:
        return ""
    data = bytes(value) if not isinstance(value, bytes) else value
    decoded = data.decode("utf-8", errors="ignore")
    # streamtyped attributed strings put the visible NSString between binary
    # control bytes. Keep all printable Unicode punctuation (including “”,
    # ASCII commas and currency symbols) instead of an allow-list that can
    # truncate legitimate court or arbitration content.
    chunks = re.findall(r"[^\x00-\x08\x0b\x0c\x0e-\x1f\x7f]{4,}", decoded)
    chunks = [re.sub(r"\s+", " ", item).strip() for item in chunks]
    chunks = [item for item in chunks if item and not any(marker in item for marker in ("NSString", "NSObject", "NSKeyedArchiver", "DDScannerResult"))]
    if not chunks:
        return ""
    def score(item: str) -> Tuple[int, int, int]:
        cjk = len(re.findall(r"[\u3400-\u9fff]", item))
        return (1 if "【" in item or cjk >= 4 else 0, cjk, min(len(item), 5000))
    return max(chunks, key=score)[:5000]


def _best_message_body(text: Any, attributed_body: Any) -> str:
    plain = _norm(text)
    rich = _attributed_text(attributed_body)
    if not plain:
        return rich
    if rich and len(rich) > len(plain) and plain in rich:
        return rich
    return plain


def _message_bodies_by_rowid(path: Path, source_rowids: Sequence[int]) -> Dict[int, str]:
    ids = sorted({int(value) for value in source_rowids if int(value) > 0})
    if not ids:
        return {}
    conn = sqlite3.connect(f"file:{path}?mode=ro", uri=True, timeout=5)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("PRAGMA query_only=ON")
        columns = {row[1] for row in conn.execute("PRAGMA table_info(message)").fetchall()}
        text_expr = "text" if "text" in columns else "NULL AS text"
        attributed_expr = "attributedBody" if "attributedBody" in columns else "NULL AS attributedBody"
        result: Dict[int, str] = {}
        for start in range(0, len(ids), 500):
            batch = ids[start:start + 500]
            placeholders = ",".join("?" for _ in batch)
            rows = conn.execute(
                f"SELECT ROWID AS source_rowid,{text_expr},{attributed_expr} FROM message WHERE ROWID IN ({placeholders})",
                batch,
            ).fetchall()
            for row in rows:
                result[int(row["source_rowid"])] = _best_message_body(row["text"], row["attributedBody"])
        return result
    finally:
        conn.close()


def _message_rows(path: Path, last_rowid: int, initial: bool, backfill_days: int, replay_rowids: Sequence[int] = ()) -> Tuple[List[Dict[str, Any]], int]:
    uri = f"file:{path}?mode=ro"
    conn = sqlite3.connect(uri, uri=True, timeout=5)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("PRAGMA query_only=ON")
        columns = {row[1] for row in conn.execute("PRAGMA table_info(message)").fetchall()}
        if not {"date", "is_from_me"}.issubset(columns):
            raise RuntimeError("无法识别当前 macOS 的信息数据库结构")
        selected = ["m.ROWID AS source_rowid"]
        for column in ("guid", "text", "attributedBody", "service", "date"):
            selected.append(f"m.{column}" if column in columns else f"NULL AS {column}")
        join = "LEFT JOIN handle h ON h.ROWID=m.handle_id" if "handle_id" in columns else ""
        selected.append("h.id AS sender" if join else "'' AS sender")
        if initial:
            query = f"SELECT {','.join(selected)} FROM message m {join} WHERE COALESCE(m.is_from_me,0)=0 ORDER BY m.ROWID DESC LIMIT 50000"
            raw = list(reversed(conn.execute(query).fetchall()))
        else:
            replay_ids = sorted({int(value) for value in replay_rowids if int(value) > 0})
            query = f"SELECT {','.join(selected)} FROM message m {join} WHERE COALESCE(m.is_from_me,0)=0 AND m.ROWID>? ORDER BY m.ROWID LIMIT 10000"
            raw = conn.execute(query, (last_rowid,)).fetchall()
            for start in range(0, len(replay_ids), 500):
                batch = replay_ids[start:start + 500]
                placeholders = ",".join("?" for _ in batch)
                replay_query = f"SELECT {','.join(selected)} FROM message m {join} WHERE COALESCE(m.is_from_me,0)=0 AND m.ROWID IN ({placeholders})"
                raw += conn.execute(replay_query, batch).fetchall()
            raw = sorted({int(row["source_rowid"]): row for row in raw}.values(), key=lambda row: row["source_rowid"])
        cutoff = dt.datetime.now() - dt.timedelta(days=max(1, backfill_days))
        rows: List[Dict[str, Any]] = []
        highwater = last_rowid
        for row in raw:
            item = dict(row)
            highwater = max(highwater, int(item.get("source_rowid") or 0))
            received_at = _apple_date(item.get("date"))
            try:
                received_dt = dt.datetime.fromisoformat(received_at)
            except ValueError:
                received_dt = dt.datetime.now()
            if received_dt < cutoff:
                continue
            body = _best_message_body(item.get("text"), item.get("attributedBody"))
            if not body:
                continue
            rows.append({
                "source_rowid": int(item.get("source_rowid") or 0),
                "source_guid": _norm(item.get("guid")),
                "sender": _norm(item.get("sender")),
                "service": _norm(item.get("service")),
                "received_at": received_at,
                "body": body[:10000],
            })
        attachments = _message_attachments_by_rowid(
            path, [int(item["source_rowid"]) for item in rows]
        )
        for item in rows:
            item["attachments"] = attachments.get(int(item["source_rowid"]), [])
        return rows, highwater
    finally:
        conn.close()


def _ensure_schema(get_db: Callable[[], Any]) -> None:
    stamp = _now()
    with get_db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS case_manager_message_settings (
                id INTEGER PRIMARY KEY CHECK(id=1), enabled INTEGER NOT NULL DEFAULT 1,
                poll_interval_seconds INTEGER NOT NULL DEFAULT 300,
                backfill_days INTEGER NOT NULL DEFAULT 90, ai_enabled INTEGER NOT NULL DEFAULT 1,
                ai_model TEXT NOT NULL DEFAULT '', hmac_secret TEXT NOT NULL DEFAULT '',
                last_source_rowid INTEGER NOT NULL DEFAULT 0, initial_scan_completed INTEGER NOT NULL DEFAULT 0,
                last_scan_started_at TEXT DEFAULT '', last_scan_completed_at TEXT DEFAULT '',
                last_success_at TEXT DEFAULT '', last_error TEXT DEFAULT '', last_db_mtime TEXT DEFAULT '',
                last_message_at TEXT DEFAULT '', recipient_phones_json TEXT DEFAULT '[]', updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS case_manager_judicial_messages (
                id INTEGER PRIMARY KEY AUTOINCREMENT, source_guid TEXT DEFAULT '', source_rowid INTEGER DEFAULT 0,
                fingerprint TEXT NOT NULL UNIQUE, sender TEXT DEFAULT '', service TEXT DEFAULT '',
                received_at TEXT NOT NULL, raw_body TEXT DEFAULT '', institution TEXT DEFAULT '',
                category TEXT DEFAULT '', extracted_json TEXT DEFAULT '{}', analysis_json TEXT DEFAULT '{}',
                candidate_cases_json TEXT DEFAULT '[]', suggested_case_id INTEGER,
                matched_case_id INTEGER, case_numbers_json TEXT DEFAULT '[]', workflow_status TEXT NOT NULL DEFAULT 'pending',
                decision_mode TEXT DEFAULT '', relevance_score INTEGER DEFAULT 0, category_confidence INTEGER DEFAULT 0,
                case_confidence INTEGER DEFAULT 0, ai_model TEXT DEFAULT '', prompt_version TEXT DEFAULT '',
                analysis_error TEXT DEFAULT '', analyzed_at TEXT DEFAULT '', confirmed_at TEXT DEFAULT '',
                imported_scan_run_id INTEGER DEFAULT 0, is_new INTEGER NOT NULL DEFAULT 0,
                is_deleted INTEGER NOT NULL DEFAULT 0, deleted_at TEXT DEFAULT '',
                created_at TEXT NOT NULL, updated_at TEXT NOT NULL,
                FOREIGN KEY(suggested_case_id) REFERENCES case_manager_cases(id) ON DELETE SET NULL,
                FOREIGN KEY(matched_case_id) REFERENCES case_manager_cases(id) ON DELETE SET NULL
            );
            CREATE UNIQUE INDEX IF NOT EXISTS idx_case_manager_jm_guid ON case_manager_judicial_messages(source_guid) WHERE source_guid<>'';
            CREATE INDEX IF NOT EXISTS idx_case_manager_jm_status_time ON case_manager_judicial_messages(workflow_status,received_at DESC);
            CREATE INDEX IF NOT EXISTS idx_case_manager_jm_case ON case_manager_judicial_messages(matched_case_id,suggested_case_id);
            CREATE TABLE IF NOT EXISTS case_manager_message_fingerprints (
                fingerprint TEXT PRIMARY KEY, source_guid TEXT DEFAULT '', source_rowid INTEGER DEFAULT 0,
                received_at TEXT DEFAULT '', disposition TEXT NOT NULL DEFAULT 'irrelevant', created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS case_manager_trusted_senders (
                sender TEXT PRIMARY KEY, institution TEXT DEFAULT '', verified INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL, updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS case_manager_message_links (
                id INTEGER PRIMARY KEY AUTOINCREMENT, message_id INTEGER NOT NULL, case_id INTEGER NOT NULL,
                target_type TEXT NOT NULL, target_id INTEGER NOT NULL, creation_mode TEXT NOT NULL,
                target_created_at TEXT NOT NULL, created_at TEXT NOT NULL,
                UNIQUE(message_id,target_type),
                FOREIGN KEY(message_id) REFERENCES case_manager_judicial_messages(id) ON DELETE CASCADE,
                FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS idx_case_manager_message_links_target ON case_manager_message_links(target_type,target_id);
            CREATE TABLE IF NOT EXISTS case_manager_message_documents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                message_id INTEGER NOT NULL,
                source_kind TEXT NOT NULL DEFAULT 'url',
                source_url TEXT DEFAULT '',
                source_path TEXT DEFAULT '',
                original_filename TEXT DEFAULT '',
                mime_type TEXT DEFAULT '',
                status TEXT NOT NULL DEFAULT 'discovered',
                case_id INTEGER,
                staging_path TEXT DEFAULT '',
                file_id INTEGER,
                final_filename TEXT DEFAULT '',
                sha256 TEXT DEFAULT '',
                attempts INTEGER NOT NULL DEFAULT 0,
                last_error TEXT DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(message_id,source_kind,source_url,source_path),
                FOREIGN KEY(message_id) REFERENCES case_manager_judicial_messages(id) ON DELETE CASCADE,
                FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE SET NULL,
                FOREIGN KEY(file_id) REFERENCES case_manager_files(id) ON DELETE SET NULL
            );
            CREATE INDEX IF NOT EXISTS idx_case_manager_message_documents
                ON case_manager_message_documents(message_id,status,created_at);
            CREATE TABLE IF NOT EXISTS case_manager_message_scan_runs (
                id INTEGER PRIMARY KEY AUTOINCREMENT, started_at TEXT NOT NULL, completed_at TEXT DEFAULT '',
                status TEXT NOT NULL DEFAULT 'running', scanned_count INTEGER DEFAULT 0, candidate_count INTEGER DEFAULT 0,
                auto_written_count INTEGER DEFAULT 0, error TEXT DEFAULT ''
            );
            """
        )
        message_columns = {row[1] for row in conn.execute("PRAGMA table_info(case_manager_judicial_messages)").fetchall()}
        added_is_new = "is_new" not in message_columns
        migrations = (
            ("imported_scan_run_id", "INTEGER DEFAULT 0"),
            ("is_new", "INTEGER NOT NULL DEFAULT 0"),
            ("is_deleted", "INTEGER NOT NULL DEFAULT 0"),
            ("deleted_at", "TEXT DEFAULT ''"),
        )
        for column, definition in migrations:
            if column not in message_columns:
                conn.execute(f"ALTER TABLE case_manager_judicial_messages ADD COLUMN {column} {definition}")
        settings_columns = {row[1] for row in conn.execute("PRAGMA table_info(case_manager_message_settings)").fetchall()}
        if "recipient_phones_json" not in settings_columns:
            conn.execute(
                "ALTER TABLE case_manager_message_settings ADD COLUMN recipient_phones_json TEXT DEFAULT '[]'"
            )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_case_manager_jm_new ON case_manager_judicial_messages(is_deleted,is_new,created_at DESC)")
        # On the first upgrade, preserve the most recent non-empty scan as the
        # visible “new messages” batch instead of making every historical row new.
        if added_is_new:
            latest_import = conn.execute(
                """SELECT id,started_at,completed_at FROM case_manager_message_scan_runs
                   WHERE status='completed' AND candidate_count>0
                   ORDER BY id DESC LIMIT 1"""
            ).fetchone()
            if latest_import:
                conn.execute(
                    """UPDATE case_manager_judicial_messages
                       SET is_new=1,imported_scan_run_id=?
                       WHERE created_at>=? AND created_at<=? AND COALESCE(is_deleted,0)=0""",
                    (latest_import["id"], latest_import["started_at"], latest_import["completed_at"]),
                )
        row = conn.execute("SELECT id FROM case_manager_message_settings WHERE id=1").fetchone()
        if not row:
            conn.execute(
                """INSERT INTO case_manager_message_settings
                   (id,enabled,poll_interval_seconds,backfill_days,ai_enabled,ai_model,hmac_secret,updated_at)
                   VALUES (1,1,300,90,1,?,?,?)""",
                (DEFAULT_AI_MODEL, secrets.token_hex(32), stamp),
            )
        existing_phones = conn.execute(
            "SELECT recipient_phones_json FROM case_manager_message_settings WHERE id=1"
        ).fetchone()
        phones = [
            value for value in _loads(existing_phones[0] if existing_phones else "[]", [])
            if re.fullmatch(r"1\d{10}", str(value))
        ]
        for default_phone in ("18600001011", "17600000902"):
            if default_phone not in phones:
                phones.append(default_phone)
        conn.execute(
            "UPDATE case_manager_message_settings SET recipient_phones_json=? WHERE id=1",
            (_json(phones),),
        )
        conn.execute(
            """INSERT OR IGNORE INTO case_manager_trusted_senders(sender,institution,verified,created_at,updated_at)
               VALUES ('12368','人民法院',0,?,?)""", (stamp, stamp)
        )
        conn.commit()


def annotate_message_sources(get_db: Callable[[], Any], items: List[Dict[str, Any]], item_type: str = "") -> List[Dict[str, Any]]:
    if not items:
        return items
    keys = [(item_type or _norm(item.get("item_type")), int(item.get("id") or 0)) for item in items]
    valid = [(kind, item_id) for kind, item_id in keys if kind and item_id]
    if not valid:
        return items
    try:
        rows = []
        with get_db() as conn:
            for start in range(0, len(valid), 300):
                batch = valid[start:start + 300]
                clauses = " OR ".join("(l.target_type=? AND l.target_id=?)" for _ in batch)
                args: List[Any] = []
                for kind, item_id in batch:
                    args.extend([kind, item_id])
                rows.extend(conn.execute(
                    f"""SELECT l.target_type,l.target_id,l.message_id,l.creation_mode,m.received_at
                         FROM case_manager_message_links l JOIN case_manager_judicial_messages m ON m.id=l.message_id
                         WHERE {clauses}""", args,
                ).fetchall())
        mapping = {(row["target_type"], int(row["target_id"])): dict(row) for row in rows}
        for item in items:
            kind = item_type or _norm(item.get("item_type"))
            source = mapping.get((kind, int(item.get("id") or 0)))
            if source:
                item["source_message_id"] = source["message_id"]
                item["message_created_mode"] = source["creation_mode"]
                item["message_source_received_at"] = source["received_at"]
        return items
    except sqlite3.OperationalError:
        return items


def register_judicial_message_routes(
    app: Any,
    get_db: Callable[[], Any],
    callbacks: Dict[str, Callable[..., Any]],
) -> None:
    _ensure_schema(get_db)
    JUDICIAL_STAGING_DIR.mkdir(parents=True, exist_ok=True)
    scheduler: Dict[str, Any] = {"task": None, "scan_lock": None}
    # Short-lived in-memory sessions for recipient verification. Phone
    # numbers, captcha answers, and SMS codes are never written to disk.
    delivery_verifications: Dict[int, Dict[str, Any]] = {}

    # Request-local transfer counters. No URLs, credentials or document content.
    download_progress: Dict[tuple[int, int], Dict[str, Any]] = {}
    current_progress: ContextVar = ContextVar("judicial_download_progress", default=None)

    def report_download(**values: Any) -> None:
        progress = current_progress.get()
        if progress is not None:
            progress.update(values)
            if progress.get('_center_id') and (set(values) != {'received_bytes'} or time.monotonic()-progress.get('_center_stamp',0)>.4):
                progress['_center_stamp']=time.monotonic()
                callbacks['download_center_progress'](progress['_center_id'],progress)

    def begin_download_file(filename: str) -> None:
        report_download(phase="downloading", filename=filename, received_bytes=0, total_bytes=0)

    def complete_download_file(failed: bool = False) -> None:
        progress = current_progress.get()
        if progress is not None:
            field = "failed_files" if failed else "completed_files"
            progress[field] += 1
            report_download()

    def verification_session(document_id: int) -> Dict[str, Any]:
        session = delivery_verifications.get(document_id)
        if not session or session.get("expires_at", 0) < dt.datetime.now().timestamp():
            delivery_verifications.pop(document_id, None)
            raise HTTPException(status_code=410, detail="验证会话已过期，请重新获取图形验证码")
        return session

    def verification_document(message_id: int, document_id: int) -> Dict[str, Any]:
        with get_db() as conn:
            row = conn.execute(
                """SELECT d.*,m.matched_case_id,m.suggested_case_id,m.case_numbers_json
                   FROM case_manager_message_documents d
                   JOIN case_manager_judicial_messages m ON m.id=d.message_id
                   WHERE d.id=? AND d.message_id=?""",
                (document_id, message_id),
            ).fetchone()
        if not row or not _gdems_delivery_id(row["source_url"] or ""):
            raise HTTPException(status_code=404, detail="这不是可内嵌验证的电子送达链接")
        return dict(row)

    async def gdems_request(
        session: Dict[str, Any], method: str, url: str, **kwargs: Any
    ) -> httpx.Response:
        _validate_public_url(url)
        client: httpx.AsyncClient = session["client"]
        try:
            response = await client.request(method, url, **kwargs)
        except httpx.HTTPError as exc:
            raise HTTPException(status_code=502, detail="电子送达平台暂时无法访问") from exc
        return response

    async def close_verification(document_id: int) -> None:
        session = delivery_verifications.pop(document_id, None)
        if session:
            await session["client"].aclose()

    def response_filename(response: httpx.Response, fallback: str = "法院文书") -> str:
        disposition = response.headers.get("content-disposition") or ""
        match = re.search(r"filename\\*?=(?:UTF-8''|[\"']?)([^\"';]+)", disposition, re.I)
        if match:
            return urllib.parse.unquote(match.group(1))
        return Path(urllib.parse.urlparse(str(response.url)).path).name or fallback

    async def collect_gdems_browser_downloads(session: Dict[str, Any]) -> List[tuple[bytes, str, str, str]]:
        """Use the verified cookies in a headless browser and collect files from UI actions/XHR."""
        try:
            from playwright.async_api import async_playwright
        except ImportError:
            return []
        collected: List[tuple[bytes, str, str, str]] = []
        seen: set[str] = set()
        diagnostics: List[str] = []

        def add(content: bytes, filename: str, content_type: str, url: str) -> None:
            if len(content) > MAX_DOCUMENT_BYTES:
                return
            valid, detected = _valid_document(content, filename, content_type)
            if not valid:
                return
            digest = hashlib.sha256(content).hexdigest()
            if digest in seen:
                return
            seen.add(digest)
            if Path(filename).suffix.lower() != detected:
                filename = f"{Path(filename).stem or '法院文书'}{detected}"
            collected.append((content, filename, content_type, url))

        try:
            async with async_playwright() as playwright:
                options: Dict[str, Any] = {"headless": True}
                executable = _browser_executable()
                if executable:
                    options["executable_path"] = executable
                diagnostics.append(f"browser executable={executable or 'bundled'}")
                browser = await playwright.chromium.launch(**options)
                try:
                    context = await browser.new_context(accept_downloads=True)
                    cookies = []
                    for cookie in session["client"].cookies.jar:
                        cookies.append({
                            "name": cookie.name,
                            "value": cookie.value,
                            "domain": cookie.domain or "sd.gdems.com",
                            "path": cookie.path or "/",
                            "secure": bool(cookie.secure),
                        })
                    diagnostics.append(f"cookies {[(item['name'], item['domain'], item['path']) for item in cookies]}")
                    if cookies:
                        await context.add_cookies(cookies)
                    pending: List[asyncio.Task[Any]] = []

                    async def proxy_gdems(route: Any, request: Any) -> None:
                        """Run delivery-page traffic through the verified HTTP session.

                        The gateway closes connections from headless Chrome, but
                        accepts the same requests from httpx. Keeping the DOM in
                        Chromium while proxying its network calls preserves the
                        site's JavaScript-generated download requests.
                        """
                        try:
                            method = request.method.upper()
                            request_headers = {
                                key: value for key, value in request.headers.items()
                                if key.lower() not in {
                                    "host", "content-length", "cookie", "accept-encoding",
                                    "connection", "sec-fetch-dest", "sec-fetch-mode",
                                    "sec-fetch-site", "sec-ch-ua", "sec-ch-ua-mobile",
                                    "sec-ch-ua-platform",
                                }
                            }
                            body = request.post_data_buffer if method not in {"GET", "HEAD"} else None
                            response = await gdems_request(
                                session,
                                method,
                                request.url,
                                headers=request_headers,
                                content=body,
                            )
                            response_headers = {
                                key: value for key, value in response.headers.items()
                                if key.lower() not in {
                                    "content-encoding", "content-length", "transfer-encoding",
                                    "connection", "set-cookie",
                                }
                            }
                            diagnostics.append(
                                f"proxy {method} {response.status_code} "
                                f"{response.headers.get('content-type','')[:50]} "
                                f"{len(response.content)} {request.url[:300]}"
                            )
                            await route.fulfill(
                                status=response.status_code,
                                headers=response_headers,
                                body=response.content,
                            )
                        except Exception as exc:
                            diagnostics.append(f"proxy-error {request.method} {request.url[:250]} {exc}")
                            await route.abort("failed")

                    await context.route("https://sd.gdems.com/**", proxy_gdems)

                    async def capture_response(response: Any) -> None:
                        try:
                            parsed_url = urllib.parse.urlparse(response.url)
                            if parsed_url.hostname != "sd.gdems.com":
                                return
                            content_type = response.headers.get("content-type", "")
                            disposition = response.headers.get("content-disposition", "")
                            resource_type = response.request.resource_type
                            interesting = (
                                resource_type in {"document", "xhr", "fetch"}
                                or bool(disposition)
                                or any(value in content_type.lower() for value in ("pdf", "word", "image/", "octet-stream", "zip"))
                            )
                            if not interesting:
                                return
                            content = await response.body()
                            diagnostics.append(
                                f"{response.status} {response.request.method} {resource_type} "
                                f"{content_type[:50]} {len(content)} {response.url[:300]}"
                            )
                            match = re.search(r"filename\\*?=(?:UTF-8''|[\"']?)([^\"';]+)", disposition, re.I)
                            filename = urllib.parse.unquote(match.group(1)) if match else Path(parsed_url.path).name
                            add(content, filename or "法院文书", content_type, response.url)
                        except Exception as exc:
                            diagnostics.append(f"response-error {response.url[:200]} {exc}")

                    async def capture_download(download: Any) -> None:
                        try:
                            path = await download.path()
                            if path:
                                add(Path(path).read_bytes(), download.suggested_filename or "法院文书", "", download.url)
                        except Exception as exc:
                            diagnostics.append(f"download-error {exc}")

                    def on_response(response: Any) -> None:
                        pending.append(asyncio.create_task(capture_response(response)))

                    def on_download(download: Any) -> None:
                        pending.append(asyncio.create_task(capture_download(download)))

                    def attach_download_listener(target_page: Any) -> None:
                        target_page.on("download", on_download)

                    context.on("response", on_response)
                    context.on("page", attach_download_listener)
                    page = await context.new_page()
                    attach_download_listener(page)
                    try:
                        await page.goto(session["source_url"], wait_until="domcontentloaded", timeout=45_000)
                    except Exception as exc:
                        # Delivery pages commonly keep polling and may never
                        # reach a quiet lifecycle state. The DOM is still usable.
                        diagnostics.append(f"goto-warning {type(exc).__name__}: {exc}")
                    await page.wait_for_timeout(1_500)
                    diagnostics.append(f"page {page.url} title={await page.title()} frames={len(page.frames)}")
                    try:
                        diagnostics.append("body " + re.sub(r"\s+", " ", (await page.locator("body").inner_text())[:2000]))
                        elements = await page.locator("a,button,input,[onclick],[role=button]").evaluate_all(
                            "els => els.slice(0,200).map(e => ({tag:e.tagName,text:(e.innerText||e.value||e.title||e.alt||'').trim(),href:e.href||'',onclick:e.getAttribute('onclick')||''}))"
                        )
                        diagnostics.append("elements " + _json(elements)[:12000])
                    except Exception as exc:
                        diagnostics.append(f"dom-error {exc}")
                    # The authenticated site uses JavaScript handlers rather than
                    # ordinary hrefs. Click only document-oriented controls.
                    selectors = (
                        "a:has-text('下载')", "button:has-text('下载')", "input[value*='下载']",
                        "a:has-text('查看文书')", "button:has-text('查看文书')",
                        "a:has-text('预览')", "button:has-text('预览')",
                        "a:has-text('查看材料')", "button:has-text('查看材料')",
                        "[title*='下载']", "[aria-label*='下载']", "[onclick*='download' i]", "[onclick*='wj' i]",
                    )
                    for frame in page.frames:
                        for selector in selectors:
                            try:
                                locator = frame.locator(selector)
                                count = min(await locator.count(), 30)
                            except Exception as exc:
                                diagnostics.append(f"selector-error {selector} {exc}")
                                continue
                            for index in range(count):
                                try:
                                    diagnostics.append(f"click {selector} #{index}")
                                    await locator.nth(index).click(timeout=4_000)
                                    await page.wait_for_timeout(700)
                                except Exception as exc:
                                    diagnostics.append(f"click-error {selector} #{index} {exc}")
                                    continue
                    await page.wait_for_timeout(2_000)
                    if pending:
                        await asyncio.gather(*pending, return_exceptions=True)
                finally:
                    await browser.close()
        except Exception as exc:
            diagnostics.append(f"collector-error {type(exc).__name__}: {exc}")
            diagnostics.append(traceback.format_exc()[-8000:])
        finally:
            session["download_diagnostics"] = diagnostics[-300:]
        return collected

    def document_rows(message_id: int) -> List[Dict[str, Any]]:
        with get_db() as conn:
            rows = conn.execute(
                """SELECT d.*,COALESCE(f.relative_path,'') AS relative_path,
                          COALESCE(f.absolute_path,'') AS absolute_path
                   FROM case_manager_message_documents d
                   LEFT JOIN case_manager_files f ON f.id=d.file_id
                   WHERE d.message_id=? ORDER BY d.id""",
                (message_id,),
            ).fetchall()
        items = [dict(row) for row in rows]
        for item in items:
            error = item.get("last_error") or ""
            if (
                "BrowserType." in error
                or "Target page, context or browser has been closed" in error
                or "Browser logs:" in error
            ):
                item["last_error"] = _friendly_browser_error(RuntimeError(error))
        return items

    def discover_documents(
        message_id: int,
        body: str,
        attachments: Sequence[Dict[str, str]] | None = None,
    ) -> List[Dict[str, Any]]:
        stamp = _now()
        with get_db() as conn:
            if not conn.execute(
                "SELECT 1 FROM case_manager_judicial_messages WHERE id=?", (message_id,)
            ).fetchone():
                raise HTTPException(status_code=404, detail="短信不存在")
            urls = extract_document_urls(body)
            gzac_params = _gzac_delivery_params(body)
            if gzac_params:
                # The GZAC search page itself is not a document.  Persist one
                # dedicated task containing the SMS lookup credentials, so the
                # downloader can verify the case and enumerate every attachment
                # through the official public interface.
                conn.execute(
                    """INSERT OR IGNORE INTO case_manager_message_documents
                       (message_id,source_kind,source_url,source_path,original_filename,status,created_at,updated_at)
                       VALUES (?,'gzac_delivery',?,?,?,'discovered',?,?)""",
                    (
                        message_id,
                        gzac_params["source_url"],
                        _json({
                            "case_year": gzac_params["case_year"],
                            "case_no": gzac_params["case_no"],
                            "query_code": gzac_params["query_code"],
                        }),
                        "广州仲裁委电子送达",
                        stamp,
                        stamp,
                    ),
                )
            for url in urls:
                if _is_court_platform_homepage(url) or _is_gzac_search_url(url):
                    continue
                name = Path(urllib.parse.urlparse(url).path).name or "法院文书"
                conn.execute(
                    """INSERT OR IGNORE INTO case_manager_message_documents
                       (message_id,source_kind,source_url,original_filename,status,created_at,updated_at)
                       VALUES (?,'url',?,?,'discovered',?,?)""",
                    (message_id, url, name[:240], stamp, stamp),
                )
            # Older versions treated the optional “也可以登录平台” homepage as
            # another document. Remove only unfinished task rows; real archived
            # files and manually imported content are never touched.
            stale_rows = conn.execute(
                """SELECT id,source_url FROM case_manager_message_documents
                   WHERE message_id=? AND source_kind='url'
                     AND status IN ('discovered','needs_login','failed','cancelled')
                     AND COALESCE(staging_path,'')='' AND file_id IS NULL""",
                (message_id,),
            ).fetchall()
            for stale in stale_rows:
                if (
                    _is_court_platform_homepage(stale["source_url"] or "")
                    or _is_gzac_search_url(stale["source_url"] or "")
                ):
                    conn.execute(
                        "DELETE FROM case_manager_message_documents WHERE id=?",
                        (int(stale["id"]),),
                    )
            for item in attachments or []:
                source_path = _norm(item.get("path"))
                if not source_path:
                    continue
                conn.execute(
                    """INSERT OR IGNORE INTO case_manager_message_documents
                       (message_id,source_kind,source_path,original_filename,mime_type,status,created_at,updated_at)
                       VALUES (?,'attachment',?,?,?,'discovered',?,?)""",
                    (
                        message_id,
                        source_path,
                        (_norm(item.get("name")) or Path(source_path).name)[:240],
                        _norm(item.get("mime_type"))[:120],
                        stamp,
                        stamp,
                    ),
                )
            conn.commit()
        return document_rows(message_id)

    def case_folder_available(case_id: int) -> bool:
        """Whether a matched case currently has a usable local folder binding."""
        if not case_id:
            return False
        getter = callbacks.get("get_folder_binding")
        if getter is None:
            return False
        try:
            binding = getter(int(case_id), required=False)
        except Exception:
            return False
        return bool(binding.get("available"))

    def stage_document(document_id: int, content: bytes, filename: str) -> str:
        suffix = Path(filename).suffix.lower() or ".bin"
        final_path = JUDICIAL_STAGING_DIR / f"{document_id}-{hashlib.sha256(content).hexdigest()[:16]}{suffix}"
        fd, temporary = tempfile.mkstemp(prefix=".judicial-", suffix=".tmp", dir=str(JUDICIAL_STAGING_DIR))
        try:
            with os.fdopen(fd, "wb") as handle:
                handle.write(content)
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temporary, final_path)
        except Exception:
            try:
                os.unlink(temporary)
            except OSError:
                pass
            raise
        return str(final_path)

    def archive_document(
        document_id: int,
        case_id: int,
        content: bytes | None = None,
        filename: str = "",
        mime_type: str = "",
    ) -> Dict[str, Any]:
        with get_db() as conn:
            row = conn.execute(
                """SELECT d.*,m.received_at,m.institution,m.category,m.case_numbers_json
                   FROM case_manager_message_documents d
                   JOIN case_manager_judicial_messages m ON m.id=d.message_id
                   WHERE d.id=?""",
                (document_id,),
            ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="文书任务不存在")
        document = dict(row)
        if content is None:
            staging = Path(document.get("staging_path") or "")
            if not staging.is_file():
                raise HTTPException(status_code=409, detail="文书尚未下载完成")
            content = staging.read_bytes()
        digest = hashlib.sha256(content).hexdigest()
        with get_db() as conn:
            duplicate = conn.execute(
                """SELECT * FROM case_manager_files
                   WHERE case_id=? AND category='court_document' AND sha256=?
                   ORDER BY id LIMIT 1""",
                (case_id, digest),
            ).fetchone()
        if duplicate:
            file_item = dict(duplicate)
        else:
            valid, detected = _valid_document(
                content,
                filename or document.get("original_filename") or "",
                mime_type or document.get("mime_type") or "",
            )
            if not valid:
                raise HTTPException(status_code=400, detail="下载内容不是可识别的法院文书")
            output_name = _document_filename(
                document,
                filename or document.get("original_filename") or "法院文书",
                detected,
                _infer_document_type(
                    content,
                    detected,
                    filename or document.get("original_filename") or "",
                ),
            )
            file_item = callbacks["store_case_file"](
                case_id,
                "court_document",
                output_name,
                content,
                mime_type=mime_type or document.get("mime_type") or "",
                source_type="judicial_message",
                source_id=int(document["message_id"]),
                original_filename=filename or document.get("original_filename") or "",
                metadata={
                    "document_id": document_id,
                    "source_url": document.get("source_url") or "",
                    "source_kind": document.get("source_kind") or "",
                },
            )
        stamp = _now()
        with get_db() as conn:
            conn.execute(
                """UPDATE case_manager_message_documents
                   SET status='archived',case_id=?,file_id=?,final_filename=?,sha256=?,
                       staging_path='',last_error='',updated_at=? WHERE id=?""",
                (
                    case_id,
                    int(file_item["id"]),
                    Path(file_item["absolute_path"]).name,
                    digest,
                    stamp,
                    document_id,
                ),
            )
            callbacks["add_audit_log"](
                conn,
                case_id,
                "judicial_document",
                document_id,
                "created",
                f"归档法院文书：{Path(file_item['absolute_path']).name}",
                {},
                {"路径": file_item["relative_path"]},
                {"路径": "Stella电子送达归档"},
            )
            conn.commit()
        staging_path = document.get("staging_path") or ""
        if staging_path:
            try:
                Path(staging_path).unlink(missing_ok=True)
            except OSError:
                pass
        return dict(next(item for item in document_rows(int(document["message_id"])) if item["id"] == document_id))

    async def fetch_url_document(url: str) -> tuple[bytes, str, str]:
        current = url
        async with httpx.AsyncClient(
            timeout=httpx.Timeout(45.0, connect=15.0),
            follow_redirects=False,
            # Some court delivery gateways return a false “data does not
            # exist” page to non-browser user agents. Use a normal browser UA
            # while retaining the same SSRF and size checks below.
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/124.0 Safari/537.36"
                ),
                "Accept": "text/html,application/xhtml+xml,application/pdf,application/octet-stream;q=0.9,*/*;q=0.8",
            },
        ) as client:
            for _ in range(6):
                _validate_public_url(current)
                async with client.stream("GET", current) as response:
                    if response.status_code in {301, 302, 303, 307, 308}:
                        location = response.headers.get("location")
                        if not location:
                            raise HTTPException(status_code=502, detail="下载地址重定向不完整")
                        current = urllib.parse.urljoin(current, location)
                        continue
                    if response.status_code >= 400:
                        raise HTTPException(status_code=502, detail=f"文书下载失败（HTTP {response.status_code}）")
                    declared = int(response.headers.get("content-length") or 0)
                    if response.headers.get("content-encoding"):
                        declared = 0  # aiter_bytes reports decoded, not wire bytes.
                    if declared > MAX_DOCUMENT_BYTES:
                        raise HTTPException(status_code=413, detail=f"文书超过 {MAX_DOCUMENT_MB}MB 限制")
                    report_download(total_bytes=declared, received_bytes=0)
                    chunks: List[bytes] = []
                    total = 0
                    async for chunk in response.aiter_bytes():
                        total += len(chunk)
                        if total > MAX_DOCUMENT_BYTES:
                            raise HTTPException(status_code=413, detail=f"文书超过 {MAX_DOCUMENT_MB}MB 限制")
                        chunks.append(chunk)
                        report_download(received_bytes=total)
                    disposition = response.headers.get("content-disposition") or ""
                    match = re.search(r"filename\\*?=(?:UTF-8''|[\"']?)([^\"';]+)", disposition, re.I)
                    name = urllib.parse.unquote(match.group(1)) if match else Path(urllib.parse.urlparse(current).path).name
                    return b"".join(chunks), name or "法院文书", response.headers.get("content-type") or ""
            raise HTTPException(status_code=502, detail="下载地址重定向次数过多")

    async def court_delivery_manifest(url: str) -> List[Dict[str, str]]:
        """Resolve every file in a public court SMS delivery link without login."""
        params = _court_delivery_params(url)
        if not params:
            return []
        parsed = urllib.parse.urlparse(url)
        endpoint = urllib.parse.urlunparse((
            parsed.scheme,
            parsed.netloc,
            "/yzw/yzw-zxfw-sdfw/api/v1/sdfw/getWsListBySdbhNew",
            "",
            "",
            "",
        ))
        _validate_public_url(endpoint)
        try:
            async with httpx.AsyncClient(
                timeout=httpx.Timeout(45.0, connect=15.0),
                follow_redirects=False,
                headers={
                    "User-Agent": "Mozilla/5.0 LegalKnowledgeBase/1.0",
                    "Origin": f"{parsed.scheme}://{parsed.netloc}",
                    "Referer": f"{parsed.scheme}://{parsed.netloc}/zxfw/",
                },
            ) as client:
                response = await client.post(endpoint, json=params)
        except httpx.HTTPError as exc:
            raise HTTPException(status_code=502, detail="法院送达页文件清单读取失败") from exc
        if response.status_code != 200:
            raise HTTPException(
                status_code=502,
                detail=f"法院送达页文件清单读取失败（HTTP {response.status_code}）",
            )
        try:
            payload = response.json()
        except ValueError as exc:
            raise HTTPException(status_code=502, detail="法院送达页未返回有效文件清单") from exc
        if int(payload.get("code") or 0) != 200 or not isinstance(payload.get("data"), list):
            raise HTTPException(
                status_code=502,
                detail=_norm(payload.get("msg") or payload.get("message")) or "法院送达页未返回文件清单",
            )
        items: List[Dict[str, str]] = []
        seen: set[str] = set()
        for index, raw in enumerate(payload["data"][:100]):
            if not isinstance(raw, dict):
                continue
            file_url = _norm(raw.get("wjlj"))
            if not file_url:
                continue
            _validate_public_url(file_url)
            extension = _norm(raw.get("c_wjgs")).lower().lstrip(".")
            stem = re.sub(r'[<>:"/\\|?*]', "_", _norm(raw.get("c_wsmc"))).strip(" ._")
            url_name = urllib.parse.unquote(Path(urllib.parse.urlparse(file_url).path).name)
            filename = f"{stem}.{extension}" if stem and extension else url_name or f"法院文书-{index + 1}"
            key = _norm(raw.get("c_wsbh")) or hashlib.sha256(
                urllib.parse.urlunsplit((*urllib.parse.urlsplit(file_url)[:3], "", "")).encode()
            ).hexdigest()
            if key in seen:
                continue
            seen.add(key)
            items.append({"key": key, "url": file_url, "filename": filename})
        if not items:
            raise HTTPException(status_code=404, detail="法院送达链接中没有可下载文书")
        return items

    async def gzac_delivery_manifest(
        document: Dict[str, Any], client: httpx.AsyncClient
    ) -> List[Dict[str, str]]:
        """Verify one GZAC SMS lookup and return its official attachment list.

        The first request establishes the website session.  The second request
        must use that same session; it returns the temporary file codes which
        the public delivery page normally renders after a user presses 查询.
        """
        params = _gzac_document_params(document)
        if not params:
            raise HTTPException(status_code=400, detail="广州仲裁委短信缺少可用的案件年份、编号或查询码")
        lookup_url = f"{GZAC_API_ORIGIN}/mod-openinterface/caseInfo/getSecretaryInfo"
        materials_url = f"{GZAC_API_ORIGIN}/mod-case/mod/delivery/getSmsAndEmailUrlData"
        _validate_public_url(lookup_url)
        _validate_public_url(materials_url)
        try:
            lookup_response = await client.post(
                lookup_url,
                json={
                    "collectNo": "",
                    "caseYear": params["case_year"],
                    "caseNo": params["case_no"],
                    "code": params["query_code"],
                },
            )
        except httpx.HTTPError as exc:
            raise HTTPException(status_code=502, detail="广州仲裁委查询服务暂时无法访问") from exc
        if lookup_response.status_code != 200:
            raise HTTPException(
                status_code=502,
                detail=f"广州仲裁委查询验证失败（HTTP {lookup_response.status_code}）",
            )
        try:
            lookup_payload = lookup_response.json()
        except ValueError as exc:
            raise HTTPException(status_code=502, detail="广州仲裁委查询服务未返回有效数据") from exc
        lookup_data = lookup_payload.get("data") if isinstance(lookup_payload, dict) else None
        delivery_cid = _norm(lookup_data.get("deliveryCid")) if isinstance(lookup_data, dict) else ""
        if not delivery_cid:
            detail = _norm(lookup_payload.get("message")) if isinstance(lookup_payload, dict) else ""
            raise HTTPException(status_code=404, detail=detail or "未查询到可下载的广州仲裁委送达材料")

        # The portal's own page posts the SMS link's ``s`` value as ``cid``
        # (and ``mc`` as ``refId``); without that matching call the materials
        # endpoint refuses to return the attachment list.  Fall back to the
        # verified lookup id for SMS variants that lack the ``s`` parameter.
        search_id = _gzac_sms_search_id(_norm(document.get("source_url")))
        cid_candidates: List[str] = []
        for candidate in (search_id, delivery_cid):
            if candidate and candidate not in cid_candidates:
                cid_candidates.append(candidate)
        ref_id = _gzac_sms_ref_id(_norm(document.get("source_url")))
        cid_for_urls = ""
        materials_error = ""
        for cid in cid_candidates:
            request_data: Dict[str, str] = {"cid": cid}
            if ref_id:
                request_data["refId"] = ref_id
            try:
                materials_response = await client.post(materials_url, json=request_data)
            except httpx.HTTPError as exc:
                raise HTTPException(status_code=502, detail="广州仲裁委送达材料清单读取失败") from exc
            if materials_response.status_code != 200:
                materials_error = f"广州仲裁委送达材料清单读取失败（HTTP {materials_response.status_code}）"
                continue
            try:
                materials_payload = materials_response.json()
            except ValueError as exc:
                raise HTTPException(status_code=502, detail="广州仲裁委送达材料页未返回有效清单") from exc
            data = materials_payload.get("data") if isinstance(materials_payload, dict) else None
            result_file = data.get("resultFile") if isinstance(data, dict) else None
            raw_items = result_file.get("materialContent") if isinstance(result_file, dict) else None
            if isinstance(raw_items, str):
                raw_items = _loads(raw_items, [])
            if isinstance(raw_items, list) and raw_items:
                cid_for_urls = cid
                break
            materials_error = _norm(materials_payload.get("message")) or "广州仲裁委查询结果中没有可下载材料"
        if not cid_for_urls:
            raise HTTPException(status_code=404, detail=materials_error)
        delivery_id = cid_for_urls

        result: List[Dict[str, str]] = []
        seen: set[str] = set()
        for index, raw in enumerate(raw_items[:100]):
            if not isinstance(raw, dict):
                continue
            # The portal passes the entry's file code (or, on current data
            # exports, its file name) as ``fileCode`` when downloading.
            file_code = _norm(
                raw.get("fileCode") or raw.get("storageCode") or _norm(raw.get("name"))
            )[:500]
            if not file_code or file_code in seen:
                continue
            seen.add(file_code)
            name = re.sub(r'[<>:"/\\|?*]', "_", _norm(raw.get("name"))).strip(" ._")
            filename = name or f"广州仲裁委文书-{index + 1}"
            download_query = {"fileCode": file_code, "deliveryId": delivery_id}
            preflight_query = dict(download_query)
            if ref_id:
                download_query["deliveryMainCid"] = ref_id
                preflight_query["deliveryMainCid"] = ref_id
            result.append({
                "key": file_code,
                "filename": filename,
                "url": f"{GZAC_API_ORIGIN}/mod-case/mod/manageFile/downloadDeliveryAttr?{urllib.parse.urlencode(download_query)}",
                "preflight_url": f"{GZAC_API_ORIGIN}/mod-case/mod/manageFile/downloadDeliveryAttrBefore?{urllib.parse.urlencode(preflight_query)}",
            })
        if not result:
            raise HTTPException(status_code=404, detail="广州仲裁委查询结果中没有可下载材料")
        return result

    async def fetch_gzac_document(
        client: httpx.AsyncClient, preflight_url: str, download_url: str
    ) -> tuple[bytes, str, str]:
        """Pass the portal's public download-limit check, then stream one file."""
        _validate_public_url(preflight_url)
        try:
            preflight = await client.get(preflight_url)
        except httpx.HTTPError as exc:
            raise HTTPException(status_code=502, detail="广州仲裁委文书下载校验失败") from exc
        if preflight.status_code >= 400:
            raise HTTPException(status_code=502, detail=f"广州仲裁委文书下载校验失败（HTTP {preflight.status_code}）")
        try:
            allowed = preflight.json().get("data")
        except (ValueError, AttributeError) as exc:
            raise HTTPException(status_code=502, detail="广州仲裁委文书下载校验未返回有效结果") from exc
        if not allowed:
            raise HTTPException(status_code=429, detail="广州仲裁委提示该文书下载次数过多，请稍后重试")

        current = download_url
        for _ in range(6):
            _validate_public_url(current)
            try:
                async with client.stream("GET", current) as response:
                    if response.status_code in {301, 302, 303, 307, 308}:
                        location = response.headers.get("location")
                        if not location:
                            raise HTTPException(status_code=502, detail="广州仲裁委文书下载重定向不完整")
                        current = urllib.parse.urljoin(current, location)
                        continue
                    if response.status_code >= 400:
                        raise HTTPException(status_code=502, detail=f"广州仲裁委文书下载失败（HTTP {response.status_code}）")
                    declared = int(response.headers.get("content-length") or 0)
                    if response.headers.get("content-encoding"):
                        declared = 0  # aiter_bytes reports decoded, not wire bytes.
                    if declared > MAX_DOCUMENT_BYTES:
                        raise HTTPException(status_code=413, detail=f"文书超过 {MAX_DOCUMENT_MB}MB 限制")
                    report_download(total_bytes=declared, received_bytes=0)
                    chunks: List[bytes] = []
                    total = 0
                    async for chunk in response.aiter_bytes():
                        total += len(chunk)
                        if total > MAX_DOCUMENT_BYTES:
                            raise HTTPException(status_code=413, detail=f"文书超过 {MAX_DOCUMENT_MB}MB 限制")
                        chunks.append(chunk)
                        report_download(received_bytes=total)
                    disposition = response.headers.get("content-disposition") or ""
                    match = re.search(r"filename\\*?=(?:UTF-8''|[\"']?)([^\"';]+)", disposition, re.I)
                    filename = urllib.parse.unquote(match.group(1)) if match else Path(urllib.parse.urlparse(current).path).name
                    return b"".join(chunks), filename or "广州仲裁委文书", response.headers.get("content-type") or ""
            except httpx.HTTPError as exc:
                raise HTTPException(status_code=502, detail="广州仲裁委文书下载失败") from exc
        raise HTTPException(status_code=502, detail="广州仲裁委文书下载重定向次数过多")

    async def process_court_delivery(
        document: Dict[str, Any],
        case_id: int = 0,
        automatic: bool = False,
    ) -> Dict[str, Any]:
        """Download, validate, and stage/archive every item in one delivery link."""
        message_id = int(document["message_id"])
        manifest = await court_delivery_manifest(document["source_url"])
        target_key = _norm(document.get("source_path")) if document.get("source_kind") == "delivery_item" else ""
        if target_key:
            manifest = [item for item in manifest if item["key"] == target_key]
            if not manifest:
                raise HTTPException(status_code=404, detail="法院送达链接中已找不到该文书")

        assignments: List[tuple[int, Dict[str, str]]] = []
        parent_status = _norm(document.get("status"))
        parent_name = Path(
            _norm(document.get("final_filename") or document.get("original_filename"))
        ).stem
        parent_match = next(
            (
                index for index, item in enumerate(manifest)
                if parent_name and Path(item["filename"]).stem == parent_name
            ),
            -1,
        )
        stamp = _now()
        with get_db() as conn:
            for index, item in enumerate(manifest):
                use_parent = (
                    document.get("source_kind") == "delivery_item"
                    or (
                        document.get("source_kind") == "url"
                        and (
                            (parent_status != "archived" and index == 0)
                            or (parent_status == "archived" and index == parent_match)
                        )
                    )
                )
                if use_parent:
                    item_id = int(document["id"])
                else:
                    conn.execute(
                        """INSERT OR IGNORE INTO case_manager_message_documents
                           (message_id,source_kind,source_url,source_path,original_filename,status,created_at,updated_at)
                           VALUES (?,'delivery_item',?,?,?,'discovered',?,?)""",
                        (
                            message_id,
                            document["source_url"],
                            item["key"],
                            item["filename"][:240],
                            stamp,
                            stamp,
                        ),
                    )
                    child = conn.execute(
                        """SELECT id FROM case_manager_message_documents
                           WHERE message_id=? AND source_kind='delivery_item'
                             AND source_url=? AND source_path=?""",
                        (message_id, document["source_url"], item["key"]),
                    ).fetchone()
                    if not child:
                        continue
                    item_id = int(child["id"])
                assignments.append((item_id, item))
            conn.commit()

        report_download(total_files=len(assignments))
        successes = 0
        failures: List[str] = []
        for item_id, item in assignments:
            begin_download_file(item["filename"])
            with get_db() as conn:
                current = conn.execute(
                    "SELECT status FROM case_manager_message_documents WHERE id=?",
                    (item_id,),
                ).fetchone()
                if current and current["status"] == "archived":
                    successes += 1
                    complete_download_file()
                    continue
                conn.execute(
                    """UPDATE case_manager_message_documents
                       SET status='downloading',attempts=attempts+1,original_filename=?,
                           last_error='',updated_at=? WHERE id=?""",
                    (item["filename"][:240], _now(), item_id),
                )
                conn.commit()
            try:
                content, fetched_name, content_type = await fetch_url_document(item["url"])
                report_download(phase="saving", received_bytes=len(content))
                filename = item["filename"] or fetched_name
                valid, detected = _valid_document(content, filename, content_type)
                if not valid:
                    raise HTTPException(status_code=400, detail="下载内容不是可识别的法院文书")
                if Path(filename).suffix.lower() != detected:
                    filename = f"{Path(filename).stem}{detected}"
                item_target = int(case_id or document.get("matched_case_id") or 0)
                message_numbers = {
                    normalize_case_no(value)
                    for value in _loads(document.get("case_numbers_json"), [])
                    if value
                }
                content_numbers = {
                    normalize_case_no(value)
                    for value in _document_case_numbers(content, detected)
                    if value
                }
                has_conflict = bool(
                    message_numbers and content_numbers and message_numbers.isdisjoint(content_numbers)
                )
                if automatic and has_conflict:
                    item_target = 0
                stage_note = ""
                if item_target and not case_folder_available(item_target):
                    stage_note = (
                        f"关联案件（#{item_target}）尚未绑定可用的本地文件夹，"
                        "文书已保存到应用暂存区；绑定后点击“确认案件并归档”完成归档"
                    )
                    item_target = 0
                if item_target:
                    archive_document(item_id, item_target, content, filename, content_type)
                else:
                    staging_path = stage_document(item_id, content, filename)
                    with get_db() as conn:
                        conn.execute(
                            """UPDATE case_manager_message_documents
                               SET status='waiting_confirmation',staging_path=?,original_filename=?,
                                   mime_type=?,sha256=?,last_error=?,updated_at=? WHERE id=?""",
                            (
                                staging_path,
                                filename[:240],
                                content_type[:120],
                                hashlib.sha256(content).hexdigest(),
                                "；".join(
                                    note
                                    for note in (
                                        "文书案号与短信不一致，请人工确认" if has_conflict else "",
                                        stage_note,
                                    )
                                    if note
                                ),
                                _now(),
                                item_id,
                            ),
                        )
                        conn.commit()
                successes += 1
                complete_download_file()
            except Exception as exc:
                complete_download_file(failed=True)
                detail = exc.detail if isinstance(exc, HTTPException) else str(exc)
                failures.append(f"{item['filename']}：{detail}")
                report_download(error=failures[0][:1000])
                with get_db() as conn:
                    conn.execute(
                        """UPDATE case_manager_message_documents
                           SET status='failed',last_error=?,updated_at=? WHERE id=?""",
                        (str(detail)[:1000], _now(), item_id),
                    )
                    conn.commit()
        if not successes:
            raise HTTPException(
                status_code=502,
                detail=failures[0] if failures else "法院送达链接中的文书下载失败",
            )
        return next(
            item for item in document_rows(message_id)
            if int(item["id"]) == int(document["id"])
        )

    async def gzac_browser_materials(
        document: Dict[str, Any], only_name: str = ""
    ) -> List[Dict[str, Any]]:
        """Load the GZAC delivery page in a browser and collect the attachments.

        The official page encrypts its lookup body (SM4 ``{"v": ...}``) and
        encrypts every query parameter, so a plain httpx call can neither
        request the material list nor download a file.  Driving the real page
        replicates the site's own flow: the SPA requests the list itself and
        each attachment download is triggered with the page's own handler.
        """
        from playwright.async_api import TimeoutError as PlaywrightTimeoutError
        from playwright.async_api import async_playwright

        params = _gzac_document_params(document)
        if not params:
            raise HTTPException(status_code=400, detail="广州仲裁委短信缺少可用的案件年份、编号或查询码")
        source_url = _norm(document.get("source_url"))
        search_id = _gzac_sms_search_id(source_url)
        ref_id = _gzac_sms_ref_id(source_url)
        if not search_id:
            lookup_url = f"{GZAC_API_ORIGIN}/mod-openinterface/caseInfo/getSecretaryInfo"
            _validate_public_url(lookup_url)
            lookup_json = {
                "collectNo": "",
                "caseYear": params["case_year"],
                "caseNo": params["case_no"],
                "code": params["query_code"],
            }
            try:
                async with httpx.AsyncClient(
                    timeout=httpx.Timeout(45.0, connect=15.0)
                ) as client:
                    response = await client.post(lookup_url, json=lookup_json)
            except httpx.HTTPError as exc:
                raise HTTPException(
                    status_code=502, detail="广州仲裁委查询服务暂时无法访问"
                ) from exc
            payload = response.json() if response.status_code == 200 else {}
            data = payload.get("data") if isinstance(payload, dict) else None
            search_id = _norm(data.get("deliveryCid")) if isinstance(data, dict) else ""
            if not search_id:
                detail = _norm(payload.get("message")) if isinstance(payload, dict) else ""
                raise HTTPException(
                    status_code=404,
                    detail=detail or "未查询到可下载的广州仲裁委送达材料",
                )
        params_pairs = [("deliveryId", search_id)]
        if ref_id:
            params_pairs += [("src", "sms"), ("mc", ref_id)]
        page_url = (
            f"{GZAC_API_ORIGIN}/secretary/showEmailAndSms?"
            + urllib.parse.urlencode(params_pairs)
        )
        holder: Dict[str, Any] = {}

        async with async_playwright() as playwright:
            launch_options: Dict[str, Any] = {
                "args": ["--no-first-run", "--no-default-browser-check"],
            }
            executable = _browser_executable()
            if executable:
                launch_options["executable_path"] = executable
            try:
                browser = await playwright.chromium.launch(**launch_options)
            except Exception:
                launch_options.pop("executable_path", None)
                browser = await playwright.chromium.launch(**launch_options)
            try:
                context = await browser.new_context(accept_downloads=True)
                page = await context.new_page()

                async def on_response(response):
                    url = response.url or ""
                    if "getSmsAndEmailUrlData" in url and "data" not in holder:
                        try:
                            body_json = await response.json()
                        except Exception:
                            return
                        holder["data"] = (
                            body_json.get("data") if isinstance(body_json, dict) else None
                        )

                page.on("response", on_response)
                try:
                    await page.goto(page_url, wait_until="domcontentloaded", timeout=60000)
                except PlaywrightTimeoutError:
                    pass
                await page.wait_for_timeout(4000)
                payload = holder.get("data")
                if not isinstance(payload, dict):
                    raise HTTPException(
                        status_code=404,
                        detail="未查询到可下载的广州仲裁委送达材料（页面会话未返回清单）",
                    )
                result_file = payload.get("resultFile") or {}
                raw_items = result_file.get("materialContent")
                if isinstance(raw_items, str):
                    raw_items = _loads(raw_items, [])
                result: List[Dict[str, Any]] = []
                for index, raw in enumerate(raw_items[:100] if isinstance(raw_items, list) else []):
                    if not isinstance(raw, dict):
                        continue
                    original_name = _norm(raw.get("name"))
                    key = _norm(
                        raw.get("fileCode") or raw.get("storageCode") or original_name
                    )[:500]
                    if not key:
                        continue
                    display_name = re.sub(
                        r'[<>:"/\\|?*]', "_", original_name
                    ).strip(" ._") or f"广州仲裁委文书-{index + 1}"
                    if only_name and original_name != only_name:
                        continue
                    result.append(
                        {"key": key, "filename": display_name, "original_name": original_name}
                    )
                if not result:
                    raise HTTPException(
                        status_code=404, detail="广州仲裁委查询结果中没有可下载材料"
                    )
                for item in result:
                    original_name = item["original_name"]
                    target = page.get_by_text(original_name or item["filename"], exact=True)
                    try:
                        async with page.expect_download(timeout=60000) as download_info:
                            await target.wait_for(state="visible", timeout=20000)
                            await target.nth(0).click()
                        download = await download_info.value
                        item["content_type"] = (
                            "application/pdf"
                            if item["filename"].lower().endswith(".pdf")
                            else "application/octet-stream"
                        )
                    except Exception as exc:
                        raise HTTPException(
                            status_code=502,
                            detail=f"广州仲裁委文书下载失败（{item['filename']}：{exc}）",
                        ) from exc
                    try:
                        download_path = await download.path()
                        if download_path and os.path.exists(download_path):
                            with open(download_path, "rb") as fetched:
                                item["content"] = fetched.read()
                    except Exception as exc:
                        raise HTTPException(
                            status_code=502,
                            detail=f"广州仲裁委文书读取失败（{item['filename']}：{exc}）",
                        ) from exc
                    if not item.get("content"):
                        raise HTTPException(
                            status_code=502,
                            detail=f"广州仲裁委文书下载失败（{item['filename']}：内容为空）",
                        )

                return result
            finally:
                if browser.is_connected():
                    await browser.close()

    async def process_gzac_delivery(
        document: Dict[str, Any],
        case_id: int = 0,
        automatic: bool = False,
    ) -> Dict[str, Any]:
        """Download every attachment from a verified GZAC SMS delivery.

        The official page encrypts both the lookup body and the per-file query
        parameters, so a plain httpx call frequently fails; whenever it does,
        fall back to driving the real page inside a browser.
        """
        message_id = int(document["message_id"])
        target_key = (
            _gzac_delivery_item_key(document)
            if document.get("source_kind") == "gzac_delivery_item"
            else ""
        )
        httpx_headers = {
            "User-Agent": (
                "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
            ),
            "Origin": GZAC_API_ORIGIN,
            "Referer": f"{GZAC_API_ORIGIN}/search/",
            "Accept": "application/json, text/plain, */*",
        }
        try:
            async with httpx.AsyncClient(
                timeout=httpx.Timeout(45.0, connect=15.0),
                follow_redirects=False,
                headers=httpx_headers,
            ) as client:
                manifest = await gzac_delivery_manifest(document, client)
                if target_key:
                    manifest = [item for item in manifest if item["key"] == target_key]
                    if not manifest:
                        raise HTTPException(
                            status_code=404, detail="广州仲裁委送达清单中已找不到该文书"
                        )
        except HTTPException as httpx_exc:
            # The site rejects plain requests; drive the official page and
            # download with its own handler instead.
            only_name = _norm(document.get("original_filename")) if target_key else ""
            manifest = await gzac_browser_materials(document, only_name=only_name)
            if not manifest:
                raise HTTPException(
                    status_code=404,
                    detail=_norm(httpx_exc.detail) or "未查询到可下载的广州仲裁委送达材料",
                )

        assignments: List[tuple[int, Dict[str, Any]]] = []
        parent_status = _norm(document.get("status"))
        parent_name = Path(
            _norm(document.get("final_filename") or document.get("original_filename"))
        ).stem
        parent_match = next(
            (
                index
                for index, item in enumerate(manifest)
                if parent_name and Path(item["filename"]).stem == parent_name
            ),
            -1,
        )
        stamp = _now()
        with get_db() as conn:
            for index, item in enumerate(manifest):
                use_parent = (
                    document.get("source_kind") == "gzac_delivery_item"
                    or (
                        document.get("source_kind") == "gzac_delivery"
                        and (
                            (parent_status != "archived" and index == 0)
                            or (parent_status == "archived" and index == parent_match)
                        )
                    )
                )
                if use_parent:
                    item_id = int(document["id"])
                else:
                    child_source_path = _json(
                        {**_gzac_document_params(document), "key": item["key"]}
                    )
                    conn.execute(
                        """INSERT OR IGNORE INTO case_manager_message_documents
                           (message_id,source_kind,source_url,source_path,original_filename,status,created_at,updated_at)
                           VALUES (?,'gzac_delivery_item',?,?,?,'discovered',?,?)""",
                        (
                            message_id,
                            document["source_url"],
                            child_source_path,
                            item["filename"][:240],
                            stamp,
                            stamp,
                        ),
                    )
                    child = conn.execute(
                        """SELECT id FROM case_manager_message_documents
                           WHERE message_id=? AND source_kind='gzac_delivery_item'
                             AND source_url=? AND source_path=?""",
                        (message_id, document["source_url"], child_source_path),
                    ).fetchone()
                    if not child:
                        continue
                    item_id = int(child["id"])
                assignments.append((item_id, item))
            conn.commit()

        report_download(total_files=len(assignments))
        successes = 0
        failures: List[str] = []
        for item_id, item in assignments:
            begin_download_file(item["filename"])
            with get_db() as conn:
                current = conn.execute(
                    "SELECT status FROM case_manager_message_documents WHERE id=?", (item_id,)
                ).fetchone()
                if current and current["status"] == "archived":
                    successes += 1
                    complete_download_file()
                    continue
                conn.execute(
                    """UPDATE case_manager_message_documents
                       SET status='downloading',attempts=attempts+1,original_filename=?,
                           last_error='',updated_at=? WHERE id=?""",
                    (item["filename"][:240], _now(), item_id),
                )
                conn.commit()
            try:
                if "content" in item:
                    content = item["content"]
                    fetched_name = ""
                    content_type = item.get("content_type") or "application/octet-stream"
                else:
                    async with httpx.AsyncClient(
                        timeout=httpx.Timeout(45.0, connect=15.0),
                        follow_redirects=False,
                        headers=httpx_headers,
                    ) as client:
                        content, fetched_name, content_type = await fetch_gzac_document(
                            client, item["preflight_url"], item["url"]
                        )
                report_download(phase="saving", received_bytes=len(content))
                filename = item["filename"] or fetched_name
                valid, detected = _valid_document(content, filename, content_type)
                if not valid:
                    raise HTTPException(status_code=400, detail="下载内容不是可识别的仲裁文书")
                if Path(filename).suffix.lower() != detected:
                    filename = f"{Path(filename).stem}{detected}"
                item_target = int(case_id or document.get("matched_case_id") or 0)
                message_numbers = {
                    normalize_case_no(value)
                    for value in _loads(document.get("case_numbers_json"), [])
                    if value
                }
                content_numbers = {
                    normalize_case_no(value)
                    for value in _document_case_numbers(content, detected)
                    if value
                }
                has_conflict = bool(
                    message_numbers and content_numbers and message_numbers.isdisjoint(content_numbers)
                )
                if automatic and has_conflict:
                    item_target = 0
                stage_note = ""
                if item_target and not case_folder_available(item_target):
                    stage_note = (
                        f"关联案件（#{item_target}）尚未绑定可用的本地文件夹，"
                        "文书已保存到应用暂存区；绑定后点击“确认案件并归档”完成归档"
                    )
                    item_target = 0
                if item_target:
                    archive_document(item_id, item_target, content, filename, content_type)
                else:
                    staging_path = stage_document(item_id, content, filename)
                    with get_db() as conn:
                        conn.execute(
                            """UPDATE case_manager_message_documents
                               SET status='waiting_confirmation',staging_path=?,original_filename=?,
                                   mime_type=?,sha256=?,last_error=?,updated_at=? WHERE id=?""",
                            (
                                staging_path,
                                filename[:240],
                                content_type[:120],
                                hashlib.sha256(content).hexdigest(),
                                "；".join(
                                    note
                                    for note in (
                                        "文书案号与短信不一致，请人工确认" if has_conflict else "",
                                        stage_note,
                                    )
                                    if note
                                ),
                                _now(),
                                item_id,
                            ),
                        )
                        conn.commit()
                successes += 1
                complete_download_file()
            except Exception as exc:
                complete_download_file(failed=True)
                detail = exc.detail if isinstance(exc, HTTPException) else str(exc)
                failures.append(f"{item['filename']}：{detail}")
                report_download(error=failures[0][:1000])
                with get_db() as conn:
                    conn.execute(
                        """UPDATE case_manager_message_documents
                           SET status='failed',last_error=?,updated_at=? WHERE id=?""",
                        (str(detail)[:1000], _now(), item_id),
                    )
                    conn.commit()
        if not successes:
            raise HTTPException(
                status_code=502,
                detail=failures[0] if failures else "广州仲裁委送达材料下载失败",
            )
        return next(
            item for item in document_rows(message_id)
            if int(item["id"]) == int(document["id"])
        )

    async def process_document(document_id: int, case_id: int = 0, automatic: bool = False) -> Dict[str, Any]:
        with get_db() as conn:
            row = conn.execute(
                """SELECT d.*,m.matched_case_id,m.suggested_case_id,m.received_at,m.institution,
                          m.category,m.case_numbers_json
                   FROM case_manager_message_documents d
                   JOIN case_manager_judicial_messages m ON m.id=d.message_id
                   WHERE d.id=?""",
                (document_id,),
            ).fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="文书任务不存在")
        document = dict(row)
        try:
            if _court_delivery_params(document.get("source_url") or ""):
                return await process_court_delivery(document, case_id, automatic)
            if document.get("source_kind") in {"gzac_delivery", "gzac_delivery_item"}:
                return await process_gzac_delivery(document, case_id, automatic)
            report_download(total_files=1)
            begin_download_file(document.get("original_filename") or "法院文书")
            with get_db() as conn:
                conn.execute(
                    """UPDATE case_manager_message_documents
                       SET status='downloading',attempts=attempts+1,last_error='',updated_at=? WHERE id=?""",
                    (_now(), document_id),
                )
                conn.commit()
            cached_path = ""
            cached_name = ""
            cached_type = ""
            if document["source_kind"] == "url" and document.get("source_url"):
                with get_db() as conn:
                    cached = conn.execute(
                        """SELECT d.staging_path,d.original_filename,d.mime_type,
                                  COALESCE(f.absolute_path,'') AS absolute_path
                           FROM case_manager_message_documents d
                           LEFT JOIN case_manager_files f ON f.id=d.file_id
                           WHERE d.id<>? AND d.source_url=? AND d.sha256<>''
                           ORDER BY CASE WHEN d.status='archived' THEN 0 ELSE 1 END,d.id LIMIT 1""",
                        (document_id, document["source_url"]),
                    ).fetchone()
                if cached:
                    candidate = cached["absolute_path"] or cached["staging_path"] or ""
                    if candidate and Path(candidate).is_file():
                        cached_path = candidate
                        cached_name = cached["original_filename"] or Path(candidate).name
                        cached_type = cached["mime_type"] or ""
            if cached_path:
                content = Path(cached_path).read_bytes()
                filename = cached_name
                content_type = cached_type
            elif document["source_kind"] == "attachment":
                path = _safe_attachment_path(document["source_path"])
                if path.stat().st_size > MAX_DOCUMENT_BYTES:
                    raise HTTPException(status_code=413, detail=f"短信附件超过 {MAX_DOCUMENT_MB}MB 限制")
                content = path.read_bytes()
                filename = document.get("original_filename") or path.name
                content_type = document.get("mime_type") or mimetypes.guess_type(filename)[0] or ""
            else:
                content, filename, content_type = await fetch_url_document(document["source_url"])
            report_download(phase="saving", received_bytes=len(content))
            valid, detected = _valid_document(content, filename, content_type)
            if not valid and detected in {"login", "verification", "interaction"}:
                access_messages = {
                    "login": "链接需要账号登录；请点击“登录并下载”完成后续操作",
                    "verification": "送达页需要核验收件手机号和短信验证码（无需账号登录）；请点击“验证并下载”",
                    "interaction": "链接是公开网页，需要在页面内选择并下载文书；请点击“打开并下载”",
                }
                with get_db() as conn:
                    conn.execute(
                        """UPDATE case_manager_message_documents
                           SET status='needs_login',last_error=?,
                               updated_at=? WHERE id=?""",
                        (access_messages[detected], _now(), document_id),
                    )
                    conn.commit()
                return next(item for item in document_rows(int(document["message_id"])) if item["id"] == document_id)
            if not valid:
                raise HTTPException(status_code=400, detail="下载内容不是 PDF、Word 或可识别图片文书")
            safe_name = filename or f"法院文书{detected}"
            if Path(safe_name).suffix.lower() != detected:
                safe_name = f"{Path(safe_name).stem}{detected}"
            target_case = int(case_id or document.get("matched_case_id") or 0)
            message_numbers = {
                normalize_case_no(item) for item in _loads(document.get("case_numbers_json"), []) if item
            }
            content_numbers = {
                normalize_case_no(item) for item in _document_case_numbers(content, detected) if item
            }
            has_conflict = bool(message_numbers and content_numbers and message_numbers.isdisjoint(content_numbers))
            if automatic and has_conflict:
                target_case = 0
            stage_note = ""
            if target_case and not case_folder_available(target_case):
                stage_note = (
                    f"关联案件（#{target_case}）尚未绑定可用的本地文件夹，"
                    "文书已保存到应用暂存区；绑定后点击“确认案件并归档”完成归档"
                )
                target_case = 0
            if target_case:
                return archive_document(document_id, target_case, content, safe_name, content_type)
            staging_path = stage_document(document_id, content, safe_name)
            with get_db() as conn:
                conn.execute(
                    """UPDATE case_manager_message_documents
                       SET status='waiting_confirmation',staging_path=?,original_filename=?,mime_type=?,
                           sha256=?,last_error=?,updated_at=? WHERE id=?""",
                    (
                        staging_path,
                        safe_name[:240],
                        content_type[:120],
                        hashlib.sha256(content).hexdigest(),
                        "；".join(
                            note
                            for note in (
                                "文书案号与短信不一致，请人工确认" if has_conflict else "",
                                stage_note,
                            )
                            if note
                        ),
                        _now(),
                        document_id,
                    ),
                )
                conn.commit()
            return next(item for item in document_rows(int(document["message_id"])) if item["id"] == document_id)
        except HTTPException as exc:
            with get_db() as conn:
                conn.execute(
                    """UPDATE case_manager_message_documents
                       SET status='failed',last_error=?,updated_at=? WHERE id=?""",
                    (str(exc.detail)[:1000], _now(), document_id),
                )
                conn.commit()
            raise
        except Exception as exc:
            with get_db() as conn:
                conn.execute(
                    """UPDATE case_manager_message_documents
                       SET status='failed',last_error=?,updated_at=? WHERE id=?""",
                    (str(exc)[:1000], _now(), document_id),
                )
                conn.commit()
            raise HTTPException(status_code=502, detail=f"文书处理失败：{exc}") from exc

    def settings() -> Dict[str, Any]:
        with get_db() as conn:
            row = conn.execute("SELECT * FROM case_manager_message_settings WHERE id=1").fetchone()
            senders = [dict(item) for item in conn.execute("SELECT * FROM case_manager_trusted_senders ORDER BY sender").fetchall()]
        data = dict(row) if row else {}
        data.pop("hmac_secret", None)
        data["recipient_phones"] = [
            value for value in _loads(data.pop("recipient_phones_json", "[]"), [])
            if re.fullmatch(r"1\d{10}", str(value))
        ]
        data["ai_model"] = effective_purpose_model("judicial_messages")
        data["trusted_senders"] = senders
        data["categories"] = CATEGORIES
        data["available_models"] = sorted(ALLOWED_MODELS)
        data["python_executable"] = sys.executable
        data["messages_database"] = str(MESSAGE_DB)
        return data

    def fingerprint(body: str, sender: str, received_at: str, guid: str = "") -> str:
        with get_db() as conn:
            secret = conn.execute("SELECT hmac_secret FROM case_manager_message_settings WHERE id=1").fetchone()[0]
        payload = "|".join([guid, sender, received_at, body]).encode("utf-8", errors="ignore")
        return hmac.new(secret.encode("utf-8"), payload, hashlib.sha256).hexdigest()

    def trusted_sender(sender: str, institution: str) -> bool:
        with get_db() as conn:
            row = conn.execute("SELECT * FROM case_manager_trusted_senders WHERE sender=? AND verified=1", (sender,)).fetchone()
        if not row:
            return False
        expected = _norm(row["institution"])
        return not expected or expected in institution or institution in expected

    def relevance_score(body: str, sender: str) -> int:
        if is_verification_code_text(body):
            return 0
        score = 0
        if any(term in body for term in INSTITUTION_TERMS): score += 45
        if any(term in body for term in EVENT_TERMS): score += 30
        if is_bankruptcy_meeting_notice(body): score += 45
        if extract_case_numbers(body): score += 25
        with get_db() as conn:
            if conn.execute("SELECT 1 FROM case_manager_trusted_senders WHERE sender=?", (sender,)).fetchone(): score += 25
        return min(score, 100)

    def case_candidates(extracted: Dict[str, Any], body: str) -> Tuple[Optional[int], int, List[Dict[str, Any]], str]:
        numbers = [normalize_case_no(value) for value in extracted.get("case_numbers") or []]
        with get_db() as conn:
            rows = conn.execute(
                """SELECT c.id,c.title,c.case_type,c.current_stage,n.case_no,n.court
                   FROM case_manager_cases c LEFT JOIN case_manager_case_numbers n ON n.case_id=c.id
                   WHERE COALESCE(c.source_status,'active')='active'"""
            ).fetchall()
            parties = conn.execute(
                """SELECT p.case_id,p.name FROM case_manager_parties p JOIN case_manager_cases c ON c.id=p.case_id
                   WHERE COALESCE(c.source_status,'active')='active' AND COALESCE(p.name,'')<>''"""
            ).fetchall()
        exact_ids = {int(row["id"]) for row in rows if normalize_case_no(row["case_no"]) in numbers and normalize_case_no(row["case_no"])}
        by_case: Dict[int, Dict[str, Any]] = {}
        for row in rows:
            item = by_case.setdefault(int(row["id"]), {"id": int(row["id"]), "title": row["title"], "case_type": row["case_type"], "current_stage": row["current_stage"], "case_numbers": [], "courts": [], "parties": [], "score": 0, "reasons": []})
            if row["case_no"]: item["case_numbers"].append(row["case_no"])
            if row["court"]: item["courts"].append(row["court"])
        for row in parties:
            if int(row["case_id"]) in by_case:
                by_case[int(row["case_id"])]["parties"].append(row["name"])
        if len(exact_ids) == 1:
            case_id = next(iter(exact_ids)); item = by_case[case_id]; item["score"] = 100; item["reasons"] = ["案号唯一精确匹配"]
            return case_id, 100, [item], "exact_unique"
        if len(exact_ids) > 1:
            candidates = [by_case[item_id] for item_id in exact_ids]
            for item in candidates: item["score"] = 70; item["reasons"] = ["案号对应多个在办案件"]
            return None, 0, candidates, "exact_conflict"
        institution = extracted.get("institution") or ""
        for item in by_case.values():
            if item["title"] and item["title"] in body:
                item["score"] += 55; item["reasons"].append("短信包含案件名称")
            party_hits = [name for name in item["parties"] if len(name) >= 2 and name in body]
            if party_hits:
                item["score"] += min(35, 18 * len(party_hits)); item["reasons"].append("短信包含当事人")
            if institution and any(institution in court or court in institution for court in item["courts"] if court):
                item["score"] += 15; item["reasons"].append("受理机构相符")
            item["score"] = min(item["score"], 95)
        candidates = sorted((item for item in by_case.values() if item["score"]), key=lambda x: (-x["score"], x["id"]))[:5]
        if candidates and candidates[0]["score"] >= 55 and (len(candidates) == 1 or candidates[0]["score"] - candidates[1]["score"] >= 20):
            return candidates[0]["id"], candidates[0]["score"], candidates, "fuzzy_unique"
        return None, 0, candidates, "unmatched"

    async def ai_analysis(body: str, extracted: Dict[str, Any], candidates: List[Dict[str, Any]], selected_model: str) -> Tuple[Dict[str, Any], str]:
        minimal = [{k: item.get(k) for k in ("id", "title", "case_numbers", "courts", "parties")} for item in candidates[:5]]
        prompt = {
            "sms": body[:5000],
            "local_extraction": extracted,
            "candidate_cases": minimal,
            "allowed_categories": CATEGORIES,
        }
        result, actual_model = await call_ai_with_fallback(
            selected_model,
            [
                {"role": "system", "content": "你是法律案件短信分类器。只能根据输入事实判断，不得补充事实。只返回JSON对象，字段为 relevance(boolean)、institution、category、case_id(integer或null)、case_confidence(0-100)、category_confidence(0-100)、event_at、location、event_type、summary、conflicts(array)、reason。case_id只能从候选案件中选择，category只能从允许类别中选择。"},
                {"role": "user", "content": _json(prompt)},
            ],
            temperature=0,
            max_tokens=1000,
        )
        content = result.get("choices", [{}])[0].get("message", {}).get("content", "")
        content = re.sub(r"^```(?:json)?\s*|\s*```$", "", content.strip(), flags=re.I)
        data = json.loads(content)
        if data.get("category") not in CATEGORIES:
            data["category"] = ""
        allowed_ids = {int(item["id"]) for item in candidates}
        try: selected_id = int(data.get("case_id") or 0)
        except (TypeError, ValueError): selected_id = 0
        if selected_id not in allowed_ids:
            data["case_id"] = None
        return data, actual_model

    def default_destinations(category: str, event_type: str = "") -> List[str]:
        if category == "开庭/庭询/听证" or (category == "调解/谈话/询问" and event_type == "询问"):
            return ["hearing", "ledger"]
        if category in {"举证/答辩/补正期限", "上诉/异议/复议期限"}:
            return ["todo"]
        return ["ledger"]

    def create_outputs(conn: Any, message: Dict[str, Any], case_id: int, category: str, destinations: Iterable[str], fields: Dict[str, Any], creation_mode: str) -> List[Dict[str, Any]]:
        case = conn.execute("SELECT * FROM case_manager_cases WHERE id=? AND COALESCE(source_status,'active')='active'", (case_id,)).fetchone()
        if not case:
            raise HTTPException(status_code=404, detail="关联案件不存在或已结案")
        stamp = _now(); links: List[Dict[str, Any]] = []
        extracted = _loads(message.get("extracted_json"), {})
        title = _norm(fields.get("title")) or _norm(case["title"])
        summary = _norm(fields.get("summary")) or extracted.get("summary") or _norm(message.get("raw_body"))[:180]
        event_at = _norm(fields.get("event_at")) or extracted.get("event_at") or ""
        location = _norm(fields.get("location")) or extracted.get("location") or ""
        event_type = _norm(fields.get("event_type")) or extracted.get("event_type") or category
        case_no = _norm(fields.get("case_no")) or next(iter(_loads(message.get("case_numbers_json"), [])), "")
        institution = _norm(fields.get("institution")) or message.get("institution") or ""
        source_key = hashlib.sha256(f"judicial-message|{message['id']}".encode()).hexdigest()
        for target_type in dict.fromkeys(destinations):
            if target_type not in {"hearing", "todo", "ledger", "memo"}:
                continue
            existing = conn.execute("SELECT target_id FROM case_manager_message_links WHERE message_id=? AND target_type=?", (message["id"], target_type)).fetchone()
            if existing:
                links.append({"target_type": target_type, "target_id": int(existing["target_id"])}); continue
            if target_type == "hearing":
                if not event_at: raise HTTPException(status_code=400, detail="创建开庭记录需要明确时间")
                cur = conn.execute(
                    """INSERT INTO case_manager_hearings(case_id,title,reminder_type,hearing_at,location,case_no,court,court_contact,party_contact,notes,created_at,updated_at,source_key)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (case_id,title,event_type or "开庭",event_at,location,case_no,institution,"","",f"来源：司法短信\n{summary}",stamp,stamp,source_key),
                ); target_id = int(cur.lastrowid)
                callbacks["sync_hearing_to_case_ledger"](conn, target_id)
            elif target_type == "todo":
                if not event_at: raise HTTPException(status_code=400, detail="创建待办需要明确提醒或截止时间")
                cur = conn.execute(
                    """INSERT INTO case_manager_todos(case_id,title,todo_title,remind_at,notes,status,created_at,updated_at,source_key)
                       VALUES (?,?,?,?,?,'pending',?,?,?)""",
                    (case_id,title,event_type or category,event_at,f"来源：司法短信\n{summary}",stamp,stamp,source_key),
                ); target_id = int(cur.lastrowid)
            elif target_type == "ledger":
                project_id = callbacks["ensure_case_research_project"](conn, case_id)
                cur = conn.execute(
                    """INSERT INTO research_case_ledgers(project_id,title,content,happened_at,source_type,source_id,created_at,updated_at)
                       VALUES (?,?,?,?,?,?,?,?)""",
                    (project_id,f"司法短信：{category}",f"机构：{institution}\n案号：{case_no}\n事项：{summary}\n\n原短信：\n{message.get('raw_body') or ''}",message.get("received_at") or stamp,"case_manager_sms",message["id"],stamp,stamp),
                ); target_id = int(cur.lastrowid)
            else:
                project_id = callbacks["ensure_case_research_project"](conn, case_id)
                markdown = f"## 事项摘要\n\n{summary}\n\n## 原短信\n\n{message.get('raw_body') or ''}\n\n- 来源机构：{institution}\n- 接收时间：{message.get('received_at') or ''}\n- 案号：{case_no}"
                cur = conn.execute(
                    """INSERT INTO research_memos(project_id,title,question,markdown,sources_json,created_at,updated_at)
                       VALUES (?,?,?,?,?,?,?)""",
                    (project_id,f"司法短信备忘：{category}",summary,markdown,_json([{"type":"judicial_message","id":message["id"]}]),stamp,stamp),
                ); target_id = int(cur.lastrowid)
            conn.execute(
                """INSERT INTO case_manager_message_links(message_id,case_id,target_type,target_id,creation_mode,target_created_at,created_at)
                   VALUES (?,?,?,?,?,?,?)""", (message["id"],case_id,target_type,target_id,creation_mode,stamp,stamp),
            )
            links.append({"target_type": target_type, "target_id": target_id})
        callbacks["add_audit_log"](conn,case_id,"judicial_message",int(message["id"]),"created",f"{('短信自动创建' if creation_mode == 'auto' else '确认司法短信')}：{category}",{}, {"短信": summary}, {"短信":"事项"})
        return links

    def strict_auto_allowed(message: Dict[str, Any], extracted: Dict[str, Any], local_case_id: Optional[int], match_mode: str, ai: Dict[str, Any]) -> bool:
        category = extracted.get("category") or ""
        event_type = extracted.get("event_type") or ""
        time_sensitive = category in {"开庭/庭询/听证", "举证/答辩/补正期限", "上诉/异议/复议期限"} or (category == "调解/谈话/询问" and event_type == "询问")
        event_at = extracted.get("event_at") or ""
        try: future = dt.datetime.fromisoformat(event_at) > dt.datetime.now() - dt.timedelta(hours=1)
        except ValueError: future = False
        return bool(
            time_sensitive and future and local_case_id and match_mode == "exact_unique"
            and trusted_sender(message.get("sender") or "", extracted.get("institution") or "")
            and ai and not (ai.get("conflicts") or [])
            and ai.get("category") == category and int(ai.get("case_id") or 0) == int(local_case_id)
            and _norm(ai.get("event_at")).replace("T", " ")[:16] == _norm(event_at).replace("T", " ")[:16]
            and int(ai.get("category_confidence") or 0) == 100 and int(ai.get("case_confidence") or 0) == 100
        )

    async def analyze_saved(message_id: int, allow_auto: bool = True) -> Dict[str, Any]:
        with get_db() as conn:
            row = conn.execute("SELECT * FROM case_manager_judicial_messages WHERE id=?", (message_id,)).fetchone()
            config = conn.execute("SELECT * FROM case_manager_message_settings WHERE id=1").fetchone()
        if not row: raise HTTPException(status_code=404, detail="短信不存在")
        config = dict(config); config["ai_model"] = effective_purpose_model("judicial_messages")
        message = dict(row); body = message.get("raw_body") or ""
        manual = str(message.get('service') or '').startswith('manual-')
        if manual:
            allow_auto = False
        if not body: return message
        discover_documents(message_id, body)
        extracted = extract_fields(body, message.get("received_at") or "", allow_inferred_year=not manual or bool(message.get('received_at')))
        local_case_id, local_confidence, candidates, match_mode = case_candidates(extracted, body)
        ai: Dict[str, Any] = {}; actual_model = ""; error = ""
        if config["ai_enabled"] and model_available(config["ai_model"]):
            try:
                ai, actual_model = await ai_analysis(body, extracted, candidates, config["ai_model"])
            except Exception as exc:
                error = str(exc)[:1000]
        elif config["ai_enabled"]:
            error = f"AI 模型未配置密钥：{config['ai_model']}"
        suggested_case_id = local_case_id
        case_confidence = local_confidence
        if not suggested_case_id and ai.get("case_id"):
            suggested_case_id = int(ai["case_id"]); case_confidence = min(int(ai.get("case_confidence") or 0), 89)
        category = extracted.get("category") or ai.get("category") or ""
        category_confidence = 95 if extracted.get("category") else min(int(ai.get("category_confidence") or 0), 89)
        if ai.get("category") == extracted.get("category") and category:
            category_confidence = min(100, max(category_confidence, int(ai.get("category_confidence") or 0)))
        if ai.get("event_at") and not extracted.get("event_at"):
            extracted["event_at"] = _norm(ai.get("event_at"))
        for field in ("location", "event_type", "summary"):
            if not extracted.get(field) and ai.get(field): extracted[field] = _norm(ai[field])
        extracted["category"] = category
        manually_confirmed = bool(message.get("matched_case_id") and message.get("decision_mode") == "manual")
        if manually_confirmed:
            suggested_case_id = int(message["matched_case_id"])
            status = message.get("workflow_status") or "confirmed"
        else:
            status = "suggested" if suggested_case_id and category else "pending"
        stamp = _now()
        with get_db() as conn:
            conn.execute(
                """UPDATE case_manager_judicial_messages SET institution=?,category=?,extracted_json=?,analysis_json=?,
                   candidate_cases_json=?,suggested_case_id=?,case_numbers_json=?,workflow_status=?,category_confidence=?,
                   case_confidence=?,ai_model=?,prompt_version=?,analysis_error=?,analyzed_at=?,updated_at=? WHERE id=?""",
                (extracted.get("institution") or "",category,_json(extracted),_json(ai),_json(candidates),suggested_case_id,
                 _json(extracted.get("case_numbers") or []),status,category_confidence,case_confidence,actual_model,PROMPT_VERSION,error,stamp,stamp,message_id),
            ); conn.commit()
            refreshed = dict(conn.execute("SELECT * FROM case_manager_judicial_messages WHERE id=?", (message_id,)).fetchone())
        if allow_auto and strict_auto_allowed(refreshed, extracted, local_case_id, match_mode, ai):
            destinations = ["hearing"] if category in {"开庭/庭询/听证", "调解/谈话/询问"} else ["todo"]
            with get_db() as conn:
                kind_row = conn.execute("SELECT case_kind FROM case_manager_cases WHERE id=?", (int(local_case_id),)).fetchone()
                if kind_row and kind_row["case_kind"] == "series_child" and "ledger" not in destinations:
                    destinations.append("ledger")
                latest = dict(conn.execute("SELECT * FROM case_manager_judicial_messages WHERE id=?", (message_id,)).fetchone())
                links = create_outputs(conn, latest, int(local_case_id), category, destinations, extracted, "auto")
                conn.execute("UPDATE case_manager_judicial_messages SET matched_case_id=?,workflow_status='auto_written',decision_mode='auto',confirmed_at=?,updated_at=? WHERE id=?", (local_case_id,stamp,stamp,message_id))
                conn.commit()
            for link in links:
                if link["target_type"] in {"hearing", "todo"}:
                    try: callbacks["auto_sync_if_enabled"](link["target_type"], link["target_id"], force=link["target_type"] == "todo")
                    except Exception: pass
            refreshed = get_message(message_id)
        if (
            allow_auto
            and local_case_id
            and match_mode == "exact_unique"
            and trusted_sender(message.get("sender") or "", extracted.get("institution") or "")
        ):
            # 案件尚未绑定本地文件夹时下载也不再失败：文书自动进入应用暂存区，
            # 绑定后可通过“确认案件并归档”完成归档。
            for document in document_rows(message_id):
                if document["status"] == "discovered":
                    try:
                        await process_document(int(document["id"]), int(local_case_id), automatic=True)
                    except Exception:
                        # The task row retains the actionable failure state.
                        pass
            refreshed = get_message(message_id)
        return refreshed

    def get_message(message_id: int) -> Dict[str, Any]:
        with get_db() as conn:
            row = conn.execute(
                """SELECT m.*,COALESCE(c.title,'') AS case_title FROM case_manager_judicial_messages m
                   LEFT JOIN case_manager_cases c ON c.id=COALESCE(m.matched_case_id,m.suggested_case_id) WHERE m.id=?""", (message_id,),
            ).fetchone()
            if not row: raise HTTPException(status_code=404, detail="短信不存在")
            links = [dict(item) for item in conn.execute("SELECT * FROM case_manager_message_links WHERE message_id=? ORDER BY id", (message_id,)).fetchall()]
        data = dict(row); data["links"] = links
        discover_documents(message_id, data.get("raw_body") or "")
        data["documents"] = document_rows(message_id)
        for key in ("extracted_json", "analysis_json", "candidate_cases_json", "case_numbers_json"):
            data[key[:-5] if key.endswith("_json") else key] = _loads(data.get(key), [] if key.endswith("cases_json") or key == "case_numbers_json" else {})
        data["status_label"] = STATUS_LABELS.get(data.get("workflow_status"), data.get("workflow_status"))
        data["default_destinations"] = default_destinations(data.get("category") or "", (data.get("extracted") or {}).get("event_type") or "")
        return data

    def repair_saved_bodies() -> List[int]:
        """Restore previously truncated pending bodies from their source rows.

        Ignored messages are intentionally excluded because ignoring a message
        is also the user's request to erase its body from the application DB.
        Fingerprints stay unchanged so the repair cannot affect deduplication.
        """
        with get_db() as conn:
            saved = [dict(row) for row in conn.execute(
                """SELECT id,source_rowid,raw_body,received_at,institution,category
                   FROM case_manager_judicial_messages
                   WHERE workflow_status IN ('pending','suggested')
                     AND source_rowid>0 AND raw_body<>''"""
            ).fetchall()]
        source_bodies = _message_bodies_by_rowid(
            MESSAGE_DB,
            [int(item["source_rowid"] or 0) for item in saved],
        )
        repaired: List[int] = []
        stamp = _now()
        with get_db() as conn:
            for item in saved:
                current = _norm(item.get("raw_body"))
                source = _norm(source_bodies.get(int(item["source_rowid"] or 0)))
                # Only extend a known prefix. This avoids replacing user-visible
                # content if Apple changes a row or the decoder selects a
                # different attributed-string field in a future schema.
                if not source or len(source) <= len(current) + 3 or current not in source:
                    continue
                extracted = extract_fields(source, item.get("received_at") or "")
                conn.execute(
                    """UPDATE case_manager_judicial_messages
                       SET raw_body=?,institution=?,category=?,extracted_json=?,case_numbers_json=?,
                           analysis_json='{}',analysis_error='',analyzed_at='',updated_at=?
                       WHERE id=? AND workflow_status IN ('pending','suggested')""",
                    (
                        source,
                        extracted.get("institution") or item.get("institution") or "",
                        extracted.get("category") or item.get("category") or "",
                        _json(extracted),
                        _json(extracted.get("case_numbers") or []),
                        stamp,
                        int(item["id"]),
                    ),
                )
                repaired.append(int(item["id"]))
            conn.commit()
        return repaired

    async def _scan_once_unlocked(manual: bool = False) -> Dict[str, Any]:
        started = _now()
        with get_db() as conn:
            config = dict(conn.execute("SELECT * FROM case_manager_message_settings WHERE id=1").fetchone())
            cur = conn.execute("INSERT INTO case_manager_message_scan_runs(started_at,status) VALUES (?,'running')", (started,)); run_id = int(cur.lastrowid)
            conn.execute("UPDATE case_manager_message_settings SET last_scan_started_at=?,last_error='',updated_at=? WHERE id=1", (started,started)); conn.commit()
        if not config["enabled"] and not manual:
            return {"ok": True, "disabled": True}
        cursor_from = int(config["last_source_rowid"] or 0)
        processed_highwater = cursor_from
        scanned = candidates_count = auto_count = repaired_count = 0; error = ""
        try:
            # The private Mac currently runs Apple Command Line Tools Python 3.8,
            # where asyncio.to_thread() is not available yet. run_in_executor()
            # provides the same non-blocking database read behavior on 3.8+.
            loop = asyncio.get_running_loop()
            repaired_ids = await loop.run_in_executor(None, repair_saved_bodies)
            repaired_count = len(repaired_ids)
            # Re-run classification and matching, but never auto-create a
            # historical hearing or todo merely because its body was repaired.
            for message_id in repaired_ids:
                await analyze_saved(message_id, False)
            with get_db() as conn:
                cutoff = (dt.datetime.now() - dt.timedelta(days=max(1, int(config["backfill_days"] or 90)))).isoformat(sep=" ")
                legacy_ignored = conn.execute(
                    """SELECT fingerprint,source_rowid FROM case_manager_message_fingerprints
                       WHERE disposition IN ('irrelevant','irrelevant-v2','irrelevant-v3') AND source_rowid>0 AND received_at>=?
                       ORDER BY source_rowid LIMIT 1000""",
                    (cutoff,),
                ).fetchall()
            replay_ids = {int(row["source_rowid"]) for row in legacy_ignored}
            rows, highwater = await loop.run_in_executor(
                None,
                _message_rows,
                MESSAGE_DB,
                int(config["last_source_rowid"] or 0),
                not bool(config["initial_scan_completed"]),
                int(config["backfill_days"] or 90),
                tuple(replay_ids),
            )
            replay_seen = {int(item["source_rowid"]) for item in rows} & replay_ids
            # Missing or expired source messages cannot be reconsidered again.
            with get_db() as conn:
                for row in legacy_ignored:
                    if int(row["source_rowid"]) not in replay_seen:
                        conn.execute(
                            "UPDATE case_manager_message_fingerprints SET disposition=? WHERE fingerprint=? AND disposition IN ('irrelevant','irrelevant-v2','irrelevant-v3')",
                            (IRRELEVANT_DISPOSITION, row["fingerprint"]),
                        )
                conn.commit()
            for item in rows:
                scanned += 1
                digest = fingerprint(item["body"], item["sender"], item["received_at"], item["source_guid"])
                with get_db() as conn:
                    known = conn.execute("SELECT disposition FROM case_manager_message_fingerprints WHERE fingerprint=?", (digest,)).fetchone()
                    if (known and not (int(item["source_rowid"]) in replay_ids and known["disposition"] in REPLAYABLE_IRRELEVANT_DISPOSITIONS)) or conn.execute("SELECT 1 FROM case_manager_judicial_messages WHERE fingerprint=? OR (source_guid<>'' AND source_guid=?)", (digest,item["source_guid"])).fetchone():
                        processed_highwater = max(processed_highwater, int(item["source_rowid"] or 0))
                        continue
                score = relevance_score(item["body"], item["sender"])
                if score < 45:
                    with get_db() as conn:
                        conn.execute("INSERT INTO case_manager_message_fingerprints(fingerprint,source_guid,source_rowid,received_at,disposition,created_at) VALUES (?,?,?,?,?,?) ON CONFLICT(fingerprint) DO UPDATE SET disposition=excluded.disposition WHERE disposition IN ('irrelevant','irrelevant-v2','irrelevant-v3')", (digest,item["source_guid"],item["source_rowid"],item["received_at"],IRRELEVANT_DISPOSITION,_now())); conn.commit()
                    processed_highwater = max(processed_highwater, int(item["source_rowid"] or 0))
                    continue
                candidates_count += 1; stamp = _now(); extracted = extract_fields(item["body"], item["received_at"])
                with get_db() as conn:
                    conn.execute("DELETE FROM case_manager_message_fingerprints WHERE fingerprint=? AND disposition IN ('irrelevant','irrelevant-v2','irrelevant-v3')", (digest,))
                    cur = conn.execute(
                        """INSERT INTO case_manager_judicial_messages(source_guid,source_rowid,fingerprint,sender,service,received_at,raw_body,institution,category,extracted_json,case_numbers_json,relevance_score,imported_scan_run_id,is_new,created_at,updated_at)
                           VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,1,?,?)""",
                        (item["source_guid"],item["source_rowid"],digest,item["sender"],item["service"],item["received_at"],item["body"],extracted["institution"],extracted["category"],_json(extracted),_json(extracted["case_numbers"]),score,run_id,stamp,stamp),
                    ); message_id = int(cur.lastrowid); conn.commit()
                discover_documents(message_id, item["body"], item.get("attachments") or [])
                analyzed = await analyze_saved(message_id, True)
                if analyzed.get("workflow_status") == "auto_written": auto_count += 1
                processed_highwater = max(processed_highwater, int(item["source_rowid"] or 0))
            with get_db() as conn:
                for row in legacy_ignored:
                    conn.execute(
                        "UPDATE case_manager_message_fingerprints SET disposition=? WHERE fingerprint=? AND disposition IN ('irrelevant','irrelevant-v2','irrelevant-v3')",
                        (IRRELEVANT_DISPOSITION, row["fingerprint"]),
                    )
                conn.commit()
            completed = _now(); mtime = ""
            try: mtime = dt.datetime.fromtimestamp(MESSAGE_DB.stat().st_mtime).replace(microsecond=0).isoformat(sep=" ")
            except OSError: pass
            with get_db() as conn:
                conn.execute("""UPDATE case_manager_message_settings SET last_source_rowid=?,initial_scan_completed=1,last_scan_completed_at=?,last_success_at=?,last_error='',last_db_mtime=?,last_message_at=COALESCE((SELECT MAX(received_at) FROM case_manager_judicial_messages),'') ,updated_at=? WHERE id=1""", (highwater,completed,completed,mtime,completed))
                conn.execute("UPDATE case_manager_message_scan_runs SET completed_at=?,status='completed',scanned_count=?,candidate_count=?,auto_written_count=? WHERE id=?", (completed,scanned,candidates_count,auto_count,run_id)); conn.commit()
            return {"ok": True, "scanned": scanned, "candidates": candidates_count, "auto_written": auto_count, "repaired_count": repaired_count, "cursor_from": cursor_from, "cursor_to": highwater}
        except Exception as exc:
            error = str(exc)[:1000]; completed = _now()
            if "authorization denied" in error.lower() or "not authorized" in error.lower(): error = f"没有读取“信息”数据库的权限。请为 {full_disk_access_target()} 授予“完全磁盘访问权限”。"
            with get_db() as conn:
                conn.execute("UPDATE case_manager_message_settings SET last_source_rowid=MAX(last_source_rowid,?),last_scan_completed_at=?,last_error=?,updated_at=? WHERE id=1", (processed_highwater,completed,error,completed))
                conn.execute("UPDATE case_manager_message_scan_runs SET completed_at=?,status='failed',error=? WHERE id=?", (completed,error,run_id)); conn.commit()
            if manual: raise HTTPException(status_code=503, detail=error)
            return {"ok": False, "error": error}

    async def scan_once(manual: bool = False) -> Dict[str, Any]:
        # Startup scheduling and a user's “立即扫描” can occur at nearly the
        # same time. Serialize them so cursor and repair work never overlap.
        if scheduler["scan_lock"] is None:
            scheduler["scan_lock"] = asyncio.Lock()
        async with scheduler["scan_lock"]:
            return await _scan_once_unlocked(manual)

    def health() -> Dict[str, Any]:
        permission = False; permission_error = ""; db_mtime = ""
        try:
            if not MESSAGE_DB.exists(): raise FileNotFoundError(str(MESSAGE_DB))
            conn = sqlite3.connect(f"file:{MESSAGE_DB}?mode=ro", uri=True, timeout=2)
            try: conn.execute("SELECT count(*) FROM message LIMIT 1").fetchone(); permission = True
            finally: conn.close()
            db_mtime = dt.datetime.fromtimestamp(MESSAGE_DB.stat().st_mtime).replace(microsecond=0).isoformat(sep=" ")
        except Exception as exc: permission_error = str(exc)
        with get_db() as conn:
            config = dict(conn.execute("SELECT * FROM case_manager_message_settings WHERE id=1").fetchone())
            counts = {row["workflow_status"]: row["count"] for row in conn.execute("SELECT workflow_status,count(*) AS count FROM case_manager_judicial_messages WHERE COALESCE(is_deleted,0)=0 GROUP BY workflow_status").fetchall()}
            new_row = conn.execute(
                """SELECT count(*) AS count,COALESCE(MAX(created_at),'') AS latest_at
                   FROM case_manager_judicial_messages
                   WHERE COALESCE(is_deleted,0)=0 AND COALESCE(is_new,0)=1"""
            ).fetchone()
        config["ai_model"] = effective_purpose_model("judicial_messages")
        ai_ready = bool(config["ai_enabled"] and model_available(config["ai_model"]))
        stale = False
        if db_mtime:
            try: stale = dt.datetime.fromisoformat(db_mtime) < dt.datetime.now() - dt.timedelta(days=1)
            except ValueError: pass
        return {"permission_granted":permission,"permission_error":permission_error,"database_exists":MESSAGE_DB.exists(),"database_mtime":db_mtime,"source_stale":stale,"ai_ready":ai_ready,"counts":counts,"new_count":int(new_row["count"] or 0),"latest_new_at":new_row["latest_at"] or "","settings":{k:v for k,v in config.items() if k != "hmac_secret"},"python_executable":full_disk_access_target(),"messages_database":str(MESSAGE_DB)}

    @app.get("/api/case-manager/judicial-messages/status")
    def message_status(): return health()

    @app.get("/api/case-manager/judicial-messages/settings")
    def get_message_settings(): return settings()

    @app.post("/api/case-manager/judicial-messages/import-preview")
    async def preview_manual_message(request: Request):
        """Read a bounded text/image input without writing messages or case records."""
        from starlette.concurrency import run_in_threadpool
        form = await request.form()
        text = str(form.get('text') or '')
        source = 'manual-text'
        upload = form.get('file')
        if upload and hasattr(upload, 'read'):
            raw = await upload.read(8 * 1024 * 1024 + 1)
            await upload.close()
            if len(raw) > 8 * 1024 * 1024:
                raise HTTPException(413, '导入文件不能超过 8MB')
            suffix = Path(upload.filename or '').suffix.lower()
            if suffix == '.txt':
                try:
                    text = raw.decode('utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig')
                except UnicodeError:
                    raise HTTPException(400, '请使用 UTF-8 或 UTF-16 编码的文本文件')
                source = 'manual-text'
            elif suffix in {'.png', '.jpg', '.jpeg', '.webp'}:
                from PIL import Image
                from io import BytesIO
                try:
                    with Image.open(BytesIO(raw)) as image:
                        if image.width * image.height > 20_000_000:
                            raise ValueError('image size')
                        image.verify()
                except Exception:
                    raise HTTPException(400, '请选择有效且不超过 2000 万像素的短信截图')
                from .court_directory import ocr_image_bytes
                text = await run_in_threadpool(ocr_image_bytes, raw)
                source = 'manual-image'
            else:
                raise HTTPException(400, '请选择 TXT 文本或 PNG、JPEG、WebP 截图')
        text = text.strip()
        if not text or len(text) > 100_000:
            raise HTTPException(400, '未识别到文字，或文字超过 10 万字；请检查后重新导入')
        return {'text': text, 'source': source, 'requires_confirmation': True}

    @app.post("/api/case-manager/judicial-messages/import")
    async def import_manual_message(request: Request):
        body = await request.json()
        text = str(body.get('text') or '').strip()
        sender = str(body.get('sender') or '').strip()
        received = str(body.get('received_at') or '').strip()
        if body.get('confirmed') is not True or not text or len(text) > 100_000 or len(sender) > 100:
            raise HTTPException(400, '请核对短信内容后确认导入')
        if received:
            try:
                received = dt.datetime.fromisoformat(received).isoformat(timespec='seconds')
            except ValueError:
                raise HTTPException(400, '接收时间格式不正确；不确定时请留空')
        source = 'manual-image' if body.get('source') == 'manual-image' else 'manual-text'
        checksum = fingerprint(text, sender, received, 'manual-v1')
        stamp = _now()
        extracted = extract_fields(text, received, allow_inferred_year=bool(received))
        with get_db() as conn:
            row = conn.execute('SELECT id FROM case_manager_judicial_messages WHERE fingerprint=?', (checksum,)).fetchone()
            if row:
                return {'ok': True, 'duplicate': True, 'id': row['id']}
            cursor = conn.execute('''INSERT OR IGNORE INTO case_manager_judicial_messages
                (source_guid,fingerprint,sender,service,received_at,raw_body,institution,category,
                 extracted_json,case_numbers_json,relevance_score,is_new,created_at,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,1,?,?)''',
                ('manual:' + checksum, checksum, sender, source, received, text,
                 extracted.get('institution', ''), extracted.get('category', ''), _json(extracted),
                 _json(extracted.get('case_numbers', [])), relevance_score(text, sender), stamp, stamp))
            if not cursor.rowcount:
                existing = conn.execute('SELECT id FROM case_manager_judicial_messages WHERE fingerprint=?', (checksum,)).fetchone()
                return {'ok': True, 'duplicate': True, 'id': existing['id']}
            message_id = cursor.lastrowid
            conn.commit()
        # Imported text and OCR never bypass the existing human confirmation workflow.
        await analyze_saved(message_id, False)
        return {'ok': True, 'duplicate': False, 'id': message_id}

    @app.put("/api/case-manager/judicial-messages/settings")
    async def save_message_settings(request: Request):
        body = await request.json(); stamp = _now()
        from .platform_capabilities import enabled
        if not enabled('messages_sync'):
            body['enabled'] = False
        interval = max(60, min(int(body.get("poll_interval_seconds") or 300), 3600))
        backfill = max(1, min(int(body.get("backfill_days") or 90), 3650))
        model = _norm(body.get("ai_model")) or effective_purpose_model("judicial_messages")
        if model not in ALLOWED_MODELS: raise HTTPException(status_code=400, detail="不支持的 AI 模型")
        recipient_phones: List[str] = []
        for raw in body.get("recipient_phones") or []:
            phone = re.sub(r"\D", "", _norm(raw))
            if phone and not re.fullmatch(r"1\d{10}", phone):
                raise HTTPException(status_code=400, detail=f"收件手机号格式不正确：{raw}")
            if phone and phone not in recipient_phones:
                recipient_phones.append(phone)
        with get_db() as conn:
            conn.execute("UPDATE case_manager_message_settings SET enabled=?,poll_interval_seconds=?,backfill_days=?,ai_enabled=?,ai_model=?,recipient_phones_json=?,updated_at=? WHERE id=1", (1 if body.get("enabled", True) else 0,interval,backfill,1 if body.get("ai_enabled", True) else 0,model,_json(recipient_phones),stamp))
            if "trusted_senders" in body:
                conn.execute("DELETE FROM case_manager_trusted_senders")
                for item in body.get("trusted_senders") or []:
                    sender = _norm(item.get("sender"))
                    if sender: conn.execute("INSERT INTO case_manager_trusted_senders(sender,institution,verified,created_at,updated_at) VALUES (?,?,?,?,?)", (sender,_norm(item.get("institution")),1 if item.get("verified") else 0,stamp,stamp))
            conn.commit()
        from .ai_config import load_ai_prefs, save_ai_prefs
        prefs=load_ai_prefs(); prefs['purpose_models']={**prefs.get('purpose_models',{}),'judicial_messages':model}
        save_ai_prefs(prefs)
        return {"ok": True, "settings": settings()}

    @app.get("/api/case-manager/judicial-messages")
    def list_messages(q: str = "", status: str = "", category: str = "", institution: str = "", case_no: str = "", case_id: int = 0, date_from: str = "", date_to: str = "", new_only: int = 0, page: int = 1, page_size: int = 50):
        where = ["COALESCE(m.is_deleted,0)=0"]; args: List[Any] = []
        if new_only: where.append("COALESCE(m.is_new,0)=1")
        if status: where.append("m.workflow_status=?"); args.append(status)
        if category: where.append("m.category=?"); args.append(category)
        if institution: where.append("m.institution LIKE ?"); args.append(f"%{institution}%")
        if case_no: where.append("m.case_numbers_json LIKE ?"); args.append(f"%{case_no}%")
        if case_id: where.append("COALESCE(m.matched_case_id,m.suggested_case_id)=?"); args.append(case_id)
        if date_from: where.append("date(m.received_at)>=date(?)"); args.append(date_from)
        if date_to: where.append("date(m.received_at)<=date(?)"); args.append(date_to)
        if q:
            where.append("(m.raw_body LIKE ? OR m.institution LIKE ? OR m.category LIKE ? OR m.case_numbers_json LIKE ? OR c.title LIKE ?)"); args.extend([f"%{q}%"]*5)
        size = max(1,min(page_size,200)); offset = max(0,page-1)*size
        sql_where = " AND ".join(where)
        with get_db() as conn:
            total = conn.execute(f"SELECT count(*) FROM case_manager_judicial_messages m LEFT JOIN case_manager_cases c ON c.id=COALESCE(m.matched_case_id,m.suggested_case_id) WHERE {sql_where}", args).fetchone()[0]
            rows = conn.execute(f"""SELECT m.*,COALESCE(c.title,'') AS case_title FROM case_manager_judicial_messages m
                 LEFT JOIN case_manager_cases c ON c.id=COALESCE(m.matched_case_id,m.suggested_case_id)
                 WHERE {sql_where} ORDER BY datetime(m.received_at) DESC,m.id DESC LIMIT ? OFFSET ?""", [*args,size,offset]).fetchall()
        items = []
        for row in rows:
            item = dict(row); item["raw_body_preview"] = (item.get("raw_body") or "")[:260]; item["status_label"] = STATUS_LABELS.get(item["workflow_status"],item["workflow_status"]); item.pop("analysis_json",None); item.pop("extracted_json",None); items.append(item)
        return {"items":items,"total":total,"page":page,"page_size":size,"categories":CATEGORIES,"statuses":STATUS_LABELS}

    @app.post("/api/case-manager/judicial-messages/read-bulk")
    async def read_messages_bulk(request: Request):
        body = await request.json()
        message_ids = list(dict.fromkeys(int(value) for value in (body.get("message_ids") or []) if int(value) > 0))
        mark_all = bool(body.get("all_new"))
        if not mark_all and not message_ids:
            raise HTTPException(status_code=400, detail="请选择需要标记为已读的短信")
        stamp = _now()
        with get_db() as conn:
            if mark_all:
                cur = conn.execute(
                    "UPDATE case_manager_judicial_messages SET is_new=0,updated_at=? WHERE COALESCE(is_deleted,0)=0 AND is_new=1",
                    (stamp,),
                )
            else:
                placeholders = ",".join("?" for _ in message_ids)
                cur = conn.execute(
                    f"UPDATE case_manager_judicial_messages SET is_new=0,updated_at=? WHERE id IN ({placeholders}) AND COALESCE(is_deleted,0)=0",
                    [stamp, *message_ids],
                )
            conn.commit()
        return {"ok":True,"read_count":int(cur.rowcount or 0)}

    @app.post("/api/case-manager/judicial-messages/delete-bulk")
    async def delete_messages_bulk(request: Request):
        body = await request.json()
        message_ids = list(dict.fromkeys(int(value) for value in (body.get("message_ids") or []) if int(value) > 0))
        if not message_ids:
            raise HTTPException(status_code=400, detail="请选择需要删除的短信")
        if len(message_ids) > 500:
            raise HTTPException(status_code=400, detail="单次最多删除500条短信")
        stamp = _now(); placeholders = ",".join("?" for _ in message_ids)
        with get_db() as conn:
            rows = conn.execute(
                f"""SELECT id,fingerprint,source_guid,source_rowid,received_at
                    FROM case_manager_judicial_messages
                    WHERE id IN ({placeholders}) AND COALESCE(is_deleted,0)=0""",
                message_ids,
            ).fetchall()
            deleted_ids = [int(row["id"]) for row in rows]
            linked_count = 0
            if deleted_ids:
                deleted_placeholders = ",".join("?" for _ in deleted_ids)
                linked_count = int(conn.execute(
                    f"SELECT count(DISTINCT message_id) FROM case_manager_message_links WHERE message_id IN ({deleted_placeholders})",
                    deleted_ids,
                ).fetchone()[0] or 0)
                for row in rows:
                    conn.execute(
                        "INSERT OR REPLACE INTO case_manager_message_fingerprints(fingerprint,source_guid,source_rowid,received_at,disposition,created_at) VALUES (?,?,?,?,?,?)",
                        (row["fingerprint"],row["source_guid"],row["source_rowid"],row["received_at"],"deleted",stamp),
                    )
                # Soft-delete the inbox record so linked hearings, todos,
                # ledgers and memos keep their source relation and audit trail.
                conn.execute(
                    f"""UPDATE case_manager_judicial_messages
                        SET raw_body='',sender='',service='',institution='',category='',
                            extracted_json='{{}}',analysis_json='{{}}',candidate_cases_json='[]',
                            case_numbers_json='[]',analysis_error='',is_new=0,is_deleted=1,
                            deleted_at=?,updated_at=?
                        WHERE id IN ({deleted_placeholders})""",
                    [stamp, stamp, *deleted_ids],
                )
            conn.commit()
        return {"ok":True,"deleted_count":len(deleted_ids),"deleted_ids":deleted_ids,"linked_count":linked_count,"skipped_count":len(message_ids)-len(deleted_ids)}

    @app.get("/api/case-manager/judicial-messages/{message_id}/documents")
    def list_message_documents(message_id: int):
        message = get_message(message_id)
        return {"items": message["documents"]}

    @app.post("/api/case-manager/judicial-messages/{message_id}/documents/discover")
    def rediscover_message_documents(message_id: int):
        message = get_message(message_id)
        return {"items": discover_documents(message_id, message.get("raw_body") or "")}

    @app.get("/api/case-manager/judicial-messages/{message_id}/documents/{document_id}/download-progress")
    async def message_document_download_progress(message_id: int, document_id: int):
        with get_db() as conn:
            row = conn.execute(
                "SELECT 1 FROM case_manager_message_documents WHERE id=? AND message_id=?",
                (document_id, message_id),
            ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="文书任务不存在")
        return dict(download_progress.get((message_id, document_id), {"status": "idle"}))

    @app.post("/api/case-manager/judicial-messages/{message_id}/documents/{document_id}/download")
    async def download_message_document(message_id: int, document_id: int, request: Request):
        body = await request.json()
        with get_db() as conn:
            row = conn.execute(
                "SELECT 1 FROM case_manager_message_documents WHERE id=? AND message_id=?",
                (document_id, message_id),
            ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="文书任务不存在")
        key = (message_id, document_id)
        if download_progress.get(key, {}).get("status") == "running":
            raise HTTPException(status_code=409, detail="文书正在下载，请查看进度")
        # Retain a bounded set of results for drawer reopen/reconnect.
        while len(download_progress) >= 64:
            finished = next((key for key, value in download_progress.items() if value["status"] != "running"), None)
            if finished is None:
                raise HTTPException(status_code=429, detail="下载任务较多，请稍后重试")
            download_progress.pop(finished)
        progress = {
            "status": "running", "phase": "preparing", "filename": "",
            "received_bytes": 0, "total_bytes": 0, "total_files": 0,
            "completed_files": 0, "failed_files": 0, "error": "",
        }
        if callbacks.get('download_center_start'):
            progress['_center_id']=callbacks['download_center_start'](document_id)
        download_progress[key] = progress
        token = current_progress.set(progress)
        try:
            item = await process_document(document_id, int(body.get("case_id") or 0))
            if item.get("status") == "needs_login":
                progress.update(status="needs_action", error=item.get("last_error") or "请完成页面验证")
            else:
                if not progress["completed_files"] and not progress["failed_files"]:
                    progress.update(total_files=1, completed_files=1)
                progress["status"] = "partial" if progress["failed_files"] else "completed"
            return {"ok": True, "document": item, "progress": dict(progress)}
        except asyncio.CancelledError:
            progress.update(status="failed", error="下载已中断，请重试")
            raise
        except Exception as exc:
            progress.update(status="failed", error=str(exc.detail if isinstance(exc, HTTPException) else exc)[:1000])
            raise
        finally:
            report_download()
            current_progress.reset(token)

    @app.post("/api/case-manager/judicial-messages/{message_id}/documents/{document_id}/archive")
    async def archive_message_document(message_id: int, document_id: int, request: Request):
        body = await request.json()
        case_id = int(body.get("case_id") or 0)
        if not case_id:
            raise HTTPException(status_code=400, detail="请选择归档案件")
        with get_db() as conn:
            row = conn.execute(
                "SELECT 1 FROM case_manager_message_documents WHERE id=? AND message_id=?",
                (document_id, message_id),
            ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="文书任务不存在")
        return {"ok": True, "document": archive_document(document_id, case_id)}

    @app.post("/api/case-manager/judicial-messages/{message_id}/documents/{document_id}/import")
    async def import_message_document(message_id: int, document_id: int, file: UploadFile = File(...)):
        filename = file.filename or "法院文书"
        content = await file.read(MAX_DOCUMENT_BYTES + 1)
        if len(content) > MAX_DOCUMENT_BYTES:
            raise HTTPException(status_code=413, detail=f"文书超过 {MAX_DOCUMENT_MB}MB 限制")
        valid, detected = _valid_document(content, filename, file.content_type or "")
        if not valid:
            raise HTTPException(status_code=400, detail="请选择 PDF、Word 或图片文书")
        if Path(filename).suffix.lower() != detected:
            filename = f"{Path(filename).stem}{detected}"
        with get_db() as conn:
            row = conn.execute(
                "SELECT 1 FROM case_manager_message_documents WHERE id=? AND message_id=?",
                (document_id, message_id),
            ).fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="文书任务不存在")
            staging_path = stage_document(document_id, content, filename)
            conn.execute(
                """UPDATE case_manager_message_documents
                   SET status='waiting_confirmation',staging_path=?,original_filename=?,mime_type=?,
                       sha256=?,last_error='',updated_at=? WHERE id=?""",
                (
                    staging_path,
                    filename[:240],
                    (file.content_type or "")[:120],
                    hashlib.sha256(content).hexdigest(),
                    _now(),
                    document_id,
                ),
            )
            conn.commit()
        return {"ok": True, "document": next(item for item in document_rows(message_id) if item["id"] == document_id)}

    @app.post("/api/case-manager/judicial-messages/{message_id}/documents/{document_id}/browser-download")
    async def browser_download_message_document(message_id: int, document_id: int):
        with get_db() as conn:
            row = conn.execute(
                """SELECT d.*,m.matched_case_id FROM case_manager_message_documents d
                   JOIN case_manager_judicial_messages m ON m.id=d.message_id
                   WHERE d.id=? AND d.message_id=?""",
                (document_id, message_id),
            ).fetchone()
        if not row or not row["source_url"]:
            raise HTTPException(status_code=404, detail="没有可接管的文书链接")
        _validate_public_url(row["source_url"])
        # Public links sent by 人民法院在线服务 expose a document-list API.
        # Resolve it directly so old “登录接管” tasks also become zero-click and
        # collect every attachment instead of closing after the first download.
        if _court_delivery_params(row["source_url"]):
            item = await process_document(
                document_id,
                int(row["matched_case_id"] or 0),
            )
            return {
                "ok": True,
                "document": item,
                "documents": document_rows(message_id),
            }
        try:
            from playwright.async_api import TimeoutError as PlaywrightTimeoutError
            from playwright.async_api import async_playwright
        except ImportError as exc:
            raise HTTPException(status_code=501, detail="当前版本未安装浏览器接管组件") from exc
        with get_db() as conn:
            conn.execute(
                "UPDATE case_manager_message_documents SET status='needs_login',last_error='',updated_at=? WHERE id=?",
                (_now(), document_id),
            )
            conn.commit()
        try:
            async with async_playwright() as playwright:
                launch_options: Dict[str, Any] = {
                    "headless": False,
                    "args": ["--no-first-run", "--no-default-browser-check"],
                }
                executable = _browser_executable()
                if executable:
                    launch_options["executable_path"] = executable
                browser = await playwright.chromium.launch(**launch_options)
                try:
                    context = await browser.new_context(accept_downloads=True)
                    loop = asyncio.get_running_loop()
                    download_future: asyncio.Future[Any] = loop.create_future()
                    disconnected_future: asyncio.Future[bool] = loop.create_future()
                    attached_pages: set[int] = set()

                    def receive_download(download: Any) -> None:
                        if not download_future.done():
                            download_future.set_result(download)

                    def attach_page(page: Any) -> None:
                        page_key = id(page)
                        if page_key in attached_pages:
                            return
                        attached_pages.add(page_key)
                        page.on("download", receive_download)

                    def browser_disconnected() -> None:
                        if not disconnected_future.done():
                            disconnected_future.set_result(True)

                    context.on("page", attach_page)
                    browser.on("disconnected", browser_disconnected)
                    page = await context.new_page()
                    attach_page(page)
                    try:
                        await page.goto(row["source_url"], wait_until="domcontentloaded", timeout=90_000)
                    except PlaywrightTimeoutError:
                        # The page can remain interactive even when a dynamic court
                        # platform never reaches the DOMContentLoaded milestone.
                        pass
                    done, _ = await asyncio.wait(
                        {download_future, disconnected_future},
                        timeout=300,
                        return_when=asyncio.FIRST_COMPLETED,
                    )
                    if download_future not in done:
                        if disconnected_future in done:
                            raise RuntimeError("Target page, context or browser has been closed")
                        raise asyncio.TimeoutError("waiting for browser download")
                    download = download_future.result()
                    failure = await download.failure()
                    if failure:
                        raise RuntimeError(f"download canceled: {failure}")
                    download_path = await download.path()
                    if not download_path:
                        raise RuntimeError("浏览器未返回下载文件")
                    downloaded_file = Path(download_path)
                    if downloaded_file.stat().st_size > MAX_DOCUMENT_BYTES:
                        raise HTTPException(status_code=413, detail=f"文书超过 {MAX_DOCUMENT_MB}MB 限制")
                    content = downloaded_file.read_bytes()
                    filename = download.suggested_filename or row["original_filename"] or "法院文书"
                finally:
                    if browser.is_connected():
                        await browser.close()
            valid, detected = _valid_document(content, filename)
            if not valid:
                raise HTTPException(status_code=400, detail="浏览器下载的内容不是可识别文书")
            if Path(filename).suffix.lower() != detected:
                filename = f"{Path(filename).stem}{detected}"
            staging_path = stage_document(document_id, content, filename)
            with get_db() as conn:
                conn.execute(
                    """UPDATE case_manager_message_documents
                       SET status='waiting_confirmation',staging_path=?,original_filename=?,sha256=?,
                           last_error='',updated_at=? WHERE id=?""",
                    (
                        staging_path,
                        filename[:240],
                        hashlib.sha256(content).hexdigest(),
                        _now(),
                        document_id,
                    ),
                )
                conn.commit()
            return {"ok": True, "document": next(item for item in document_rows(message_id) if item["id"] == document_id)}
        except HTTPException:
            raise
        except Exception as exc:
            friendly_error = _friendly_browser_error(exc)
            with get_db() as conn:
                conn.execute(
                    "UPDATE case_manager_message_documents SET status='needs_login',last_error=?,updated_at=? WHERE id=?",
                    (friendly_error, _now(), document_id),
                )
                conn.commit()
            raise HTTPException(status_code=408, detail=friendly_error) from exc

    @app.post("/api/case-manager/judicial-messages/{message_id}/documents/{document_id}/verification/start")
    async def start_document_verification(message_id: int, document_id: int):
        document = verification_document(message_id, document_id)
        await close_verification(document_id)
        client = httpx.AsyncClient(
            timeout=httpx.Timeout(45.0, connect=15.0),
            follow_redirects=True,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
                ),
                "Referer": document["source_url"],
            },
        )
        session = {
            "client": client,
            "source_url": document["source_url"],
            "delivery_id": _gdems_delivery_id(document["source_url"]),
            "expires_at": dt.datetime.now().timestamp() + 10 * 60,
        }
        delivery_verifications[document_id] = session
        page = await gdems_request(session, "GET", document["source_url"])
        if page.status_code >= 400:
            await close_verification(document_id)
            raise HTTPException(status_code=502, detail=f"送达页打开失败（HTTP {page.status_code}）")
        captcha_url = urllib.parse.urljoin(str(page.url), "/servlet/validateCodeServlet")
        captcha = await gdems_request(session, "GET", captcha_url)
        if captcha.status_code != 200 or not captcha.content:
            await close_verification(document_id)
            raise HTTPException(status_code=502, detail="图形验证码获取失败")
        media_type = (captcha.headers.get("content-type") or "image/jpeg").split(";", 1)[0]
        return {
            "ok": True,
            "captcha": f"data:{media_type};base64,{base64.b64encode(captcha.content).decode('ascii')}",
            "expires_in": 600,
        }

    @app.post("/api/case-manager/judicial-messages/{message_id}/documents/{document_id}/verification/send-code")
    async def send_document_verification_code(message_id: int, document_id: int, request: Request):
        verification_document(message_id, document_id)
        session = verification_session(document_id)
        body = await request.json()
        phone = re.sub(r"\D", "", _norm(body.get("phone")))
        captcha = _norm(body.get("captcha"))
        if not re.fullmatch(r"1\d{10}", phone):
            raise HTTPException(status_code=400, detail="请输入正确的 11 位收件手机号")
        if not captcha:
            raise HTTPException(status_code=400, detail="请输入图形验证码")
        endpoint = urllib.parse.urljoin(session["source_url"], "/v3/dzsd/getCode")
        response = await gdems_request(
            session,
            "POST",
            endpoint,
            data={"id": session["delivery_id"], "photoCode": captcha, "number": phone},
        )
        try:
            payload = response.json()
        except ValueError as exc:
            raise HTTPException(status_code=502, detail="送达平台未返回有效的验证码结果") from exc
        if str(payload.get("result")) != "1":
            raise HTTPException(status_code=400, detail=_norm(payload.get("errorDetail")) or "验证码发送失败")
        session["expires_at"] = dt.datetime.now().timestamp() + 10 * 60
        return {"ok": True, "message": "短信验证码已发送", "retry_after": 120}

    @app.post("/api/case-manager/judicial-messages/{message_id}/documents/{document_id}/verification/complete")
    async def complete_document_verification(message_id: int, document_id: int, request: Request):
        document = verification_document(message_id, document_id)
        session = verification_session(document_id)
        body = await request.json()
        code = _norm(body.get("code"))
        case_id = int(body.get("case_id") or document.get("matched_case_id") or document.get("suggested_case_id") or 0)
        if not code and not session.get("verified"):
            raise HTTPException(status_code=400, detail="请输入短信验证码")
        if not session.get("verified"):
            endpoint = urllib.parse.urljoin(session["source_url"], "/v3/dzsd/validateCode")
            response = await gdems_request(
                session,
                "POST",
                endpoint,
                data={"id": session["delivery_id"], "code": code},
            )
            try:
                payload = response.json()
            except ValueError as exc:
                raise HTTPException(status_code=502, detail="送达平台未返回有效的验证结果") from exc
            if str(payload.get("result")) != "1":
                raise HTTPException(status_code=400, detail=_norm(payload.get("errorDetail")) or "短信验证码不正确")
            session["verified"] = True
            session["expires_at"] = dt.datetime.now().timestamp() + 10 * 60

        page = await gdems_request(session, "GET", session["source_url"])
        if page.status_code >= 400:
            raise HTTPException(status_code=502, detail="验证成功，但文书页面打开失败")
        try:
            (JUDICIAL_STAGING_DIR / f"gdems-page-{document_id}.html").write_text(
                page.text, encoding="utf-8"
            )
        except OSError:
            pass
        gdems_manifest = _gdems_document_manifest(page.text, str(page.url))
        candidates = [item["url"] for item in gdems_manifest]
        candidate_names = {item["url"]: item["filename"] for item in gdems_manifest}
        if not candidates:
            candidates = _html_download_candidates(page.text, str(page.url))
            candidate_names = {}
        downloads: List[tuple[bytes, str, str, str]] = []
        errors: List[str] = []
        for url in candidates:
            try:
                response = await gdems_request(session, "GET", url)
                if response.status_code >= 400 or len(response.content) > MAX_DOCUMENT_BYTES:
                    continue
                disposition = response.headers.get("content-disposition") or ""
                match = re.search(r"filename\\*?=(?:UTF-8''|[\"']?)([^\"';]+)", disposition, re.I)
                filename = (
                    urllib.parse.unquote(match.group(1))
                    if match else candidate_names.get(url) or Path(urllib.parse.urlparse(str(response.url)).path).name
                )
                valid, detected = _valid_document(response.content, filename, response.headers.get("content-type") or "")
                if not valid:
                    continue
                if Path(filename).suffix.lower() != detected:
                    filename = f"{Path(filename).stem or '法院文书'}{detected}"
                digest = hashlib.sha256(response.content).hexdigest()
                if digest not in {hashlib.sha256(item[0]).hexdigest() for item in downloads}:
                    downloads.append((response.content, filename, response.headers.get("content-type") or "", url))
            except HTTPException as exc:
                errors.append(str(exc.detail))
        if not downloads:
            try:
                downloads = await collect_gdems_browser_downloads(session)
            except Exception as exc:
                errors.append(str(exc))
        if not downloads:
            diagnostic_path = JUDICIAL_STAGING_DIR / f"gdems-diagnostic-{document_id}.log"
            try:
                diagnostic_path.write_text(
                    "\n".join(session.get("download_diagnostics") or []), encoding="utf-8"
                )
            except OSError:
                pass
            raise HTTPException(
                status_code=502,
                detail="收件验证已通过，自动下载仍未取得文书；可直接点击“重试自动下载”，无需再次获取验证码",
            )

        result_ids: List[int] = []
        stamp = _now()
        for index, (content, filename, content_type, source_url) in enumerate(downloads):
            item_id = document_id
            if index:
                key = hashlib.sha256(source_url.encode()).hexdigest()
                with get_db() as conn:
                    conn.execute(
                        """INSERT OR IGNORE INTO case_manager_message_documents
                           (message_id,source_kind,source_url,source_path,original_filename,status,created_at,updated_at)
                           VALUES (?,'delivery_item',?,?,?,'discovered',?,?)""",
                        (message_id, document["source_url"], key, filename[:240], stamp, stamp),
                    )
                    row = conn.execute(
                        """SELECT id FROM case_manager_message_documents
                           WHERE message_id=? AND source_kind='delivery_item' AND source_url=? AND source_path=?""",
                        (message_id, document["source_url"], key),
                    ).fetchone()
                    conn.commit()
                    item_id = int(row["id"])
            staging_path = stage_document(item_id, content, filename)
            with get_db() as conn:
                conn.execute(
                    """UPDATE case_manager_message_documents
                       SET status='waiting_confirmation',staging_path=?,original_filename=?,mime_type=?,
                           sha256=?,last_error='',updated_at=? WHERE id=?""",
                    (staging_path, filename[:240], content_type[:120], hashlib.sha256(content).hexdigest(), stamp, item_id),
                )
                conn.commit()
            if case_id:
                archive_document(item_id, case_id)
            result_ids.append(item_id)
        await close_verification(document_id)
        return {
            "ok": True,
            "archived": bool(case_id),
            "downloaded_count": len(result_ids),
            "documents": document_rows(message_id),
        }

    @app.post("/api/case-manager/judicial-messages/{message_id}/documents/{document_id}/open")
    def open_message_document(message_id: int, document_id: int):
        item = next(
            (row for row in document_rows(message_id) if int(row["id"]) == document_id),
            None,
        )
        if not item:
            raise HTTPException(status_code=404, detail="文书任务不存在")
        path = item.get("absolute_path") or item.get("staging_path") or ""
        if not path or not Path(path).is_file():
            raise HTTPException(status_code=404, detail="文书文件当前不可用")
        if sys.platform != "darwin":
            raise HTTPException(status_code=501, detail="当前系统不支持 Finder 打开")
        subprocess.Popen(["open", "-R", path])
        return {"ok": True}

    @app.get("/api/case-manager/judicial-messages/{message_id}")
    def message_detail(message_id: int): return get_message(message_id)

    @app.post("/api/case-manager/judicial-messages/scan")
    async def scan_messages(): return await scan_once(True)

    @app.post("/api/case-manager/judicial-messages/{message_id}/reanalyze")
    async def reanalyze_message(message_id: int): return await analyze_saved(message_id, True)

    @app.post("/api/case-manager/judicial-messages/{message_id}/confirm")
    async def confirm_message(message_id: int, request: Request):
        body = await request.json(); case_id = int(body.get("case_id") or 0); category = _norm(body.get("category")); destinations = body.get("destinations") or []
        if not case_id or category not in CATEGORIES or not destinations: raise HTTPException(status_code=400, detail="请选择案件、类别和至少一个写入目标")
        stamp = _now()
        with get_db() as conn:
            row = conn.execute("SELECT * FROM case_manager_judicial_messages WHERE id=?", (message_id,)).fetchone()
            if not row: raise HTTPException(status_code=404, detail="短信不存在")
            message = dict(row); links = create_outputs(conn,message,case_id,category,destinations,body.get("fields") or {},"manual")
            conn.execute("UPDATE case_manager_judicial_messages SET matched_case_id=?,category=?,workflow_status='confirmed',decision_mode='manual',confirmed_at=?,updated_at=? WHERE id=?", (case_id,category,stamp,stamp,message_id)); conn.commit()
        for link in links:
            if link["target_type"] in {"hearing","todo"}:
                try: callbacks["auto_sync_if_enabled"](link["target_type"],link["target_id"],force=link["target_type"]=="todo")
                except Exception: pass
        for document in list(document_rows(message_id)):
            if (
                _court_delivery_params(document.get("source_url") or "")
                and document.get("source_kind") == "url"
            ):
                try:
                    await process_document(int(document["id"]), case_id)
                except Exception:
                    # Keep the confirmed SMS result; the individual document
                    # row carries a retryable error if one file failed.
                    pass
            elif document["status"] in {"discovered", "failed", "needs_login"}:
                try:
                    await process_document(int(document["id"]), case_id)
                except Exception:
                    pass
            elif document["status"] in {"waiting_confirmation", "downloaded", "verified"}:
                try:
                    archive_document(int(document["id"]), case_id)
                except Exception:
                    pass
        return {"ok":True,"message":get_message(message_id)}

    @app.post("/api/case-manager/judicial-messages/{message_id}/ignore")
    def ignore_message(message_id: int):
        stamp = _now()
        with get_db() as conn:
            row = conn.execute("SELECT fingerprint,source_guid,source_rowid,received_at,workflow_status FROM case_manager_judicial_messages WHERE id=?", (message_id,)).fetchone()
            if not row: raise HTTPException(status_code=404, detail="短信不存在")
            if row["workflow_status"] not in {"pending", "suggested"}: raise HTTPException(status_code=409, detail="已处理或已写入案件的短信不能直接忽略")
            conn.execute("INSERT OR REPLACE INTO case_manager_message_fingerprints(fingerprint,source_guid,source_rowid,received_at,disposition,created_at) VALUES (?,?,?,?,?,?)", (row["fingerprint"],row["source_guid"],row["source_rowid"],row["received_at"],"ignored",stamp))
            conn.execute("UPDATE case_manager_judicial_messages SET raw_body='',workflow_status='ignored',decision_mode='manual',updated_at=? WHERE id=?", (stamp,message_id)); conn.commit()
        return {"ok":True}

    @app.post("/api/case-manager/judicial-messages/ignore-bulk")
    async def ignore_messages_bulk(request: Request):
        body = await request.json()
        message_ids = list(dict.fromkeys(int(value) for value in (body.get("message_ids") or []) if int(value) > 0))
        if not message_ids:
            raise HTTPException(status_code=400, detail="请选择需要忽略的短信")
        if len(message_ids) > 500:
            raise HTTPException(status_code=400, detail="单次最多处理500条短信")
        stamp = _now(); placeholders = ",".join("?" for _ in message_ids)
        with get_db() as conn:
            rows = conn.execute(
                f"""SELECT id,fingerprint,source_guid,source_rowid,received_at
                    FROM case_manager_judicial_messages
                    WHERE id IN ({placeholders}) AND workflow_status IN ('pending','suggested')""",
                message_ids,
            ).fetchall()
            eligible_ids = [int(row["id"]) for row in rows]
            for row in rows:
                conn.execute(
                    "INSERT OR REPLACE INTO case_manager_message_fingerprints(fingerprint,source_guid,source_rowid,received_at,disposition,created_at) VALUES (?,?,?,?,?,?)",
                    (row["fingerprint"],row["source_guid"],row["source_rowid"],row["received_at"],"ignored",stamp),
                )
            if eligible_ids:
                eligible_placeholders = ",".join("?" for _ in eligible_ids)
                conn.execute(
                    f"UPDATE case_manager_judicial_messages SET raw_body='',workflow_status='ignored',decision_mode='manual',updated_at=? WHERE id IN ({eligible_placeholders})",
                    [stamp, *eligible_ids],
                )
            conn.commit()
        return {"ok":True,"ignored_count":len(eligible_ids),"ignored_ids":eligible_ids,"skipped_count":len(message_ids)-len(eligible_ids)}

    @app.post("/api/case-manager/judicial-messages/{message_id}/undo")
    def undo_auto_message(message_id: int):
        with get_db() as conn:
            message = conn.execute("SELECT * FROM case_manager_judicial_messages WHERE id=?", (message_id,)).fetchone()
            if not message or message["workflow_status"] != "auto_written": raise HTTPException(status_code=400, detail="只能撤销短信自动创建且尚未修改的记录")
            links = [dict(row) for row in conn.execute("SELECT * FROM case_manager_message_links WHERE message_id=? AND creation_mode='auto'", (message_id,)).fetchall()]
            tables = {"hearing":"case_manager_hearings","todo":"case_manager_todos","ledger":"research_case_ledgers"}
            for link in links:
                table = tables.get(link["target_type"])
                if not table: raise HTTPException(status_code=409, detail="该短信包含不能自动撤销的记录")
                target = conn.execute(f"SELECT updated_at FROM {table} WHERE id=?", (link["target_id"],)).fetchone()
                if target and _norm(target["updated_at"]) != _norm(link["target_created_at"]): raise HTTPException(status_code=409, detail="自动创建的记录已被人工修改，请在对应列表中处理")
            for link in links:
                callbacks["delete_calendar_sync_record"](link["target_type"],link["target_id"])
                if link["target_type"] == "hearing": conn.execute("DELETE FROM research_case_ledgers WHERE source_type='case_manager_hearing' AND source_id=?", (link["target_id"],))
                conn.execute(f"DELETE FROM {tables[link['target_type']]} WHERE id=?", (link["target_id"],))
            conn.execute("DELETE FROM case_manager_message_links WHERE message_id=?", (message_id,))
            conn.execute("UPDATE case_manager_judicial_messages SET workflow_status='suggested',decision_mode='',matched_case_id=NULL,confirmed_at='',updated_at=? WHERE id=?", (_now(),message_id)); conn.commit()
        return {"ok":True,"message":get_message(message_id)}

    async def scheduler_loop() -> None:
        await asyncio.sleep(5)
        while True:
            try: await scan_once(False)
            except Exception as exc: print(f"[judicial-messages] scan failed: {exc}")
            with get_db() as conn: interval = int(conn.execute("SELECT poll_interval_seconds FROM case_manager_message_settings WHERE id=1").fetchone()[0] or 300)
            await asyncio.sleep(max(60,min(interval,3600)))

    @app.on_event("startup")
    async def start_message_scheduler():
        from .platform_capabilities import enabled
        if not enabled('messages_sync'): return
        if os.environ.get("JUDICIAL_SMS_SCHEDULER_DISABLED") == "1": return
        if scheduler["task"] is None or scheduler["task"].done(): scheduler["task"] = asyncio.create_task(scheduler_loop())

    @app.on_event("shutdown")
    async def stop_message_scheduler():
        if scheduler["task"]: scheduler["task"].cancel(); scheduler["task"] = None
