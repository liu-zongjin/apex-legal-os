"""Readable chronology table with indexed original citations, without changing event text."""
from .citation_locations import display_locator
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _text(paragraph, text, size=10.5, bold=False, color='202327'):
    run=paragraph.add_run(text)
    run.font.name='Arial'
    run.font.size=Pt(size)
    run.font.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)
    fonts=run._element.get_or_add_rPr().rFonts
    fonts.set(qn('w:eastAsia'),'SimSun')
    paragraph.paragraph_format.space_after=Pt(3)
    paragraph.paragraph_format.line_spacing=1.18
    return paragraph


def _cell_paragraph(cell, text, first=False, **kwargs):
    paragraph=cell.paragraphs[0] if first else cell.add_paragraph()
    return _text(paragraph,text,**kwargs)


def write_document(events, destination, review_labels=None, export_options=None):
    from .timeline_presentation import presentation,confirmed_position
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from docx.text.paragraph import Paragraph
    from lxml import etree
    rows,references,settings=presentation(events,export_options)
    doc=Document();section=doc.sections[0]
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Mm(18)
    section.footer_distance=Mm(9)
    for name,size in [('Normal',10.5),('Title',21),('Heading 1',14),('Heading 2',11)]:
        style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(6)
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    _text(doc.add_paragraph(style='Title'),'案件事实经过表' if settings['style']=='table' else '案件事实时间轴',size=21)
    _text(doc.add_paragraph(),f'共 {len(rows)} 项事件    日期保留原记录精度，事件间距不表示时长。',size=9,color='606770')
    cited=set()
    def cite(paragraph,numbers):
        for offset,number in enumerate(numbers):
            if offset:paragraph.add_run(',').font.superscript=True
            if number not in cited:
                start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(number));start.set(qn('w:name'),'timeline_note_'+str(number));paragraph._p.append(start)
                run=paragraph.add_run();run.font.superscript=True
                note=OxmlElement('w:endnoteReference');note.set(qn('w:id'),str(number));run._r.append(note)
                end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(number));paragraph._p.append(end);cited.add(number)
            else:
                field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),' NOTEREF timeline_note_'+str(number)+r' \h ')
                run=OxmlElement('w:r');prop=OxmlElement('w:rPr');vert=OxmlElement('w:vertAlign');vert.set(qn('w:val'),'superscript');prop.append(vert);run.append(prop)
                text=OxmlElement('w:t');text.text=str(number);run.append(text);field.append(run);paragraph._p.append(field)
    def content(cell,event):
        paragraph=_cell_paragraph(cell,event['description'],first=True);cite(paragraph,event['note_numbers'])
        if event['actors']:_cell_paragraph(cell,'主体：'+event['actors'],size=9,color='606770')
    def table_settings(table,widths,border='insideH'):
        table.autofit=False
        for column,width in zip(table.columns,widths):column.width=Mm(width)
        props=table._tbl.tblPr;edges=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single' if edge==border else 'nil');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'DDE1E6');edges.append(item)
        props.append(edges)
        margins=OxmlElement('w:tblCellMar')
        for edge,value in [('top',100),('bottom',130),('left',100),('right',100)]:
            item=OxmlElement('w:'+edge);item.set(qn('w:w'),str(value));item.set(qn('w:type'),'dxa');margins.append(item)
        props.append(margins)
    style=settings['style']
    if style in {'table','axis'}:
        widths=[11,37,126] if style=='table' else [39,135]
        table=doc.add_table(rows=1 if style=='table' else 0,cols=len(widths));table_settings(table,widths,'insideH' if style=='table' else 'insideV')
        if style=='table':
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for cell,width,label in zip(table.rows[0].cells,widths,['序号','日期','事实经过']):
                cell.width=Mm(width);_cell_paragraph(cell,label,first=True,size=9,bold=True)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F4F7');cell._tc.get_or_add_tcPr().append(shade)
        for index,event in enumerate(rows,1):
            cells=table.add_row().cells
            for cell,width in zip(cells,widths):cell.width=Mm(width)
            if style=='table':_cell_paragraph(cells[0],str(index),first=True,size=9,color='606770')
            _cell_paragraph(cells[-2],event['date_text'] or '日期不明',first=True,size=10,bold=True)
            content(cells[-1],event)
            if sum(len(p.text) for cell in cells for p in cell.paragraphs)<450:
                keep=OxmlElement('w:cantSplit');table.rows[-1]._tr.get_or_add_trPr().append(keep)
    else:
        for event in rows:
            table=doc.add_table(rows=1,cols=1);table_settings(table,[174],'bottom')
            cell=table.cell(0,0);cell.width=Mm(174)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F6F7F9');cell._tc.get_or_add_tcPr().append(shade)
            heading=_cell_paragraph(cell,event['date_text'] or '日期不明',first=True,size=12,bold=True);heading.paragraph_format.keep_with_next=True
            paragraph=_cell_paragraph(cell,event['description']);cite(paragraph,event['note_numbers'])
            if event['actors']:_cell_paragraph(cell,'主体：'+event['actors'],size=9,color='606770')
            if len(event['description'])+len(event['actors'])<450:
                keep=OxmlElement('w:cantSplit');table.rows[0]._tr.get_or_add_trPr().append(keep)
            spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(4);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=Pt(2)
    if not rows:_text(doc.add_paragraph(),'没有纳入导出的事件。')
    if references:
        root=OxmlElement('w:endnotes')
        for identifier,kind,tag in [(-1,'separator','separator'),(0,'continuationSeparator','continuationSeparator')]:
            endnote=OxmlElement('w:endnote');endnote.set(qn('w:id'),str(identifier));endnote.set(qn('w:type'),kind)
            p=OxmlElement('w:p');r=OxmlElement('w:r');r.append(OxmlElement('w:'+tag));p.append(r);endnote.append(p);root.append(endnote)
        for number,ref in enumerate(references,1):
            endnote=OxmlElement('w:endnote');endnote.set(qn('w:id'),str(number))
            p=OxmlElement('w:p');paragraph=Paragraph(p,doc._body)
            run=paragraph.add_run();run.font.superscript=True;run._r.append(OxmlElement('w:endnoteRef'))
            location=confirmed_position(ref)
            _text(paragraph,' '+ref['path']+(' · '+location if location else '')+'\n'+ref['quote'],size=9)
            endnote.append(p);root.append(endnote)
        part=Part(PackURI('/word/endnotes.xml'),'application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml',etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True),doc.part.package)
        doc.part.relate_to(part,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes')
        endnote_properties=OxmlElement('w:endnotePr');position=OxmlElement('w:pos');position.set(qn('w:val'),'docEnd');endnote_properties.append(position);number_format=OxmlElement('w:numFmt');number_format.set(qn('w:val'),'decimal');endnote_properties.append(number_format);section._sectPr.insert_element_before(endnote_properties,'w:type','w:pgSz','w:pgMar','w:cols','w:docGrid')
        defaults=doc.settings.element.find(qn('w:endnotePr'))
        if defaults is not None:
            existing=defaults.find(qn('w:numFmt'))
            if existing is not None:existing.set(qn('w:val'),'decimal')
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;_text(footer,'第 ',size=8,color='606770')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);_text(footer,' 页',size=8,color='606770')
    doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.save(destination)
