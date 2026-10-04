"""Local case archive workbench.

The criminal checklist and classification approach are adapted from
FachuanHybridSystem (Elastic License 2.0); adapted for Stella on 2026-09-30.
See third_party/fachuan-archive-LICENSE.txt. Source files are never modified.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import uuid
import zipfile
from pathlib import Path
from typing import Any, Callable

from fastapi import HTTPException, Request
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse

from .case_files import _bound_root, _resolve_inside


OUTPUT_DIR = "APEX案卷归档"
SUPPORTED = {".pdf", ".png", ".jpg", ".jpeg", ".doc", ".docx"}
INTERNAL_NAMES = ("结案提成表", "卷内目录", ".ds_store")
COMPILED_NAMES = ("合并打印", "结案材料")

LITIGATION = [
    "委托代理协议、终止协议", "授权委托书、所函", "收费凭证", "阅卷笔录",
    "会见当事人谈话笔录", "调查材料（证人证言、书证）", "起诉状、答辩状、上诉状",
    "诉讼保全申请书、证据保全申请书、先行给付申请书和法院裁定书",
    "承办律师意见", "集体讨论记录", "代理词", "出庭通知书", "庭审笔录",
    "裁定书、判决书、调解书、上诉书", "办案小结", "其他",
]
NON_LITIGATION = [
    "委托代理合同", "收费凭证", "与委托人谈话笔录", "委托人提供的证据材料",
    "调查材料", "律师出具的法律意见、或草拟的法律文书，办理具体法律事务活动的记录等",
    "工作小结", "其他",
]
CRIMINAL = [
    "案卷封面", "结案归档登记表", "案卷目录", "委托代理合同、风险告知书",
    "授权委托证明材料", "收费凭证", "律师办案工作日记",
    "会见被告人、委托人、证人笔录", "调查材料等案件关联材料",
    "承办人提出的辩护或代理意见", "集体讨论记录", "起诉书、上诉书",
    "辩护词或代理词", "出庭通知书", "裁定书、判决书", "上诉书、抗诉书",
    "办案服务质量监督卡", "办案小结",
]
DEFAULTS = {"litigation": LITIGATION, "non_litigation": NON_LITIGATION, "criminal": CRIMINAL}
TEMPLATE_DIR = Path(__file__).resolve().parent / "archive_templates"
AUTO_SUMMARY = {"litigation": 15, "non_litigation": 7}
DEFAULT_NA = {"litigation": {4, 5, 8, 9, 10, 11, 12, 13, 16},
              "non_litigation": {3, 4, 5, 8}}

KEYWORDS = {
    "litigation": {
        1: ("委托代理协议", "委托合同", "四方协议", "补充协议", "终止协议"),
        2: ("授权委托", "授权书", "所函"), 3: ("发票", "收费凭证", "缴费"),
        4: ("阅卷笔录",), 5: ("谈话笔录", "会见笔录"),
        6: ("证据", "调查材料", "证人证言"),
        7: ("起诉状", "起诉书", "答辩状", "上诉状", "诉状"),
        8: ("保全申请", "保全裁定"), 9: ("律师意见", "代理意见"),
        10: ("集体讨论",), 11: ("代理词",), 12: ("出庭通知",),
        13: ("庭审笔录",), 14: ("判决书", "裁定书", "调解书", "撤诉"),
        15: ("办案小结", "工作小结"),
    },
    "non_litigation": {
        1: ("委托合同", "委托代理合同", "委托代理协议"),
        2: ("发票", "收费凭证"), 3: ("谈话笔录",),
        4: ("委托人提供", "客户材料"), 5: ("调查材料",),
        6: ("法律意见", "律师函", "情况说明", "通知书", "领受物品清单"),
        7: ("办案小结", "工作小结"),
    },
    "criminal": {
        4: ("委托合同", "风险告知"), 5: ("授权委托", "授权书"),
        6: ("发票", "收费凭证"), 7: ("工作日记",),
        8: ("会见笔录", "谈话笔录"), 9: ("调查材料", "证据"),
        10: ("辩护意见", "代理意见"), 11: ("集体讨论",),
        12: ("起诉书",), 13: ("辩护词", "代理词"),
        14: ("出庭通知",), 15: ("判决书", "裁定书"),
        16: ("上诉书", "抗诉书"), 17: ("监督卡",), 18: ("办案小结",),
    },
}


def _stamp() -> str:
    return dt.datetime.now().astimezone().isoformat(timespec="seconds")


def _category(case_type: str) -> str:
    value = (case_type or "").lower()
    if any(word in value for word in ("刑", "criminal")):
        return "criminal"
    if any(word in value for word in ("非诉", "顾问", "专项", "special", "advisor")):
        return "non_litigation"
    return "litigation"


def _default_checklist(category: str) -> list[dict[str, Any]]:
    return [{"code": f"{category[:2]}_{index}", "name": name,
             "applicability": "na" if index in DEFAULT_NA.get(category, set()) else "review"}
            for index, name in enumerate(DEFAULTS[category], 1)]


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _keyword(name: str) -> str:
    stem = Path(name).stem.lower()
    return re.sub(r"^[#\s\d.、_-]+", "", stem).strip()[:100]


def _classify(relative: str, category: str, learned: dict[str, tuple[str, int, int]]) -> tuple[str, str]:
    name = Path(relative).name
    key = _keyword(name)
    if key in learned:
        code, _hits, conflict = learned[key]
        if conflict:
            return "", "ambiguous"
        return code, "learned"
    haystack = relative.lower()
    matches = []
    for index, words in KEYWORDS[category].items():
        if any(word in haystack for word in words):
            matches.append(f"{category[:2]}_{index}")
    if len(matches) == 1:
        return matches[0], "suggested"
    return "", "ambiguous" if matches else "unmatched"


def _case(get_db: Callable[[], Any], case_id: int) -> dict[str, Any]:
    with get_db() as conn:
        available = {column[1] for column in conn.execute("PRAGMA table_info(case_manager_cases)")}
        columns = [name for name in ("id", "title", "case_type", "source_status", "case_kind",
                                    "series_code", "series_display_name", "cause", "project_code",
                                    "accepted_at", "closed_at") if name in available]
        row = conn.execute(f"SELECT {','.join(columns)} FROM case_manager_cases WHERE id=?", (case_id,)).fetchone()
        clients = []
        if row and conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='case_manager_parties'").fetchone():
            clients = [entry[0] for entry in conn.execute(
                "SELECT name FROM case_manager_parties WHERE case_id=? AND party_group='client' ORDER BY sort_order,id", (case_id,))]
    if not row or row["source_status"] == "deleted":
        raise HTTPException(status_code=404, detail="案件不存在")
    case = dict(row)
    display = case.get("series_display_name") if case.get("case_kind") == "series_child" else case["title"]
    if case.get("case_kind") == "series_child" and (not display or display == case.get("series_code") or display.startswith("新子案")):
        display = case["title"].rsplit("·", 1)[-1]
    case["display_name"] = display or case["title"]
    case["client_name"] = "、".join(dict.fromkeys(name for name in clients if name))
    return case


def _project(get_db: Callable[[], Any], case_id: int, case_type: str) -> tuple[str, list[dict[str, Any]]]:
    with get_db() as conn:
        row = conn.execute("SELECT category,checklist_json FROM case_archive_projects WHERE case_id=?", (case_id,)).fetchone()
    if row:
        return row["category"], json.loads(row["checklist_json"])
    category = _category(case_type)
    return category, _default_checklist(category)


def _walk(root: Path) -> list[Path]:
    results: list[Path] = []
    for directory, dirs, files in os.walk(root, followlinks=False):
        dirs[:] = sorted(d for d in dirs if d != OUTPUT_DIR and not d.startswith(".")
                         and not (Path(directory) / d).is_symlink())
        for name in sorted(files):
            source = Path(directory) / name
            if source.is_symlink() or name.startswith((".", "~$")) or source.suffix.lower() not in SUPPORTED:
                continue
            if any(part in name.lower() for part in INTERNAL_NAMES) or "办案小结" in name or "工作小结" in name:
                continue
            results.append(source)
            if len(results) > 3000:
                raise HTTPException(status_code=409, detail="案件文件超过 3000 个，请先缩小绑定文件夹范围")
    return results


def _workspace(get_db: Callable[[], Any], case_id: int, *, scan: bool = False) -> dict[str, Any]:
    case = _case(get_db, case_id)
    category, checklist = _project(get_db, case_id, case["case_type"])
    try:
        root = _bound_root(get_db, case_id)
    except HTTPException as exc:
        if exc.status_code == 409:
            return {"case": case, "category": category, "checklist": checklist, "bound": False,
                    "materials": [], "exports": [], "saves": [], "message": exc.detail}
        raise
    with get_db() as conn:
        stored = {row["relative_path"]: dict(row) for row in conn.execute(
            "SELECT * FROM case_archive_materials WHERE case_id=?", (case_id,))}
        learned = {row["keyword"]: (row["item_code"], row["confirmations"], row["conflicted"])
                   for row in conn.execute("SELECT * FROM case_archive_rules WHERE category=?", (category,))}
        exports = [dict(row) for row in conn.execute(
            "SELECT id,category,relative_directory,created_at FROM case_archive_exports WHERE case_id=? ORDER BY id DESC LIMIT 30", (case_id,))]
        saves = [dict(row) for row in conn.execute(
            "SELECT id,category,created_at,manifest_json FROM case_archive_saves WHERE case_id=? ORDER BY id DESC LIMIT 30", (case_id,))]
    for saved in saves:
        manifest = json.loads(saved.pop("manifest_json"))
        saved["material_count"] = len(manifest.get("materials", []))
    for export in exports:
        export["summary_available"] = (root / export["relative_directory"] / "办案小结.docx").is_file()
    seen_hashes: dict[str, str] = {}
    materials: list[dict[str, Any]] = []
    updates: list[tuple[Any, ...]] = []
    for source in _walk(root):
        relative = source.relative_to(root).as_posix()
        prior = stored.pop(relative, None)
        stat = source.stat()
        fingerprint = (prior["fingerprint"] if prior and not scan and
                       prior["size_bytes"] == stat.st_size and prior["mtime_ns"] == stat.st_mtime_ns
                       else _file_hash(source))
        duplicate_of = seen_hashes.get(fingerprint, "")
        seen_hashes.setdefault(fingerprint, relative)
        aggregate = any(word in source.stem for word in COMPILED_NAMES)
        template = "模板" in source.stem or "空白" in source.stem
        suggested_code, suggestion = _classify(relative, category, learned)
        if prior:
            changed = prior["fingerprint"] != fingerprint
            state = "changed" if changed and prior["review_state"] == "confirmed" else prior["review_state"]
            code = prior["item_code"]
            position = prior["position"]
        else:
            state = "template" if template else "review" if aggregate or duplicate_of else suggestion
            code = "" if aggregate or duplicate_of else suggested_code
            position = 0
            changed = False
        item = {"path": relative, "name": source.name, "fingerprint": fingerprint,
                "size": stat.st_size, "item_code": code, "position": position,
                "state": state, "duplicate_of": duplicate_of, "aggregate": aggregate,
                "changed": changed, "suggestion": suggestion}
        materials.append(item)
        if scan and (not prior or changed):
            updates.append((case_id, relative, fingerprint, stat.st_size, stat.st_mtime_ns, code, position, state, _stamp()))
    for relative, prior in stored.items():
        materials.append({"path": relative, "name": Path(relative).name, "fingerprint": prior["fingerprint"],
                          "item_code": prior["item_code"], "position": prior["position"],
                          "state": "missing", "duplicate_of": "", "aggregate": False, "changed": False})
    if scan:
        with get_db() as conn:
            conn.execute("INSERT OR IGNORE INTO case_archive_projects(case_id,category,checklist_json,updated_at) VALUES(?,?,?,?)",
                         (case_id, category, json.dumps(checklist, ensure_ascii=False), _stamp()))
            conn.executemany("""INSERT INTO case_archive_materials
                (case_id,relative_path,fingerprint,size_bytes,mtime_ns,item_code,position,review_state,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(case_id,relative_path) DO UPDATE SET
                fingerprint=excluded.fingerprint,item_code=excluded.item_code,
                size_bytes=excluded.size_bytes,mtime_ns=excluded.mtime_ns,
                position=excluded.position,review_state=excluded.review_state,updated_at=excluded.updated_at""", updates)
            conn.commit()
    counts = {item["code"]: sum(1 for material in materials if material["item_code"] == item["code"]
                                 and material["state"] == "confirmed") for item in checklist}
    return {"case": case, "category": category, "checklist": checklist, "bound": True,
            "materials": materials, "counts": counts, "exports": exports, "saves": saves}


def _convert_word(source: Path, output_dir: Path) -> Path:
    if source.suffix.lower() == ".docx":
        if not zipfile.is_zipfile(source):
            raise RuntimeError(f"不是有效的 Word 文档：{source.name}")
        with zipfile.ZipFile(source) as package:
            if "word/document.xml" not in package.namelist():
                raise RuntimeError(f"不是有效的 Word 文档：{source.name}")
    else:
        with source.open("rb") as word_file:
            if word_file.read(8) != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
                raise RuntimeError(f"不是有效的旧版 Word 文档：{source.name}")
    soffice = os.environ.get("LEGAL_KB_SOFFICE") or shutil.which("soffice")
    if not soffice:
        app = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        soffice = str(app) if app.is_file() else ""
    if not soffice:
        raise RuntimeError(f"未安装 LibreOffice，无法转换：{source.name}")
    with tempfile.TemporaryDirectory(prefix="apex-archive-lo-") as profile:
        environment = os.environ.copy()
        if sys.platform == "darwin":
            font_config = Path(profile) / "fonts.conf"
            font_config.write_text('<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd">'
                                   '<fontconfig><dir>/System/Library/Fonts</dir><dir>/Library/Fonts</dir>'
                                   f'<cachedir>{profile}/font-cache</cachedir></fontconfig>', encoding="utf-8")
            environment["FONTCONFIG_FILE"] = str(font_config)
        result = subprocess.run([soffice, f"-env:UserInstallation=file://{profile}", "--headless",
                                 "--convert-to", "pdf", "--outdir", str(output_dir), str(source)],
                                capture_output=True, text=True, timeout=120, check=False, env=environment)
    converted = output_dir / (source.stem + ".pdf")
    if result.returncode or not converted.is_file():
        raise RuntimeError(f"Word 转换失败：{source.name}；{(result.stderr or result.stdout)[-180:]}")
    return converted


def _set_cell(cell: Any, value: str) -> None:
    """Change text in a firm template without replacing its cell or paragraph styling."""
    first = cell.paragraphs[0]
    if first.runs:
        first.runs[0].text = value
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.add_run(value)
    for paragraph in cell.paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


def _readable_template_fonts(doc: Any) -> None:
    """WPS samples carry fonts that LibreOffice renders as missing-glyph boxes."""
    from docx.oxml.ns import qn
    for paragraph in [*doc.paragraphs, *(p for table in doc.tables for row in table.rows
                                         for cell in row.cells for p in cell.paragraphs)]:
        for run in paragraph.runs:
            run.font.name = "SimSun"
            rpr = run._element.get_or_add_rPr()
            fonts = rpr.rFonts
            if fonts is not None:
                fonts.set(qn("w:eastAsia"), "SimSun")


def _summary_text(category: str, case: dict[str, Any], entries: list[dict[str, Any]],
                  checklist: list[dict[str, Any]]) -> str:
    names = {item["code"]: item["name"] for item in checklist}
    kinds = list(dict.fromkeys(names[e["item_code"]] for e in entries if e["item_code"] in names))
    documents = list(dict.fromkeys(re.sub(r"^[\d\s._、：:-]+", "", Path(e["path"]).stem)[:24]
                                   for e in entries if e.get("path") and not e.get("generated")))
    documents = [name for name in documents if name][:3]
    examples = "、".join(f"《{name}》" for name in documents)
    scope = examples or "、".join(kinds[:3])
    if category == "non_litigation":
        return "本所律师接受委托后，围绕本项目整理并核对了" + scope + "等材料，办理了相关非诉法律事务。具体工作成果以本卷所附材料为准。"
    return "本所律师接受委托后，研究了" + scope + "等材料，办理了相关诉讼或仲裁代理工作。案件程序及结果以本卷所附文书为准。"


def _make_summary(destination: Path, category: str, case: dict[str, Any],
                  entries: list[dict[str, Any]], checklist: list[dict[str, Any]]) -> Path | None:
    if category not in AUTO_SUMMARY:
        return None
    from docx import Document
    template = TEMPLATE_DIR / f"{category}_summary.docx"
    doc = Document(template)
    table = doc.tables[0]
    _set_cell(table.cell(0, 2), str(case.get("project_code") or ""))
    _set_cell(table.cell(0, 6), str(case.get("client_name") or ""))
    _set_cell(table.cell(1, 2), str(case.get("cause") or case.get("display_name") or ""))
    _set_cell(table.cell(2, 2), str(case.get("accepted_at") or ""))
    _set_cell(table.cell(2, 5), str(case.get("closed_at") or ""))
    _set_cell(table.cell(3, 1), _summary_text(category, case, entries, checklist))
    if category == "non_litigation":
        _set_cell(table.cell(2, 8), "")
        for paragraph in doc.paragraphs:
            if "诉讼类" in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace("诉讼类", "非诉类")
            if "结案日期：" in paragraph.text:
                for run in paragraph.runs:
                    run.text = ""
    _readable_template_fonts(doc)
    output = destination / "办案小结.docx"
    doc.save(output)
    return output


def _make_catalog(destination: Path, category: str, checklist: list[dict[str, Any]],
                  entries: list[dict[str, Any]], pages: int) -> None:
    from docx import Document
    if category in {"litigation", "non_litigation"}:
        doc = Document(TEMPLATE_DIR / f"{category}_catalog.docx")
        if category == "litigation":
            from docx.shared import Cm
            doc.sections[0].top_margin = Cm(2.0)
            doc.sections[0].bottom_margin = Cm(1.2)
        table = doc.tables[0]
        for index, item in enumerate(checklist):
            row = table.rows[index + 1] if index + 1 < len(table.rows) else table.add_row()
            found = [entry for entry in entries if entry["item_code"] == item["code"]]
            values = (str(index + 1), item["name"], "", str(len(found)) if found else "",
                      "、".join(f"{e['start']}-{e['end']}" if e["start"] != e["end"] else str(e["start"]) for e in found),
                      "不适用" if item.get("applicability") == "na" else "")
            for cell, value in zip(row.cells, values):
                _set_cell(cell, value)
        for index in range(len(table.rows) - 1, len(checklist), -1):
            table._tbl.remove(table.rows[index]._tr)
        _set_cell(doc.tables[1].cell(0, 0), f"本卷宗共 {pages} 页              附实物材料（如光盘等）    /   袋")
    else:
        doc = Document()
        doc.add_paragraph("卷内目录（刑事卷）")
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, ("序号", "材料名称", "原件(份)", "复印件(份)", "页码", "备注")):
            cell.text = value
        for index, item in enumerate(checklist, 1):
            found = [entry for entry in entries if entry["item_code"] == item["code"]]
            row = table.add_row().cells
            for cell, value in zip(row, (str(index), item["name"], "", str(len(found)) if found else "",
                                           "、".join(f"{e['start']}-{e['end']}" for e in found), "")):
                cell.text = value
        doc.add_paragraph(f"本卷宗共 {pages} 页")
    _readable_template_fonts(doc)
    doc.save(destination / "卷内目录.docx")


def _compile(root: Path, case: dict[str, Any], category: str, checklist: list[dict[str, Any]],
             materials: list[dict[str, Any]], destination: Path) -> dict[str, Any]:
    import fitz

    selected = [m for m in materials if m["state"] == "confirmed"]
    order = {item["code"]: index for index, item in enumerate(checklist)}
    selected.sort(key=lambda m: (order.get(m["item_code"], 9999), m["position"], m["path"]))
    if not selected:
        raise HTTPException(status_code=409, detail="请先确认至少一份入卷材料")
    if any(m["item_code"] not in order for m in selected):
        raise HTTPException(status_code=409, detail="入卷材料关联的清单项目已被删除，请重新分类")
    merged = fitz.open()
    entries: list[dict[str, Any]] = []
    work = destination / "转换缓存"
    work.mkdir()
    try:
        for material in selected:
            source = _resolve_inside(root, material["path"], allow_root=False)
            if source.is_symlink() or not source.is_file() or _file_hash(source) != material["fingerprint"]:
                raise HTTPException(status_code=409, detail=f"源文件已变化，请重新扫描：{material['name']}")
            ext = source.suffix.lower()
            try:
                if ext in {".doc", ".docx"}:
                    word_work = work / uuid.uuid4().hex
                    word_work.mkdir()
                    converted = _convert_word(source, word_work)
                    document = fitz.open(str(converted))
                else:
                    document = fitz.open(str(source))
            except Exception as exc:
                raise HTTPException(status_code=409, detail=f"无法读取或转换材料：{material['name']}；{str(exc)[:150]}") from exc
            try:
                if document.needs_pass:
                    raise HTTPException(status_code=409, detail=f"文件有密码，无法入卷：{material['name']}")
                start = len(merged) + 1
                try:
                    if ext in {".png", ".jpg", ".jpeg"}:
                        page = merged.new_page(width=document[0].rect.width, height=document[0].rect.height)
                        page.insert_image(page.rect, filename=str(source))
                    else:
                        merged.insert_pdf(document)
                except Exception as exc:
                    raise HTTPException(status_code=409, detail=f"无法合并材料：{material['name']}；{str(exc)[:150]}") from exc
                end = len(merged)
                if end < start:
                    raise HTTPException(status_code=409, detail=f"空文件不能入卷：{material['name']}")
                entries.append({"item_code": material["item_code"], "path": material["path"],
                                "sha256": material["fingerprint"], "start": start, "end": end})
            finally:
                document.close()
        summary_path = _make_summary(destination, category, case, entries, checklist)
        if summary_path:
            word_work = work / "summary"
            word_work.mkdir()
            converted = _convert_word(summary_path, word_work)
            summary_pdf = fitz.open(str(converted))
            try:
                start = len(merged) + 1
                merged.insert_pdf(summary_pdf)
                entries.append({"item_code": f"{category[:2]}_{AUTO_SUMMARY[category]}",
                                "path": "办案小结.docx", "generated": True,
                                "start": start, "end": len(merged)})
            finally:
                summary_pdf.close()
        case_label = case["display_name"][:28]
        for index, page in enumerate(merged, 1):
            page.insert_text((36, 24), f"{case_label} 结案归档材料  第{index}页",
                             fontname="china-s", fontsize=9, color=(0.22, 0.22, 0.22), overlay=True)
        pdf_path = destination / "结案材料.pdf"
        merged.save(str(pdf_path), garbage=3, deflate=True)
    finally:
        merged.close()
        shutil.rmtree(work, ignore_errors=True)

    _make_catalog(destination, category, checklist, entries, entries[-1]["end"])
    return {"pages": entries[-1]["end"], "entries": entries}


def register_case_archive_routes(app: Any, get_db: Callable[[], Any], frontend_dir: Path,
                                 require_action: Callable[[Request], Any]) -> None:
    @app.get("/case-archive")
    def page():
        return HTMLResponse((frontend_dir / "case-archive.html").read_text(encoding="utf-8"),
                            headers={"Cache-Control": "no-store"})

    @app.get("/api/case-archive/cases/{case_id}")
    def workspace(case_id: int):
        return _workspace(get_db, case_id)

    @app.get("/api/case-archive/templates/{category}")
    def template(category: str):
        if category not in DEFAULTS:
            raise HTTPException(status_code=404, detail="没有对应的案卷清单")
        return {"category": category, "items": _default_checklist(category)}

    @app.post("/api/case-archive/cases/{case_id}/scan")
    def scan(case_id: int, request: Request):
        require_action(request)
        return _workspace(get_db, case_id, scan=True)

    @app.put("/api/case-archive/cases/{case_id}/checklist")
    async def save_checklist(case_id: int, request: Request):
        require_action(request)
        _case(get_db, case_id)
        body = await request.json()
        category = body.get("category")
        items = body.get("items")
        if category not in DEFAULTS or not isinstance(items, list) or not 1 <= len(items) <= 50:
            raise HTTPException(status_code=400, detail="请选择有效的案卷类别和清单")
        codes = [str(item.get("code", "")) for item in items]
        if len(set(codes)) != len(codes) or any(not re.fullmatch(r"[a-z0-9_]{2,50}", code) for code in codes):
            raise HTTPException(status_code=400, detail="清单编号无效或重复")
        cleaned = []
        for item in items:
            name = str(item.get("name", "")).strip()
            applicability = item.get("applicability", "review")
            if not name or len(name) > 120 or applicability not in {"review", "na"}:
                raise HTTPException(status_code=400, detail="清单名称或状态无效")
            cleaned.append({"code": item["code"], "name": name, "applicability": applicability})
        if category in AUTO_SUMMARY:
            generated = next((item for item in cleaned if item["code"] == f"{category[:2]}_{AUTO_SUMMARY[category]}"), None)
            if not generated or generated["name"] != DEFAULTS[category][AUTO_SUMMARY[category] - 1] or generated["applicability"] != "review":
                raise HTTPException(status_code=400, detail="自动生成的小结项目不能删除或设为不适用")
        with get_db() as conn:
            old = conn.execute("SELECT category FROM case_archive_projects WHERE case_id=?", (case_id,)).fetchone()
            conn.execute("""INSERT INTO case_archive_projects(case_id,category,checklist_json,updated_at)
                VALUES(?,?,?,?) ON CONFLICT(case_id) DO UPDATE SET category=excluded.category,
                checklist_json=excluded.checklist_json,updated_at=excluded.updated_at""",
                (case_id, category, json.dumps(cleaned, ensure_ascii=False), _stamp()))
            valid = {item["code"] for item in cleaned if item["applicability"] != "na"}
            if old and old["category"] != category:
                conn.execute("""UPDATE case_archive_materials SET item_code='',review_state='review',updated_at=?
                    WHERE case_id=? AND review_state='confirmed'""", (_stamp(), case_id))
            else:
                for row in conn.execute("SELECT relative_path,item_code FROM case_archive_materials WHERE case_id=? AND review_state='confirmed'", (case_id,)):
                    if row["item_code"] not in valid:
                        conn.execute("""UPDATE case_archive_materials SET item_code='',review_state='review',updated_at=?
                            WHERE case_id=? AND relative_path=?""", (_stamp(), case_id, row["relative_path"]))
            conn.commit()
        return _workspace(get_db, case_id)

    @app.put("/api/case-archive/cases/{case_id}/materials")
    async def save_material(case_id: int, request: Request):
        require_action(request)
        case = _case(get_db, case_id)
        category, checklist = _project(get_db, case_id, case["case_type"])
        body = await request.json()
        relative = str(body.get("path", ""))
        state = body.get("state")
        code = str(body.get("item_code", ""))
        valid_codes = {i["code"] for i in checklist if i["applicability"] != "na"}
        if state not in {"confirmed", "excluded", "review"} or (state == "confirmed" and
                (code not in valid_codes or code == f"{category[:2]}_{AUTO_SUMMARY.get(category, 0)}")):
            raise HTTPException(status_code=400, detail="请选择有效的归档类别和状态")
        root = _bound_root(get_db, case_id)
        source = _resolve_inside(root, relative, allow_root=False)
        if source.is_symlink() or not source.is_file() or source.suffix.lower() not in SUPPORTED:
            raise HTTPException(status_code=400, detail="材料不是可归档文件")
        fingerprint = _file_hash(source)
        stat = source.stat()
        position = max(0, min(int(body.get("position") or 0), 100000))
        with get_db() as conn:
            conn.execute("""INSERT INTO case_archive_materials
                (case_id,relative_path,fingerprint,size_bytes,mtime_ns,item_code,position,review_state,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(case_id,relative_path) DO UPDATE SET
                fingerprint=excluded.fingerprint,item_code=excluded.item_code,position=excluded.position,
                size_bytes=excluded.size_bytes,mtime_ns=excluded.mtime_ns,
                review_state=excluded.review_state,updated_at=excluded.updated_at""",
                (case_id, relative, fingerprint, stat.st_size, stat.st_mtime_ns,
                 code if state == "confirmed" else "", position, state, _stamp()))
            if state == "confirmed":
                key = _keyword(source.name)
                if key:
                    prior = conn.execute("SELECT item_code FROM case_archive_rules WHERE category=? AND keyword=?",
                                         (category, key)).fetchone()
                    if prior and prior["item_code"] != code:
                        conn.execute("UPDATE case_archive_rules SET conflicted=1 WHERE category=? AND keyword=?", (category, key))
                    else:
                        conn.execute("""INSERT INTO case_archive_rules(category,keyword,item_code)
                            VALUES(?,?,?) ON CONFLICT(category,keyword) DO UPDATE SET
                            confirmations=confirmations+1""", (category, key, code))
            conn.commit()
        return _workspace(get_db, case_id)

    @app.put("/api/case-archive/cases/{case_id}/materials/batch")
    async def save_materials_batch(case_id: int, request: Request):
        require_action(request)
        case = _case(get_db, case_id)
        category, checklist = _project(get_db, case_id, case["case_type"])
        body = await request.json()
        paths = body.get("paths")
        code = str(body.get("item_code", ""))
        valid = {item["code"] for item in checklist if item["applicability"] != "na"}
        if (not isinstance(paths, list) or not 1 <= len(paths) <= 500 or
                len(set(paths)) != len(paths) or code not in valid or
                code == f"{category[:2]}_{AUTO_SUMMARY.get(category, 0)}"):
            raise HTTPException(status_code=400, detail="请选择有效的材料和归档类别")
        root = _bound_root(get_db, case_id)
        rows = []
        for relative in paths:
            source = _resolve_inside(root, relative, allow_root=False)
            if source.is_symlink() or not source.is_file() or source.suffix.lower() not in SUPPORTED:
                raise HTTPException(status_code=400, detail="材料不是可归档文件")
            stat = source.stat()
            rows.append((case_id, relative, _file_hash(source), stat.st_size, stat.st_mtime_ns,
                         code, 0, "confirmed", _stamp()))
        with get_db() as conn:
            conn.executemany("""INSERT INTO case_archive_materials
                (case_id,relative_path,fingerprint,size_bytes,mtime_ns,item_code,position,review_state,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(case_id,relative_path) DO UPDATE SET
                fingerprint=excluded.fingerprint,item_code=excluded.item_code,
                size_bytes=excluded.size_bytes,mtime_ns=excluded.mtime_ns,
                review_state=excluded.review_state,updated_at=excluded.updated_at""", rows)
            conn.commit()
        return _workspace(get_db, case_id)

    @app.put("/api/case-archive/cases/{case_id}/order")
    async def save_order(case_id: int, request: Request):
        require_action(request)
        _case(get_db, case_id)
        body = await request.json()
        paths = body.get("paths")
        if not isinstance(paths, list) or len(paths) > 3000 or len(set(paths)) != len(paths):
            raise HTTPException(status_code=400, detail="材料顺序无效")
        root = _bound_root(get_db, case_id)
        for path in paths:
            source = _resolve_inside(root, path, allow_root=False)
            if not source.is_file() or source.is_symlink():
                raise HTTPException(status_code=400, detail="材料路径无效")
        with get_db() as conn:
            for position, path in enumerate(paths, 1):
                conn.execute("UPDATE case_archive_materials SET position=?,updated_at=? WHERE case_id=? AND relative_path=?",
                             (position, _stamp(), case_id, path))
            conn.commit()
        return _workspace(get_db, case_id)

    @app.post("/api/case-archive/cases/{case_id}/saves")
    def save_manifest(case_id: int, request: Request):
        require_action(request)
        state = _workspace(get_db, case_id)
        if not state["bound"]:
            raise HTTPException(status_code=409, detail=state["message"])
        manifest = {
            "case_id": case_id,
            "case_name": state["case"]["display_name"],
            "category": state["category"],
            "created_at": _stamp(),
            "checklist": state["checklist"],
            "materials": [{
                "path": item["path"], "name": item["name"], "state": item["state"],
                "item_code": item["item_code"], "position": item["position"],
                "fingerprint": item["fingerprint"],
            } for item in state["materials"]],
        }
        with get_db() as conn:
            cursor = conn.execute("""INSERT INTO case_archive_saves
                (case_id,category,manifest_json,created_at) VALUES(?,?,?,?)""",
                (case_id, state["category"], json.dumps(manifest, ensure_ascii=False), manifest["created_at"]))
            conn.commit()
            save_id = cursor.lastrowid
        return {"id": save_id, "created_at": manifest["created_at"],
                "material_count": sum(item["state"] == "confirmed" for item in manifest["materials"])}

    @app.get("/api/case-archive/cases/{case_id}/saves/{save_id}")
    def download_saved_manifest(case_id: int, save_id: int):
        _case(get_db, case_id)
        with get_db() as conn:
            row = conn.execute("SELECT manifest_json FROM case_archive_saves WHERE id=? AND case_id=?",
                               (save_id, case_id)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="归档清单保存记录不存在")
        return JSONResponse(
            json.loads(row["manifest_json"]),
            headers={"Content-Disposition": f'attachment; filename="archive-list-{save_id}.json"'},
        )

    @app.post("/api/case-archive/cases/{case_id}/exports")
    def create_export(case_id: int, request: Request):
        require_action(request)
        state = _workspace(get_db, case_id)
        if not state["bound"]:
            raise HTTPException(status_code=409, detail=state["message"])
        stale = [item["name"] for item in state["materials"] if item["state"] in {"changed", "missing"} and item["item_code"]]
        if stale:
            raise HTTPException(status_code=409, detail=f"已入卷材料发生变化，请重新核对：{stale[0]}")
        case = state["case"]
        root = _bound_root(get_db, case_id, writable=True)
        output_root = root / OUTPUT_DIR
        output_root.mkdir(exist_ok=True)
        version = dt.datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + uuid.uuid4().hex[:6]
        staging = Path(tempfile.mkdtemp(prefix=".apex-archive-", dir=output_root))
        try:
            result = _compile(root, case, state["category"], state["checklist"], state["materials"], staging)
            manifest = {"case_id": case_id, "category": state["category"], "created_at": _stamp(),
                        "pages": result["pages"], "entries": result["entries"],
                        "checklist": state["checklist"]}
            (staging / "材料清单.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
            final = output_root / version
            os.replace(staging, final)
            try:
                with get_db() as conn:
                    cursor = conn.execute("""INSERT INTO case_archive_exports
                        (case_id,category,relative_directory,manifest_json,created_at) VALUES(?,?,?,?,?)""",
                        (case_id, state["category"], final.relative_to(root).as_posix(),
                         json.dumps(manifest, ensure_ascii=False), manifest["created_at"]))
                    conn.commit()
                    export_id = cursor.lastrowid
            except Exception:
                shutil.rmtree(final, ignore_errors=True)
                raise
            return {"id": export_id, "pages": result["pages"], "directory": final.relative_to(root).as_posix()}
        except Exception:
            shutil.rmtree(staging, ignore_errors=True)
            raise

    @app.get("/api/case-archive/cases/{case_id}/exports/{export_id}/{kind}")
    def download(case_id: int, export_id: int, kind: str):
        if kind not in {"pdf", "catalog", "summary", "manifest"}:
            raise HTTPException(status_code=404, detail="导出文件不存在")
        _case(get_db, case_id)
        with get_db() as conn:
            row = conn.execute("SELECT relative_directory FROM case_archive_exports WHERE id=? AND case_id=?",
                               (export_id, case_id)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="导出版本不存在")
        root = _bound_root(get_db, case_id)
        filename = {"pdf": "结案材料.pdf", "catalog": "卷内目录.docx",
                    "summary": "办案小结.docx", "manifest": "材料清单.json"}[kind]
        path = _resolve_inside(root, f"{row['relative_directory']}/{filename}", allow_root=False)
        if not path.is_file():
            raise HTTPException(status_code=404, detail="导出文件不存在")
        return FileResponse(path, filename=filename)
