from __future__ import annotations

import datetime as dt
import base64
from decimal import Decimal, InvalidOperation
import hashlib
import json
import mimetypes
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
import uuid
import zipfile
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable

from fastapi import File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse


# Judicial-message attachments are kept separate from other case material.  Keep
# this constant as the single source of truth for both new bindings and writes.
COURT_DOCUMENT_SUBDIR = "#APEX电子送达归档"
LEGACY_COURT_DOCUMENT_SUBDIR = "#法院文书"
CONTRACT_SUBDIR = "#委托合同及授权"
DATA_DIR = Path(os.environ.get("LEGAL_KB_DATA_DIR", str(Path(__file__).resolve().parent)))
TEMPLATE_DIR = DATA_DIR / "contract_templates"
PREVIEW_DIR = DATA_DIR / "contract_previews"
MAX_TEMPLATE_BYTES = 20 * 1024 * 1024

DEFAULT_FIRM = {
    "name": "示例律师事务所（虚构）",
    "address": "示例地址（虚构）",
    "credit_code": "990000199001010110",
    "bank": "示例地址（虚构）",
    "account": "0000 0000 0000 0000 00",
}

AUTHORIZATION_OPTIONS = [
    {"id": "investigate", "label": "调查、收集证据", "keyword": "调查取证"},
    {"id": "file_case", "label": "提起诉讼或仲裁", "keyword": "提起诉讼"},
    {"id": "respond", "label": "应诉、答辩", "keyword": "应诉答辩"},
    {"id": "change_claims", "label": "变更、放弃诉讼请求", "keyword": "变更或者放弃诉讼请求"},
    {"id": "admit_claims", "label": "承认对方诉讼请求", "keyword": "承认诉讼请求"},
    {"id": "counterclaim", "label": "提起反诉", "keyword": "提起反诉"},
    {"id": "settle", "label": "参与调解、和解并签署相关文书", "keyword": "调解或者和解"},
    {"id": "appeal", "label": "提起上诉", "keyword": "提起上诉"},
    {"id": "enforce", "label": "申请执行及参与执行程序", "keyword": "申请执行"},
    {"id": "documents", "label": "签署、送达、接受法律文书", "keyword": "签署、送达、接受法律文书"},
]

ENTRUST_TYPE_DEFINITIONS = [
    {"id": "plaintiff", "label": "一审原告", "stage": "一审",
     "matter": "在本案一审程序中担任原告的委托代理人",
     "authority_lines": [
         "代为申请财产保全；", "代为起诉、出庭；", "代为收集、提供证据，进行辩论；",
         "代为查阅本案有关材料；", "代为提出回避申请、管辖权异议申请；",
         "代为签署、递交、接收法律文书；", "代为上诉，进行和解，请求调解，提起反诉。",
     ]},
    {"id": "defendant", "label": "一审被告", "stage": "一审",
     "matter": "在本案一审程序中担任被告的委托代理人",
     "authority_lines": [
         "代为应诉、答辩、出庭；", "代为收集、提供证据，进行辩论；",
         "代为查阅本案有关材料；", "代为提出回避申请、管辖权异议申请；",
         "代为签署、递交、接收法律文书；", "代为提起反诉或上诉；",
         "代为进行和解，请求调解。",
     ]},
    {"id": "appellant", "label": "二审上诉人", "stage": "二审",
     "matter": "在本案二审程序中担任上诉人的委托代理人",
     "authority_lines": [
         "代为上诉、答辩、出庭；", "代为收集、提供证据，进行辩论；",
         "代为查阅本案有关材料；", "代为提出回避申请、管辖权异议申请；",
         "代为签署、递交、接收法律文书；",
         "代为承认、变更、放弃上诉请求，代为进行和解，请求调解。",
     ]},
    {"id": "appellee", "label": "二审被上诉人", "stage": "二审",
     "matter": "在本案二审程序中担任被上诉人的委托代理人",
     "authority_lines": [
         "代为答辩、出庭；", "代为收集、提供证据，进行辩论；",
         "代为查阅本案有关材料；", "代为提出回避申请、管辖权异议申请；",
         "代为签署、递交、接收法律文书；", "代为进行和解，请求调解。",
     ]},
    {"id": "arbApplicant", "label": "仲裁申请人", "stage": "仲裁",
     "matter": "在本案仲裁程序中担任申请人的委托代理人",
     "authority_lines": [
         "代为提起仲裁、出庭；", "代为选择仲裁员；",
         "代为收集、提供证据，进行辩论；", "代为查阅本案有关材料；",
         "代为提出回避申请、主管权异议申请；", "代为签署、递交、接收法律文书；",
         "代为和解，请求调解。",
     ]},
    {"id": "arbRespondent", "label": "仲裁被申请人", "stage": "仲裁",
     "matter": "在本案仲裁程序中担任被申请人的委托代理人",
     "authority_lines": [
         "代为应诉（裁）、出庭；", "代为选择仲裁员；",
         "代为收集、提供证据，进行辩论；", "代为查阅本案有关材料；",
         "代为提出回避申请；", "代为签署、递交、接收法律文书；",
         "代为和解，请求调解。",
     ]},
    {"id": "execApplicant", "label": "申请执行人", "stage": "强制执行",
     "matter": "在本案强制执行程序中担任申请执行人的委托代理人",
     "authority_lines": [
         "代为申请强制执行、参与执行程序、代为协助法院执行；",
         "代为调解或执行和解；", "代为签收相关法律文件；", "代为领取案款；",
         "代为投诉、提出申诉、控告；", "代为参与执行相关的活动；代为查阅、复制档案等。",
     ]},
    {"id": "execRespondent", "label": "被执行人", "stage": "强制执行",
     "matter": "在本案强制执行程序中担任被执行人的委托代理人",
     "authority_lines": [
         "代为调查取证；", "代为提出回避申请；", "代为参加执行和解、执行调解；",
         "代为提出执行异议及复议；", "代为提出申诉、投诉、举报、控告；",
         "代为签署、递交、接收法律文书。",
     ]},
    {"id": "execObjection", "label": "执行异议", "stage": "执行异议",
     "matter": "在本案执行异议程序中担任申请人的委托代理人",
     "authority_lines": [
         "代为提出执行异议申请（申请排除执行）；", "代为出庭接受询问；",
         "代为收集、提供证据，进行辩论；", "代为查阅本案有关材料；",
         "代为承认、变更、放弃异议申请；", "代为签署、递交、接收法律文书。",
     ]},
    {"id": "bankruptcy", "label": "破产清算", "stage": "破产清算",
     "matter": "在本案破产程序中担任债权人的委托代理人",
     "authority_lines": [
         "申报债权、对债权审核意见提出异议；", "进行调解、和解；",
         "参加债权人会议、行使表决权；",
         "审查/表决重整计划（草案）、财产处置方案、分配方案等各项方案议案，提出异议或意见；",
         "参与财产分配，领取分配所得；", "提交、签收有关法律文件；",
         "其他与破产有关的事宜。",
     ]},
]

FEE_PLAN_DEFINITIONS = [
    {"id": "fixed", "label": "固定收费", "description": "约定一个固定律师服务费"},
    {"id": "staged_fixed", "label": "分阶段固定收费", "description": "自由选择一个或多个代理阶段"},
    {"id": "base_risk", "label": "基础费＋责任减免风险费", "description": "基础费加责任减免结果收费"},
    {"id": "recovery_tiered", "label": "基础费＋分段回款风险费", "description": "自由设置一档或多档回款费率"},
]

RISK_METHOD_DEFINITIONS = [
    {"id": "recovery_rate", "label": "按实际回款比例", "description": "以实际收回款项为基数乘以约定比例"},
    {"id": "reduction_rate", "label": "按减免金额比例", "description": "以对方主张金额与最终承担金额之差为基数"},
    {"id": "success_fixed", "label": "达到特定结果收固定成功费", "description": "达到无需承担责任等明确结果时收取固定金额"},
    {"id": "recovery_tiered", "label": "按回款金额分段计费", "description": "根据实际回款落入的区间适用不同比例"},
]


def _authority_ids_from_lines(lines: Iterable[str]) -> list[str]:
    text = "\n".join(str(item) for item in lines)
    matches: list[str] = []
    keyword_groups = {
        "investigate": ("调查", "收集、提供证据", "查阅"),
        "file_case": ("起诉", "提起仲裁", "执行异议申请"),
        "respond": ("应诉", "答辩"),
        "change_claims": ("变更", "放弃"),
        "admit_claims": ("承认",),
        "counterclaim": ("反诉",),
        "settle": ("调解", "和解"),
        "appeal": ("上诉",),
        "enforce": ("强制执行", "执行程序"),
        "documents": ("法律文书", "法律文件", "签收"),
    }
    for item in AUTHORIZATION_OPTIONS:
        if any(keyword in text for keyword in keyword_groups[item["id"]]):
            matches.append(item["id"])
    return matches


def _chinese_year_month(value: dt.date) -> str:
    digits = "〇一二三四五六七八九"
    year = "".join(digits[int(char)] for char in str(value.year))
    month = (
        digits[value.month]
        if value.month < 10
        else ("十" if value.month == 10 else f"十{digits[value.month - 10]}")
    )
    return f"{year}年{month}月"


SEMANTIC_FIELDS: dict[str, dict[str, Any]] = {
    "case.title": {"label": "案件名称", "patterns": ("案件名称", "案件", "案名")},
    "case.type": {"label": "案由/案件类型", "patterns": ("案由", "案件类型", "纠纷类型")},
    "case.stage": {"label": "代理阶段", "patterns": ("代理阶段", "委托阶段", "诉讼阶段")},
    "case.number": {"label": "案号", "patterns": ("案号",)},
    "case.court": {"label": "法院/仲裁机构", "patterns": ("受理法院", "法院", "仲裁委", "受理机构")},
    "case.subject_amount": {"label": "争议标的额", "patterns": ("争议标的", "标的额")},
    "client.name": {"label": "委托人", "patterns": ("委托人", "甲方")},
    "client.identity_code": {"label": "委托人证件号码", "patterns": ("统一社会信用代码", "身份证号码", "证件号码")},
    "client.address": {"label": "委托人地址", "patterns": ("委托人地址", "甲方地址", "住所地", "住所")},
    "client.email": {"label": "委托人电子邮箱", "patterns": ("委托人邮箱", "甲方邮箱", "电子邮箱", "Email")},
    "client.legal_representative": {"label": "法定代表人", "patterns": ("法定代表人", "负责人")},
    "client.phone": {"label": "委托人联系电话", "patterns": ("委托人电话", "甲方电话", "联系电话")},
    "opponents.names": {"label": "对方当事人", "patterns": ("对方当事人", "相对方", "乙方")},
    "firm.name": {"label": "律师事务所", "patterns": ("律师事务所", "承办律所")},
    "firm.address": {"label": "律所地址", "patterns": ("律所地址", "事务所地址")},
    "firm.credit_code": {"label": "律所统一社会信用代码", "patterns": ("律所统一社会信用代码",)},
    "firm.bank": {"label": "开户银行", "patterns": ("开户银行", "开户行")},
    "firm.account": {"label": "银行账号", "patterns": ("银行账号", "收款账号", "账户")},
    "engagement.lawyers": {"label": "代理律师", "patterns": ("代理律师", "承办律师", "受委托人")},
    "engagement.scope": {"label": "委托范围", "patterns": ("委托范围", "委托事项", "代理权限")},
    "engagement.fee_type": {"label": "收费方式", "patterns": ("收费方式", "收费模式")},
    "engagement.fee_amount": {"label": "律师费", "patterns": ("律师费", "代理费", "服务费")},
    "engagement.payment_terms": {"label": "付款安排", "patterns": ("付款安排", "付款方式", "支付方式")},
    "engagement.authorization": {"label": "代理权限", "patterns": ("代理权限", "授权范围")},
    "engagement.fee_clause": {"label": "收费条款", "patterns": ("收费条款", "律师费方案")},
    "engagement.copies": {"label": "合同份数", "patterns": ("合同份数", "一式")},
    "engagement.sign_date": {"label": "签署日期", "patterns": ("签署日期", "签订日期", "日期")},
}


def _now() -> str:
    return dt.datetime.now().replace(microsecond=0).isoformat(sep=" ")


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _loads(value: Any, default: Any) -> Any:
    try:
        return json.loads(value or "")
    except (TypeError, ValueError, json.JSONDecodeError):
        return default


def _row(row: Any) -> dict[str, Any]:
    return dict(row) if row else {}


def safe_filename(value: str, fallback: str = "文件") -> str:
    text = re.sub(r"[\x00-\x1f<>:\"/\\|?*]+", "_", str(value or "")).strip(" ._")
    text = re.sub(r"\s+", " ", text)
    return (text[:180] or fallback)


def _hash_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def ensure_case_automation_schema(get_db: Callable[[], Any]) -> None:
    with get_db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS case_manager_folder_bindings (
                case_id INTEGER PRIMARY KEY,
                folder_path TEXT NOT NULL,
                real_path TEXT NOT NULL,
                device_id INTEGER,
                inode INTEGER,
                status TEXT NOT NULL DEFAULT 'active',
                subdirs_json TEXT NOT NULL DEFAULT '{}',
                last_verified_at TEXT DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
            );
            CREATE TABLE IF NOT EXISTS case_manager_files (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                case_id INTEGER NOT NULL,
                category TEXT NOT NULL,
                relative_path TEXT NOT NULL,
                absolute_path TEXT NOT NULL,
                original_filename TEXT DEFAULT '',
                sha256 TEXT NOT NULL,
                size_bytes INTEGER NOT NULL DEFAULT 0,
                mime_type TEXT DEFAULT '',
                source_type TEXT DEFAULT '',
                source_id INTEGER,
                metadata_json TEXT DEFAULT '{}',
                created_at TEXT NOT NULL,
                UNIQUE(case_id, relative_path),
                FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS idx_case_manager_files_case
                ON case_manager_files(case_id, category, created_at DESC);
            CREATE INDEX IF NOT EXISTS idx_case_manager_files_hash
                ON case_manager_files(case_id, sha256);

            CREATE TABLE IF NOT EXISTS contract_templates (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                original_filename TEXT NOT NULL,
                stored_path TEXT NOT NULL,
                sha256 TEXT NOT NULL UNIQUE,
                version INTEGER NOT NULL DEFAULT 1,
                active INTEGER NOT NULL DEFAULT 1,
                analysis_status TEXT NOT NULL DEFAULT 'pending',
                structure_json TEXT DEFAULT '{}',
                warnings_json TEXT DEFAULT '[]',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS contract_template_mappings (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                template_id INTEGER NOT NULL,
                semantic_key TEXT NOT NULL,
                label TEXT NOT NULL,
                locator_json TEXT NOT NULL,
                fill_type TEXT NOT NULL DEFAULT 'replace_token',
                token_text TEXT DEFAULT '',
                required INTEGER NOT NULL DEFAULT 0,
                confidence REAL NOT NULL DEFAULT 0,
                confirmed INTEGER NOT NULL DEFAULT 0,
                sort_order INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(template_id) REFERENCES contract_templates(id) ON DELETE CASCADE
            );
            CREATE INDEX IF NOT EXISTS idx_contract_template_mappings
                ON contract_template_mappings(template_id, sort_order, id);
            CREATE TABLE IF NOT EXISTS case_manager_contract_profiles (
                case_id INTEGER PRIMARY KEY,
                template_id INTEGER,
                lawyers_json TEXT DEFAULT '[]',
                entrust_type TEXT DEFAULT '',
                representation_stage TEXT DEFAULT '',
                engagement_scope TEXT DEFAULT '',
                background_matter TEXT DEFAULT '',
                background_stage TEXT DEFAULT '',
                matter_kind TEXT DEFAULT 'case',
                fee_type TEXT DEFAULT '',
                fee_amount TEXT DEFAULT '',
                payment_terms TEXT DEFAULT '',
                contract_copies TEXT DEFAULT '',
                sign_date TEXT DEFAULT '',
                firm_json TEXT DEFAULT '{}',
                custom_values_json TEXT DEFAULT '{}',
                fee_plan TEXT DEFAULT '',
                fee_config_json TEXT DEFAULT '{}',
                authorization_json TEXT DEFAULT '[]',
                authority_lines_json TEXT DEFAULT '[]',
                subject_amount TEXT DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE,
                FOREIGN KEY(template_id) REFERENCES contract_templates(id) ON DELETE SET NULL
            );
            CREATE TABLE IF NOT EXISTS contract_profile_presets (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                source_case_id INTEGER,
                profile_json TEXT NOT NULL DEFAULT '{}',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(source_case_id) REFERENCES case_manager_cases(id) ON DELETE SET NULL
            );
            CREATE INDEX IF NOT EXISTS idx_contract_profile_presets_updated
                ON contract_profile_presets(updated_at DESC, id DESC);
            CREATE TABLE IF NOT EXISTS case_manager_contract_previews (
                token TEXT PRIMARY KEY,
                case_id INTEGER NOT NULL,
                template_id INTEGER NOT NULL,
                file_path TEXT NOT NULL,
                profile_fingerprint TEXT DEFAULT '',
                values_json TEXT NOT NULL,
                missing_json TEXT NOT NULL DEFAULT '[]',
                created_at TEXT NOT NULL,
                FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE,
                FOREIGN KEY(template_id) REFERENCES contract_templates(id) ON DELETE CASCADE
            );
            """
        )
        profile_columns = {
            row[1] for row in conn.execute("PRAGMA table_info(case_manager_contract_profiles)").fetchall()
        }
        migrations = {
            "entrust_type": "TEXT DEFAULT ''",
            "background_matter": "TEXT DEFAULT ''",
            "background_stage": "TEXT DEFAULT ''",
            "matter_kind": "TEXT DEFAULT 'case'",
            "fee_plan": "TEXT DEFAULT ''",
            "fee_config_json": "TEXT DEFAULT '{}'",
            "authorization_json": "TEXT DEFAULT '[]'",
            "authority_lines_json": "TEXT DEFAULT '[]'",
            "subject_amount": "TEXT DEFAULT ''",
        }
        for column, declaration in migrations.items():
            if column not in profile_columns:
                conn.execute(
                    f"ALTER TABLE case_manager_contract_profiles ADD COLUMN {column} {declaration}"
                )
        preview_columns = {
            row[1] for row in conn.execute("PRAGMA table_info(case_manager_contract_previews)").fetchall()
        }
        if "profile_fingerprint" not in preview_columns:
            conn.execute(
                "ALTER TABLE case_manager_contract_previews ADD COLUMN profile_fingerprint TEXT DEFAULT ''"
            )
        conn.commit()


def _case_exists(conn: Any, case_id: int) -> bool:
    return bool(conn.execute("SELECT 1 FROM case_manager_cases WHERE id=?", (case_id,)).fetchone())


def safe_subfolder_name(name: str) -> str:
    """Strip path separators so a requested folder name cannot escape its root."""
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(name or "")).strip(" ._")
    return cleaned[:120]


def series_child_folder_name(parent_title: str, child_title: str) -> str:
    """Derive a subfolder name for a series child from its display title."""
    child = str(child_title or "").strip()
    parent = str(parent_title or "").strip()
    if parent and child.startswith(parent):
        child = child[len(parent):].lstrip("·•\u00b7\u2022 -_—").strip()
    return child or parent or "子案"


def bind_case_folder(
    get_db: Callable[[], Any],
    case_id: int,
    folder_path: str,
    real: Path,
    stat_result: os.stat_result,
) -> None:
    """Insert or refresh one case's folder binding row (shared by all bind paths)."""
    stamp = _now()
    subdirs = {"court_document": COURT_DOCUMENT_SUBDIR, "contract": CONTRACT_SUBDIR}
    with get_db() as conn:
        previous = conn.execute(
            "SELECT real_path FROM case_manager_folder_bindings WHERE case_id=?", (case_id,)
        ).fetchone()
        conn.execute(
            """INSERT INTO case_manager_folder_bindings
               (case_id,folder_path,real_path,device_id,inode,status,subdirs_json,
                last_verified_at,created_at,updated_at)
               VALUES (?,?,?,?,?,'active',?,?,?,?)
               ON CONFLICT(case_id) DO UPDATE SET
                 folder_path=excluded.folder_path,real_path=excluded.real_path,
                 device_id=excluded.device_id,inode=excluded.inode,status='active',
                 subdirs_json=excluded.subdirs_json,last_verified_at=excluded.last_verified_at,
                 updated_at=excluded.updated_at""",
            (
                case_id, str(folder_path or real), str(real), int(stat_result.st_dev), int(stat_result.st_ino),
                _json(subdirs), stamp, stamp, stamp,
            ),
        )
        if previous and str(previous["real_path"] or "") != str(real):
            try:
                conn.execute("DELETE FROM case_manager_local_file_index WHERE case_id=?", (case_id,))
            except sqlite3.OperationalError:
                pass
        conn.commit()


def validate_case_folder(path_text: str) -> tuple[Path, os.stat_result]:
    if not path_text or "\x00" in path_text:
        raise HTTPException(status_code=400, detail="请选择案件文件夹")
    path = Path(path_text).expanduser()
    try:
        real = path.resolve(strict=True)
    except (OSError, RuntimeError) as exc:
        raise HTTPException(status_code=400, detail=f"文件夹不存在或当前不可访问：{exc}") from exc
    home = Path.home().resolve()
    forbidden = {Path("/"), home, home.parent}
    if real in forbidden or len(real.parts) < 3:
        raise HTTPException(status_code=400, detail="请选择具体的案件文件夹，不能绑定系统根目录或用户主目录")
    if not real.is_dir():
        raise HTTPException(status_code=400, detail="所选路径不是文件夹")
    if not os.access(real, os.R_OK | os.W_OK | os.X_OK):
        raise HTTPException(status_code=400, detail="当前程序没有该文件夹的读写权限")
    return real, real.stat()


def get_folder_binding(get_db: Callable[[], Any], case_id: int, required: bool = False) -> dict[str, Any]:
    with get_db() as conn:
        row = conn.execute(
            "SELECT * FROM case_manager_folder_bindings WHERE case_id=?", (case_id,)
        ).fetchone()
    if not row:
        if required:
            raise HTTPException(status_code=409, detail="请先为案件绑定本地文件夹")
        return {"case_id": case_id, "bound": False, "available": False}
    item = dict(row)
    path = Path(item["real_path"])
    available = path.is_dir() and os.access(path, os.R_OK | os.W_OK | os.X_OK)
    status = "offline"
    if available:
        try:
            stat = path.stat()
            stored_inode = int(item.get("inode") or 0)
            stored_device = int(item.get("device_id") or 0)
            # st_dev can change after a macOS volume is remounted or its mount
            # topology changes.  The directory inode remains the useful identity
            # signal in that situation; requiring both values would incorrectly
            # take every otherwise valid binding offline.  Older bindings without
            # an inode still fall back to the stored device identifier.
            same_identity = (
                int(stat.st_ino) == stored_inode
                if stored_inode
                else (not stored_device or int(stat.st_dev) == stored_device)
            )
            available = same_identity
            status = "active" if same_identity else "replaced"
            if same_identity:
                item["device_id"] = int(stat.st_dev)
                item["inode"] = int(stat.st_ino)
        except OSError:
            available = False
    stamp = _now()
    try:
        with get_db() as conn:
            conn.execute(
                """UPDATE case_manager_folder_bindings
                   SET device_id=?,inode=?,status=?,subdirs_json=?,last_verified_at=? WHERE case_id=?""",
                (
                    item.get("device_id"), item.get("inode"),
                    status if available or status == "replaced" else "offline",
                    _json({"court_document": COURT_DOCUMENT_SUBDIR, "contract": CONTRACT_SUBDIR}),
                    stamp, case_id,
                ),
            )
            conn.commit()
    except sqlite3.Error:
        pass
    item["status"] = status if available or status == "replaced" else "offline"
    item["last_verified_at"] = stamp
    item["bound"] = True
    item["available"] = available
    if available:
        _migrate_court_document_directory(get_db, case_id, path)
    item["subdirs"] = _loads(item.pop("subdirs_json", "{}"), {})
    item["subdirs"]["court_document"] = COURT_DOCUMENT_SUBDIR
    item["subdirs"].setdefault("contract", CONTRACT_SUBDIR)
    if required and not available:
        raise HTTPException(status_code=409, detail="案件文件夹当前离线或不可写，请重新绑定")
    return item


def _unique_destination(directory: Path, filename: str) -> Path:
    stem = Path(filename).stem
    suffix = Path(filename).suffix
    candidate = directory / filename
    index = 2
    while candidate.exists():
        candidate = directory / f"{stem}_{index}{suffix}"
        index += 1
    return candidate


def _migrate_court_document_directory(get_db: Callable[[], Any], case_id: int, root: Path) -> None:
    """Rename the legacy judicial-delivery folder without losing colliding files.

    A user can already have a new-named folder (for example after manually
    preparing a case folder).  In that case, move legacy entries in one at a
    time and give only colliding entries a suffix, rather than overwriting them.
    File-index rows are then adjusted to the actual destination paths.
    """
    legacy = root / LEGACY_COURT_DOCUMENT_SUBDIR
    target = root / COURT_DOCUMENT_SUBDIR
    legacy_exists = legacy.is_dir() and not legacy.is_symlink()
    if not legacy_exists and not target.is_dir():
        return

    moved: list[tuple[Path, Path]] = []
    try:
        if legacy_exists and not target.exists():
            legacy.rename(target)
            moved.append((legacy, target))
        elif legacy_exists and target.is_dir() and not target.is_symlink():
            for source in list(legacy.iterdir()):
                destination = _unique_destination(target, source.name)
                shutil.move(str(source), str(destination))
                moved.append((source, destination))
            legacy.rmdir()
        elif legacy_exists:
            # Do not replace a file or symlink bearing the target folder name.
            return
    except OSError:
        # A later binding read or application start can retry a temporarily
        # unavailable folder; normal document archiving remains unaffected.
        return

    with get_db() as conn:
        rows = conn.execute(
            "SELECT id,relative_path,absolute_path FROM case_manager_files "
            "WHERE case_id=? AND category='court_document'",
            (case_id,),
        ).fetchall()
        for row in rows:
            old_path = Path(str(row["absolute_path"] or ""))
            new_path: Path | None = None
            for source, destination in moved:
                try:
                    suffix = old_path.relative_to(source)
                except ValueError:
                    continue
                new_path = destination / suffix
                break
            if new_path is None:
                # The row may predate file indexing; still update the standard
                # legacy prefix when its physical directory was renamed.
                relative = str(row["relative_path"] or "")
                if relative == LEGACY_COURT_DOCUMENT_SUBDIR or relative.startswith(LEGACY_COURT_DOCUMENT_SUBDIR + os.sep):
                    new_path = root / COURT_DOCUMENT_SUBDIR / Path(relative).relative_to(LEGACY_COURT_DOCUMENT_SUBDIR)
            if new_path is not None:
                conn.execute(
                    "UPDATE case_manager_files SET relative_path=?,absolute_path=? WHERE id=?",
                    (str(new_path.relative_to(root)), str(new_path), int(row["id"])),
                )
        conn.execute(
            "UPDATE case_manager_folder_bindings SET subdirs_json=?,updated_at=? WHERE case_id=?",
            (_json({"court_document": COURT_DOCUMENT_SUBDIR, "contract": CONTRACT_SUBDIR}), _now(), case_id),
        )
        conn.commit()


def store_case_file(
    get_db: Callable[[], Any],
    case_id: int,
    category: str,
    filename: str,
    content: bytes,
    *,
    mime_type: str = "",
    source_type: str = "",
    source_id: int | None = None,
    original_filename: str = "",
    metadata: dict[str, Any] | None = None,
) -> dict[str, Any]:
    binding = get_folder_binding(get_db, case_id, required=True)
    subdir = COURT_DOCUMENT_SUBDIR if category == "court_document" else CONTRACT_SUBDIR
    root = Path(binding["real_path"])
    directory = root / subdir
    directory.mkdir(parents=False, exist_ok=True)
    if directory.resolve().parent != root.resolve():
        raise HTTPException(status_code=400, detail="非法的案件子目录")
    final_path = _unique_destination(directory, safe_filename(filename))
    fd, temporary = tempfile.mkstemp(prefix=".legal-kb-", suffix=".tmp", dir=str(directory))
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, final_path)
    except Exception:
        try:
            os.unlink(temporary)
        except OSError:
            pass
        raise
    relative = str(final_path.relative_to(root))
    stamp = _now()
    digest = _hash_bytes(content)
    with get_db() as conn:
        cur = conn.execute(
            """INSERT INTO case_manager_files
               (case_id,category,relative_path,absolute_path,original_filename,sha256,size_bytes,
                mime_type,source_type,source_id,metadata_json,created_at)
               VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                case_id, category, relative, str(final_path), original_filename, digest, len(content),
                mime_type, source_type, source_id, _json(metadata or {}), stamp,
            ),
        )
        conn.commit()
        file_id = int(cur.lastrowid)
    return {
        "id": file_id,
        "case_id": case_id,
        "category": category,
        "relative_path": relative,
        "absolute_path": str(final_path),
        "sha256": digest,
        "size_bytes": len(content),
        "created_at": stamp,
    }


def _iter_doc_paragraphs(doc: Any) -> Iterable[tuple[dict[str, Any], Any]]:
    for index, paragraph in enumerate(doc.paragraphs):
        yield {"scope": "body", "kind": "paragraph", "paragraph": index}, paragraph
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            seen: set[int] = set()
            for col_index, cell in enumerate(row.cells):
                cell_key = id(cell._tc)
                if cell_key in seen:
                    continue
                seen.add(cell_key)
                for paragraph_index, paragraph in enumerate(cell.paragraphs):
                    yield {
                        "scope": "body", "kind": "cell", "table": table_index,
                        "row": row_index, "col": col_index, "paragraph": paragraph_index,
                    }, paragraph
    for section_index, section in enumerate(doc.sections):
        for scope, container in (("header", section.header), ("footer", section.footer)):
            for paragraph_index, paragraph in enumerate(container.paragraphs):
                yield {
                    "scope": scope, "kind": "paragraph", "section": section_index,
                    "paragraph": paragraph_index,
                }, paragraph
            for table_index, table in enumerate(container.tables):
                for row_index, row in enumerate(table.rows):
                    seen: set[int] = set()
                    for col_index, cell in enumerate(row.cells):
                        cell_key = id(cell._tc)
                        if cell_key in seen:
                            continue
                        seen.add(cell_key)
                        for paragraph_index, paragraph in enumerate(cell.paragraphs):
                            yield {
                                "scope": scope, "kind": "cell", "section": section_index,
                                "table": table_index, "row": row_index, "col": col_index,
                                "paragraph": paragraph_index,
                            }, paragraph


def _semantic_for_text(text: str) -> tuple[str, float]:
    normalized = re.sub(r"\s+", "", text or "")
    best_key = ""
    best_score = 0.0
    for key, info in SEMANTIC_FIELDS.items():
        for pattern in info["patterns"]:
            if pattern in normalized:
                score = min(0.98, 0.72 + len(pattern) / 30)
                if score > best_score:
                    best_key, best_score = key, score
    return best_key, best_score


def analyze_docx(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]], list[str]]:
    from docx import Document

    doc = Document(str(path))
    structure: dict[str, Any] = {"paragraphs": [], "tables": len(doc.tables), "sections": len(doc.sections)}
    mappings: list[dict[str, Any]] = []
    warnings: list[str] = []
    token_pattern = re.compile(r"(\{\{\s*([^{}]+?)\s*\}\}|【\s*([^【】]+?)\s*】)")
    for order, (locator, paragraph) in enumerate(_iter_doc_paragraphs(doc)):
        text = paragraph.text or ""
        if not text.strip():
            continue
        structure["paragraphs"].append({"locator": locator, "text": text[:1000]})
        tokens = list(token_pattern.finditer(text))
        if tokens:
            for token in tokens:
                token_text = token.group(1)
                token_label = token.group(2) or token.group(3) or ""
                key, confidence = _semantic_for_text(token_label)
                if not key and token_label in SEMANTIC_FIELDS:
                    key, confidence = token_label, 0.99
                mappings.append({
                    "semantic_key": key,
                    "label": SEMANTIC_FIELDS.get(key, {}).get("label", token_label),
                    "locator": locator,
                    "fill_type": "replace_token",
                    "token_text": token_text,
                    "required": bool(key in {"client.name", "case.title"}),
                    "confidence": confidence,
                    "confirmed": False,
                    "sort_order": order,
                })
            continue
        key, confidence = _semantic_for_text(text)
        has_blank = bool(re.search(r"_{3,}|〔\s*〕|（\s*）|:\s*$|：\s*$", text))
        if key and has_blank:
            mappings.append({
                "semantic_key": key,
                "label": SEMANTIC_FIELDS[key]["label"],
                "locator": locator,
                "fill_type": "fill_blank",
                "token_text": "",
                "required": bool(key in {"client.name", "case.title"}),
                "confidence": confidence,
                "confirmed": False,
                "sort_order": order,
            })
    with zipfile.ZipFile(path) as archive:
        raw = b"\n".join(
            archive.read(name) for name in archive.namelist() if name.startswith("word/") and name.endswith(".xml")
        )
    if b"w:txbxContent" in raw:
        warnings.append("模板包含文本框；第一版不会自动填充文本框中的字段")
    if b"w:sdt" in raw:
        warnings.append("模板包含内容控件或复杂域；请重点核对预览，不能稳定映射的位置不会自动填写")
    if any(marker in raw for marker in ("□".encode(), "☐".encode(), "☑".encode())):
        warnings.append("模板包含复选框；只有设置了明确占位符并确认映射的复选框文字才会填充，系统不会自行勾选条款")
    if not mappings:
        warnings.append("未检测到明确占位符或可填写空白，请在字段映射中手动添加")
    return structure, mappings, warnings


def _get_paragraph(doc: Any, locator: dict[str, Any]) -> Any:
    scope = locator.get("scope", "body")
    if scope in {"header", "footer"}:
        section = doc.sections[int(locator.get("section", 0))]
        container = section.header if scope == "header" else section.footer
        if locator.get("kind") == "cell":
            cell = container.tables[int(locator["table"])].rows[int(locator["row"])].cells[int(locator["col"])]
            return cell.paragraphs[int(locator.get("paragraph", 0))]
        return container.paragraphs[int(locator["paragraph"])]
    if locator.get("kind") == "cell":
        cell = doc.tables[int(locator["table"])].rows[int(locator["row"])].cells[int(locator["col"])]
        return cell.paragraphs[int(locator.get("paragraph", 0))]
    return doc.paragraphs[int(locator["paragraph"])]


def _replace_across_runs(paragraph: Any, old: str, new: str) -> bool:
    if not old:
        return False
    replaced = False
    search_from = 0
    while True:
        runs = list(paragraph.runs)
        full = "".join(run.text for run in runs)
        start = full.find(old, search_from)
        if start < 0:
            return replaced
        end = start + len(old)
        offsets: list[tuple[int, int]] = []
        cursor = 0
        for run in runs:
            offsets.append((cursor, cursor + len(run.text)))
            cursor += len(run.text)
        start_index = next(index for index, (_, stop) in enumerate(offsets) if start < stop)
        end_index = next(index for index, (_, stop) in enumerate(offsets) if end <= stop)
        start_run_start = offsets[start_index][0]
        end_run_start = offsets[end_index][0]
        prefix = runs[start_index].text[: start - start_run_start]
        suffix = runs[end_index].text[end - end_run_start :]
        runs[start_index].text = prefix + new + (suffix if start_index == end_index else "")
        for index in range(start_index + 1, end_index):
            runs[index].text = ""
        if end_index != start_index:
            runs[end_index].text = suffix
        replaced = True
        # Do not search inside the replacement text. In particular, replacing
        # an unchecked checkbox "☐" with "☐" must not loop forever.
        search_from = start + len(new)


def _replace_first_across_runs(paragraph: Any, old: str, new: str) -> bool:
    if not old:
        return False
    runs = list(paragraph.runs)
    full = "".join(run.text for run in runs)
    start = full.find(old)
    if start < 0:
        return False
    end = start + len(old)
    offsets: list[tuple[int, int]] = []
    cursor = 0
    for run in runs:
        offsets.append((cursor, cursor + len(run.text)))
        cursor += len(run.text)
    start_index = next(index for index, (_, stop) in enumerate(offsets) if start < stop)
    end_index = next(index for index, (_, stop) in enumerate(offsets) if end <= stop)
    start_run_start = offsets[start_index][0]
    end_run_start = offsets[end_index][0]
    prefix = runs[start_index].text[: start - start_run_start]
    suffix = runs[end_index].text[end - end_run_start :]
    runs[start_index].text = prefix + new + (suffix if start_index == end_index else "")
    for index in range(start_index + 1, end_index):
        runs[index].text = ""
    if end_index != start_index:
        runs[end_index].text = suffix
    return True


def _fill_blank(paragraph: Any, value: str) -> None:
    text = paragraph.text or ""
    match = re.search(r"_{3,}|〔\s*〕|（\s*）", text)
    if match and _replace_across_runs(paragraph, match.group(0), value):
        return
    if paragraph.runs:
        paragraph.runs[-1].text += value
    else:
        paragraph.add_run(value)


def _money(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return "【待填写】"
    if re.fullmatch(r"[\d,.]+", text):
        return f"{text}元"
    return text


def build_fee_clause(profile: dict[str, Any]) -> str:
    plan = str(profile.get("fee_plan") or profile.get("fee_type") or "fixed")
    config = profile.get("fee_config") or {}
    payment = str(
        config.get("payment_terms")
        or profile.get("payment_terms")
        or "本合同签订后3日内支付"
    ).strip()
    base_amount = _money(config.get("base_amount") or profile.get("fee_amount"))
    if plan == "fixed":
        return f"本案律师服务费采用固定收费方式，共计{base_amount}；{payment}。"
    if plan == "staged_fixed":
        first = _money(config.get("first_amount"))
        second = _money(config.get("second_amount"))
        return (
            f"本案按代理阶段固定收费：一审或仲裁阶段律师服务费为{first}；"
            f"如进入二审阶段，律师服务费为{second}；{payment}。"
        )

    prefix = ""
    if plan == "base_risk":
        prefix = f"本案采用基础费加风险代理费方式，基础律师服务费为{base_amount}；"
    elif plan == "pure_risk":
        prefix = "本案采用纯风险代理方式，不预先收取基础律师服务费；"
    method = str(config.get("risk_method") or "recovery_rate")
    rate = str(config.get("rate") or "【待填写】").rstrip("%")
    if method == "reduction_rate":
        risk = (
            f"以对方主张金额与生效裁判或经确认和解方案中委托人最终承担金额之差额"
            f"作为计费基数，按{rate}%计收风险代理费"
        )
    elif method == "success_fixed":
        result = str(config.get("success_condition") or "委托人无需承担付款责任")
        risk = f"如实现“{result}”的约定结果，收取成功代理费{_money(config.get('success_amount'))}"
    elif method == "recovery_tiered":
        tiers = config.get("tiers") or []
        rendered = []
        for item in tiers:
            threshold = str(item.get("threshold") or "").strip()
            tier_rate = str(item.get("rate") or "").strip().rstrip("%")
            if threshold and tier_rate:
                rendered.append(f"实际回款达到{threshold}时按{tier_rate}%计收")
        risk = "；".join(rendered) or "按双方确认的实际回款分段比例计收风险代理费"
    else:
        risk = f"以委托人实际收回的款项为计费基数，按{rate}%计收风险代理费"
    return f"{prefix}{risk}；{payment}。"


def contract_context(conn: Any, case_id: int, profile: dict[str, Any] | None = None) -> dict[str, Any]:
    profile = profile or {}
    values = build_contract_values(conn, case_id, profile)
    case_numbers = [
        dict(row) for row in conn.execute(
            """SELECT case_no,court,number_type,sort_order
               FROM case_manager_case_numbers
               WHERE case_id=? ORDER BY sort_order,id""",
            (case_id,),
        ).fetchall()
    ]
    lawyers = [
        dict(row) for row in conn.execute(
            "SELECT id,name,phone FROM suohan_lawyers ORDER BY id"
        ).fetchall()
    ]
    stages = [
        dict(row) for row in conn.execute(
            "SELECT id,name FROM suohan_stages ORDER BY id"
        ).fetchall()
    ]
    courts = [
        dict(row) for row in conn.execute(
            "SELECT id,name,province,city FROM suohan_courts ORDER BY id"
        ).fetchall()
    ]
    authority_templates: dict[str, list[str]] = {}
    for row in conn.execute(
        "SELECT entrust_type,authorities_json FROM power_attorney_authority_templates"
    ).fetchall():
        authority_templates[str(row["entrust_type"])] = _loads(row["authorities_json"], [])
    entrust_types = []
    for definition in ENTRUST_TYPE_DEFINITIONS:
        item = dict(definition)
        item["authority_lines"] = authority_templates.get(
            item["id"], list(definition.get("authority_lines") or [])
        )
        item["authorization_ids"] = _authority_ids_from_lines(item["authority_lines"])
        entrust_types.append(item)
    return {
        "case": {
            "title": values["case.title"],
            "type": values["case.type"],
            "stage": values["case.stage"],
            "number": values["case.number"],
            "court": values["case.court"],
            "subject_amount": values["case.subject_amount"],
        },
        "client": {
            "name": values["client.name"],
            "identity_code": values["client.identity_code"],
            "address": values["client.address"],
            "email": values["client.email"],
            "legal_representative": values["client.legal_representative"],
            "legal_representative_title": values["client.legal_representative_title"],
            "contact_person": values["client.contact_person"],
            "phone": values["client.phone"],
        },
        "opponents": {"names": values["opponents.names"]},
        "defaults": {
            "firm": DEFAULT_FIRM,
            "authorizations": [],
            "selected_lawyers": profile.get("lawyers") or [],
            "contract_copies": "3",
        },
        "case_numbers": case_numbers,
        "lawyers": lawyers,
        "stages": stages,
        "courts": courts,
        "entrust_types": entrust_types,
        "authorization_options": AUTHORIZATION_OPTIONS,
        "fee_plans": FEE_PLAN_DEFINITIONS,
        "risk_methods": RISK_METHOD_DEFINITIONS,
    }


def build_contract_values(conn: Any, case_id: int, profile: dict[str, Any]) -> dict[str, str]:
    case = conn.execute("SELECT * FROM case_manager_cases WHERE id=?", (case_id,)).fetchone()
    if not case:
        raise HTTPException(status_code=404, detail="案件不存在")
    parties = [dict(row) for row in conn.execute(
        """SELECT p.*,c.identity_code AS client_identity_code,c.address AS client_address,
                  c.legal_representative AS client_legal_representative,c.phone AS client_phone,
                  c.legal_representative_title AS client_legal_representative_title,
                  c.contact_person AS client_contact_person
           FROM case_manager_parties p
           LEFT JOIN case_manager_clients c ON c.id=p.client_id
           WHERE p.case_id=? ORDER BY p.sort_order,p.id""", (case_id,)
    ).fetchall()]
    numbers = [dict(row) for row in conn.execute(
        "SELECT * FROM case_manager_case_numbers WHERE case_id=? ORDER BY sort_order,id", (case_id,)
    ).fetchall()]
    clients = [item for item in parties if item.get("party_group") == "client"]
    opponents = [item for item in parties if item.get("party_group") != "client"]
    client = clients[0] if clients else {}
    stage = str(profile.get("representation_stage") or case["current_stage"] or "")
    number = next(
        (
            item for item in reversed(numbers)
            if stage and (
                stage in str(item.get("number_type") or "")
                or str(item.get("number_type") or "") in stage
            )
        ),
        numbers[-1] if numbers else {},
    )
    firm = dict(DEFAULT_FIRM)
    firm.update({
        key: value for key, value in (profile.get("firm") or {}).items()
        if str(value or "").strip()
    })
    lawyers = profile.get("lawyers") or []
    entrust_type = str(profile.get("entrust_type") or "")
    definition = next(
        (item for item in ENTRUST_TYPE_DEFINITIONS if item["id"] == entrust_type), {}
    )
    authority_lines = (
        profile.get("authority_lines")
        or definition.get("authority_lines")
        or []
    )
    authorizations = (
        profile.get("authorizations")
        or _authority_ids_from_lines(authority_lines)
    )
    authorization_labels = [
        item["label"] for item in AUTHORIZATION_OPTIONS if item["id"] in authorizations
    ]
    values = {
        "case.title": str(case["title"] or ""),
        "case.type": str(case["case_type"] or ""),
        "case.stage": stage,
        "case.number": str(number.get("case_no") or ""),
        "case.court": str(number.get("court") or ""),
        "case.subject_amount": str(profile.get("subject_amount") or ""),
        "client.name": "、".join(item.get("name") or "" for item in clients).strip("、"),
        "client.identity_code": str(client.get("client_identity_code") or client.get("credit_code") or ""),
        "client.address": str(client.get("client_address") or ""),
        "client.email": str((profile.get("custom_values") or {}).get("client.email") or ""),
        "client.legal_representative": str(client.get("client_legal_representative") or ""),
        "client.legal_representative_title": str(client.get("client_legal_representative_title") or ""),
        "client.contact_person": str(client.get("client_contact_person") or ""),
        "client.phone": str(client.get("client_phone") or ""),
        "opponents.names": "、".join(item.get("name") or "" for item in opponents).strip("、"),
        "firm.name": str(firm.get("name") or ""),
        "firm.address": str(firm.get("address") or ""),
        "firm.credit_code": str(firm.get("credit_code") or ""),
        "firm.bank": str(firm.get("bank") or ""),
        "firm.account": str(firm.get("account") or ""),
        "engagement.lawyers": "、".join(
            str(item.get("name") if isinstance(item, dict) else item) for item in lawyers
        ).strip("、"),
        "engagement.scope": str(profile.get("engagement_scope") or ""),
        "engagement.entrust_type": entrust_type,
        "engagement.background_matter": str(
            profile.get("background_matter")
            or (f"{case['case_type']}一案" if case["case_type"] else "")
        ),
        "engagement.background_stage": str(
            profile.get("background_stage") or stage
        ),
        "engagement.matter_kind": (
            "matter" if str(profile.get("matter_kind") or "case") == "matter" else "case"
        ),
        "engagement.fee_type": str(profile.get("fee_type") or ""),
        "engagement.fee_amount": str(profile.get("fee_amount") or ""),
        "engagement.payment_terms": str(profile.get("payment_terms") or ""),
        "engagement.authorization": "；".join(authorization_labels),
        "engagement.authority_lines": "\n".join(str(item) for item in authority_lines),
        "engagement.authorization_ids": ",".join(str(item) for item in authorizations),
        "engagement.fee_clause": build_fee_clause(profile),
        "engagement.fee_plan": str(profile.get("fee_plan") or "fixed"),
        "engagement.fee_config_json": _json(profile.get("fee_config") or {}),
        "engagement.copies": str(profile.get("contract_copies") or "3"),
        "engagement.sign_date": str(profile.get("sign_date") or dt.date.today().isoformat()),
        "contract.current_year": str(dt.date.today().year),
        "contract.cover_month": _chinese_year_month(dt.date.today()),
    }
    for key, value in (profile.get("custom_values") or {}).items():
        values[str(key)] = str(value or "")
    return values


def _iter_all_paragraphs(doc: Any) -> Iterable[Any]:
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _fill_known_contract_labels(doc: Any, values: dict[str, str]) -> None:
    replacements = {
        "【委托人】": values.get("client.name", ""),
        "【对方当事人】": values.get("opponents.names", ""),
        "【案由】": values.get("case.type", ""),
        "【法院/仲裁机构】": values.get("case.court", ""),
        "【受理机构】": values.get("case.court", ""),
        "【案号】": values.get("case.number", ""),
        "【代理律师】": values.get("engagement.lawyers", ""),
    }
    label_values = {
        "法定代表人": values.get("client.legal_representative", ""),
        "统一社会信用代码": values.get("client.identity_code", ""),
        "住所": values.get("client.address", ""),
        "争议标的": values.get("case.subject_amount", ""),
    }
    for paragraph in _iter_all_paragraphs(doc):
        for token, value in replacements.items():
            if value:
                _replace_across_runs(paragraph, token, value)
        text = paragraph.text or ""
        direct_labels = (
            ("法定代表人：", values.get("client.legal_representative", "")),
            ("统一社会信用代码：", values.get("client.identity_code", "")),
            ("对方当事人名称或者姓名：", values.get("opponents.names", "")),
            ("案由：", values.get("case.type", "")),
            ("受理机关：", values.get("case.court", "")),
            ("涉及标的：", values.get("case.subject_amount", "")),
        )
        for label, value in direct_labels:
            if value and label in text and not text.split(label, 1)[1].strip():
                if paragraph.runs:
                    paragraph.runs[-1].text += value
                else:
                    paragraph.add_run(value)
                text = paragraph.text or ""
                break
        if values.get("engagement.background_matter") and "甲方因" in text and "根据中华" in text:
            matched = re.search(r"(?<=甲方因).*?(?=，?\s*根据中华)", text)
            if matched:
                _replace_across_runs(
                    paragraph,
                    matched.group(0),
                    f" {values['engagement.background_matter']} ",
                )
                text = paragraph.text or ""
        if values.get("engagement.background_stage") and "聘请乙方" in text and "阶段的委托代理人" in text:
            matched = re.search(r"(?<=作为).*?(?=阶段的委托代理人)", text)
            if matched:
                _replace_across_runs(
                    paragraph,
                    matched.group(0),
                    f" {values['engagement.background_stage']} ",
                )
                text = paragraph.text or ""
        if "在下列" in text and "担任甲方的委托代理人" in text:
            kind = "事项" if values.get("engagement.matter_kind") == "matter" else "案件"
            fixed_clause = f"乙方接受甲方委托，委派律师在下列{kind}中担任甲方的委托代理人："
            _replace_across_runs(
                paragraph,
                text,
                fixed_clause,
            )
            text = paragraph.text or ""
        if values.get("engagement.lawyers") and "作为上述案件" in text and "乙方委派" in text:
            matched = re.search(r"委派\s*.+?\s*律师(?:团队)?作为上述案件", text)
            if matched:
                lawyer_names = [
                    item.strip()
                    for item in values["engagement.lawyers"].split("、")
                    if item.strip()
                ]
                designation = (
                    "示上某 律师团队"
                    if lawyer_names == ["示上某"]
                    else f"{'、'.join(lawyer_names)} 律师"
                )
                _replace_across_runs(
                    paragraph,
                    matched.group(0),
                    f"委派 {designation}作为上述案件",
                )
                text = paragraph.text or ""
        if (
            values.get("client.name")
            and text.strip().startswith("甲方：")
            and "乙方：" in text
        ):
            replaced = re.sub(
                r"甲方：.*?(?=\s+乙方：)",
                f"甲方：{values['client.name']}",
                text,
                count=1,
            )
            if replaced != text:
                _replace_across_runs(paragraph, text, replaced)
                text = paragraph.text or ""
        contract_year = values.get("contract.current_year", str(dt.date.today().year))
        year_match = re.search(r"（\d{4}）粤明思律民字第", text)
        if year_match:
            _replace_across_runs(
                paragraph,
                year_match.group(0),
                f"（{contract_year}）粤明思律民字第",
            )
            text = paragraph.text or ""
        if re.fullmatch(
            r"[二〇零一二三四五六七八九十]+年[一二三四五六七八九十]+月",
            text.strip(),
        ):
            _replace_across_runs(
                paragraph, text, values.get("contract.cover_month", "")
            )
            text = paragraph.text or ""
        for label, value in label_values.items():
            if value and label in text and re.search(r"_{3,}|〔\s*〕|（\s*）", text):
                _fill_blank(paragraph, value)
                break
    address = values.get("client.address", "").strip()
    email = values.get("client.email", "").strip()
    for index, paragraph in enumerate(doc.paragraphs[:-1]):
        if not (paragraph.text or "").strip().startswith("统一社会信用代码："):
            continue
        target = doc.paragraphs[index + 1]
        if (target.text or "").strip():
            continue
        _set_paragraph_text(
            target,
            f"送达地址：{address}    电子邮箱：{email}",
            style_source=paragraph,
        )
        break


def _apply_authorization_checks(doc: Any, selected: set[str]) -> None:
    for paragraph in _iter_all_paragraphs(doc):
        text = paragraph.text or ""
        for item in AUTHORIZATION_OPTIONS:
            if item.get("keyword", item["label"][:4]) not in text:
                continue
            checked = item["id"] in selected
            marker = "☑" if checked else "☐"
            if not re.sub(r"^[\s□☐☑√×]+", "", text).strip():
                continue
            for old in ("□", "☐", "☑", "√", "×"):
                if old in paragraph.text:
                    _replace_across_runs(paragraph, old, marker)
                    break


def _replace_authority_lines(doc: Any, authority_lines: Iterable[str]) -> bool:
    """Make the contract permission list exactly match the current editor selection."""
    paragraphs = list(doc.paragraphs)
    start = next(
        (
            index for index, paragraph in enumerate(paragraphs)
            if "特别授权" in (paragraph.text or "")
            and "选择项" in (paragraph.text or "")
        ),
        None,
    )
    if start is None:
        return False
    end = next(
        (
            index for index in range(start + 1, len(paragraphs))
            if "第三条" in (paragraphs[index].text or "")
        ),
        None,
    )
    if end is None:
        return False
    destinations = list(paragraphs[start + 1:end])
    if not destinations:
        return False
    lines = [
        re.sub(r"^[\s□☐☑√×]+", "", str(line or "")).strip()
        for line in authority_lines
    ]
    lines = [line for line in lines if line]
    while len(destinations) < len(lines):
        destinations.append(_clone_paragraph_after(destinations[-1]))
    for index, line in enumerate(lines):
        _set_paragraph_text(destinations[index], f"☑{line}")
    for paragraph in destinations[len(lines):]:
        _remove_paragraph(paragraph)
    return True


def _decimal_amount(value: Any) -> Decimal | None:
    text = str(value or "").strip().replace(",", "").replace("￥", "").replace("¥", "")
    multiplier = Decimal("10000") if "万" in text else Decimal("1")
    matched = re.search(r"-?\d+(?:\.\d+)?", text)
    if not matched:
        return None
    try:
        return Decimal(matched.group(0)) * multiplier
    except InvalidOperation:
        return None


def _cn_integer(value: int) -> str:
    if value == 0:
        return "零"
    digits = "零壹贰叁肆伍陆柒捌玖"
    small_units = ("", "拾", "佰", "仟")
    big_units = ("", "万", "亿", "兆")
    groups: list[int] = []
    while value:
        groups.append(value % 10000)
        value //= 10000
    rendered: list[str] = []
    pending_zero = False
    for group_index in range(len(groups) - 1, -1, -1):
        group = groups[group_index]
        if group == 0:
            pending_zero = bool(rendered)
            continue
        chars: list[str] = []
        inner_zero = False
        for unit_index in range(3, -1, -1):
            divisor = 10 ** unit_index
            digit = group // divisor % 10
            if digit:
                if (pending_zero or inner_zero) and (not chars or chars[-1] != "零"):
                    chars.append("零")
                chars.append(digits[digit] + small_units[unit_index])
                pending_zero = False
                inner_zero = False
            elif chars:
                inner_zero = True
        rendered.append("".join(chars).rstrip("零") + big_units[group_index])
        if group < 1000 and group_index:
            pending_zero = True
    return "".join(rendered).replace("零零", "零").rstrip("零")


def _amount_yuan(value: Any) -> str:
    amount = _decimal_amount(value)
    if amount is None:
        return str(value or "").strip()
    if amount == amount.to_integral():
        return str(int(amount))
    return format(amount.normalize(), "f")


def _amount_upper_yuan(value: Any) -> str:
    amount = _decimal_amount(value)
    if amount is None:
        return ""
    return _cn_integer(int(amount))


def _amount_upper_wan(value: Any) -> str:
    amount = _decimal_amount(value)
    if amount is None:
        return ""
    wan = amount / Decimal("10000")
    if wan == wan.to_integral():
        return _cn_integer(int(wan))
    return format(wan.normalize(), "f")


def _replace_fee_blanks(paragraph: Any, replacements: Iterable[Any]) -> None:
    for value in replacements:
        text = paragraph.text or ""
        matched = re.search(r"【\s*】", text)
        if not matched:
            return
        value_text = str(value or "").strip()
        if value_text:
            _replace_first_across_runs(paragraph, matched.group(0), value_text)


def _set_paragraph_text(paragraph: Any, text: str, style_source: Any | None = None) -> None:
    original = paragraph.text or ""
    if original:
        _replace_across_runs(paragraph, original, text)
    else:
        source = style_source or paragraph
        source_run = next((run for run in source.runs if run.text), None)
        run = paragraph.add_run(text)
        if source_run is not None and source_run._r.rPr is not None:
            if run._r.rPr is not None:
                run._r.remove(run._r.rPr)
            run._r.insert(0, deepcopy(source_run._r.rPr))


def _clone_paragraph_after(paragraph: Any) -> Any:
    from docx.text.paragraph import Paragraph

    element = deepcopy(paragraph._element)
    paragraph._element.addnext(element)
    return Paragraph(element, paragraph._parent)


def _remove_paragraph(paragraph: Any) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _strip_fee_brackets(paragraph: Any) -> None:
    text = paragraph.text or ""
    for matched in list(re.finditer(r"【([^【】]*)】", text))[::-1]:
        _replace_across_runs(paragraph, matched.group(0), matched.group(1).strip())
    text = paragraph.text or ""
    matched = re.match(r"方案\s*[A-DＡ-Ｄ]\s*[.．、]?\s*", text)
    if matched:
        _replace_first_across_runs(paragraph, matched.group(0), "")


def _stage_fee_items(config: dict[str, Any]) -> list[dict[str, str]]:
    items = []
    for item in config.get("stage_fees") or []:
        name = str(item.get("stage") or item.get("name") or "").strip()
        amount = str(item.get("amount") or "").strip()
        if name and amount:
            items.append({"stage": name, "amount": amount})
    if items:
        return items
    for name, key in (("一审", "first_amount"), ("二审", "second_amount")):
        amount = str(config.get(key) or "").strip()
        if amount:
            items.append({"stage": name, "amount": amount})
    return items


def _risk_tier_items(config: dict[str, Any]) -> list[dict[str, str]]:
    items = []
    for item in config.get("tiers") or []:
        rate = str(item.get("rate") or "").strip().rstrip("%")
        if not rate:
            continue
        items.append({
            "lower": str(item.get("lower") or "").strip(),
            "upper": str(item.get("upper") or "").strip(),
            "rate": rate,
        })
    if items:
        return items
    legacy = (
        ("", "tier1_upper_wan", "tier1_rate"),
        ("tier2_lower_wan", "tier2_upper_wan", "tier2_rate"),
        ("tier3_lower_wan", "tier3_upper_wan", "tier3_rate"),
        ("tier4_lower_wan", "", "tier4_rate"),
    )
    for lower_key, upper_key, rate_key in legacy:
        rate = str(config.get(rate_key) or "").strip().rstrip("%")
        if rate:
            items.append({
                "lower": str(config.get(lower_key) or "").strip() if lower_key else "",
                "upper": str(config.get(upper_key) or "").strip() if upper_key else "",
                "rate": rate,
            })
    return items


def _risk_tier_clause(item: dict[str, str]) -> str:
    lower, upper, rate = item["lower"], item["upper"], item["rate"]
    if lower and upper:
        bound = f"回款金额 {lower} 万元以上不足 {upper} 万元的部分"
    elif upper:
        bound = f"回款金额 {upper} 万元以下的部分"
    elif lower:
        bound = f"回款金额 {lower} 万元以上的部分"
    else:
        bound = "全部回款金额"
    return f"{bound}，按 {rate}% 计收风险代理费；"


def _replace_fee_section(doc: Any, values: dict[str, str]) -> None:
    paragraphs = list(doc.paragraphs)
    markers = {}
    for label in ("方案A", "方案B", "方案C", "方案D"):
        markers[label] = next(
            (
                index for index, paragraph in enumerate(paragraphs)
                if label in (paragraph.text or "")
            ),
            None,
        )
    if any(markers[label] is None for label in markers):
        return
    end = next(
        (
            index for index in range(markers["方案A"] + 1, len(paragraphs))
            if "收款账号" in (paragraphs[index].text or "")
            or "收款账户" in (paragraphs[index].text or "")
        ),
        None,
    )
    if end is None:
        return
    plan = values.get("engagement.fee_plan") or "fixed"
    selected_label = {
        "fixed": "方案A",
        "staged_fixed": "方案B",
        "base_risk": "方案C",
        "recovery_tiered": "方案D",
        "pure_risk": "方案D",
    }.get(plan, "方案A")
    ordered = ["方案A", "方案B", "方案C", "方案D"]
    selected_start = markers[selected_label]
    marker_position = ordered.index(selected_label)
    selected_end = (
        markers[ordered[marker_position + 1]]
        if marker_position + 1 < len(ordered)
        else end
    )
    selected_paragraphs = paragraphs[selected_start:selected_end]
    try:
        config = json.loads(values.get("engagement.fee_config_json") or "{}")
    except json.JSONDecodeError:
        config = {}

    if selected_label == "方案A":
        _replace_fee_blanks(
            selected_paragraphs[0],
            (
                _amount_yuan(config.get("base_amount")),
                _amount_upper_yuan(config.get("base_amount")),
            ),
        )
    elif selected_label == "方案B":
        stage_fees = _stage_fee_items(config)
        destinations = list(selected_paragraphs[1:3])
        while len(destinations) < len(stage_fees):
            destinations.append(_clone_paragraph_after(destinations[-1]))
        for index, item in enumerate(stage_fees):
            stage = item["stage"]
            amount = item["amount"]
            if index == 0:
                payment = "甲方应于本合同签订之日起三日内向乙方一次性支付"
            elif "二审" in stage:
                payment = "甲方应在收到二审法院开庭传票之日起三日内向乙方一次性支付"
            else:
                payment = f"甲方应在进入{stage}阶段之日起三日内向乙方一次性支付"
            _set_paragraph_text(
                destinations[index],
                f"{stage}代理费人民币 {_amount_yuan(amount)} 元"
                f"（大写人民币 {_amount_upper_wan(amount)} 万元整），{payment}"
                f"{'。' if index == len(stage_fees) - 1 else '；'}",
            )
        for paragraph in destinations[len(stage_fees):]:
            _remove_paragraph(paragraph)
    elif selected_label == "方案C":
        _replace_fee_blanks(
            selected_paragraphs[1],
            (
                _amount_yuan(config.get("base_amount")),
                _amount_upper_wan(config.get("base_amount")),
            ),
        )
        _replace_fee_blanks(
            selected_paragraphs[3],
            (
                config.get("liability_action"),
                config.get("payee_name") or values.get("opponents.names"),
                config.get("claim_amount_wan"),
                config.get("claim_amount_wan"),
                config.get("baseline_amount_wan"),
                str(config.get("rate") or "").rstrip("%"),
            ),
        )
        _replace_fee_blanks(
            selected_paragraphs[4],
            (
                config.get("no_liability_action"),
                config.get("liability_examples"),
                config.get("success_amount_wan"),
            ),
        )
    else:
        _replace_fee_blanks(
            selected_paragraphs[1],
            (
                _amount_yuan(config.get("base_amount")),
                _amount_upper_wan(config.get("base_amount")),
            ),
        )
        guarantee = selected_paragraphs[2]
        if config.get("guarantee_threshold_wan") and config.get("guarantee_fee_wan"):
            _replace_fee_blanks(
                guarantee,
                (
                    config.get("guarantee_threshold_wan"),
                    config.get("guarantee_fee_wan"),
                    config.get("guarantee_threshold_wan"),
                ),
            )
        else:
            _remove_paragraph(guarantee)
        tiers = _risk_tier_items(config)
        destinations = list(selected_paragraphs[3:7])
        while len(destinations) < len(tiers):
            destinations.append(_clone_paragraph_after(destinations[-1]))
        for index, item in enumerate(tiers):
            clause = _risk_tier_clause(item)
            if index == len(tiers) - 1:
                clause = clause[:-1] + "。"
            _set_paragraph_text(destinations[index], clause)
        for paragraph in destinations[len(tiers):]:
            _remove_paragraph(paragraph)

    for paragraph in selected_paragraphs:
        if paragraph._element.getparent() is not None:
            _strip_fee_brackets(paragraph)
    for index, paragraph in enumerate(paragraphs):
        if markers["方案A"] <= index < end and not (
            selected_start <= index < selected_end
        ):
            _remove_paragraph(paragraph)


def _adjust_cover_floating_box(doc: Any) -> None:
    """Keep the template artwork while leaving room for the cover month."""
    for shape in doc._element.xpath(".//*[local-name()='shape']"):
        style = str(shape.get("style") or "")
        if "position:absolute" not in style or "z-index:" not in style:
            continue
        style = re.sub(r"z-index:\s*-?\d+", "z-index:251659264", style)
        # The uploaded cover's law-firm introduction box extends over the date
        # paragraph in WPS/QuickLook. Its text needs less height than the legacy box.
        style = re.sub(r"height:\s*230(?:\.75)?pt", "height:180pt", style)
        shape.set("style", style)
    for element in doc._element.xpath(
        ".//*[local-name()='extent' or local-name()='ext']"
    ):
        if element.get("cy") == "2930525":
            element.set("cy", "2286000")


def _append_word_field(paragraph: Any, instruction: str, cached: str = "1") -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    begin_run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run.append(begin)
    paragraph._p.append(begin_run)

    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} \\* MERGEFORMAT "
    instr_run.append(instr)
    paragraph._p.append(instr_run)

    separate_run = OxmlElement("w:r")
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run.append(separate)
    paragraph._p.append(separate_run)

    value_run = OxmlElement("w:r")
    value_text = OxmlElement("w:t")
    value_text.text = cached
    value_run.append(value_text)
    paragraph._p.append(value_run)

    end_run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    paragraph._p.append(end_run)


def _ensure_dynamic_page_fields(doc: Any) -> None:
    """Replace the template's fixed '共11页' footer with PAGE/NUMPAGES fields."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    footer_roots: list[Any] = []
    seen: set[int] = set()
    for section in doc.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            root = footer._element
            if id(root) not in seen:
                seen.add(id(root))
                footer_roots.append(root)
    for root in footer_roots:
        for paragraph_element in root.xpath(".//*[local-name()='p']"):
            text = "".join(paragraph_element.xpath(".//*[local-name()='t']/text()"))
            if "页 共" not in text or "页" not in text:
                continue
            paragraph = Paragraph(paragraph_element, paragraph_element.getparent())
            for child in list(paragraph_element):
                if child.tag != qn("w:pPr"):
                    paragraph_element.remove(child)
            paragraph.add_run("第 ")
            _append_word_field(paragraph, "PAGE")
            paragraph.add_run(" 页 共 ")
            _append_word_field(paragraph, "NUMPAGES")
            paragraph.add_run(" 页")
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def fill_contract_docx(
    template_path: Path,
    mappings: list[dict[str, Any]],
    values: dict[str, str],
) -> tuple[bytes, list[str]]:
    from docx import Document

    doc = Document(str(template_path))
    missing: list[str] = []
    for mapping in mappings:
        key = mapping.get("semantic_key") or ""
        value = values.get(key, "")
        if mapping.get("required") and not value:
            missing.append(mapping.get("label") or key)
        if not value:
            continue
        paragraph = _get_paragraph(doc, mapping.get("locator") or {})
        if mapping.get("fill_type") == "replace_token":
            _replace_across_runs(paragraph, mapping.get("token_text") or "", value)
        else:
            _fill_blank(paragraph, value)
    _fill_known_contract_labels(doc, values)
    authority_lines = [
        line for line in values.get("engagement.authority_lines", "").splitlines()
        if line.strip()
    ]
    if not authority_lines or not _replace_authority_lines(doc, authority_lines):
        _apply_authorization_checks(
            doc, set(filter(None, values.get("engagement.authorization_ids", "").split(",")))
        )
    _replace_fee_section(doc, values)
    _adjust_cover_floating_box(doc)
    _ensure_dynamic_page_fields(doc)
    output = BytesIO()
    doc.save(output)
    return output.getvalue(), list(dict.fromkeys(missing))


def validate_contract_profile(values: dict[str, str], profile: dict[str, Any]) -> list[str]:
    missing: list[str] = []
    required_values = (
        ("委托人", values.get("client.name")),
        ("代理律师", values.get("engagement.lawyers")),
        ("代理阶段", values.get("case.stage")),
        ("代理权限", values.get("engagement.authorization")),
    )
    for label, value in required_values:
        if not str(value or "").strip():
            missing.append(label)
    plan = str(profile.get("fee_plan") or "fixed")
    config = profile.get("fee_config") or {}
    if plan == "fixed" and not (config.get("base_amount") or profile.get("fee_amount")):
        missing.append("固定律师费")
    elif plan == "staged_fixed" and not _stage_fee_items(config):
        missing.append("至少一个代理阶段及对应律师费")
    elif plan == "base_risk":
        if not (config.get("base_amount") or profile.get("fee_amount")):
            missing.append("基础律师费")
        if not config.get("rate"):
            missing.append("风险代理费比例")
        if not config.get("success_amount_wan"):
            missing.append("无需承担付款责任时的风险代理费")
    elif plan in {"recovery_tiered", "pure_risk"}:
        if plan == "recovery_tiered" and not (
            config.get("base_amount") or profile.get("fee_amount")
        ):
            missing.append("基础律师费")
        if not _risk_tier_items(config):
            missing.append("至少一档回款风险费率")
    return missing


def _contract_profile_fingerprint(case_id: int, template_id: int, profile: dict[str, Any]) -> str:
    relevant_keys = (
        "template_id", "lawyers", "entrust_type", "representation_stage",
        "engagement_scope", "background_matter", "background_stage", "matter_kind",
        "fee_type", "fee_amount", "payment_terms",
        "contract_copies", "sign_date", "firm", "custom_values", "fee_plan",
        "fee_config", "authorizations", "authority_lines", "subject_amount",
    )
    payload = {
        "case_id": int(case_id),
        "template_id": int(template_id),
        "profile": {key: profile.get(key) for key in relevant_keys},
    }
    encoded = json.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _render_contract_pdf(docx_path: Path) -> Path:
    pdf_path = docx_path.with_suffix(".pdf")
    render_docx = docx_path.with_name(f"{docx_path.stem}-render.docx")
    # LibreOffice needs the installed PostScript family aliases to render
    # several Word-only Chinese font names without dropping glyphs.
    replacements = {
        "仿宋": "Noto Sans CJK SC",
        "华文楷体": "Noto Sans CJK SC",
        "宋体": "Noto Sans CJK SC",
    }
    with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(
        render_docx, "w", zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename.endswith(".xml"):
                text = content.decode("utf-8")
                for old, new in replacements.items():
                    text = text.replace(old, new)
                content = text.encode("utf-8")
            target.writestr(info, content)
    rendered_pdf = render_docx.with_suffix(".pdf")
    candidates = [
        Path(os.environ.get("LEGAL_KB_SOFFICE", "")),
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
        Path("/Users/liuzongjin/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice"),
    ]
    soffice = next((path for path in candidates if str(path) and path.is_file()), None)
    if not soffice:
        resolved = shutil.which("soffice")
        soffice = Path(resolved) if resolved else None
    if not soffice:
        raise RuntimeError("未找到合同预览转换组件")
    with tempfile.TemporaryDirectory(prefix="legal-kb-lo-") as profile_dir:
        result = subprocess.run(
            [
                str(soffice),
                f"-env:UserInstallation=file://{profile_dir}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(docx_path.parent),
                str(render_docx),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=45,
            check=False,
        )
    if result.returncode != 0 or not rendered_pdf.is_file():
        render_docx.unlink(missing_ok=True)
        detail = (result.stderr or result.stdout or "转换失败").strip()
        raise RuntimeError(detail[-500:])
    os.replace(rendered_pdf, pdf_path)
    render_docx.unlink(missing_ok=True)
    return pdf_path


def _render_contract_html(docx_path: Path) -> Path:
    html_path = docx_path.with_suffix(".html")
    with tempfile.TemporaryDirectory(prefix="legal-kb-ql-") as output_dir:
        result = subprocess.run(
            ["/usr/bin/qlmanage", "-p", "-o", output_dir, str(docx_path)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            timeout=30,
            check=False,
        )
        previews = list(Path(output_dir).glob("*.qlpreview/Preview.html"))
        if result.returncode != 0 or not previews:
            detail = (result.stderr or result.stdout or "预览转换失败").strip()
            raise RuntimeError(detail[-500:])
        source_html = previews[0]
        content = source_html.read_text(encoding="utf-8")
        content = re.sub(
            r"line-height:-([0-9]+(?:\.[0-9]+)?)",
            lambda match: f"line-height:{float(match.group(1)) * 10:g}px",
            content,
        )
        content = content.replace(
            '<div class="s1">',
            '<div class="s1 contract-preview-page">',
            1,
        )
        content = re.sub(
            r'<span([^>]*)>\f</span></p>',
            '</p></div><div class="s1 contract-preview-page">',
            content,
        )
        bundle = source_html.parent
        for attachment in bundle.iterdir():
            if not attachment.is_file() or attachment.name == "Preview.html":
                continue
            mime = mimetypes.guess_type(attachment.name)[0] or "application/octet-stream"
            encoded = base64.b64encode(attachment.read_bytes()).decode("ascii")
            content = content.replace(
                f'src="{attachment.name}"',
                f'src="data:{mime};base64,{encoded}"',
            )
        preview_style = """
        <style>
          html { background:#e8ebf0; scroll-behavior:auto; }
          body { margin:0!important; padding:14px 0 40px!important; background:#e8ebf0; }
          body, body * { font-synthesis:none; text-rendering:geometricPrecision; }
          .contract-preview-page { box-sizing:border-box!important; width:595px!important;
            height:842px!important; min-height:842px!important; margin:0 auto 18px!important;
            background:#fff; box-shadow:0 3px 18px rgba(15,23,42,.18);
            overflow:hidden!important; position:relative!important; }
          .contract-preview-page.contract-flow-page { padding-top:68px!important;
            padding-bottom:142px!important; }
          .contract-running-header { position:absolute; z-index:20; top:13px; left:70px; right:70px;
            height:43px; border-bottom:1px solid #222; display:flex; align-items:center; }
          .contract-running-header img { display:block; width:250px!important; max-height:38px!important;
            object-fit:contain; object-position:left center; }
          .contract-running-footer { position:absolute; z-index:20; bottom:14px; left:0; right:0;
            color:#111; font-family:SimSun,Songti SC,serif; font-size:11px; text-align:center; }
          [data-contract-section] { transition:background-color .18s ease, outline-color .18s ease; }
          [data-contract-section].contract-focus { background:#fff4a8!important;
            outline:3px solid rgba(190,122,18,.38); outline-offset:3px; border-radius:2px; }
          .contract-cover-date { position:relative!important; z-index:251659265!important; }
        </style>
        <script>
        (() => {
          const sectionRules = [
            ['cover', text => text.includes('民事委托代理合同')],
            ['case', text => text.trim().startsWith('甲方：')],
            ['client-contact', text => text.includes('送达地址') && text.includes('电子邮箱')],
            ['background', text => text.includes('甲方因') && text.includes('聘请乙方')],
            ['matter', text => text.includes('第一条') && text.includes('委托代理事项')],
            ['matter-court', text => text.includes('受理机关：')],
            ['matter-subject', text => text.includes('涉及标的：')],
            ['authority', text => text.includes('第二条') && text.includes('委托代理权限')],
            ['lawyer', text => text.includes('乙方委派') && text.includes('作为上述案件')],
            ['fee', text => text.includes('第五条') && (text.includes('收费') || text.includes('律师代理费'))],
            ['sign', text => text.includes('第十条') && text.includes('合同的生效')]
          ];
          function removeLegacyFooters() {
            for (const element of document.querySelectorAll('p,div')) {
              const text = (element.textContent || '').replace(/\\s+/g, ' ').trim();
              if (/^第\\s*\\d+\\s*页\\s*共\\s*\\d+\\s*页$/.test(text) && element.children.length < 8) {
                element.remove();
              }
            }
          }
          function paginate() {
            const initialPages = [...document.querySelectorAll('.contract-preview-page')];
            if (!initialPages.length) return;
            const logo = document.querySelector('.contract-preview-page img');
            for (let pageIndex = 1; pageIndex < initialPages.length; pageIndex++) {
              let source = initialPages[pageIndex];
              const children = [...source.children].filter(node =>
                !node.classList?.contains('contract-running-header') &&
                !node.classList?.contains('contract-running-footer'));
              source.replaceChildren();
              source.classList.add('contract-flow-page');
              let current = source;
              let guard = 0;
              for (const child of children) {
                if (++guard > 2500) break;
                current.appendChild(child);
                if (current.scrollHeight > current.clientHeight && current.children.length > 1) {
                  current.removeChild(child);
                  const next = source.cloneNode(false);
                  current.after(next);
                  current = next;
                  current.appendChild(child);
                  if (document.querySelectorAll('.contract-preview-page').length > 50) break;
                }
              }
            }
            const pages = [...document.querySelectorAll('.contract-preview-page')];
            pages.forEach((page,index) => {
              if (index > 0 && logo) {
                const header = document.createElement('div');
                header.className = 'contract-running-header';
                header.appendChild(logo.cloneNode(true));
                page.appendChild(header);
              }
              if (index > 0) {
                const footer = document.createElement('div');
                footer.className = 'contract-running-footer';
                footer.textContent = `第 ${index + 1} 页 共 ${pages.length} 页`;
                page.appendChild(footer);
              }
            });
          }
          function prepare() {
            const claimed = new Set();
            for (const element of document.querySelectorAll(
              '.contract-preview-page p, .contract-preview-page div.s16, .contract-preview-page div.s33'
            )) {
              const text = (element.textContent || '').replace(/\\s+/g, ' ');
              if (/^[二〇零一二三四五六七八九十]+年[一二三四五六七八九十]+月$/.test(text.trim())) {
                element.classList.add('contract-cover-date');
              }
              for (const [name, match] of sectionRules) {
                if (!claimed.has(name) && match(text)) {
                  element.dataset.contractSection = name;
                  claimed.add(name);
                  break;
                }
              }
            }
          }
          function focusSection(name) {
            document.querySelectorAll('.contract-focus').forEach(node =>
              node.classList.remove('contract-focus'));
            const target = document.querySelector(`[data-contract-section="${name}"]`);
            if (!target) return;
            target.classList.add('contract-focus');
            const top = target.getBoundingClientRect().top + window.scrollY - 58;
            window.scrollTo({top:Math.max(0, top), behavior:'auto'});
          }
          addEventListener('DOMContentLoaded', () => {
            removeLegacyFooters();
            paginate();
            prepare();
            parent.postMessage({type:'contract-preview-ready'}, '*');
          });
          addEventListener('message', event => {
            if (event.data && event.data.type === 'contract-focus') {
              focusSection(event.data.section);
            }
          });
        })();
        </script>
        """
        content = content.replace("</head>", f"{preview_style}</head>", 1)
        html_path.write_text(content, encoding="utf-8")
    return html_path


def _render_contract_exact_html(docx_path: Path) -> Path:
    """Render the generated DOCX page-for-page, then add non-destructive focus overlays."""
    import fitz

    pdf_path = _render_contract_pdf(docx_path)
    html_path = docx_path.with_suffix(".html")
    document = fitz.open(pdf_path)
    section_needles = (
        ("cover", ("民事委托代理合同",)),
        ("case", ("甲方：",)),
        ("background", ("甲方因",)),
        ("matter", ("第一条", "委托代理事项")),
        ("authority", ("第二条", "委托代理权限")),
        ("lawyer", ("乙方委派",)),
        ("fee", ("第五条",)),
        ("sign", ("第十条",)),
    )
    claimed: set[str] = set()
    pages: list[str] = []
    scale = 1.65
    matrix = fitz.Matrix(scale, scale)
    for page_index, page in enumerate(document):
        pixmap = page.get_pixmap(matrix=matrix, alpha=False)
        encoded = base64.b64encode(pixmap.tobytes("png")).decode("ascii")
        overlays: list[str] = []
        page_text = page.get_text("text")
        for section, needles in section_needles:
            if section in claimed or not all(needle in page_text for needle in needles):
                continue
            rectangles = page.search_for(needles[0])
            if not rectangles:
                continue
            rect = rectangles[0]
            claimed.add(section)
            left = max(0.0, rect.x0 / page.rect.width * 100)
            top = max(0.0, rect.y0 / page.rect.height * 100)
            width = min(100.0 - left, max(14.0, rect.width / page.rect.width * 100))
            height = max(2.2, rect.height / page.rect.height * 100)
            overlays.append(
                f'<div class="contract-focus-target" data-contract-section="{section}" '
                f'style="left:{left:.4f}%;top:{top:.4f}%;width:{width:.4f}%;height:{height:.4f}%"></div>'
            )
        pages.append(
            '<section class="contract-rendered-page">'
            f'<img alt="合同第{page_index + 1}页" src="data:image/png;base64,{encoded}">'
            f'{"".join(overlays)}</section>'
        )
    document.close()
    content = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8">
<style>
html,body{{margin:0;background:#e8ebf0;scroll-behavior:auto}}
body{{padding:14px 0 38px}}
.contract-rendered-page{{position:relative;width:min(94vw,794px);margin:0 auto 18px;
  background:#fff;box-shadow:0 3px 18px rgba(15,23,42,.18);line-height:0}}
.contract-rendered-page img{{display:block;width:100%;height:auto}}
.contract-focus-target{{position:absolute;z-index:3;pointer-events:none;border-radius:3px;
  opacity:0;transition:opacity .15s ease,background-color .15s ease}}
.contract-focus-target.contract-focus{{opacity:1;background:rgba(255,225,64,.42);
  outline:3px solid rgba(190,122,18,.42);outline-offset:2px}}
</style></head><body>{"".join(pages)}
<script>
function focusSection(name){{
  document.querySelectorAll('.contract-focus').forEach(node=>node.classList.remove('contract-focus'));
  const target=document.querySelector(`[data-contract-section="${{name}}"]`);
  if(!target)return;
  target.classList.add('contract-focus');
  const top=target.getBoundingClientRect().top+window.scrollY
    -Math.max(18,(window.innerHeight-target.offsetHeight)/2);
  window.scrollTo({{top:Math.max(0,top),behavior:'auto'}});
}}
addEventListener('DOMContentLoaded',()=>parent.postMessage({{type:'contract-preview-ready'}},'*'));
addEventListener('message',event=>{{
  if(event.data&&event.data.type==='contract-focus')focusSection(event.data.section);
}});
</script></body></html>"""
    html_path.write_text(content, encoding="utf-8")
    return html_path


def register_case_automation_routes(
    app: Any,
    get_db: Callable[[], Any],
    *,
    add_audit_log: Callable[..., Any] | None = None,
) -> None:
    ensure_case_automation_schema(get_db)
    # Apply the directory-name migration eagerly for every currently available
    # binding.  Offline bindings are retried the next time they come online.
    with get_db() as conn:
        bound_case_ids = [int(row["case_id"]) for row in conn.execute(
            "SELECT case_id FROM case_manager_folder_bindings"
        ).fetchall()]
    for bound_case_id in bound_case_ids:
        get_folder_binding(get_db, bound_case_id)
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)

    @app.get("/api/case-manager/cases/{case_id}/folder-binding")
    def read_folder_binding(case_id: int):
        return get_folder_binding(get_db, case_id)

    def require_series_child_binding(case_id: int, unbound_detail: str):
        """Validate a series child and return its row plus the parent binding."""
        with get_db() as conn:
            case_row = conn.execute(
                """SELECT c.id,c.title,c.case_kind,c.parent_series_id,p.title AS parent_title
                   FROM case_manager_cases c LEFT JOIN case_manager_cases p ON p.id=c.parent_series_id
                   WHERE c.id=?""",
                (case_id,),
            ).fetchone()
        if not case_row:
            raise HTTPException(status_code=404, detail="案件不存在")
        if case_row["case_kind"] != "series_child" or not int(case_row["parent_series_id"] or 0):
            raise HTTPException(
                status_code=400,
                detail="只有系列案件下的子案才能在系列文件夹内创建或选择文件夹",
            )
        try:
            parent_binding = get_folder_binding(get_db, int(case_row["parent_series_id"]), required=True)
        except HTTPException as exc:
            raise HTTPException(status_code=exc.status_code, detail=unbound_detail) from exc
        if not parent_binding.get("available"):
            raise HTTPException(status_code=409, detail=unbound_detail)
        return case_row, parent_binding

    @app.put("/api/case-manager/cases/{case_id}/folder-binding")
    async def save_folder_binding(case_id: int, request: Request):
        body = await request.json()
        with get_db() as conn:
            if not _case_exists(conn, case_id):
                raise HTTPException(status_code=404, detail="案件不存在")
        real, stat = validate_case_folder(str(body.get("folder_path") or ""))
        bind_case_folder(get_db, case_id, str(body.get("folder_path") or real), real, stat)
        return get_folder_binding(get_db, case_id)

    @app.post("/api/case-manager/cases/{case_id}/folder-binding/create-in-series")
    async def create_series_child_folder_binding(case_id: int, request: Request):
        """Create a subfolder named after the child inside the series folder and bind it.

        Lets every sub-case of a series case keep its own delivery archive without
        asking the user to pick the same parent tree again in the folder dialog.
        """
        body = await request.json()
        case_row, parent_binding = require_series_child_binding(
            case_id, "请先为系列案件绑定可用的本地文件夹，再为子案创建文件夹"
        )
        name = (
            safe_subfolder_name(str(body.get("name") or ""))
            or series_child_folder_name(case_row["parent_title"], case_row["title"])
            or f"子案-{case_id}"
        )
        root = Path(parent_binding["real_path"])
        directory = _unique_destination(root, name)
        try:
            directory.mkdir(parents=True)
        except OSError as exc:
            raise HTTPException(status_code=500, detail=f"创建子案文件夹失败：{exc}") from exc
        real = directory.resolve()
        bind_case_folder(get_db, case_id, str(directory), real, real.stat())
        return get_folder_binding(get_db, case_id)

    @app.get("/api/case-manager/cases/{case_id}/folder-binding/series-options")
    def series_folder_options(case_id: int):
        """List existing first-level folders inside the series parent's folder."""
        with get_db() as conn:
            case_row = conn.execute(
                "SELECT case_kind,parent_series_id FROM case_manager_cases WHERE id=?",
                (case_id,),
            ).fetchone()
        if not case_row:
            raise HTTPException(status_code=404, detail="案件不存在")
        if case_row["case_kind"] != "series_child" or not int(case_row["parent_series_id"] or 0):
            raise HTTPException(status_code=400, detail="只有系列案件下的子案可以查看系列文件夹")
        parent = get_folder_binding(get_db, int(case_row["parent_series_id"]), required=False)
        if not (parent.get("bound") and parent.get("available")):
            return {"bound": False, "real_path": "", "folders": []}
        root = Path(parent["real_path"])
        reserved = {COURT_DOCUMENT_SUBDIR, CONTRACT_SUBDIR}
        folders: list[dict[str, str]] = []
        try:
            for entry in sorted(root.iterdir()):
                if entry.name in reserved or entry.name.startswith("."):
                    continue
                # Only real directories; symlinks could point outside the series folder.
                if not entry.is_dir() or entry.is_symlink():
                    continue
                folders.append({"name": entry.name, "path": str(entry)})
        except OSError:
            folders = []
        return {"bound": True, "real_path": str(root), "folders": folders[:200]}

    @app.post("/api/case-manager/cases/{case_id}/folder-binding/bind-in-series")
    async def bind_existing_series_folder(case_id: int, request: Request):
        """Bind one of the existing folders inside the series folder to this child."""
        body = await request.json()
        case_row, parent_binding = require_series_child_binding(
            case_id, "请先为系列案件绑定可用的本地文件夹，再为子案选择文件夹"
        )
        name = safe_subfolder_name(str(body.get("name") or ""))
        if not name:
            raise HTTPException(status_code=400, detail="请选择总文件夹内的文件夹")
        root = Path(parent_binding["real_path"])
        target = root / name
        if not target.is_dir() or target.is_symlink():
            raise HTTPException(status_code=404, detail="总文件夹内找不到该子文件夹")
        real = target.resolve()
        bind_case_folder(get_db, case_id, str(target), real, real.stat())
        return get_folder_binding(get_db, case_id)

    @app.delete("/api/case-manager/cases/{case_id}/folder-binding")
    def delete_folder_binding(case_id: int):
        with get_db() as conn:
            cur = conn.execute("DELETE FROM case_manager_folder_bindings WHERE case_id=?", (case_id,))
            try:
                conn.execute("DELETE FROM case_manager_local_file_index WHERE case_id=?", (case_id,))
            except sqlite3.OperationalError:
                pass
            conn.commit()
        return {"ok": True, "deleted": bool(cur.rowcount)}

    @app.post("/api/case-manager/cases/{case_id}/folder-binding/select")
    def select_folder_binding(case_id: int):
        with get_db() as conn:
            if not _case_exists(conn, case_id):
                raise HTTPException(status_code=404, detail="案件不存在")
        if shutil.which("osascript") is None:
            raise HTTPException(status_code=501, detail="当前系统不支持系统文件夹选择器，请手动粘贴路径")
        script = 'POSIX path of (choose folder with prompt "选择案件文件夹")'
        result = subprocess.run(
            ["osascript", "-e", script], text=True, capture_output=True, timeout=120
        )
        if result.returncode != 0:
            if "-128" in result.stderr:
                return {"ok": False, "cancelled": True}
            raise HTTPException(status_code=500, detail=result.stderr.strip() or "选择文件夹失败")
        return {"ok": True, "folder_path": result.stdout.strip().rstrip("/")}

    @app.post("/api/case-manager/cases/{case_id}/folder-binding/open")
    def open_folder_binding(case_id: int):
        binding = get_folder_binding(get_db, case_id, required=True)
        if shutil.which("open") is None:
            raise HTTPException(status_code=501, detail="当前系统不支持 Finder")
        subprocess.run(["open", binding["real_path"]], check=True)
        return {"ok": True}

    @app.get("/api/case-manager/cases/{case_id}/files")
    def list_case_files(case_id: int):
        with get_db() as conn:
            rows = conn.execute(
                "SELECT * FROM case_manager_files WHERE case_id=? ORDER BY datetime(created_at) DESC,id DESC",
                (case_id,),
            ).fetchall()
        items = []
        for row in rows:
            item = dict(row)
            item["metadata"] = _loads(item.pop("metadata_json", "{}"), {})
            item["exists"] = Path(item["absolute_path"]).is_file()
            items.append(item)
        return {"items": items}

    @app.post("/api/contract-templates")
    async def upload_contract_template(name: str, file: UploadFile = File(...)):
        filename = file.filename or ""
        if Path(filename).suffix.lower() != ".docx":
            raise HTTPException(status_code=400, detail="仅支持未加密的 .docx 模板")
        content = await file.read()
        if len(content) > MAX_TEMPLATE_BYTES:
            raise HTTPException(status_code=413, detail="模板文件不能超过20MB")
        digest = _hash_bytes(content)
        with get_db() as conn:
            existing = conn.execute("SELECT * FROM contract_templates WHERE sha256=?", (digest,)).fetchone()
            if existing:
                return dict(existing)
        stored_path = TEMPLATE_DIR / f"{uuid.uuid4().hex}.docx"
        stored_path.write_bytes(content)
        try:
            from docx import Document
            Document(str(stored_path))
        except Exception as exc:
            stored_path.unlink(missing_ok=True)
            raise HTTPException(status_code=400, detail="模板损坏、加密或不是有效的 DOCX 文件") from exc
        stamp = _now()
        with get_db() as conn:
            version = int(conn.execute(
                "SELECT COALESCE(MAX(version),0)+1 FROM contract_templates WHERE name=?", (name,)
            ).fetchone()[0])
            conn.execute("UPDATE contract_templates SET active=0 WHERE name=?", (name,))
            cur = conn.execute(
                """INSERT INTO contract_templates
                   (name,original_filename,stored_path,sha256,version,active,created_at,updated_at)
                   VALUES (?,?,?,?,?,1,?,?)""",
                (name.strip() or Path(filename).stem, filename, str(stored_path), digest, version, stamp, stamp),
            )
            conn.commit()
            template_id = int(cur.lastrowid)
        return {"id": template_id, "name": name, "version": version, "analysis_status": "pending"}

    @app.get("/api/contract-templates")
    def list_contract_templates():
        with get_db() as conn:
            rows = conn.execute(
                """SELECT t.*,(SELECT COUNT(*) FROM contract_template_mappings m
                   WHERE m.template_id=t.id) AS mapping_count
                   FROM contract_templates t ORDER BY t.active DESC,t.name,t.version DESC"""
            ).fetchall()
        return {"items": [dict(row) for row in rows]}

    @app.post("/api/contract-templates/{template_id}/analyze")
    def analyze_contract_template(template_id: int):
        with get_db() as conn:
            template = conn.execute("SELECT * FROM contract_templates WHERE id=?", (template_id,)).fetchone()
        if not template:
            raise HTTPException(status_code=404, detail="合同模板不存在")
        structure, mappings, warnings = analyze_docx(Path(template["stored_path"]))
        stamp = _now()
        with get_db() as conn:
            conn.execute("DELETE FROM contract_template_mappings WHERE template_id=?", (template_id,))
            for item in mappings:
                conn.execute(
                    """INSERT INTO contract_template_mappings
                       (template_id,semantic_key,label,locator_json,fill_type,token_text,required,
                        confidence,confirmed,sort_order,created_at,updated_at)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        template_id, item["semantic_key"], item["label"], _json(item["locator"]),
                        item["fill_type"], item["token_text"], int(item["required"]),
                        item["confidence"], 0, item["sort_order"], stamp, stamp,
                    ),
                )
            conn.execute(
                """UPDATE contract_templates SET analysis_status='review',structure_json=?,
                   warnings_json=?,updated_at=? WHERE id=?""",
                (_json(structure), _json(warnings), stamp, template_id),
            )
            conn.commit()
        return {"template_id": template_id, "mappings": mappings, "warnings": warnings}

    @app.get("/api/contract-templates/{template_id}/mappings")
    def get_template_mappings(template_id: int):
        with get_db() as conn:
            rows = conn.execute(
                "SELECT * FROM contract_template_mappings WHERE template_id=? ORDER BY sort_order,id",
                (template_id,),
            ).fetchall()
        items = []
        for row in rows:
            item = dict(row)
            item["locator"] = _loads(item.pop("locator_json"), {})
            items.append(item)
        return {"items": items, "semantic_fields": SEMANTIC_FIELDS}

    @app.put("/api/contract-templates/{template_id}/mappings")
    async def save_template_mappings(template_id: int, request: Request):
        body = await request.json()
        items = body.get("items") or []
        stamp = _now()
        with get_db() as conn:
            if not conn.execute("SELECT 1 FROM contract_templates WHERE id=?", (template_id,)).fetchone():
                raise HTTPException(status_code=404, detail="合同模板不存在")
            conn.execute("DELETE FROM contract_template_mappings WHERE template_id=?", (template_id,))
            for order, item in enumerate(items):
                key = str(item.get("semantic_key") or "")
                if not key:
                    continue
                conn.execute(
                    """INSERT INTO contract_template_mappings
                       (template_id,semantic_key,label,locator_json,fill_type,token_text,required,
                        confidence,confirmed,sort_order,created_at,updated_at)
                       VALUES (?,?,?,?,?,?,?,?,1,?,?,?)""",
                    (
                        template_id, key, str(item.get("label") or SEMANTIC_FIELDS.get(key, {}).get("label") or key),
                        _json(item.get("locator") or {}), str(item.get("fill_type") or "replace_token"),
                        str(item.get("token_text") or ""), int(bool(item.get("required"))),
                        float(item.get("confidence") or 1), order, stamp, stamp,
                    ),
                )
            conn.execute(
                "UPDATE contract_templates SET analysis_status='ready',updated_at=? WHERE id=?",
                (stamp, template_id),
            )
            conn.commit()
        return get_template_mappings(template_id)

    def _read_profile(conn: Any, case_id: int) -> dict[str, Any]:
        row = conn.execute("SELECT * FROM case_manager_contract_profiles WHERE case_id=?", (case_id,)).fetchone()
        if not row:
            return {
                "case_id": case_id,
                "lawyers": [],
                "entrust_type": "",
                "background_matter": "",
                "background_stage": "",
                "matter_kind": "case",
                "firm": dict(DEFAULT_FIRM),
                "custom_values": {},
                "fee_plan": "fixed",
                "fee_config": {},
                "authorizations": [],
                "authority_lines": [],
                "contract_copies": "3",
                "sign_date": dt.date.today().isoformat(),
            }
        item = dict(row)
        item["lawyers"] = _loads(item.pop("lawyers_json"), [])
        item["firm"] = {**DEFAULT_FIRM, **_loads(item.pop("firm_json"), {})}
        item["custom_values"] = _loads(item.pop("custom_values_json"), {})
        item["fee_config"] = _loads(item.pop("fee_config_json", "{}"), {})
        item["authorizations"] = _loads(item.pop("authorization_json", "[]"), [])
        item["authority_lines"] = _loads(item.pop("authority_lines_json", "[]"), [])
        item["fee_plan"] = item.get("fee_plan") or "fixed"
        item["matter_kind"] = (
            "matter" if item.get("matter_kind") == "matter" else "case"
        )
        item["contract_copies"] = item.get("contract_copies") or "3"
        return item

    @app.get("/api/case-manager/cases/{case_id}/contract-profile")
    def get_contract_profile(case_id: int):
        with get_db() as conn:
            if not _case_exists(conn, case_id):
                raise HTTPException(status_code=404, detail="案件不存在")
            profile = _read_profile(conn, case_id)
            profile["context"] = contract_context(conn, case_id, profile)
            return profile

    @app.get("/api/contract-fee-plans")
    def get_contract_fee_plans():
        return {
            "fee_plans": FEE_PLAN_DEFINITIONS,
            "risk_methods": RISK_METHOD_DEFINITIONS,
            "authorization_options": AUTHORIZATION_OPTIONS,
        }

    @app.get("/api/contract-profile-presets")
    def get_contract_profile_presets():
        with get_db() as conn:
            rows = conn.execute(
                """SELECT id,name,source_case_id,profile_json,created_at,updated_at
                   FROM contract_profile_presets ORDER BY updated_at DESC,id DESC"""
            ).fetchall()
        return {
            "items": [
                {
                    **{key: value for key, value in dict(row).items() if key != "profile_json"},
                    "profile": _loads(row["profile_json"], {}),
                }
                for row in rows
            ]
        }

    @app.post("/api/contract-profile-presets")
    async def save_contract_profile_preset(request: Request):
        body = await request.json()
        name = str(body.get("name") or "").strip()
        if not name:
            raise HTTPException(status_code=400, detail="请填写合同方案名称")
        profile = body.get("profile") or {}
        preset_id = int(body.get("id") or 0)
        source_case_id = int(body.get("source_case_id") or 0) or None
        stamp = _now()
        with get_db() as conn:
            if preset_id:
                row = conn.execute(
                    "SELECT 1 FROM contract_profile_presets WHERE id=?", (preset_id,)
                ).fetchone()
                if not row:
                    raise HTTPException(status_code=404, detail="合同方案不存在")
                conn.execute(
                    """UPDATE contract_profile_presets
                       SET name=?,source_case_id=?,profile_json=?,updated_at=? WHERE id=?""",
                    (name, source_case_id, _json(profile), stamp, preset_id),
                )
            else:
                cur = conn.execute(
                    """INSERT INTO contract_profile_presets
                       (name,source_case_id,profile_json,created_at,updated_at)
                       VALUES (?,?,?,?,?)""",
                    (name, source_case_id, _json(profile), stamp, stamp),
                )
                preset_id = int(cur.lastrowid)
            conn.commit()
        return {"id": preset_id, "name": name, "source_case_id": source_case_id, "profile": profile}

    @app.delete("/api/contract-profile-presets/{preset_id}")
    def delete_contract_profile_preset(preset_id: int):
        with get_db() as conn:
            conn.execute("DELETE FROM contract_profile_presets WHERE id=?", (preset_id,))
            conn.commit()
        return {"ok": True}

    @app.get("/api/case-manager/cases/{case_id}/authorization-material-context")
    def get_authorization_material_context(case_id: int):
        with get_db() as conn:
            if not _case_exists(conn, case_id):
                raise HTTPException(status_code=404, detail="案件不存在")
            profile = _read_profile(conn, case_id)
            return contract_context(conn, case_id, profile)

    @app.put("/api/case-manager/cases/{case_id}/contract-profile")
    async def save_contract_profile(case_id: int, request: Request):
        body = await request.json()
        stamp = _now()
        with get_db() as conn:
            if not _case_exists(conn, case_id):
                raise HTTPException(status_code=404, detail="案件不存在")
            conn.execute(
                """INSERT INTO case_manager_contract_profiles
                   (case_id,template_id,lawyers_json,entrust_type,representation_stage,engagement_scope,
                    background_matter,background_stage,matter_kind,
                    fee_type,fee_amount,payment_terms,contract_copies,sign_date,firm_json,
                    custom_values_json,fee_plan,fee_config_json,authorization_json,authority_lines_json,subject_amount,
                    created_at,updated_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(case_id) DO UPDATE SET
                    template_id=excluded.template_id,lawyers_json=excluded.lawyers_json,
                    entrust_type=excluded.entrust_type,
                    representation_stage=excluded.representation_stage,
                    engagement_scope=excluded.engagement_scope,fee_type=excluded.fee_type,
                    background_matter=excluded.background_matter,
                    background_stage=excluded.background_stage,
                    matter_kind=excluded.matter_kind,
                    fee_amount=excluded.fee_amount,payment_terms=excluded.payment_terms,
                    contract_copies=excluded.contract_copies,sign_date=excluded.sign_date,
                    firm_json=excluded.firm_json,custom_values_json=excluded.custom_values_json,
                    fee_plan=excluded.fee_plan,fee_config_json=excluded.fee_config_json,
                    authorization_json=excluded.authorization_json,
                    authority_lines_json=excluded.authority_lines_json,
                    subject_amount=excluded.subject_amount,
                    updated_at=excluded.updated_at""",
                (
                    case_id, body.get("template_id"), _json(body.get("lawyers") or []),
                    str(body.get("entrust_type") or ""),
                    str(body.get("representation_stage") or ""), str(body.get("engagement_scope") or ""),
                    str(body.get("background_matter") or ""),
                    str(body.get("background_stage") or ""),
                    "matter" if str(body.get("matter_kind") or "case") == "matter" else "case",
                    str(body.get("fee_type") or ""), str(body.get("fee_amount") or ""),
                    str(body.get("payment_terms") or ""), str(body.get("contract_copies") or ""),
                    str(body.get("sign_date") or ""), _json(body.get("firm") or {}),
                    _json(body.get("custom_values") or {}),
                    str(body.get("fee_plan") or "fixed"),
                    _json(body.get("fee_config") or {}),
                    _json(body.get("authorizations") or []),
                    _json(body.get("authority_lines") or []),
                    str(body.get("subject_amount") or ""),
                    stamp, stamp,
                ),
            )
            conn.commit()
            return _read_profile(conn, case_id)

    def _template_and_mappings(conn: Any, template_id: int) -> tuple[dict[str, Any], list[dict[str, Any]]]:
        template = conn.execute("SELECT * FROM contract_templates WHERE id=?", (template_id,)).fetchone()
        if not template:
            raise HTTPException(status_code=404, detail="合同模板不存在")
        rows = conn.execute(
            """SELECT * FROM contract_template_mappings
               WHERE template_id=? AND confirmed=1 ORDER BY sort_order,id""", (template_id,)
        ).fetchall()
        mappings = []
        for row in rows:
            item = dict(row)
            item["locator"] = _loads(item.pop("locator_json"), {})
            mappings.append(item)
        return dict(template), mappings

    @app.post("/api/case-manager/cases/{case_id}/contracts/preview")
    async def preview_contract(case_id: int, request: Request):
        body = await request.json()
        template_id = int(body.get("template_id") or 0)
        supersedes_token = str(body.get("supersedes_token") or "")
        with get_db() as conn:
            profile = _read_profile(conn, case_id)
            profile.update(body.get("profile") or {})
            template, mappings = _template_and_mappings(conn, template_id)
            values = build_contract_values(conn, case_id, profile)
        content, missing = fill_contract_docx(Path(template["stored_path"]), mappings, values)
        missing = list(dict.fromkeys([*missing, *validate_contract_profile(values, profile)]))
        token = uuid.uuid4().hex
        preview_path = PREVIEW_DIR / f"{token}.docx"
        preview_path.write_bytes(content)
        try:
            html_path = _render_contract_html(preview_path)
        except Exception as exc:
            preview_path.unlink(missing_ok=True)
            raise HTTPException(status_code=500, detail=f"无法生成合同实时预览：{exc}") from exc
        fingerprint = _contract_profile_fingerprint(case_id, template_id, profile)
        stamp = _now()
        with get_db() as conn:
            conn.execute(
                """INSERT INTO case_manager_contract_previews
                   (token,case_id,template_id,file_path,profile_fingerprint,values_json,missing_json,created_at)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (
                    token, case_id, template_id, str(preview_path), fingerprint,
                    _json(values), _json(missing), stamp,
                ),
            )
            if supersedes_token and supersedes_token != token:
                old = conn.execute(
                    """SELECT file_path FROM case_manager_contract_previews
                       WHERE token=? AND case_id=?""",
                    (supersedes_token, case_id),
                ).fetchone()
                if old:
                    old_path = Path(old["file_path"])
                    old_path.unlink(missing_ok=True)
                    old_path.with_suffix(".pdf").unlink(missing_ok=True)
                    old_path.with_suffix(".html").unlink(missing_ok=True)
                    conn.execute(
                        "DELETE FROM case_manager_contract_previews WHERE token=?",
                        (supersedes_token,),
                    )
            conn.commit()
        return {
            "preview_token": token,
            "download_url": f"/api/case-manager/cases/{case_id}/contracts/preview/{token}",
            "pdf_url": f"/api/case-manager/cases/{case_id}/contracts/preview/{token}/pdf",
            "html_url": f"/api/case-manager/cases/{case_id}/contracts/preview/{token}/html",
            "profile_fingerprint": fingerprint,
            "values": values,
            "missing": missing,
            "can_export": not missing,
        }

    @app.get("/api/case-manager/cases/{case_id}/contracts/preview/{token}")
    def download_contract_preview(case_id: int, token: str):
        with get_db() as conn:
            row = conn.execute(
                "SELECT * FROM case_manager_contract_previews WHERE token=? AND case_id=?",
                (token, case_id),
            ).fetchone()
        if not row or not Path(row["file_path"]).is_file():
            raise HTTPException(status_code=404, detail="预览文件不存在或已经失效")
        return FileResponse(
            row["file_path"],
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename="委托代理合同_预览.docx",
        )

    @app.get("/api/case-manager/cases/{case_id}/contracts/preview/{token}/pdf")
    def view_contract_preview_pdf(case_id: int, token: str):
        with get_db() as conn:
            row = conn.execute(
                "SELECT * FROM case_manager_contract_previews WHERE token=? AND case_id=?",
                (token, case_id),
            ).fetchone()
        pdf_path = Path(row["file_path"]).with_suffix(".pdf") if row else Path("")
        if not row or not pdf_path.is_file():
            raise HTTPException(status_code=404, detail="合同实时预览不存在或已经失效")
        return FileResponse(pdf_path, media_type="application/pdf")

    @app.get("/api/case-manager/cases/{case_id}/contracts/preview/{token}/html")
    def view_contract_preview_html(case_id: int, token: str):
        with get_db() as conn:
            row = conn.execute(
                "SELECT * FROM case_manager_contract_previews WHERE token=? AND case_id=?",
                (token, case_id),
            ).fetchone()
        html_path = Path(row["file_path"]).with_suffix(".html") if row else Path("")
        if not row or not html_path.is_file():
            raise HTTPException(status_code=404, detail="合同实时预览不存在或已经失效")
        return FileResponse(html_path, media_type="text/html; charset=utf-8")

    @app.post("/api/case-manager/cases/{case_id}/contracts/export")
    async def export_contract(case_id: int, request: Request):
        body = await request.json()
        token = str(body.get("preview_token") or "")
        with get_db() as conn:
            preview = conn.execute(
                "SELECT * FROM case_manager_contract_previews WHERE token=? AND case_id=?",
                (token, case_id),
            ).fetchone()
            case = conn.execute("SELECT title FROM case_manager_cases WHERE id=?", (case_id,)).fetchone()
            current_profile = _read_profile(conn, case_id)
        if not preview or not case:
            raise HTTPException(status_code=404, detail="合同预览不存在")
        current_profile.update(body.get("profile") or {})
        current_fingerprint = _contract_profile_fingerprint(
            case_id, int(preview["template_id"]), current_profile
        )
        if not preview["profile_fingerprint"] or preview["profile_fingerprint"] != current_fingerprint:
            raise HTTPException(
                status_code=409,
                detail="左侧内容已经变化，请等待右侧预览更新完成后再归档",
            )
        missing = _loads(preview["missing_json"], [])
        if missing:
            raise HTTPException(status_code=409, detail=f"以下必填字段尚未填写：{'、'.join(missing)}")
        preview_path = Path(preview["file_path"])
        if not preview_path.is_file():
            raise HTTPException(status_code=404, detail="合同预览文件已经失效，请重新预览")
        filename = f"{dt.date.today():%Y%m%d}_委托代理合同_{safe_filename(case['title'], '案件')}.docx"
        item = store_case_file(
            get_db, case_id, "contract", filename, preview_path.read_bytes(),
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            source_type="contract_template", source_id=int(preview["template_id"]),
            original_filename=filename,
            metadata={"template_id": int(preview["template_id"]), "values": _loads(preview["values_json"], {})},
        )
        preview_path.unlink(missing_ok=True)
        preview_path.with_suffix(".pdf").unlink(missing_ok=True)
        preview_path.with_suffix(".html").unlink(missing_ok=True)
        with get_db() as conn:
            conn.execute("DELETE FROM case_manager_contract_previews WHERE token=?", (token,))
            conn.commit()
        return item

    @app.get("/api/case-manager/cases/{case_id}/contracts")
    def list_case_contracts(case_id: int):
        with get_db() as conn:
            rows = conn.execute(
                """SELECT * FROM case_manager_files
                   WHERE case_id=? AND category='contract'
                   ORDER BY datetime(created_at) DESC,id DESC""", (case_id,)
            ).fetchall()
        return {"items": [dict(row) for row in rows]}
