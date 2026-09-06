from __future__ import annotations

import asyncio
import datetime as dt
import hashlib
import json
import os
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple
from urllib.parse import parse_qs, urlparse
from urllib.parse import quote

import httpx
import openpyxl
from fastapi import HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


EXPENSE_CATEGORIES = [
    "交通出行",
    "快递邮寄/跑腿",
    "打印复印/装订",
    "餐饮/业务招待",
    "差旅住宿",
    "培训/书刊资料",
    "调查取证/查档",
    "诉讼仲裁/公证/保全",
    "办公用品/耗材",
    "其他",
]
FEISHU_FINGERPRINT_VERSION = 2
LEGACY_EXPENSE_CATEGORIES = [
    "差旅交通", "餐饮", "快递邮寄", "打印复印/取证",
    "诉讼仲裁/公证", "办公用品/书刊", "其他",
]
ENTRY_TYPES = {"expense": "团队支出", "advance": "团队备用金"}
ENTRY_TYPE_ALIASES = {
    "expense": "expense",
    "团队支出": "expense",
    "支出": "expense",
    "advance": "advance",
    "团队备用金": "advance",
    "团队预支": "advance",
    "团队预付": "advance",
    "预支": "advance",
    "预付": "advance",
}
DEFAULT_TEXT_PRESETS = {
    "expense": [
        "文件快递费（顺丰、EMS）", "打印复印/装订费", "交通出行费", "差旅住宿费",
        "餐饮/业务招待费", "调查取证/查档费", "诉讼仲裁/公证费", "办公用品费",
    ],
    "advance": ["团队转入备用金", "补充团队备用金", "备用金调整"],
}
FEISHU_FIELDS = {
    "entry_type": "流水类型",
    "occurred_on": "发生日期",
    "project_name": "对应项目",
    "category": "支出分类",
    "description": "支出内容/说明",
    "amount": "金额",
    "sync_status": "同步状态",
    "local_id": "本地记录编号",
    "synced_at": "同步时间",
    "sync_note": "同步说明",
}
REQUIRED_FEISHU_FIELDS = set(FEISHU_FIELDS.values())
CONFIG_PATH = Path(__file__).resolve().parent.parent / "local_feishu_expense_config.json"
SHANGHAI_TZ = dt.timezone(dt.timedelta(hours=8), name="Asia/Shanghai")


def _now_iso() -> str:
    return dt.datetime.now().astimezone().isoformat(timespec="seconds")


def _today_iso() -> str:
    return dt.datetime.now(SHANGHAI_TZ).date().isoformat()


def _row_dict(row: Any) -> Dict[str, Any]:
    return dict(row) if row else {}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (int, float, Decimal)):
        return str(value)
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        return "、".join(filter(None, (_text(item) for item in value)))
    if isinstance(value, dict):
        for key in ("text", "name", "value", "full_name"):
            if key in value:
                return _text(value.get(key))
        return "、".join(filter(None, (_text(item) for item in value.values())))
    return str(value).strip()


def _amount_to_cents(value: Any) -> int:
    try:
        amount = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, AttributeError, ValueError):
        raise HTTPException(status_code=400, detail="金额格式不正确")
    if not amount.is_finite() or amount <= 0:
        raise HTTPException(status_code=400, detail="金额必须大于 0")
    quantized = amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if amount != quantized:
        raise HTTPException(status_code=400, detail="金额最多保留两位小数")
    cents = int(quantized * 100)
    if cents > 9_999_999_999_99:
        raise HTTPException(status_code=400, detail="金额超出允许范围")
    return cents


def _normalize_entry_type(value: Any) -> str:
    normalized = ENTRY_TYPE_ALIASES.get(_text(value).lower()) or ENTRY_TYPE_ALIASES.get(_text(value))
    if not normalized:
        raise HTTPException(status_code=400, detail="流水类型必须为“团队支出”或“团队备用金”")
    return normalized


def _normalize_date(value: Any, fallback: str = "") -> str:
    if value in (None, ""):
        return fallback or _today_iso()
    if isinstance(value, (int, float)) or (isinstance(value, str) and value.strip().isdigit()):
        stamp = float(value)
        if stamp > 10_000_000_000:
            stamp /= 1000
        try:
            return dt.datetime.fromtimestamp(stamp, SHANGHAI_TZ).date().isoformat()
        except (OverflowError, OSError, ValueError):
            pass
    raw = _text(value).replace("/", "-").replace("年", "-").replace("月", "-").replace("日", "")
    raw = raw[:10]
    try:
        return dt.date.fromisoformat(raw).isoformat()
    except ValueError:
        raise HTTPException(status_code=400, detail="发生日期格式不正确")


def _normalize_payload(
    payload: Dict[str, Any],
    fallback_date: str = "",
    allowed_categories: Optional[List[str]] = None,
) -> Dict[str, Any]:
    entry_type = _normalize_entry_type(payload.get("entry_type"))
    category = _text(payload.get("category"))
    if category and allowed_categories is not None and category not in allowed_categories:
        raise HTTPException(status_code=400, detail="支出分类不在允许范围内")
    description = _text(payload.get("description"))
    if entry_type == "expense" and not description:
        raise HTTPException(status_code=400, detail="支出记录必须填写支出内容")
    return {
        "occurred_on": _normalize_date(payload.get("occurred_on"), fallback_date),
        "entry_type": entry_type,
        "project_name": _text(payload.get("project_name"))[:300] if entry_type == "expense" else "",
        "category": category if entry_type == "expense" else "",
        "description": description[:2000],
        "amount_cents": _amount_to_cents(payload.get("amount")),
    }


def _serialize_expense(row: Any) -> Dict[str, Any]:
    item = _row_dict(row)
    cents = int(item.get("amount_cents") or 0)
    item["amount"] = cents / 100
    item["entry_type_label"] = ENTRY_TYPES.get(item.get("entry_type"), item.get("entry_type") or "")
    item["source_label"] = "来自飞书" if item.get("source_type") == "feishu" else "本地录入"
    if "running_balance_cents" in item:
        item["running_balance"] = int(item.get("running_balance_cents") or 0) / 100
    return item


def ensure_prepaid_expense_schema(get_db: Callable[[], Any]) -> None:
    with get_db() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS case_manager_prepaid_expenses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                occurred_on TEXT NOT NULL,
                entry_type TEXT NOT NULL,
                case_id INTEGER,
                project_name TEXT DEFAULT '',
                category TEXT DEFAULT '',
                description TEXT DEFAULT '',
                amount_cents INTEGER NOT NULL,
                source_type TEXT NOT NULL DEFAULT 'local',
                source_record_id TEXT,
                source_payload_hash TEXT DEFAULT '',
                source_created_at TEXT DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE SET NULL
            );
            CREATE UNIQUE INDEX IF NOT EXISTS idx_case_manager_prepaid_source
                ON case_manager_prepaid_expenses(source_type, source_record_id)
                WHERE source_record_id IS NOT NULL AND source_record_id<>'';
            CREATE INDEX IF NOT EXISTS idx_case_manager_prepaid_date
                ON case_manager_prepaid_expenses(occurred_on DESC, id DESC);
            CREATE INDEX IF NOT EXISTS idx_case_manager_prepaid_case
                ON case_manager_prepaid_expenses(case_id);
            CREATE TABLE IF NOT EXISTS case_manager_prepaid_sync_state (
                id INTEGER PRIMARY KEY CHECK(id=1),
                last_attempt_at TEXT DEFAULT '',
                last_success_at TEXT DEFAULT '',
                last_result_json TEXT DEFAULT '{}',
                last_error TEXT DEFAULT '',
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS case_manager_prepaid_sync_issues (
                source_record_id TEXT PRIMARY KEY,
                source_payload_hash TEXT NOT NULL,
                error_message TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS case_manager_prepaid_category_settings (
                id INTEGER PRIMARY KEY CHECK(id=1),
                categories_json TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS case_manager_prepaid_text_presets (
                kind TEXT PRIMARY KEY CHECK(kind IN ('expense','advance')),
                values_json TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS case_manager_prepaid_project_sync_state (
                id INTEGER PRIMARY KEY CHECK(id=1),
                case_fingerprint TEXT DEFAULT '',
                last_synced_at TEXT DEFAULT '',
                last_error TEXT DEFAULT '',
                updated_at TEXT NOT NULL
            );
            """
        )
        conn.execute(
            """INSERT OR IGNORE INTO case_manager_prepaid_sync_state
               (id,last_attempt_at,last_success_at,last_result_json,last_error,updated_at)
               VALUES (1,'','','{}','',?)""",
            (_now_iso(),),
        )
        conn.execute(
            """INSERT OR IGNORE INTO case_manager_prepaid_category_settings
               (id,categories_json,updated_at) VALUES (1,?,?)""",
            (json.dumps(EXPENSE_CATEGORIES, ensure_ascii=False), _now_iso()),
        )
        for preset_kind, preset_values in DEFAULT_TEXT_PRESETS.items():
            conn.execute(
                """INSERT OR IGNORE INTO case_manager_prepaid_text_presets
                   (kind,values_json,updated_at) VALUES (?,?,?)""",
                (preset_kind, json.dumps(preset_values, ensure_ascii=False), _now_iso()),
            )
        conn.execute(
            """INSERT OR IGNORE INTO case_manager_prepaid_project_sync_state
               (id,case_fingerprint,last_synced_at,last_error,updated_at) VALUES (1,'','','',?)""",
            (_now_iso(),),
        )
        category_row = conn.execute(
            "SELECT categories_json FROM case_manager_prepaid_category_settings WHERE id=1"
        ).fetchone()
        try:
            stored_categories = json.loads(category_row["categories_json"] or "[]")
        except (TypeError, ValueError, KeyError):
            stored_categories = []
        if stored_categories == LEGACY_EXPENSE_CATEGORIES:
            conn.execute(
                "UPDATE case_manager_prepaid_category_settings SET categories_json=?,updated_at=? WHERE id=1",
                (json.dumps(EXPENSE_CATEGORIES, ensure_ascii=False), _now_iso()),
            )
        conn.commit()


def _expense_categories(get_db: Callable[[], Any]) -> List[str]:
    with get_db() as conn:
        row = conn.execute(
            "SELECT categories_json FROM case_manager_prepaid_category_settings WHERE id=1"
        ).fetchone()
    try:
        values = json.loads(row["categories_json"] if row else "[]")
    except (TypeError, ValueError, KeyError):
        values = []
    normalized: List[str] = []
    for value in values:
        label = _text(value)[:60]
        if label and label not in normalized:
            normalized.append(label)
    return normalized or list(EXPENSE_CATEGORIES)


def _save_expense_categories(get_db: Callable[[], Any], values: Any) -> List[str]:
    if not isinstance(values, list):
        raise HTTPException(status_code=400, detail="分类列表格式不正确")
    categories: List[str] = []
    for value in values:
        label = _text(value)[:60]
        if label and label not in categories:
            categories.append(label)
    if not categories:
        raise HTTPException(status_code=400, detail="请至少保留一个支出分类")
    if len(categories) > 30:
        raise HTTPException(status_code=400, detail="支出分类最多设置 30 个")
    with get_db() as conn:
        conn.execute(
            """INSERT INTO case_manager_prepaid_category_settings (id,categories_json,updated_at)
               VALUES (1,?,?) ON CONFLICT(id) DO UPDATE SET
               categories_json=excluded.categories_json,updated_at=excluded.updated_at""",
            (json.dumps(categories, ensure_ascii=False), _now_iso()),
        )
        conn.commit()
    return categories


def _text_presets(get_db: Callable[[], Any]) -> Dict[str, List[str]]:
    result: Dict[str, List[str]] = {}
    with get_db() as conn:
        rows = conn.execute(
            "SELECT kind,values_json FROM case_manager_prepaid_text_presets"
        ).fetchall()
    stored = {row["kind"]: row["values_json"] for row in rows}
    for kind in ("expense", "advance"):
        try:
            raw_values = json.loads(stored.get(kind) or "[]")
        except (TypeError, ValueError):
            raw_values = []
        values: List[str] = []
        for raw_value in raw_values:
            value = _text(raw_value)[:200]
            if value and value not in values:
                values.append(value)
        result[kind] = values
    return result


def _save_text_presets(get_db: Callable[[], Any], payload: Any) -> Dict[str, List[str]]:
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="常用内容格式不正确")
    saved: Dict[str, List[str]] = {}
    with get_db() as conn:
        for kind in ("expense", "advance"):
            raw_values = payload.get(kind)
            if not isinstance(raw_values, list):
                raise HTTPException(status_code=400, detail="团队支出和团队备用金的常用内容都必须是列表")
            values: List[str] = []
            for raw_value in raw_values:
                value = _text(raw_value)[:200]
                if value and value not in values:
                    values.append(value)
            if len(values) > 50:
                raise HTTPException(status_code=400, detail="每类常用内容最多保存 50 项")
            conn.execute(
                """INSERT INTO case_manager_prepaid_text_presets (kind,values_json,updated_at)
                   VALUES (?,?,?) ON CONFLICT(kind) DO UPDATE SET
                   values_json=excluded.values_json,updated_at=excluded.updated_at""",
                (kind, json.dumps(values, ensure_ascii=False), _now_iso()),
            )
            saved[kind] = values
        conn.commit()
    return saved


def _case_project_titles(get_db: Callable[[], Any]) -> List[str]:
    with get_db() as conn:
        rows = conn.execute(
            """SELECT title FROM case_manager_cases
               WHERE COALESCE(source_status,'active')='active' AND TRIM(COALESCE(title,''))<>''
               ORDER BY updated_at DESC,id DESC"""
        ).fetchall()
    titles: List[str] = []
    for row in rows:
        title = _text(row["title"])[:300]
        if title and title not in titles:
            titles.append(title)
    return titles


def _read_config() -> Dict[str, str]:
    if not CONFIG_PATH.exists():
        return {}
    try:
        data = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
        return {str(key): _text(value) for key, value in data.items()}
    except Exception:
        return {}


def _write_config(data: Dict[str, Any]) -> None:
    CONFIG_PATH.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    os.chmod(CONFIG_PATH, 0o600)


def _mask_secret(value: str) -> str:
    if not value:
        return ""
    if len(value) <= 8:
        return "••••••••"
    return f"{value[:3]}••••••{value[-3:]}"


def _parse_base_url(value: str) -> Tuple[str, str]:
    value = _text(value)
    if not value:
        return "", ""
    parsed = urlparse(value if "://" in value else f"https://{value}")
    segments = [segment for segment in parsed.path.split("/") if segment]
    if segments[:3] == ["share", "base", "view"]:
        raise HTTPException(
            status_code=400,
            detail="这里不能使用 /share/base/view/ 分享链接。请打开多维表格的编辑页面，复制地址栏中同时包含 /base/ 和 table=tbl… 的完整链接",
        )
    app_token = ""
    if "base" in segments:
        index = segments.index("base")
        if index + 1 < len(segments):
            app_token = segments[index + 1]
    elif "wiki" in segments:
        raise HTTPException(status_code=400, detail="暂不支持知识库中的多维表格，请使用以 /base/ 开头的链接")
    table_id = (parse_qs(parsed.query).get("table") or [""])[0]
    if not app_token or not table_id:
        raise HTTPException(
            status_code=400,
            detail="链接无法识别。请从多维表格编辑页面复制地址栏完整链接，确认其中同时包含 /base/ 和 table=tbl…",
        )
    return app_token, table_id


def _settings_response(config: Optional[Dict[str, str]] = None) -> Dict[str, Any]:
    config = config or _read_config()
    return {
        "configured": bool(config.get("app_id") and config.get("app_secret") and config.get("app_token") and config.get("table_id")),
        "app_id": config.get("app_id", ""),
        "app_secret_configured": bool(config.get("app_secret")),
        "app_secret_masked": _mask_secret(config.get("app_secret", "")),
        "base_url": config.get("base_url", ""),
        "form_url": config.get("form_url", ""),
        "app_token": config.get("app_token", ""),
        "table_id": config.get("table_id", ""),
    }


def _sync_status(get_db: Callable[[], Any]) -> Dict[str, Any]:
    with get_db() as conn:
        row = conn.execute("SELECT * FROM case_manager_prepaid_sync_state WHERE id=1").fetchone()
        project_row = conn.execute(
            "SELECT last_synced_at,last_error FROM case_manager_prepaid_project_sync_state WHERE id=1"
        ).fetchone()
        issues = conn.execute("SELECT COUNT(*) FROM case_manager_prepaid_sync_issues").fetchone()[0]
    data = _row_dict(row)
    try:
        last_result = json.loads(data.get("last_result_json") or "{}")
    except Exception:
        last_result = {}
    return {
        "configured": _settings_response().get("configured", False),
        "last_attempt_at": data.get("last_attempt_at") or "",
        "last_success_at": data.get("last_success_at") or "",
        "last_error": data.get("last_error") or "",
        "last_result": last_result,
        "issue_count": int(issues or 0),
        "project_options_last_synced_at": project_row["last_synced_at"] if project_row else "",
        "project_options_error": project_row["last_error"] if project_row else "",
    }


def _summary(get_db: Callable[[], Any]) -> Dict[str, Any]:
    with get_db() as conn:
        row = conn.execute(
            """SELECT
                   COALESCE(SUM(CASE WHEN entry_type='expense' THEN amount_cents ELSE 0 END),0) AS expense_cents,
                   COALESCE(SUM(CASE WHEN entry_type='advance' THEN amount_cents ELSE 0 END),0) AS advance_cents,
                   COUNT(*) AS record_count
               FROM case_manager_prepaid_expenses"""
        ).fetchone()
    expense_cents = int(row["expense_cents"] or 0)
    advance_cents = int(row["advance_cents"] or 0)
    return {
        "expense": expense_cents / 100,
        "advance": advance_cents / 100,
        "balance": (advance_cents - expense_cents) / 100,
        "expense_cents": expense_cents,
        "advance_cents": advance_cents,
        "balance_cents": advance_cents - expense_cents,
        "record_count": int(row["record_count"] or 0),
    }


def _query_expenses(
    get_db: Callable[[], Any],
    *,
    date_from: str = "",
    date_to: str = "",
    entry_type: str = "",
    project: str = "",
    category: str = "",
    q: str = "",
    limit: Optional[int] = 1000,
) -> List[Dict[str, Any]]:
    where = ["1=1"]
    args: List[Any] = []
    if date_from:
        where.append("e.occurred_on>=?")
        args.append(_normalize_date(date_from))
    if date_to:
        where.append("e.occurred_on<=?")
        args.append(_normalize_date(date_to))
    if entry_type:
        where.append("e.entry_type=?")
        args.append(_normalize_entry_type(entry_type))
    if project:
        where.append("e.project_name LIKE ?")
        args.append(f"%{_text(project)}%")
    if category:
        where.append("e.category=?")
        args.append(_text(category))
    if q:
        token = f"%{_text(q)}%"
        where.append("(e.project_name LIKE ? OR e.category LIKE ? OR e.description LIKE ?)")
        args.extend([token, token, token])
    sql = f"""SELECT e.*,c.title AS case_title,
                     (SELECT COALESCE(SUM(CASE WHEN x.entry_type='advance' THEN x.amount_cents ELSE -x.amount_cents END),0)
                        FROM case_manager_prepaid_expenses x
                       WHERE x.occurred_on<e.occurred_on
                          OR (x.occurred_on=e.occurred_on AND x.id<=e.id)) AS running_balance_cents
              FROM case_manager_prepaid_expenses e
              LEFT JOIN case_manager_cases c ON c.id=e.case_id
              WHERE {' AND '.join(where)}
              ORDER BY e.occurred_on DESC,e.id DESC"""
    if limit is not None:
        sql += " LIMIT ?"
        args.append(max(1, min(int(limit), 5000)))
    with get_db() as conn:
        rows = conn.execute(sql, args).fetchall()
    return [_serialize_expense(row) for row in rows]


def prepaid_expense_offline_snapshot(get_db: Callable[[], Any]) -> Dict[str, Any]:
    return {
        "items": _query_expenses(get_db, limit=5000),
        "summary": _summary(get_db),
        "categories": _expense_categories(get_db),
        "presets": _text_presets(get_db),
        "sync": _sync_status(get_db),
    }


def _match_case_id(conn: Any, project_name: str) -> Optional[int]:
    if not project_name:
        return None
    rows = conn.execute(
        "SELECT id FROM case_manager_cases WHERE title=? ORDER BY id LIMIT 2",
        (project_name,),
    ).fetchall()
    return int(rows[0]["id"]) if len(rows) == 1 else None


def _save_expense(
    get_db: Callable[[], Any],
    payload: Dict[str, Any],
    *,
    item_id: int = 0,
    source_type: str = "local",
    source_record_id: str = "",
    source_payload_hash: str = "",
    source_created_at: str = "",
    fallback_date: str = "",
) -> Dict[str, Any]:
    allowed_categories = None if source_type == "feishu" else _expense_categories(get_db)
    normalized = _normalize_payload(payload, fallback_date, allowed_categories)
    stamp = _now_iso()
    with get_db() as conn:
        case_id = payload.get("case_id")
        if case_id not in (None, ""):
            row = conn.execute("SELECT id,title FROM case_manager_cases WHERE id=?", (int(case_id),)).fetchone()
            if not row:
                raise HTTPException(status_code=400, detail="关联案件不存在")
            case_id = int(row["id"])
            if not normalized["project_name"]:
                normalized["project_name"] = row["title"]
        else:
            case_id = _match_case_id(conn, normalized["project_name"])
        if item_id:
            existing = conn.execute("SELECT * FROM case_manager_prepaid_expenses WHERE id=?", (item_id,)).fetchone()
            if not existing:
                raise HTTPException(status_code=404, detail="预支费用记录不存在")
            conn.execute(
                """UPDATE case_manager_prepaid_expenses
                   SET occurred_on=?,entry_type=?,case_id=?,project_name=?,category=?,description=?,
                       amount_cents=?,updated_at=? WHERE id=?""",
                (
                    normalized["occurred_on"], normalized["entry_type"], case_id,
                    normalized["project_name"], normalized["category"], normalized["description"],
                    normalized["amount_cents"], stamp, item_id,
                ),
            )
            saved_id = item_id
        else:
            try:
                cursor = conn.execute(
                    """INSERT INTO case_manager_prepaid_expenses
                       (occurred_on,entry_type,case_id,project_name,category,description,amount_cents,
                        source_type,source_record_id,source_payload_hash,source_created_at,created_at,updated_at)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (
                        normalized["occurred_on"], normalized["entry_type"], case_id,
                        normalized["project_name"], normalized["category"], normalized["description"],
                        normalized["amount_cents"], source_type, source_record_id or None,
                        source_payload_hash, source_created_at, stamp, stamp,
                    ),
                )
                saved_id = int(cursor.lastrowid)
            except Exception as exc:
                if source_record_id:
                    existing = conn.execute(
                        """SELECT id FROM case_manager_prepaid_expenses
                           WHERE source_type=? AND source_record_id=?""",
                        (source_type, source_record_id),
                    ).fetchone()
                    if existing:
                        saved_id = int(existing["id"])
                    else:
                        raise exc
                else:
                    raise exc
        conn.commit()
        row = conn.execute(
            """SELECT e.*,c.title AS case_title FROM case_manager_prepaid_expenses e
               LEFT JOIN case_manager_cases c ON c.id=e.case_id WHERE e.id=?""",
            (saved_id,),
        ).fetchone()
    if source_type == "feishu" and normalized["category"]:
        categories = _expense_categories(get_db)
        if normalized["category"] not in categories and len(categories) < 30:
            _save_expense_categories(get_db, [*categories, normalized["category"]])
    return _serialize_expense(row)


def _export_workbook(items: List[Dict[str, Any]]) -> BytesIO:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "团队备用金使用记录"
    sheet.sheet_view.showGridLines = False
    dates = [item.get("occurred_on") for item in items if item.get("occurred_on")]
    if dates:
        earliest, latest = min(dates), max(dates)
        start_date = dt.date.fromisoformat(earliest)
        end_date = dt.date.fromisoformat(latest)
        title = (
            "团队备用金支出明细表（"
            f"{start_date.year}年{start_date.month}月{start_date.day}日—"
            f"{end_date.year}年{end_date.month}月{end_date.day}日）"
        )
    else:
        title = f"团队备用金支出明细表（{_today_iso()}）"
    sheet.merge_cells("A1:J1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="宋体", size=16, bold=True, color="243247")
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34
    headers = ["收支时间", "收支类型", "对应项目", "支出分类", "支出内容", "支出金额", "累计支出", "团队备用金", "累计备用金", "余额"]
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(2, column, header)
        cell.font = Font(name="宋体", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7F1D1D")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(bottom=thin)
    last_row = 2 + len(items)
    for index, item in enumerate(items, 3):
        occurred = dt.date.fromisoformat(item["occurred_on"])
        amount = int(item.get("amount_cents") or 0) / 100
        values = [
            occurred,
            item.get("entry_type_label") or "",
            item.get("project_name") or "",
            item.get("category") or "",
            item.get("description") or "",
            amount if item.get("entry_type") == "expense" else None,
            None,
            amount if item.get("entry_type") == "advance" else None,
            None,
            None,
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(index, column, value)
            cell.font = Font(name="宋体", size=10.5, color="243247")
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if column in (3, 4, 5) else "center",
                vertical="center",
                wrap_text=column in (3, 5),
            )
        sheet.cell(index, 7, f"=SUM(F{index}:F${last_row})")
        sheet.cell(index, 9, f"=SUM(H{index}:H${last_row})")
        sheet.cell(index, 10, f"=I{index}-G{index}")
        sheet.row_dimensions[index].height = 26
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:J{max(last_row, 2)}"
    widths = {"A": 15, "B": 13, "C": 25, "D": 18, "E": 40, "F": 14, "G": 14, "H": 14, "I": 14, "J": 14}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for cell in sheet["A"][2:]:
        cell.number_format = 'yyyy"年"m"月"d"日"'
    money_format = '#,##0.00;[Red]-#,##0.00'
    for column in (6, 7, 8, 9, 10):
        for row in range(3, last_row + 1):
            sheet.cell(row, column).number_format = money_format
    if last_row >= 3:
        red_font = Font(color="B91C1C")
        sheet.conditional_formatting.add(
            f"J3:J{last_row}",
            CellIsRule(operator="lessThan", formula=["0"], font=red_font),
        )
    sheet.print_title_rows = "1:2"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    if hasattr(workbook, "calculation"):
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _feishu_response_payload(response: httpx.Response, fallback: str) -> Dict[str, Any]:
    try:
        payload = response.json()
    except Exception:
        payload = {}
    code = payload.get("code")
    if response.status_code >= 400 or code not in (None, 0):
        message = _text(payload.get("msg"))
        if code == 99991672:
            raise RuntimeError(
                "飞书应用身份权限未生效（错误码 99991672）。请在飞书开放平台为应用身份开通"
                "“查看、评论、编辑和管理多维表格（bitable:app）”权限，发布新版本并完成管理员审批后重试"
            )
        if code not in (None, 0):
            raise RuntimeError(f"飞书错误 {code}：{message or fallback}")
        raise RuntimeError(f"飞书接口返回 HTTP {response.status_code}：{message or fallback}")
    return payload


class FeishuExpenseSync:
    def __init__(self, get_db: Callable[[], Any]):
        self.get_db = get_db
        self.lock: Optional[asyncio.Lock] = None
        self.project_lock: Optional[asyncio.Lock] = None
        self._token = ""
        self._token_expires_at = 0.0

    async def _fields(
        self, client: httpx.AsyncClient, token: str, config: Dict[str, str]
    ) -> List[Dict[str, Any]]:
        headers = {"Authorization": f"Bearer {token}"}
        url = f"https://open.feishu.cn/open-apis/bitable/v1/apps/{config['app_token']}/tables/{config['table_id']}/fields"
        response = await client.get(url, headers=headers, params={"page_size": 100})
        payload = _feishu_response_payload(response, "无法读取多维表格字段")
        return payload.get("data", {}).get("items", [])

    @staticmethod
    def _select_options(
        field: Dict[str, Any], desired_names: List[str], aliases: Optional[Dict[str, str]] = None
    ) -> List[Dict[str, Any]]:
        aliases = aliases or {}
        existing = (field.get("property") or {}).get("options") or []
        by_name = {_text(option.get("name")): option for option in existing}
        for old_name, new_name in aliases.items():
            if new_name not in by_name and old_name in by_name:
                by_name[new_name] = by_name[old_name]
        result: List[Dict[str, Any]] = []
        for index, name in enumerate(desired_names):
            option: Dict[str, Any] = {"name": name}
            previous = by_name.get(name) or {}
            if previous.get("id"):
                option["id"] = previous["id"]
            option["color"] = int(previous.get("color", index % 10))
            result.append(option)
        return result

    async def _update_select_field(
        self,
        client: httpx.AsyncClient,
        token: str,
        config: Dict[str, str],
        field: Dict[str, Any],
        desired_names: List[str],
        aliases: Optional[Dict[str, str]] = None,
    ) -> bool:
        current_names = [
            _text(option.get("name")) for option in ((field.get("property") or {}).get("options") or [])
        ]
        if int(field.get("type") or 0) == 3 and current_names == desired_names:
            return False
        headers = {"Authorization": f"Bearer {token}"}
        url = (
            f"https://open.feishu.cn/open-apis/bitable/v1/apps/{config['app_token']}"
            f"/tables/{config['table_id']}/fields/{field['field_id']}"
        )
        body = {
            "field_name": field["field_name"],
            "type": 3,
            "property": {"options": self._select_options(field, desired_names, aliases)},
        }
        response = await client.put(url, headers=headers, json=body)
        _feishu_response_payload(response, f"无法更新飞书字段“{field['field_name']}”")
        return True

    async def _sync_project_options_with_client(
        self, client: httpx.AsyncClient, token: str, config: Dict[str, str], force: bool = False
    ) -> Dict[str, Any]:
        titles = _case_project_titles(self.get_db)
        fingerprint = hashlib.sha256(
            json.dumps({"version": 2, "titles": titles}, ensure_ascii=False).encode("utf-8")
        ).hexdigest()
        with self.get_db() as conn:
            state_row = conn.execute(
                "SELECT case_fingerprint FROM case_manager_prepaid_project_sync_state WHERE id=1"
            ).fetchone()
        if not force and state_row and state_row["case_fingerprint"] == fingerprint:
            return {"ok": True, "changed": False, "project_count": len(titles)}
        try:
            fields = await self._fields(client, token, config)
            by_name = {field.get("field_name"): field for field in fields}
            project_field = by_name.get(FEISHU_FIELDS["project_name"])
            entry_type_field = by_name.get(FEISHU_FIELDS["entry_type"])
            if not project_field:
                raise RuntimeError("飞书多维表格缺少“对应项目”字段")
            if not entry_type_field:
                raise RuntimeError("飞书多维表格缺少“流水类型”字段")
            project_changed = await self._update_select_field(
                client, token, config, project_field, titles
            )
            type_changed = await self._update_select_field(
                client,
                token,
                config,
                entry_type_field,
                [ENTRY_TYPES["expense"], ENTRY_TYPES["advance"]],
                {"支出": ENTRY_TYPES["expense"], "团队预支": ENTRY_TYPES["advance"]},
            )
            stamp = _now_iso()
            with self.get_db() as conn:
                conn.execute(
                    """UPDATE case_manager_prepaid_project_sync_state SET
                       case_fingerprint=?,last_synced_at=?,last_error='',updated_at=? WHERE id=1""",
                    (fingerprint, stamp, stamp),
                )
                conn.commit()
            return {
                "ok": True,
                "changed": project_changed or type_changed,
                "project_count": len(titles),
                "last_synced_at": stamp,
            }
        except Exception as exc:
            message = _friendly_sync_error(exc)
            with self.get_db() as conn:
                conn.execute(
                    """UPDATE case_manager_prepaid_project_sync_state SET
                       last_error=?,updated_at=? WHERE id=1""",
                    (message, _now_iso()),
                )
                conn.commit()
            raise RuntimeError(message) from exc

    async def sync_project_options(self, force: bool = False) -> Dict[str, Any]:
        if self.project_lock is None:
            self.project_lock = asyncio.Lock()
        if self.project_lock.locked():
            return {"ok": True, "busy": True, "message": "案件选项正在同步"}
        async with self.project_lock:
            config = _read_config()
            if not _settings_response(config)["configured"]:
                return {"ok": False, "configured": False, "message": "尚未配置飞书同步"}
            async with httpx.AsyncClient(timeout=30.0) as client:
                token = await self._access_token(client, config)
                return await self._sync_project_options_with_client(client, token, config, force)

    async def _access_token(self, client: httpx.AsyncClient, config: Dict[str, str]) -> str:
        loop_time = asyncio.get_running_loop().time()
        if self._token and loop_time < self._token_expires_at - 300:
            return self._token
        response = await client.post(
            "https://open.feishu.cn/open-apis/auth/v3/tenant_access_token/internal",
            json={"app_id": config.get("app_id"), "app_secret": config.get("app_secret")},
        )
        payload = _feishu_response_payload(response, "飞书身份验证失败")
        self._token = payload.get("tenant_access_token") or ""
        if not self._token:
            raise RuntimeError("飞书未返回访问凭证")
        self._token_expires_at = loop_time + int(payload.get("expire") or 7200)
        return self._token

    async def verify(self, config: Optional[Dict[str, str]] = None) -> Dict[str, Any]:
        config = config or _read_config()
        if not _settings_response(config)["configured"]:
            raise HTTPException(status_code=400, detail="请完整填写飞书 App ID、App Secret 和多维表格链接")
        async with httpx.AsyncClient(timeout=20.0) as client:
            token = await self._access_token(client, config)
            headers = {"Authorization": f"Bearer {token}"}
            items = await self._fields(client, token, config)
            names = {item.get("field_name") for item in items}
            missing = sorted(REQUIRED_FEISHU_FIELDS - names)
            if missing:
                raise HTTPException(status_code=400, detail=f"多维表格缺少字段：{'、'.join(missing)}")
            expected_types = {
                FEISHU_FIELDS["entry_type"]: {3},
                FEISHU_FIELDS["occurred_on"]: {5},
                FEISHU_FIELDS["project_name"]: {1, 3},
                FEISHU_FIELDS["category"]: {3},
                FEISHU_FIELDS["description"]: {1},
                FEISHU_FIELDS["amount"]: {2},
                FEISHU_FIELDS["sync_status"]: {3},
                FEISHU_FIELDS["local_id"]: {1, 2},
                FEISHU_FIELDS["synced_at"]: {5},
                FEISHU_FIELDS["sync_note"]: {1},
            }
            wrong_types = [
                item.get("field_name")
                for item in items
                if item.get("field_name") in expected_types
                and int(item.get("type") or 0) not in expected_types[item.get("field_name")]
            ]
            if wrong_types:
                raise HTTPException(status_code=400, detail=f"以下字段类型不正确：{'、'.join(sorted(wrong_types))}")
            search_url = f"https://open.feishu.cn/open-apis/bitable/v1/apps/{config['app_token']}/tables/{config['table_id']}/records/search"
            search = await client.post(search_url, headers=headers, params={"page_size": 1}, json={"automatic_fields": True})
            search_payload = _feishu_response_payload(search, "无法读取多维表格记录")
        return {"ok": True, "message": "连接成功，字段与读取权限检查通过；回写权限将在首次同步时确认", "field_count": len(names)}

    async def _records(self, client: httpx.AsyncClient, token: str, config: Dict[str, str]) -> List[Dict[str, Any]]:
        headers = {"Authorization": f"Bearer {token}"}
        url = f"https://open.feishu.cn/open-apis/bitable/v1/apps/{config['app_token']}/tables/{config['table_id']}/records/search"
        page_token = ""
        records: List[Dict[str, Any]] = []
        while True:
            params: Dict[str, Any] = {"page_size": 500}
            if page_token:
                params["page_token"] = page_token
            response = await client.post(url, headers=headers, params=params, json={"automatic_fields": True})
            payload = _feishu_response_payload(response, "读取飞书记录失败")
            data = payload.get("data", {})
            records.extend(data.get("items") or [])
            if not data.get("has_more"):
                break
            page_token = data.get("page_token") or ""
            if not page_token:
                break
        return records

    async def _writeback(
        self,
        client: httpx.AsyncClient,
        token: str,
        config: Dict[str, str],
        updates: List[Dict[str, Any]],
    ) -> None:
        if not updates:
            return
        headers = {"Authorization": f"Bearer {token}"}
        url = f"https://open.feishu.cn/open-apis/bitable/v1/apps/{config['app_token']}/tables/{config['table_id']}/records/batch_update"
        for start in range(0, len(updates), 500):
            response = await client.post(url, headers=headers, json={"records": updates[start:start + 500]})
            _feishu_response_payload(response, "飞书同步状态回写失败")

    def _business_payload(self, record: Dict[str, Any]) -> Tuple[Dict[str, Any], str, str]:
        fields = record.get("fields") or {}
        business = {
            "entry_type": _field_value(fields.get(FEISHU_FIELDS["entry_type"])),
            "occurred_on": _field_value(fields.get(FEISHU_FIELDS["occurred_on"])),
            "project_name": _field_value(fields.get(FEISHU_FIELDS["project_name"])),
            "category": _field_value(fields.get(FEISHU_FIELDS["category"])),
            "description": _field_value(fields.get(FEISHU_FIELDS["description"])),
            "amount": _field_value(fields.get(FEISHU_FIELDS["amount"])),
        }
        encoded = json.dumps(
            {"validation_version": FEISHU_FINGERPRINT_VERSION, "business": business},
            ensure_ascii=False,
            sort_keys=True,
            default=str,
        )
        fingerprint = hashlib.sha256(encoded.encode("utf-8")).hexdigest()
        status = _text(_field_value(fields.get(FEISHU_FIELDS["sync_status"])))
        return business, fingerprint, status

    async def sync(self) -> Dict[str, Any]:
        if self.lock is None:
            self.lock = asyncio.Lock()
        if self.lock.locked():
            return {"ok": True, "busy": True, "message": "飞书同步正在进行中"}
        async with self.lock:
            config = _read_config()
            if not _settings_response(config)["configured"]:
                return {"ok": False, "configured": False, "message": "尚未配置飞书同步"}
            attempted_at = _now_iso()
            with self.get_db() as conn:
                conn.execute(
                    "UPDATE case_manager_prepaid_sync_state SET last_attempt_at=?,updated_at=? WHERE id=1",
                    (attempted_at, attempted_at),
                )
                conn.commit()
            result = {"imported": 0, "existing": 0, "invalid": 0, "skipped": 0, "writeback": 0}
            try:
                async with httpx.AsyncClient(timeout=30.0) as client:
                    token = await self._access_token(client, config)
                    try:
                        result["project_options"] = await self._sync_project_options_with_client(
                            client, token, config, False
                        )
                    except Exception as project_exc:
                        result["project_options"] = {
                            "ok": False,
                            "message": _friendly_sync_error(project_exc),
                        }
                    records = await self._records(client, token, config)
                    remote_record_ids = {
                        _text(record.get("record_id")) for record in records if _text(record.get("record_id"))
                    }
                    with self.get_db() as conn:
                        stale_issues = conn.execute(
                            "SELECT source_record_id FROM case_manager_prepaid_sync_issues"
                        ).fetchall()
                        for stale_issue in stale_issues:
                            if stale_issue["source_record_id"] not in remote_record_ids:
                                conn.execute(
                                    "DELETE FROM case_manager_prepaid_sync_issues WHERE source_record_id=?",
                                    (stale_issue["source_record_id"],),
                                )
                        conn.commit()
                    updates: List[Dict[str, Any]] = []
                    for record in records:
                        record_id = _text(record.get("record_id"))
                        if not record_id:
                            continue
                        business, fingerprint, status = self._business_payload(record)
                        fields = record.get("fields") or {}
                        local_marker = _text(_field_value(fields.get(FEISHU_FIELDS["local_id"])))
                        with self.get_db() as conn:
                            existing = conn.execute(
                                """SELECT id FROM case_manager_prepaid_expenses
                                   WHERE source_type='feishu' AND source_record_id=?""",
                                (record_id,),
                            ).fetchone()
                            issue = conn.execute(
                                "SELECT source_payload_hash,error_message FROM case_manager_prepaid_sync_issues WHERE source_record_id=?",
                                (record_id,),
                            ).fetchone()
                        if existing:
                            result["existing"] += 1
                            if status != "已同步" or local_marker != str(existing["id"]):
                                updates.append(_success_update(record_id, int(existing["id"])))
                            continue
                        # A numeric marker is only written by this system. If the corresponding
                        # local row was deliberately deleted, keep the one-way deletion choice
                        # and do not import it again. Human/example markers such as LZ2026... are
                        # not trusted as proof that this computer has imported the record.
                        if status == "已同步" and local_marker.isdigit():
                            result["skipped"] += 1
                            continue
                        if issue and issue["source_payload_hash"] == fingerprint:
                            if status != "需修正":
                                updates.append(_failure_update(record_id, issue["error_message"] or "请修正表单内容"))
                            result["skipped"] += 1
                            continue
                        try:
                            fallback_date = _normalize_date(record.get("created_time"), _today_iso())
                            saved = _save_expense(
                                self.get_db,
                                business,
                                source_type="feishu",
                                source_record_id=record_id,
                                source_payload_hash=fingerprint,
                                source_created_at=_text(record.get("created_time")),
                                fallback_date=fallback_date,
                            )
                            with self.get_db() as conn:
                                conn.execute("DELETE FROM case_manager_prepaid_sync_issues WHERE source_record_id=?", (record_id,))
                                conn.commit()
                            result["imported"] += 1
                            updates.append(_success_update(record_id, int(saved["id"])))
                        except HTTPException as exc:
                            message = str(exc.detail)
                            with self.get_db() as conn:
                                conn.execute(
                                    """INSERT INTO case_manager_prepaid_sync_issues
                                       (source_record_id,source_payload_hash,error_message,updated_at)
                                       VALUES (?,?,?,?)
                                       ON CONFLICT(source_record_id) DO UPDATE SET
                                         source_payload_hash=excluded.source_payload_hash,
                                         error_message=excluded.error_message,updated_at=excluded.updated_at""",
                                    (record_id, fingerprint, message, _now_iso()),
                                )
                                conn.commit()
                            result["invalid"] += 1
                            updates.append(_failure_update(record_id, message))
                    await self._writeback(client, token, config, updates)
                    result["writeback"] = len(updates)
                success_at = _now_iso()
                with self.get_db() as conn:
                    conn.execute(
                        """UPDATE case_manager_prepaid_sync_state
                           SET last_success_at=?,last_result_json=?,last_error='',updated_at=? WHERE id=1""",
                        (success_at, json.dumps(result, ensure_ascii=False), success_at),
                    )
                    conn.commit()
                return {"ok": True, **result, "last_success_at": success_at}
            except Exception as exc:
                message = _friendly_sync_error(exc)
                with self.get_db() as conn:
                    conn.execute(
                        """UPDATE case_manager_prepaid_sync_state
                           SET last_result_json=?,last_error=?,updated_at=? WHERE id=1""",
                        (json.dumps(result, ensure_ascii=False), message, _now_iso()),
                    )
                    conn.commit()
                return {"ok": False, **result, "message": message}


def _field_value(value: Any) -> Any:
    if isinstance(value, list):
        if len(value) == 1:
            return _field_value(value[0])
        return "、".join(_text(_field_value(item)) for item in value)
    if isinstance(value, dict):
        for key in ("value", "text", "name"):
            if key in value:
                return _field_value(value[key])
    return value


def _success_update(record_id: str, local_id: int) -> Dict[str, Any]:
    return {
        "record_id": record_id,
        "fields": {
            FEISHU_FIELDS["sync_status"]: "已同步",
            FEISHU_FIELDS["local_id"]: str(local_id),
            FEISHU_FIELDS["synced_at"]: int(dt.datetime.now().timestamp() * 1000),
            FEISHU_FIELDS["sync_note"]: "",
        },
    }


def _failure_update(record_id: str, message: str) -> Dict[str, Any]:
    return {
        "record_id": record_id,
        "fields": {
            FEISHU_FIELDS["sync_status"]: "需修正",
            FEISHU_FIELDS["sync_note"]: message[:500],
        },
    }


def _friendly_sync_error(exc: Exception) -> str:
    if isinstance(exc, httpx.TimeoutException):
        return "连接飞书超时，请检查网络后重试"
    if isinstance(exc, httpx.HTTPStatusError):
        try:
            _feishu_response_payload(exc.response, "请检查应用权限和多维表格授权")
        except RuntimeError as parsed:
            return str(parsed)[:500]
        return f"飞书接口返回 HTTP {exc.response.status_code}，请检查应用权限和多维表格授权"
    if isinstance(exc, HTTPException):
        return str(exc.detail)
    return str(exc)[:500] or "飞书同步失败"


def register_prepaid_expense_routes(app: Any, get_db: Callable[[], Any]) -> None:
    ensure_prepaid_expense_schema(get_db)
    syncer = FeishuExpenseSync(get_db)
    scheduler_task: Optional[asyncio.Task] = None

    @app.get("/api/case-manager/prepaid-expenses")
    def list_prepaid_expenses(
        date_from: str = "",
        date_to: str = "",
        entry_type: str = "",
        project: str = "",
        category: str = "",
        q: str = "",
        limit: int = 1000,
    ):
        return {
            "items": _query_expenses(
                get_db,
                date_from=date_from,
                date_to=date_to,
                entry_type=entry_type,
                project=project,
                category=category,
                q=q,
                limit=limit,
            ),
            "summary": _summary(get_db),
            "categories": _expense_categories(get_db),
            "presets": _text_presets(get_db),
            "sync": _sync_status(get_db),
        }

    @app.get("/api/case-manager/prepaid-expenses/categories")
    def get_prepaid_expense_categories():
        return {"categories": _expense_categories(get_db)}

    @app.put("/api/case-manager/prepaid-expenses/categories")
    async def put_prepaid_expense_categories(request: Request):
        body = await request.json()
        return {"categories": _save_expense_categories(get_db, body.get("categories"))}

    @app.get("/api/case-manager/prepaid-expenses/presets")
    def get_prepaid_expense_presets():
        return {"presets": _text_presets(get_db)}

    @app.put("/api/case-manager/prepaid-expenses/presets")
    async def put_prepaid_expense_presets(request: Request):
        return {"presets": _save_text_presets(get_db, await request.json())}

    @app.get("/api/case-manager/prepaid-expenses/feishu-settings")
    def get_feishu_expense_settings():
        return {**_settings_response(), "sync": _sync_status(get_db)}

    @app.put("/api/case-manager/prepaid-expenses/feishu-settings")
    async def put_feishu_expense_settings(request: Request):
        body = await request.json()
        existing = _read_config()
        base_url = _text(body.get("base_url"))
        app_token, table_id = _parse_base_url(base_url)
        secret = _text(body.get("app_secret"))
        if not secret or "•" in secret:
            secret = existing.get("app_secret", "")
        config = {
            "app_id": _text(body.get("app_id")),
            "app_secret": secret,
            "base_url": base_url,
            "form_url": _text(body.get("form_url")),
            "app_token": app_token,
            "table_id": table_id,
        }
        if not config["app_id"] or not config["app_secret"]:
            raise HTTPException(status_code=400, detail="请填写飞书 App ID 和 App Secret")
        _write_config(config)
        syncer._token = ""
        try:
            verified = await syncer.verify(config) if body.get("verify", True) else {"ok": True, "message": "设置已保存"}
            if body.get("verify", True):
                verified["project_options"] = await syncer.sync_project_options(force=True)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=400, detail=_friendly_sync_error(exc)) from exc
        return {**_settings_response(config), "verification": verified}

    @app.post("/api/case-manager/prepaid-expenses")
    async def create_prepaid_expense(request: Request):
        return _save_expense(get_db, await request.json())

    @app.put("/api/case-manager/prepaid-expenses/{item_id}")
    async def update_prepaid_expense(item_id: int, request: Request):
        return _save_expense(get_db, await request.json(), item_id=item_id)

    @app.delete("/api/case-manager/prepaid-expenses/{item_id}")
    def delete_prepaid_expense(item_id: int):
        with get_db() as conn:
            existing = conn.execute("SELECT id FROM case_manager_prepaid_expenses WHERE id=?", (item_id,)).fetchone()
            if not existing:
                raise HTTPException(status_code=404, detail="预支费用记录不存在")
            conn.execute("DELETE FROM case_manager_prepaid_expenses WHERE id=?", (item_id,))
            conn.commit()
        return {"ok": True}

    @app.get("/api/case-manager/prepaid-expenses/export")
    def export_prepaid_expenses(
        date_from: str = "",
        date_to: str = "",
        entry_type: str = "",
        project: str = "",
        category: str = "",
        q: str = "",
    ):
        items = _query_expenses(
            get_db,
            date_from=date_from,
            date_to=date_to,
            entry_type=entry_type,
            project=project,
            category=category,
            q=q,
            limit=None,
        )
        filename = f"团队备用金支出明细表（{_today_iso()}）.xlsx"
        return StreamingResponse(
            _export_workbook(items),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
        )

    @app.post("/api/case-manager/prepaid-expenses/feishu-sync")
    async def sync_feishu_expenses():
        return await syncer.sync()

    @app.post("/api/case-manager/prepaid-expenses/feishu-project-sync")
    async def sync_feishu_project_options():
        try:
            return await syncer.sync_project_options(force=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=_friendly_sync_error(exc)) from exc

    async def scheduler_loop() -> None:
        await asyncio.sleep(5)
        next_expense_sync = 0.0
        while True:
            try:
                if _settings_response().get("configured"):
                    await syncer.sync_project_options()
                    loop_time = asyncio.get_running_loop().time()
                    if loop_time >= next_expense_sync:
                        await syncer.sync()
                        next_expense_sync = loop_time + 300
            except asyncio.CancelledError:
                raise
            except Exception:
                pass
            await asyncio.sleep(15)

    @app.on_event("startup")
    async def start_prepaid_expense_scheduler():
        nonlocal scheduler_task
        if scheduler_task is None or scheduler_task.done():
            scheduler_task = asyncio.create_task(scheduler_loop())

    @app.on_event("shutdown")
    async def stop_prepaid_expense_scheduler():
        nonlocal scheduler_task
        if scheduler_task:
            scheduler_task.cancel()
            try:
                await scheduler_task
            except asyncio.CancelledError:
                pass
            scheduler_task = None
