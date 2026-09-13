from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import re
import uuid
from decimal import Decimal, InvalidOperation
from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple
from urllib.parse import quote

import openpyxl
from fastapi import File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries


SERIES_KINDS = {"ordinary", "series", "series_child"}
INHERIT_MODES = {"inherit", "supplement", "replace", "not_applicable"}
CLAIM_MODES = {"common", "inherit", "replace", "supplement", "exclude", "add"}
MILESTONE_STATUSES = {"planned", "completed", "cancelled"}
MILESTONE_TYPES = ["立案/受理", "送达", "开庭", "裁判", "执行", "回款", "结案", "自定义"]

DEFAULT_SERIES_COLUMNS = [
    {"key": "series_code", "label": "编号", "field_type": "text", "source": "builtin", "width": 92},
    {"key": "display_name", "label": "子案名称", "field_type": "text", "source": "builtin", "width": 210},
    {"key": "clients", "label": "委托人", "field_type": "lookup", "source": "clients", "width": 170},
    {"key": "opponents", "label": "对方当事人", "field_type": "lookup", "source": "opponents", "width": 190},
    {"key": "case_type", "label": "案由", "field_type": "lookup", "source": "case_types", "width": 170},
    {"key": "case_no", "label": "案号", "field_type": "text", "source": "builtin", "width": 165},
    {"key": "court", "label": "法院/仲裁委", "field_type": "lookup", "source": "courts", "width": 180},
    {"key": "current_stage", "label": "阶段", "field_type": "lookup", "source": "stages", "width": 120},
    {"key": "latest_milestone", "label": "最近节点", "field_type": "readonly", "source": "derived", "width": 135},
    {"key": "next_milestone", "label": "下一节点", "field_type": "readonly", "source": "derived", "width": 135},
    {"key": "claim_totals", "label": "请求金额", "field_type": "readonly", "source": "derived", "width": 135},
    {"key": "other_differences", "label": "差异", "field_type": "textarea", "source": "builtin", "width": 190},
    {"key": "case_focus", "label": "侧重点", "field_type": "textarea", "source": "builtin", "width": 180},
    {"key": "memo", "label": "备忘", "field_type": "textarea", "source": "builtin", "width": 180},
]

BUILTIN_COLUMN_KEYS = {item["key"] for item in DEFAULT_SERIES_COLUMNS} | {
    "court_contact", "party_contact", "facts_difference", "contract_difference", "source_status"
}
EDITABLE_CASE_FIELDS = {
    "series_code", "display_name", "case_type", "current_stage", "court_contact", "party_contact",
    "memo", "facts_difference", "contract_difference", "other_differences", "case_focus",
}

HEADER_ALIASES = {
    "series_code": {"编号", "序号", "子案编号", "案件编号", "序列号"},
    "display_name": {"子案名称", "案件名称", "名称", "项目名称"},
    "clients": {"委托人", "委托方", "客户", "我方", "申请人（我方）", "原告（我方）"},
    "opponents": {"对方", "对方当事人", "对方主体", "相对方"},
    "case_type": {"案件类型", "案由", "纠纷类型"},
    "case_no": {"案号", "案件案号", "法院案号", "仲裁案号"},
    "court": {"受理机构", "法院", "仲裁委", "法院/仲裁委", "审理机构"},
    "current_stage": {"当前阶段", "阶段", "案件阶段", "进度", "当前进度"},
    "other_differences": {"其他差异", "差异", "个案差异", "特殊情况"},
    "case_focus": {"侧重点", "案件侧重点", "重点", "处理重点"},
    "memo": {"备忘", "备注", "案件备忘", "说明"},
}

AMOUNT_HEADER_TOKENS = (
    "标的金额", "诉讼标的", "仲裁标的", "请求金额", "涉案金额", "争议金额", "合同金额",
    "合同价款", "欠款金额", "本金金额", "赔偿金额", "总金额", "合计金额", "金额", "价款",
)
DEFAULT_TABLE_SETTINGS = {"freeze_columns": 2, "freeze_rows": 0}


def _text(value: Any) -> str:
    return str(value or "").strip()


def _now() -> str:
    return dt.datetime.now().isoformat(timespec="seconds")


def _row(row: Any) -> Dict[str, Any]:
    return dict(row) if row else {}


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False)


def _case_columns(conn: Any) -> set[str]:
    return {r[1] for r in conn.execute("PRAGMA table_info(case_manager_cases)").fetchall()}


def _default_columns_json() -> str:
    return _json(DEFAULT_SERIES_COLUMNS)


def _profile_columns(profile: Dict[str, Any]) -> List[Dict[str, Any]]:
    try:
        raw = json.loads(profile.get("table_columns_json") or "[]")
    except (TypeError, ValueError):
        raw = []
    return _sanitize_columns(raw) if isinstance(raw, list) and raw else [dict(item) for item in DEFAULT_SERIES_COLUMNS]


def _profile_settings(profile: Dict[str, Any]) -> Dict[str, int]:
    try:
        raw = json.loads(profile.get("table_settings_json") or "{}")
    except (TypeError, ValueError):
        raw = {}
    if not isinstance(raw, dict):
        raw = {}
    try:
        freeze_columns = max(0, min(int(raw.get("freeze_columns", 2)), 30))
    except (TypeError, ValueError):
        freeze_columns = 2
    try:
        freeze_rows = max(0, min(int(raw.get("freeze_rows", 0)), 20))
    except (TypeError, ValueError):
        freeze_rows = 0
    return {"freeze_columns": freeze_columns, "freeze_rows": freeze_rows}


def _profile_custom_fields(profile: Dict[str, Any]) -> List[Dict[str, str]]:
    try:
        raw = json.loads(profile.get("custom_fields_json") or "[]")
    except (TypeError, ValueError, json.JSONDecodeError):
        raw = []
    fields: List[Dict[str, str]] = []
    seen: set[str] = set()
    for index, item in enumerate(raw if isinstance(raw, list) else []):
        if not isinstance(item, dict):
            continue
        label = _text(item.get("label"))
        key = _text(item.get("key")) or f"custom_{index + 1}"
        if not label or key in seen:
            continue
        seen.add(key)
        fields.append({"key": key, "label": label, "value": _text(item.get("value"))})
    return fields


def _profile_public_field_order(profile: Dict[str, Any]) -> List[str]:
    try:
        raw = json.loads(profile.get("public_field_order_json") or "[]")
    except (TypeError, ValueError, json.JSONDecodeError):
        raw = []
    if not isinstance(raw, list):
        return []
    return [value for value in (_text(item) for item in raw) if value]


def _profile_public_field_settings(profile: Dict[str, Any]) -> List[Dict[str, Any]]:
    try:
        raw = json.loads(profile.get("public_field_settings_json") or "[]")
    except (TypeError, ValueError, json.JSONDecodeError):
        raw = []
    if not isinstance(raw, list):
        return []
    result, seen = [], set()
    for item in raw:
        if not isinstance(item, dict):
            continue
        key = _text(item.get("key"))
        if not key or key in seen:
            continue
        seen.add(key)
        result.append({"key": key, "label": _text(item.get("label")), "hidden": bool(item.get("hidden"))})
    return result


DETAIL_MAPPING_TARGETS = {
    "parties.other", "contacts", "memo", "case_type", "current_stage", "case_no", "court",
}


def _profile_detail_field_mappings(profile: Dict[str, Any]) -> List[Dict[str, str]]:
    """Return stable, safe display mappings for a series' child-case detail view."""
    try:
        raw = json.loads(profile.get("detail_field_mappings_json") or "[]")
    except (TypeError, ValueError, json.JSONDecodeError):
        raw = []
    if not isinstance(raw, list):
        return []
    result: List[Dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for item in raw:
        if not isinstance(item, dict):
            continue
        source_key = _text(item.get("source_key"))
        target_key = _text(item.get("target_key"))
        marker = (source_key, target_key)
        if not source_key or target_key not in DETAIL_MAPPING_TARGETS or marker in seen:
            continue
        seen.add(marker)
        result.append({"source_key": source_key, "target_key": target_key})
    return result


def _custom_key(label: str) -> str:
    return "custom_" + hashlib.sha1(label.encode("utf-8")).hexdigest()[:12]


def _header_builtin(label: str) -> str:
    normalized = _text(label).replace(" ", "").replace("\n", "")
    for key, aliases in HEADER_ALIASES.items():
        if normalized in {item.replace(" ", "") for item in aliases}:
            return key
    return ""


def _is_amount_header(label: str) -> bool:
    normalized = _text(label).replace(" ", "").replace("\n", "")
    return any(token in normalized for token in AMOUNT_HEADER_TOKENS)


def _sanitize_columns(items: Iterable[Dict[str, Any]]) -> List[Dict[str, Any]]:
    result: List[Dict[str, Any]] = []
    seen: set[str] = set()
    for index, raw in enumerate(items or []):
        label = _text(raw.get("label"))
        key = _text(raw.get("key")) or (_header_builtin(label) or _custom_key(label or f"列{index + 1}"))
        if not label or key in seen:
            continue
        seen.add(key)
        source = _text(raw.get("source")) or ("builtin" if key in BUILTIN_COLUMN_KEYS else "custom")
        field_type = _text(raw.get("field_type")) or ("lookup" if source in {"clients", "opponents", "parties", "case_types", "stages", "courts", "party_roles", "number_types"} else "text")
        semantic_type = _text(raw.get("semantic_type")) if "semantic_type" in raw else ("amount" if _is_amount_header(label) and key != "claim_totals" else "")
        currency = _text(raw.get("currency")) or "CNY"
        if key in {"latest_milestone", "next_milestone", "claim_totals", "source_status"}:
            field_type, source = "readonly", "derived"
        elif semantic_type == "amount":
            field_type = "amount"
        try:
            width = max(80, min(int(raw.get("width") or 150), 420))
        except (TypeError, ValueError):
            width = 150
        result.append({"key": key, "label": label, "field_type": field_type, "source": source, "width": width,
                       "semantic_type": semantic_type, "currency": currency,
                       "manual_width": bool(raw.get("manual_width")),
                       "hidden": bool(raw.get("hidden"))})
    return result or [dict(item) for item in DEFAULT_SERIES_COLUMNS]


def ensure_series_schema(get_db: Callable[[], Any]) -> None:
    """Idempotent additive migration. Existing rows remain ordinary cases."""
    with get_db() as conn:
        columns = _case_columns(conn)
        additions = {
            "case_kind": "TEXT NOT NULL DEFAULT 'ordinary'",
            "parent_series_id": "INTEGER",
            "series_code": "TEXT DEFAULT ''",
            "series_order": "INTEGER DEFAULT 0",
            "series_display_name": "TEXT DEFAULT ''",
            "facts_mode": "TEXT DEFAULT 'inherit'",
            "facts_difference": "TEXT DEFAULT ''",
            "contract_mode": "TEXT DEFAULT 'inherit'",
            "contract_difference": "TEXT DEFAULT ''",
            "other_differences": "TEXT DEFAULT ''",
            "case_focus": "TEXT DEFAULT ''",
            "pre_delete_status": "TEXT DEFAULT ''",
        }
        for name, declaration in additions.items():
            if name not in columns:
                conn.execute(f"ALTER TABLE case_manager_cases ADD COLUMN {name} {declaration}")
        conn.execute("UPDATE case_manager_cases SET case_kind='ordinary' WHERE COALESCE(case_kind, '')='' OR case_kind NOT IN ('ordinary','series','series_child')")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_cases_kind_parent ON case_manager_cases(case_kind,parent_series_id,series_order,id)")
        conn.execute(
            """CREATE TABLE IF NOT EXISTS case_manager_series_profiles (
                series_case_id INTEGER PRIMARY KEY,
                common_facts TEXT DEFAULT '',
                common_contract_terms TEXT DEFAULT '',
                series_memo TEXT DEFAULT '',
                custom_fields_json TEXT DEFAULT '[]',
                public_field_order_json TEXT DEFAULT '[]',
                public_field_settings_json TEXT DEFAULT '[]',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(series_case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
            )"""
        )
        profile_columns = {r[1] for r in conn.execute("PRAGMA table_info(case_manager_series_profiles)").fetchall()}
        if "table_columns_json" not in profile_columns:
            conn.execute("ALTER TABLE case_manager_series_profiles ADD COLUMN table_columns_json TEXT DEFAULT '[]'")
        if "table_settings_json" not in profile_columns:
            conn.execute("ALTER TABLE case_manager_series_profiles ADD COLUMN table_settings_json TEXT DEFAULT '{}'")
        if "custom_fields_json" not in profile_columns:
            conn.execute("ALTER TABLE case_manager_series_profiles ADD COLUMN custom_fields_json TEXT DEFAULT '[]'")
        if "public_field_order_json" not in profile_columns:
            conn.execute("ALTER TABLE case_manager_series_profiles ADD COLUMN public_field_order_json TEXT DEFAULT '[]'")
        if "public_field_settings_json" not in profile_columns:
            conn.execute("ALTER TABLE case_manager_series_profiles ADD COLUMN public_field_settings_json TEXT DEFAULT '[]'")
        if "detail_field_mappings_json" not in profile_columns:
            conn.execute("ALTER TABLE case_manager_series_profiles ADD COLUMN detail_field_mappings_json TEXT DEFAULT '[]'")
        conn.execute(
            "UPDATE case_manager_series_profiles SET table_columns_json=? WHERE COALESCE(table_columns_json,'') IN ('','[]')",
            (_default_columns_json(),),
        )
        conn.execute(
            "UPDATE case_manager_series_profiles SET table_settings_json=? WHERE COALESCE(table_settings_json,'') IN ('','{}')",
            (_json(DEFAULT_TABLE_SETTINGS),),
        )
        conn.execute(
            """CREATE TABLE IF NOT EXISTS case_manager_series_cell_values (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                series_case_id INTEGER NOT NULL,
                child_case_id INTEGER NOT NULL,
                column_key TEXT NOT NULL,
                value_text TEXT DEFAULT '',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(child_case_id,column_key),
                FOREIGN KEY(series_case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE,
                FOREIGN KEY(child_case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
            )"""
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_series_cells_parent_child ON case_manager_series_cell_values(series_case_id,child_case_id,column_key)")
        conn.execute(
            """CREATE TABLE IF NOT EXISTS case_manager_series_cell_formats (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                series_case_id INTEGER NOT NULL,
                child_case_id INTEGER NOT NULL,
                column_key TEXT NOT NULL,
                format_json TEXT NOT NULL DEFAULT '{}',
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                UNIQUE(child_case_id,column_key),
                FOREIGN KEY(series_case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE,
                FOREIGN KEY(child_case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
            )"""
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_series_cell_formats_parent_child ON case_manager_series_cell_formats(series_case_id,child_case_id,column_key)")
        conn.execute(
            """CREATE TABLE IF NOT EXISTS case_manager_claims (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                case_id INTEGER NOT NULL,
                parent_claim_id INTEGER,
                claim_mode TEXT NOT NULL DEFAULT 'add',
                claim_code TEXT DEFAULT '',
                claim_text TEXT NOT NULL DEFAULT '',
                amount_text TEXT DEFAULT '',
                currency TEXT DEFAULT 'CNY',
                calculation_basis TEXT DEFAULT '',
                legal_basis TEXT DEFAULT '',
                notes TEXT DEFAULT '',
                sort_order INTEGER DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE,
                FOREIGN KEY(parent_claim_id) REFERENCES case_manager_claims(id) ON DELETE CASCADE
            )"""
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_claims_case_sort ON case_manager_claims(case_id,sort_order,id)")
        conn.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_claim_override_unique ON case_manager_claims(case_id,parent_claim_id) WHERE parent_claim_id IS NOT NULL")
        conn.execute(
            """CREATE TABLE IF NOT EXISTS case_manager_milestones (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                case_id INTEGER NOT NULL,
                milestone_type TEXT NOT NULL DEFAULT '自定义',
                title TEXT NOT NULL DEFAULT '',
                milestone_status TEXT NOT NULL DEFAULT 'completed',
                happened_at TEXT DEFAULT '',
                notes TEXT DEFAULT '',
                sort_order INTEGER DEFAULT 0,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
            )"""
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_milestones_case_time ON case_manager_milestones(case_id,happened_at,id)")
        conn.execute(
            """CREATE TABLE IF NOT EXISTS case_manager_series_batches (
                batch_id TEXT PRIMARY KEY,
                series_case_id INTEGER NOT NULL,
                operation TEXT NOT NULL,
                summary_json TEXT DEFAULT '{}',
                created_at TEXT NOT NULL,
                FOREIGN KEY(series_case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
            )"""
        )
        conn.execute(
            """CREATE TABLE IF NOT EXISTS case_manager_series_history (
                history_id INTEGER PRIMARY KEY AUTOINCREMENT,
                series_case_id INTEGER NOT NULL,
                operation TEXT NOT NULL DEFAULT 'excel_import',
                summary_json TEXT NOT NULL DEFAULT '{}',
                snapshot_json TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY(series_case_id) REFERENCES case_manager_cases(id) ON DELETE CASCADE
            )"""
        )
        conn.execute("CREATE INDEX IF NOT EXISTS idx_series_history_series_time ON case_manager_series_history(series_case_id,created_at DESC,history_id DESC)")
        conn.commit()


def register_series_manager_routes(app: Any, get_db: Callable[[], Any]) -> None:
    ensure_series_schema(get_db)
    import_previews: Dict[str, Dict[str, Any]] = {}
    SERIES_HISTORY_LIMIT = 10

    def require_case(conn: Any, case_id: int, kind: str = "") -> Dict[str, Any]:
        item = _row(conn.execute("SELECT * FROM case_manager_cases WHERE id=?", (case_id,)).fetchone())
        if not item or (kind and item.get("case_kind") != kind):
            raise HTTPException(status_code=404, detail="系列案件或子案件不存在")
        return item

    def rows_for_case(conn: Any, table: str, case_ids: Iterable[int]) -> List[Dict[str, Any]]:
        ids = [int(value) for value in case_ids]
        if not ids:
            return []
        marks = ",".join("?" for _ in ids)
        return [_row(row) for row in conn.execute(
            f"SELECT * FROM {table} WHERE case_id IN ({marks}) ORDER BY id", ids
        ).fetchall()]

    def capture_series_snapshot(conn: Any, series_id: int) -> Dict[str, Any]:
        """Capture every field touched by an Excel import, including soft-deleted rows."""
        series = require_case(conn, series_id, "series")
        profile = _row(conn.execute(
            "SELECT * FROM case_manager_series_profiles WHERE series_case_id=?", (series_id,)
        ).fetchone())
        children = [_row(row) for row in conn.execute(
            "SELECT * FROM case_manager_cases WHERE parent_series_id=? AND case_kind='series_child' ORDER BY series_order,id",
            (series_id,),
        ).fetchall()]
        case_ids = [series_id, *[int(item["id"]) for item in children]]
        return {
            "version": 1,
            "series": series,
            "profile": profile,
            "children": children,
            "parties": rows_for_case(conn, "case_manager_parties", case_ids),
            "case_numbers": rows_for_case(conn, "case_manager_case_numbers", case_ids),
            "claims": rows_for_case(conn, "case_manager_claims", case_ids),
            "milestones": rows_for_case(conn, "case_manager_milestones", case_ids),
            "cell_values": [_row(row) for row in conn.execute(
                "SELECT * FROM case_manager_series_cell_values WHERE series_case_id=? ORDER BY id", (series_id,)
            ).fetchall()],
            "cell_formats": [_row(row) for row in conn.execute(
                "SELECT * FROM case_manager_series_cell_formats WHERE series_case_id=? ORDER BY id", (series_id,)
            ).fetchall()],
        }

    def create_series_history(conn: Any, series_id: int, operation: str, summary: Dict[str, Any]) -> int:
        snapshot = capture_series_snapshot(conn, series_id)
        cur = conn.execute(
            """INSERT INTO case_manager_series_history
               (series_case_id,operation,summary_json,snapshot_json,created_at)
               VALUES (?,?,?,?,?)""",
            (series_id, _text(operation) or "excel_import", _json(summary or {}), _json(snapshot), _now()),
        )
        history_ids = [int(row[0]) for row in conn.execute(
            """SELECT history_id FROM case_manager_series_history
               WHERE series_case_id=? ORDER BY datetime(created_at) DESC,history_id DESC""",
            (series_id,),
        ).fetchall()]
        for history_id in history_ids[SERIES_HISTORY_LIMIT:]:
            conn.execute("DELETE FROM case_manager_series_history WHERE history_id=?", (history_id,))
        return int(cur.lastrowid)

    def restore_rows_with_ids(conn: Any, table: str, rows: Iterable[Dict[str, Any]], case_ids: Iterable[int]) -> None:
        ids = [int(value) for value in case_ids]
        if ids:
            marks = ",".join("?" for _ in ids)
            conn.execute(f"DELETE FROM {table} WHERE case_id IN ({marks})", ids)
        allowed_case_ids = set(ids)
        columns = {row[1] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()}
        for row in rows:
            if row.get("case_id") is not None and int(row.get("case_id") or 0) not in allowed_case_ids:
                continue
            values = {key: value for key, value in row.items() if key in columns}
            if "id" not in values:
                continue
            names = list(values)
            marks = ",".join("?" for _ in names)
            conn.execute(
                f"INSERT INTO {table} ({','.join(names)}) VALUES ({marks})",
                [values[name] for name in names],
            )

    def restore_series_snapshot(conn: Any, series_id: int, snapshot: Dict[str, Any]) -> None:
        if int((snapshot.get("series") or {}).get("id") or 0) != series_id:
            raise HTTPException(status_code=400, detail="历史版本与当前系列不匹配")
        require_case(conn, series_id, "series")
        snapshot_children = snapshot.get("children") if isinstance(snapshot.get("children"), list) else []
        snapshot_by_id = {int(item["id"]): item for item in snapshot_children if item.get("id")}
        current_children = [_row(row) for row in conn.execute(
            "SELECT * FROM case_manager_cases WHERE parent_series_id=? AND case_kind='series_child' ORDER BY id", (series_id,)
        ).fetchall()]
        current_ids = {int(item["id"]) for item in current_children}
        child_ids = sorted(current_ids | set(snapshot_by_id))

        case_columns = {row[1] for row in conn.execute("PRAGMA table_info(case_manager_cases)").fetchall()}
        series_row = snapshot.get("series") or {}
        for row in [series_row, *snapshot_children]:
            row_id = int(row.get("id") or 0)
            if not row_id or (row_id != series_id and row_id not in current_ids):
                continue
            values = {key: value for key, value in row.items() if key in case_columns and key != "id"}
            if values:
                assignments = ",".join(f"{key}=?" for key in values)
                conn.execute(f"UPDATE case_manager_cases SET {assignments} WHERE id=?", [*values.values(), row_id])

        stamp = _now()
        for child in current_children:
            child_id = int(child["id"])
            if child_id not in snapshot_by_id:
                conn.execute(
                    "UPDATE case_manager_cases SET pre_delete_status=COALESCE(NULLIF(pre_delete_status,''),source_status),source_status='deleted',deleted_at=?,updated_at=? WHERE id=?",
                    (stamp, stamp, child_id),
                )

        profile = snapshot.get("profile") or {}
        profile_columns = {row[1] for row in conn.execute("PRAGMA table_info(case_manager_series_profiles)").fetchall()}
        profile_values = {key: value for key, value in profile.items() if key in profile_columns and key != "series_case_id"}
        if profile_values:
            assignments = ",".join(f"{key}=?" for key in profile_values)
            conn.execute(f"UPDATE case_manager_series_profiles SET {assignments} WHERE series_case_id=?", [*profile_values.values(), series_id])

        all_case_ids = [series_id, *child_ids]
        restore_rows_with_ids(conn, "case_manager_parties", snapshot.get("parties") or [], all_case_ids)
        restore_rows_with_ids(conn, "case_manager_case_numbers", snapshot.get("case_numbers") or [], all_case_ids)
        # Clearing common claims cascades to child overrides. Restore common
        # claims first, then restore the child rows that point at them.
        restore_rows_with_ids(conn, "case_manager_claims", snapshot.get("claims") or [], [series_id])
        restore_rows_with_ids(conn, "case_manager_claims", snapshot.get("claims") or [], child_ids)
        restore_rows_with_ids(conn, "case_manager_milestones", snapshot.get("milestones") or [], all_case_ids)

        for table, key in (("case_manager_series_cell_values", "cell_values"), ("case_manager_series_cell_formats", "cell_formats")):
            conn.execute(f"DELETE FROM {table} WHERE series_case_id=?", (series_id,))
            table_columns = {row[1] for row in conn.execute(f"PRAGMA table_info({table})").fetchall()}
            for row in snapshot.get(key) or []:
                values = {name: value for name, value in row.items() if name in table_columns}
                names = list(values)
                if not names:
                    continue
                marks = ",".join("?" for _ in names)
                conn.execute(f"INSERT INTO {table} ({','.join(names)}) VALUES ({marks})", [values[name] for name in names])

        restored_series = require_case(conn, series_id, "series")
        for child in snapshot_children:
            child_id = int(child.get("id") or 0)
            if child_id:
                sync_child_title(conn, restored_series, child_id, _text(child.get("series_display_name")), _text(child.get("series_code")), stamp)

    def audit(conn: Any, case_id: int, action: str, summary: str, changes: Optional[List[Dict[str, Any]]] = None, batch_id: str = "") -> None:
        stamp = _now()
        if batch_id:
            summary = f"[{batch_id[:8]}] {summary}"
        conn.execute(
            """INSERT INTO case_manager_audit_logs
               (case_id,entity_type,entity_id,action,summary,changes_json,actor,created_at)
               VALUES (?,'series',?,?,?,?,'本机用户',?)""",
            (case_id, case_id, action, summary, _json(changes or []), stamp),
        )

    def custom_values(conn: Any, child_id: int) -> Dict[str, str]:
        return {
            _text(row["column_key"]): _text(row["value_text"])
            for row in conn.execute(
                "SELECT column_key,value_text FROM case_manager_series_cell_values WHERE child_case_id=?",
                (child_id,),
            ).fetchall()
        }

    CELL_ALIGNMENTS = {"left", "center", "right"}

    def sanitize_cell_format(raw: Any) -> Dict[str, Any]:
        if not isinstance(raw, dict):
            return {}
        sanitized: List[Dict[str, Any]] = []
        for item in raw.get("runs") or []:
            if not isinstance(item, dict):
                continue
            try:
                start, end = int(item.get("start")), int(item.get("end"))
            except (TypeError, ValueError):
                continue
            if start < 0 or end <= start or end - start > 50000:
                continue
            run: Dict[str, Any] = {"start": start, "end": end}
            for key in ("bold", "italic", "underline"):
                if item.get(key):
                    run[key] = True
            if item.get("color") == "red":
                run["color"] = "red"
            if item.get("highlight") == "yellow":
                run["highlight"] = "yellow"
            if len(run) > 2:
                sanitized.append(run)
            if len(sanitized) >= 1000:
                break
        align = raw.get("align")
        result: Dict[str, Any] = {"runs": sanitized} if sanitized else {}
        if align in CELL_ALIGNMENTS:
            result["align"] = align
        return result

    def cell_formats(conn: Any, child_id: int) -> Dict[str, Dict[str, Any]]:
        result: Dict[str, Dict[str, Any]] = {}
        for row in conn.execute(
            "SELECT column_key,format_json FROM case_manager_series_cell_formats WHERE child_case_id=?",
            (child_id,),
        ).fetchall():
            try:
                value = json.loads(row["format_json"] or "{}")
            except (TypeError, ValueError, json.JSONDecodeError):
                value = {}
            normalized = sanitize_cell_format(value)
            if normalized:
                result[_text(row["column_key"])] = normalized
        return result

    def replace_cell_format(conn: Any, series_id: int, child_id: int, field: str, raw: Any) -> Dict[str, Any]:
        normalized = sanitize_cell_format(raw)
        if not normalized:
            conn.execute(
                "DELETE FROM case_manager_series_cell_formats WHERE child_case_id=? AND column_key=?",
                (child_id, field),
            )
            return {}
        stamp = _now()
        conn.execute(
            """INSERT INTO case_manager_series_cell_formats
               (series_case_id,child_case_id,column_key,format_json,created_at,updated_at)
               VALUES (?,?,?,?,?,?)
               ON CONFLICT(child_case_id,column_key) DO UPDATE SET format_json=excluded.format_json,updated_at=excluded.updated_at""",
            (series_id, child_id, field, _json(normalized), stamp, stamp),
        )
        return normalized

    def replace_custom_values(conn: Any, series_id: int, child_id: int, values: Dict[str, Any], merge: bool = True) -> None:
        if not merge:
            conn.execute("DELETE FROM case_manager_series_cell_values WHERE child_case_id=?", (child_id,))
        stamp = _now()
        for key, value in (values or {}).items():
            column_key = _text(key)
            if not column_key or column_key in BUILTIN_COLUMN_KEYS:
                continue
            conn.execute(
                """INSERT INTO case_manager_series_cell_values
                   (series_case_id,child_case_id,column_key,value_text,created_at,updated_at)
                   VALUES (?,?,?,?,?,?)
                   ON CONFLICT(child_case_id,column_key) DO UPDATE SET value_text=excluded.value_text,updated_at=excluded.updated_at""",
                (series_id, child_id, column_key, _text(value), stamp, stamp),
            )

    def replace_parties(conn: Any, case_id: int, parties: Iterable[Dict[str, Any]]) -> None:
        stamp = _now()
        conn.execute("DELETE FROM case_manager_parties WHERE case_id=?", (case_id,))
        for index, party in enumerate(parties or []):
            group = "client" if _text(party.get("party_group")) == "client" else "other"
            name = _text(party.get("name"))
            role = _text(party.get("role"))
            credit_code = _text(party.get("credit_code"))
            if not (name or role):
                continue
            client_id = None
            opponent_id = None
            if name and group == "client":
                conn.execute(
                    """INSERT OR IGNORE INTO case_manager_clients
                       (client_type,name,identity_code,address,contact_person,phone,invoice_info,created_at,updated_at)
                       VALUES ('单位',?,?, '', '', '', '', ?,?)""",
                    (name, credit_code, stamp, stamp),
                )
                row = conn.execute("SELECT id FROM case_manager_clients WHERE name=?", (name,)).fetchone()
                client_id = row[0] if row else None
            elif name:
                conn.execute(
                    """INSERT OR IGNORE INTO case_manager_opponents
                       (opponent_type,name,identity_code,address,contact_person,phone,notes,created_at,updated_at)
                       VALUES ('单位',?,?, '', '', '', '', ?,?)""",
                    (name, credit_code, stamp, stamp),
                )
                row = conn.execute("SELECT id FROM case_manager_opponents WHERE name=?", (name,)).fetchone()
                opponent_id = row[0] if row else None
            conn.execute(
                """INSERT INTO case_manager_parties
                   (case_id,party_group,role,name,credit_code,client_id,opponent_id,sort_order,created_at,updated_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?)""",
                (case_id, group, role, name, credit_code, client_id, opponent_id, index, stamp, stamp),
            )

    def replace_case_numbers(conn: Any, case_id: int, numbers: Iterable[Dict[str, Any]]) -> None:
        stamp = _now()
        conn.execute("DELETE FROM case_manager_case_numbers WHERE case_id=?", (case_id,))
        for index, item in enumerate(numbers or []):
            values = {
                "number_type": _text(item.get("number_type")),
                "case_no": _text(item.get("case_no")),
                "court": _text(item.get("court")),
                "handler_contact": _text(item.get("handler_contact")),
                "first_filing_at": _text(item.get("first_filing_at"))[:10],
                "first_judgment_served_at": _text(item.get("first_judgment_served_at"))[:10],
                "second_filing_at": _text(item.get("second_filing_at"))[:10],
                "second_judgment_effective_at": _text(item.get("second_judgment_effective_at"))[:10],
            }
            if not any(values.values()):
                continue
            conn.execute(
                """INSERT INTO case_manager_case_numbers
                   (case_id,number_type,case_no,court,handler_contact,sort_order,first_filing_at,
                    first_judgment_served_at,second_filing_at,second_judgment_effective_at,created_at,updated_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?)""",
                (case_id, values["number_type"], values["case_no"], values["court"], values["handler_contact"], index,
                 values["first_filing_at"], values["first_judgment_served_at"], values["second_filing_at"],
                 values["second_judgment_effective_at"], stamp, stamp),
            )

    def replace_claims(conn: Any, case_id: int, claims: Iterable[Dict[str, Any]], parent_kind: bool = False) -> None:
        stamp = _now()
        claims = list(claims or [])
        if parent_kind:
            # Preserve stable common-claim ids so child inherit/replace/exclude
            # links survive edits to the common template.
            existing = [_row(r) for r in conn.execute(
                "SELECT * FROM case_manager_claims WHERE case_id=? ORDER BY sort_order,id", (case_id,)
            ).fetchall()]
            by_code = {_text(item.get("claim_code")): item for item in existing if _text(item.get("claim_code"))}
            retained: set[int] = set()
            for index, claim in enumerate(claims):
                code = _text(claim.get("claim_code")) or f"R{index + 1:03d}"
                current = by_code.get(code)
                values = (
                    code, _text(claim.get("claim_text")), _text(claim.get("amount_text")),
                    _text(claim.get("currency")) or "CNY", _text(claim.get("calculation_basis")),
                    _text(claim.get("legal_basis")), _text(claim.get("notes")), index, stamp,
                )
                if current:
                    retained.add(int(current["id"]))
                    conn.execute(
                        """UPDATE case_manager_claims SET claim_mode='common',claim_code=?,claim_text=?,amount_text=?,currency=?,
                           calculation_basis=?,legal_basis=?,notes=?,sort_order=?,updated_at=? WHERE id=?""",
                        (*values, current["id"]),
                    )
                else:
                    cur = conn.execute(
                        """INSERT INTO case_manager_claims
                           (case_id,parent_claim_id,claim_mode,claim_code,claim_text,amount_text,currency,
                            calculation_basis,legal_basis,notes,sort_order,created_at,updated_at)
                           VALUES (?,NULL,'common',?,?,?,?,?,?,?,?,?,?)""",
                        (case_id, *values[:-1], stamp, stamp),
                    )
                    retained.add(int(cur.lastrowid))
            for item in existing:
                if int(item["id"]) not in retained:
                    conn.execute("DELETE FROM case_manager_claims WHERE id=?", (item["id"],))
            return
        conn.execute("DELETE FROM case_manager_claims WHERE case_id=?", (case_id,))
        for index, claim in enumerate(claims):
            mode = _text(claim.get("claim_mode")) or ("common" if parent_kind else "add")
            if parent_kind:
                mode = "common"
            elif mode not in CLAIM_MODES - {"common"}:
                mode = "add"
            parent_claim_id = claim.get("parent_claim_id") or None
            conn.execute(
                """INSERT INTO case_manager_claims
                   (case_id,parent_claim_id,claim_mode,claim_code,claim_text,amount_text,currency,
                    calculation_basis,legal_basis,notes,sort_order,created_at,updated_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (case_id, parent_claim_id, mode, _text(claim.get("claim_code")), _text(claim.get("claim_text")),
                 _text(claim.get("amount_text")), _text(claim.get("currency")) or "CNY",
                 _text(claim.get("calculation_basis")), _text(claim.get("legal_basis")), _text(claim.get("notes")),
                 index, stamp, stamp),
            )

    def replace_milestones(conn: Any, case_id: int, milestones: Iterable[Dict[str, Any]]) -> None:
        stamp = _now()
        conn.execute("DELETE FROM case_manager_milestones WHERE case_id=?", (case_id,))
        for index, item in enumerate(milestones or []):
            mtype = _text(item.get("milestone_type")) or "自定义"
            status = _text(item.get("milestone_status")) or "completed"
            if status not in MILESTONE_STATUSES:
                status = "completed"
            conn.execute(
                """INSERT INTO case_manager_milestones
                   (case_id,milestone_type,title,milestone_status,happened_at,notes,sort_order,created_at,updated_at)
                   VALUES (?,?,?,?,?,?,?,?,?)""",
                (case_id, mtype, _text(item.get("title")) or mtype, status, _text(item.get("happened_at")),
                 _text(item.get("notes")), index, stamp, stamp),
            )

    def effective_claims(conn: Any, child_id: int, series_id: int) -> List[Dict[str, Any]]:
        common = [_row(r) for r in conn.execute(
            "SELECT * FROM case_manager_claims WHERE case_id=? ORDER BY sort_order,id", (series_id,)
        ).fetchall()]
        own = [_row(r) for r in conn.execute(
            "SELECT * FROM case_manager_claims WHERE case_id=? ORDER BY sort_order,id", (child_id,)
        ).fetchall()]
        overrides = {int(c["parent_claim_id"]): c for c in own if c.get("parent_claim_id")}
        result: List[Dict[str, Any]] = []
        for base in common:
            override = overrides.get(int(base["id"]))
            if not override or override.get("claim_mode") == "inherit":
                result.append({**base, "source": "common", "claim_mode": "inherit"})
            elif override.get("claim_mode") == "exclude":
                continue
            elif override.get("claim_mode") == "supplement":
                merged = dict(base)
                for key in ("claim_text", "calculation_basis", "legal_basis", "notes"):
                    extra = _text(override.get(key))
                    if extra:
                        merged[key] = "\n".join(x for x in [_text(base.get(key)), extra] if x)
                if _text(override.get("amount_text")):
                    merged["amount_text"] = _text(override.get("amount_text"))
                merged.update({"source": "supplement", "claim_mode": "supplement", "override_id": override["id"]})
                result.append(merged)
            else:
                result.append({**override, "source": "replace", "claim_mode": "replace"})
        result.extend({**claim, "source": "child"} for claim in own if not claim.get("parent_claim_id") and claim.get("claim_mode") != "exclude")
        return result

    def amount_totals(claims: Iterable[Dict[str, Any]]) -> Dict[str, str]:
        totals: Dict[str, Decimal] = {}
        for claim in claims:
            raw = _text(claim.get("amount_text")).replace(",", "").replace("，", "")
            if not raw:
                continue
            try:
                amount = Decimal(raw)
            except InvalidOperation:
                continue
            currency = _text(claim.get("currency")) or "CNY"
            totals[currency] = totals.get(currency, Decimal("0")) + amount
        return {key: format(value, "f") for key, value in totals.items()}

    def parse_table_amount(value: Any, default_currency: str = "CNY") -> Tuple[str, Optional[Decimal]]:
        raw = _text(value)
        if not raw:
            return default_currency or "CNY", None
        currency = _text(default_currency) or "CNY"
        upper = raw.upper()
        if any(mark in upper for mark in ("HK$", "HKD", "港币", "港元")):
            currency = "HKD"
        elif any(mark in upper for mark in ("US$", "USD", "美元")) or ("$" in raw and "HK$" not in upper):
            currency = "USD"
        elif any(mark in raw for mark in ("€", "欧元")) or "EUR" in upper:
            currency = "EUR"
        elif any(mark in raw for mark in ("￥", "¥", "人民币")) or "RMB" in upper:
            currency = "CNY"
        negative = raw.startswith("(") and raw.endswith(")")
        cleaned = raw.replace(",", "").replace("，", "").replace(" ", "")
        matches = re.findall(r"[-+]?\d+(?:\.\d+)?", cleaned)
        if not matches:
            return currency, None
        try:
            amount = Decimal(matches[0])
        except InvalidOperation:
            return currency, None
        return currency, -amount if negative and amount > 0 else amount

    def table_amount_totals(conn: Any, children: Iterable[Dict[str, Any]], columns: Iterable[Dict[str, Any]]) -> Tuple[Dict[str, str], int, List[str]]:
        amount_columns = [item for item in columns if item.get("semantic_type") == "amount"]
        totals: Dict[str, Decimal] = {}
        parsed_count = 0
        for child in children:
            values = custom_values(conn, int(child["id"]))
            for column in amount_columns:
                key = _text(column.get("key"))
                raw = child.get(key) if key in child else values.get(key, "")
                currency, amount = parse_table_amount(raw, _text(column.get("currency")) or "CNY")
                if amount is None:
                    continue
                parsed_count += 1
                totals[currency] = totals.get(currency, Decimal("0")) + amount
        return ({key: format(value, "f") for key, value in totals.items()}, parsed_count,
                [_text(item.get("label")) for item in amount_columns])

    def merged_text(common: str, mode: str, difference: str) -> str:
        if mode == "not_applicable":
            return ""
        if mode == "replace":
            return difference
        if mode == "supplement":
            return "\n\n".join(x for x in [common, difference] if x)
        return common

    def reminder_stats(conn: Any, case_ids: List[int]) -> Dict[str, int]:
        if not case_ids:
            return {"overdue": 0, "next_7_days": 0, "reminder_total": 0}
        marks = ",".join("?" for _ in case_ids)
        today = dt.date.today()
        end = today + dt.timedelta(days=7)
        union = f"""
            SELECT remind_at AS due_at FROM case_manager_todos WHERE case_id IN ({marks}) AND COALESCE(status,'pending')<>'completed'
            UNION ALL SELECT hearing_at FROM case_manager_hearings WHERE case_id IN ({marks})
            UNION ALL SELECT deadline_at FROM case_manager_preservations WHERE case_id IN ({marks})
            UNION ALL SELECT deadline_at FROM case_manager_appeals WHERE case_id IN ({marks})
        """
        args = case_ids * 4
        rows = [r[0] for r in conn.execute(union, args).fetchall() if _text(r[0])]
        dates = []
        for value in rows:
            try:
                dates.append(dt.date.fromisoformat(_text(value)[:10]))
            except ValueError:
                pass
        return {
            "overdue": sum(1 for value in dates if value < today),
            "next_7_days": sum(1 for value in dates if today <= value <= end),
            "reminder_total": len(dates),
        }

    def child_detail(conn: Any, child_id: int) -> Dict[str, Any]:
        child = require_case(conn, child_id, "series_child")
        series_id = int(child.get("parent_series_id") or 0)
        series = require_case(conn, series_id, "series")
        profile = _row(conn.execute("SELECT * FROM case_manager_series_profiles WHERE series_case_id=?", (series_id,)).fetchone())
        child["display_name"] = child.get("series_display_name") or child.get("title")
        child["series_id"] = series_id
        child["series_title"] = series.get("title") or ""
        child["parties"] = [_row(r) for r in conn.execute("SELECT * FROM case_manager_parties WHERE case_id=? ORDER BY sort_order,id", (child_id,)).fetchall()]
        child["case_numbers"] = [_row(r) for r in conn.execute("SELECT * FROM case_manager_case_numbers WHERE case_id=? ORDER BY sort_order,id", (child_id,)).fetchall()]
        child["claims"] = [_row(r) for r in conn.execute("SELECT * FROM case_manager_claims WHERE case_id=? ORDER BY sort_order,id", (child_id,)).fetchall()]
        child["effective_claims"] = effective_claims(conn, child_id, series_id)
        child["claim_totals"] = amount_totals(child["effective_claims"])
        child["custom_values"] = custom_values(conn, child_id)
        child["cell_formats"] = cell_formats(conn, child_id)
        child["milestones"] = [_row(r) for r in conn.execute("SELECT * FROM case_manager_milestones WHERE case_id=? ORDER BY datetime(happened_at),sort_order,id", (child_id,)).fetchall()]
        for table, key in (("case_manager_todos", "todos"), ("case_manager_hearings", "hearings"), ("case_manager_preservations", "preservations"), ("case_manager_appeals", "appeals")):
            child[key] = [_row(r) for r in conn.execute(f"SELECT * FROM {table} WHERE case_id=? ORDER BY id DESC", (child_id,)).fetchall()]
        child["effective_facts"] = merged_text(_text(profile.get("common_facts")), _text(child.get("facts_mode")) or "inherit", _text(child.get("facts_difference")))
        child["effective_contract_terms"] = merged_text(_text(profile.get("common_contract_terms")), _text(child.get("contract_mode")) or "inherit", _text(child.get("contract_difference")))
        child["audit_logs"] = [_row(r) for r in conn.execute("SELECT * FROM case_manager_audit_logs WHERE case_id=? ORDER BY datetime(created_at) DESC,id DESC LIMIT 200", (child_id,)).fetchall()]
        summary = child_summary(conn, child)
        for key in ("clients", "opponents", "case_no", "court", "latest_milestone", "next_milestone", "risk", "reminder_stats"):
            child[key] = summary.get(key)
        return child

    def child_summary(conn: Any, child: Dict[str, Any]) -> Dict[str, Any]:
        cid = int(child["id"])
        parties = [_row(r) for r in conn.execute("SELECT party_group,role,name FROM case_manager_parties WHERE case_id=? ORDER BY sort_order,id", (cid,)).fetchall()]
        numbers = [_row(r) for r in conn.execute("SELECT case_no,court FROM case_manager_case_numbers WHERE case_id=? ORDER BY sort_order,id", (cid,)).fetchall()]
        milestones = [_row(r) for r in conn.execute("SELECT * FROM case_manager_milestones WHERE case_id=? ORDER BY datetime(happened_at),id", (cid,)).fetchall()]
        complete = [m for m in milestones if m.get("milestone_status") == "completed"]
        planned = [m for m in milestones if m.get("milestone_status") == "planned"]
        claims = effective_claims(conn, cid, int(child["parent_series_id"]))
        stats = reminder_stats(conn, [cid])
        return {
            **child,
            "display_name": child.get("series_display_name") or child.get("title"),
            "clients": "、".join(p["name"] for p in parties if p.get("party_group") == "client" and p.get("name")),
            "opponents": "、".join(p["name"] for p in parties if p.get("party_group") != "client" and p.get("name")),
            "case_no": "\n".join(n["case_no"] for n in numbers if n.get("case_no")),
            "court": "\n".join(n["court"] for n in numbers if n.get("court")),
            "latest_milestone": complete[-1] if complete else None,
            "next_milestone": planned[0] if planned else None,
            "claim_totals": amount_totals(claims),
            "risk": "overdue" if stats["overdue"] else "",
            "reminder_stats": stats,
            "custom_values": custom_values(conn, cid),
            "cell_formats": cell_formats(conn, cid),
        }

    def split_party_names(value: Any) -> List[str]:
        names: List[str] = []
        for part in re.split(r"[、，,；;\n]+", _text(value)):
            name = _text(part)
            if name and name not in names:
                names.append(name)
        return names

    def child_counterparty_short_name(conn: Any, series_id: int, child_id: int, display: str = "", code: str = "") -> str:
        parties = [_row(row) for row in conn.execute(
            "SELECT party_group,role,name FROM case_manager_parties WHERE case_id=? ORDER BY sort_order,id",
            (child_id,),
        ).fetchall()]
        opponents: List[str] = []
        clients: List[str] = []
        for party in parties:
            name = _text(party.get("name"))
            if not name:
                continue
            target = clients if _text(party.get("party_group")) == "client" else opponents
            if name not in target:
                target.append(name)
        if opponents:
            return "、".join(opponents)

        profile = _row(conn.execute(
            "SELECT * FROM case_manager_series_profiles WHERE series_case_id=?", (series_id,)
        ).fetchone())
        party_tokens = ("对方", "申请人", "被申请人", "原告", "被告", "上诉人", "被上诉人", "债权人", "债务人", "业主", "购房人", "承租人", "出租人", "当事人")
        columns = [
            column for column in _profile_columns(profile)
            if _text(column.get("key")).startswith("custom_")
            and (_text(column.get("source")) in {"parties", "clients", "opponents"}
                 or any(token in _text(column.get("label")) for token in party_tokens))
        ]
        columns.sort(key=lambda column: next(
            (index for index, token in enumerate(party_tokens) if token in _text(column.get("label"))),
            len(party_tokens),
        ))
        values = custom_values(conn, child_id)
        for column in columns:
            candidates = split_party_names(values.get(_text(column.get("key"))))
            candidates = [name for name in candidates if name not in clients]
            if candidates:
                return "、".join(candidates)
        fallback = _text(display)
        if not fallback or fallback == _text(code):
            fallback = _text(code) or fallback
        return fallback or "子案"

    def sync_child_title(conn: Any, series: Dict[str, Any], child_id: int, display: str = "", code: str = "", stamp: str = "") -> str:
        child = require_case(conn, child_id, "series_child")
        child_code = _text(code) or _text(child.get("series_code"))
        child_display = _text(display) or _text(child.get("series_display_name"))
        short_name = child_counterparty_short_name(conn, int(series["id"]), child_id, child_display, child_code)
        suffix = short_name if short_name.endswith("案") else f"{short_name}案"
        base_title = f"{_text(series.get('title'))}·{suffix}"
        title = base_title
        duplicate = conn.execute("SELECT id FROM case_manager_cases WHERE title=? AND id<>?", (title, child_id)).fetchone()
        if duplicate:
            title = f"{base_title}（{child_code or child_id}）"
        duplicate = conn.execute("SELECT id FROM case_manager_cases WHERE title=? AND id<>?", (title, child_id)).fetchone()
        if duplicate:
            title = f"{base_title}（{child_id}）"
        if title != _text(child.get("title")):
            changed_at = stamp or _now()
            conn.execute("UPDATE case_manager_cases SET title=?,updated_at=? WHERE id=?", (title, changed_at, child_id))
            for table in ("case_manager_todos", "case_manager_hearings", "case_manager_preservations", "case_manager_appeals"):
                conn.execute(f"UPDATE {table} SET title=?,updated_at=? WHERE case_id=?", (title, changed_at, child_id))
        return title

    def series_detail(conn: Any, series_id: int, include_children: bool = True) -> Dict[str, Any]:
        series = require_case(conn, series_id, "series")
        profile = _row(conn.execute("SELECT * FROM case_manager_series_profiles WHERE series_case_id=?", (series_id,)).fetchone())
        profile["table_columns"] = _profile_columns(profile)
        profile["table_settings"] = _profile_settings(profile)
        profile["custom_fields"] = _profile_custom_fields(profile)
        profile["public_field_order"] = _profile_public_field_order(profile)
        profile["public_field_settings"] = _profile_public_field_settings(profile)
        profile["detail_field_mappings"] = _profile_detail_field_mappings(profile)
        series["profile"] = profile
        client_columns = {row[1] for row in conn.execute("PRAGMA table_info(case_manager_clients)").fetchall()}
        client_field = lambda column, alias: f"c.{column} AS {alias}" if column in client_columns else f"'' AS {alias}"
        series["parties"] = [_row(r) for r in conn.execute(
            f"""SELECT p.*,c.client_type,c.contact_person AS client_contact_person,c.phone AS client_phone,
                      c.identity_code AS client_identity_code,c.address AS client_address,
                      {client_field('legal_representative', 'client_legal_representative')},
                      {client_field('registered_capital', 'client_registered_capital')},
                      {client_field('registration_status', 'client_registration_status')}
               FROM case_manager_parties p LEFT JOIN case_manager_clients c
                 ON c.id=p.client_id OR (p.party_group='client' AND c.name=p.name)
               WHERE p.case_id=? ORDER BY p.sort_order,p.id""", (series_id,)
        ).fetchall()]
        series["common_claims"] = [_row(r) for r in conn.execute("SELECT * FROM case_manager_claims WHERE case_id=? ORDER BY sort_order,id", (series_id,)).fetchall()]
        children = [_row(r) for r in conn.execute(
            "SELECT * FROM case_manager_cases WHERE parent_series_id=? AND case_kind='series_child' ORDER BY series_order,id", (series_id,)
        ).fetchall()]
        visible_children = [child for child in children if child.get("source_status") != "deleted"]
        stages: Dict[str, int] = {}
        totals: Dict[str, Decimal] = {}
        for child in visible_children:
            stage = _text(child.get("current_stage")) or "待补充"
            stages[stage] = stages.get(stage, 0) + 1
            for currency, amount in amount_totals(effective_claims(conn, int(child["id"]), series_id)).items():
                totals[currency] = totals.get(currency, Decimal("0")) + Decimal(amount)
        child_ids = [int(item["id"]) for item in visible_children]
        custom_totals, custom_amount_count, amount_column_labels = table_amount_totals(
            conn, [item for item in children if item.get("source_status") != "deleted"], profile["table_columns"]
        )
        structured_totals = {key: format(value, "f") for key, value in totals.items()}
        series["stats"] = {
            "total": len(visible_children),
            "active": sum(1 for item in visible_children if item.get("source_status") == "active"),
            "archived": sum(1 for item in visible_children if item.get("source_status") == "archived"),
            "deleted": sum(1 for item in children if item.get("source_status") == "deleted"),
            "stage_distribution": stages,
            "claim_totals": custom_totals if custom_amount_count else structured_totals,
            "structured_claim_totals": structured_totals,
            "table_amount_totals": custom_totals,
            "amount_source": "table_columns" if custom_amount_count else "structured_claims",
            "amount_columns": amount_column_labels,
            "amount_cells_count": custom_amount_count,
            **reminder_stats(conn, child_ids + [series_id]),
        }
        if include_children:
            series["children"] = [child_summary(conn, child) for child in children]
        series["audit_logs"] = [_row(r) for r in conn.execute("SELECT * FROM case_manager_audit_logs WHERE case_id=? ORDER BY datetime(created_at) DESC,id DESC LIMIT 200", (series_id,)).fetchall()]
        return series

    def validate_mode(value: Any) -> str:
        mode = _text(value) or "inherit"
        return mode if mode in INHERIT_MODES else "inherit"

    def create_child_record(conn: Any, series: Dict[str, Any], body: Dict[str, Any], order: Optional[int] = None) -> int:
        stamp = _now()
        code = _text(body.get("series_code"))
        display = _text(body.get("display_name") or body.get("series_display_name"))
        if not display:
            raise HTTPException(status_code=400, detail="子案件名称不能为空")
        if not code:
            next_number = conn.execute("SELECT count(*)+1 FROM case_manager_cases WHERE parent_series_id=?", (series["id"],)).fetchone()[0]
            code = f"{int(next_number):03d}"
        duplicate = conn.execute(
            "SELECT 1 FROM case_manager_cases WHERE parent_series_id=? AND series_code=? AND source_status<>'deleted'",
            (series["id"], code),
        ).fetchone()
        if duplicate:
            raise HTTPException(status_code=409, detail=f"系列内编号“{code}”已存在")
        storage_title = f"{series['title']} · {code} · {display}"
        if conn.execute("SELECT 1 FROM case_manager_cases WHERE title=?", (storage_title,)).fetchone():
            storage_title += f" · {uuid.uuid4().hex[:6]}"
        if order is not None:
            sort_order = max(0, int(order))
            conn.execute(
                """UPDATE case_manager_cases SET series_order=series_order+1
                   WHERE parent_series_id=? AND case_kind='series_child' AND series_order>=?""",
                (series["id"], sort_order),
            )
        else:
            sort_order = int(conn.execute("SELECT COALESCE(MAX(series_order),-1)+1 FROM case_manager_cases WHERE parent_series_id=?", (series["id"],)).fetchone()[0])
        cur = conn.execute(
            """INSERT INTO case_manager_cases
               (title,case_type,current_stage,court_contact,party_contact,memo,source_status,deleted_at,archived_at,pinned,
                created_at,updated_at,case_kind,parent_series_id,series_code,series_order,series_display_name,
                facts_mode,facts_difference,contract_mode,contract_difference,other_differences,case_focus,pre_delete_status)
               VALUES (?,?,?,?,?,?,'active','','',0,?,?,'series_child',?,?,?,?,?,?,?,?,?,?,?)""",
            (storage_title, _text(body.get("case_type")) or _text(series.get("case_type")), _text(body.get("current_stage")),
             _text(body.get("court_contact")), _text(body.get("party_contact")), _text(body.get("memo")), stamp, stamp,
             series["id"], code, sort_order, display, validate_mode(body.get("facts_mode")), _text(body.get("facts_difference")),
             validate_mode(body.get("contract_mode")), _text(body.get("contract_difference")),
             _text(body.get("other_differences")), _text(body.get("case_focus")), ""),
        )
        child_id = int(cur.lastrowid)
        parties = body.get("parties")
        if parties is None:
            parties = [_row(r) for r in conn.execute("SELECT party_group,role,name,credit_code FROM case_manager_parties WHERE case_id=? ORDER BY sort_order,id", (series["id"],)).fetchall()]
        replace_parties(conn, child_id, parties)
        replace_case_numbers(conn, child_id, body.get("case_numbers") or [])
        replace_claims(conn, child_id, body.get("claims") or [])
        replace_milestones(conn, child_id, body.get("milestones") or [])
        replace_custom_values(conn, int(series["id"]), child_id, body.get("custom_values") or {})
        sync_child_title(conn, series, child_id, display, code, stamp)
        audit(conn, child_id, "created", "创建系列子案")
        return child_id

    def update_child_record(conn: Any, child_id: int, body: Dict[str, Any]) -> None:
        child = require_case(conn, child_id, "series_child")
        series = require_case(conn, int(child["parent_series_id"]), "series")
        stamp = _now()
        code = _text(body.get("series_code")) or _text(child.get("series_code"))
        display = _text(body.get("display_name") or body.get("series_display_name")) or _text(child.get("series_display_name"))
        duplicate = conn.execute(
            """SELECT 1 FROM case_manager_cases
               WHERE parent_series_id=? AND series_code=? AND id<>? AND source_status<>'deleted'""",
            (series["id"], code, child_id),
        ).fetchone()
        if duplicate:
            raise HTTPException(status_code=409, detail=f"系列内编号“{code}”已存在")
        def next_text(key: str) -> str:
            return _text(body.get(key)) if key in body else _text(child.get(key))
        conn.execute(
            """UPDATE case_manager_cases SET title=?,case_type=?,current_stage=?,court_contact=?,party_contact=?,memo=?,
               series_code=?,series_display_name=?,facts_mode=?,facts_difference=?,contract_mode=?,contract_difference=?,
               other_differences=?,case_focus=?,updated_at=? WHERE id=?""",
            (_text(child.get("title")), next_text("case_type") or _text(series.get("case_type")), next_text("current_stage"),
             next_text("court_contact"), next_text("party_contact"), next_text("memo"), code, display,
             validate_mode(body.get("facts_mode") if "facts_mode" in body else child.get("facts_mode")), next_text("facts_difference"),
             validate_mode(body.get("contract_mode") if "contract_mode" in body else child.get("contract_mode")),
             next_text("contract_difference"), next_text("other_differences"), next_text("case_focus"), stamp, child_id),
        )
        if "parties" in body:
            replace_parties(conn, child_id, body.get("parties") or [])
        if "case_numbers" in body:
            replace_case_numbers(conn, child_id, body.get("case_numbers") or [])
        if "claims" in body:
            replace_claims(conn, child_id, body.get("claims") or [])
        if "milestones" in body:
            replace_milestones(conn, child_id, body.get("milestones") or [])
        if "custom_values" in body:
            replace_custom_values(conn, int(series["id"]), child_id, body.get("custom_values") or {})
        sync_child_title(conn, series, child_id, display, code, stamp)
        audit(conn, child_id, "updated", "编辑系列子案")

    def update_inline_value(conn: Any, series_id: int, child_id: int, field: str, value: Any) -> None:
        child = require_case(conn, child_id, "series_child")
        if int(child.get("parent_series_id") or 0) != series_id:
            raise HTTPException(status_code=404, detail="子案件不属于该系列")
        text_value = _text(value)
        if field in EDITABLE_CASE_FIELDS:
            if field == "display_name" and not text_value:
                raise HTTPException(status_code=400, detail="子案名称不能为空")
            update_child_record(conn, child_id, {field: text_value})
            return
        if field in {"clients", "opponents"}:
            group = "client" if field == "clients" else "other"
            existing = [_row(r) for r in conn.execute(
                "SELECT party_group,role,name,credit_code FROM case_manager_parties WHERE case_id=? ORDER BY sort_order,id",
                (child_id,),
            ).fetchall()]
            kept = [item for item in existing if item.get("party_group") != group]
            names = [name.strip() for name in text_value.replace("，", "、").replace(",", "、").split("、") if name.strip()]
            role = "委托人" if group == "client" else "对方当事人"
            replace_parties(conn, child_id, [*kept, *({"party_group": group, "role": role, "name": name} for name in names)])
            sync_child_title(conn, require_case(conn, series_id, "series"), child_id)
            audit(conn, child_id, "updated", f"表格内编辑：{'委托人' if group == 'client' else '对方当事人'}")
            return
        if field in {"case_no", "court"}:
            numbers = [_row(r) for r in conn.execute(
                "SELECT * FROM case_manager_case_numbers WHERE case_id=? ORDER BY sort_order,id", (child_id,)
            ).fetchall()]
            if not numbers:
                numbers = [{}]
            numbers[0][field] = text_value
            replace_case_numbers(conn, child_id, numbers)
            audit(conn, child_id, "updated", f"表格内编辑：{'案号' if field == 'case_no' else '受理机构'}")
            return
        profile = _row(conn.execute("SELECT * FROM case_manager_series_profiles WHERE series_case_id=?", (series_id,)).fetchone())
        column_keys = {item["key"] for item in _profile_columns(profile)}
        if field not in column_keys:
            raise HTTPException(status_code=400, detail="该列不存在或已被删除")
        replace_custom_values(conn, series_id, child_id, {field: text_value})
        sync_child_title(conn, require_case(conn, series_id, "series"), child_id)
        audit(conn, child_id, "updated", f"表格内编辑：{field}")

    # Idempotently upgrade existing series-child display titles without changing substantive case data.
    with get_db() as conn:
        existing_series = [_row(row) for row in conn.execute(
            "SELECT * FROM case_manager_cases WHERE case_kind='series'"
        ).fetchall()]
        for series in existing_series:
            child_rows = conn.execute(
                "SELECT id,series_code,series_display_name FROM case_manager_cases WHERE parent_series_id=? AND case_kind='series_child' ORDER BY series_order,id",
                (series["id"],),
            ).fetchall()
            for child in child_rows:
                sync_child_title(conn, series, int(child["id"]), _text(child["series_display_name"]), _text(child["series_code"]))
        conn.commit()

    @app.get("/api/case-manager/series-options")
    def series_options(kind: str, q: str = "", limit: int = 80):
        query = _text(q).lower()
        max_items = max(10, min(int(limit or 80), 300))
        with get_db() as conn:
            if kind == "clients":
                rows = conn.execute("SELECT id,name,identity_code,client_type,contact_person,phone FROM case_manager_clients ORDER BY name LIMIT 1000").fetchall()
                items = [{"id": row["id"], "value": _text(row["name"]), "label": _text(row["name"]), "identity_code": _text(row["identity_code"]), "meta": " · ".join(value for value in [_text(row["client_type"]), _text(row["contact_person"]), _text(row["phone"]), _text(row["identity_code"])] if value)} for row in rows]
            elif kind == "opponents":
                rows = conn.execute("SELECT id,name,identity_code FROM case_manager_opponents ORDER BY name LIMIT 1000").fetchall()
                items = [{"id": row["id"], "value": _text(row["name"]), "label": _text(row["name"]), "identity_code": _text(row["identity_code"]), "meta": _text(row["identity_code"])} for row in rows]
            elif kind == "parties":
                rows = conn.execute(
                    """SELECT id,name,identity_code,'委托人库' source FROM case_manager_clients
                       UNION ALL SELECT id,name,identity_code,'对方库' source FROM case_manager_opponents
                       ORDER BY name LIMIT 1600"""
                ).fetchall()
                items = [{"id": row["id"], "value": _text(row["name"]), "label": _text(row["name"]), "meta": " · ".join(x for x in [_text(row["source"]), _text(row["identity_code"])] if x)} for row in rows]
            else:
                sql_by_kind = {
                    "case_types": "SELECT DISTINCT case_type value FROM case_manager_cases WHERE case_type<>'' ORDER BY value",
                    "stages": "SELECT DISTINCT current_stage value FROM case_manager_cases WHERE current_stage<>'' ORDER BY value",
                    "party_roles": "SELECT DISTINCT role value FROM case_manager_parties WHERE role<>'' ORDER BY value",
                    "courts": "SELECT DISTINCT court value FROM case_manager_case_numbers WHERE court<>'' ORDER BY value",
                    "number_types": "SELECT DISTINCT number_type value FROM case_manager_case_numbers WHERE number_type<>'' ORDER BY value",
                }
                defaults = {
                    "stages": ["待补充", "立案", "一审", "二审", "再审审查", "执行", "仲裁", "调解", "已结案"],
                    "party_roles": ["委托人", "原告", "被告", "第三人", "申请人", "被申请人", "申请执行人", "被执行人"],
                    "number_types": ["一审", "二审", "再审", "执行", "仲裁"],
                }
                if kind not in sql_by_kind:
                    raise HTTPException(status_code=400, detail="不支持的候选数据类型")
                values = [_text(row[0]) for row in conn.execute(sql_by_kind[kind]).fetchall()]
                values.extend(defaults.get(kind, []))
                seen: set[str] = set()
                items = []
                for value in values:
                    if value and value not in seen:
                        seen.add(value)
                        items.append({"value": value, "label": value, "meta": ""})
        if query:
            items = [item for item in items if query in (item.get("label") or "").lower() or query in (item.get("meta") or "").lower()]
        return {"items": items[:max_items]}

    @app.post("/api/case-manager/series")
    async def create_series(request: Request):
        body = await request.json()
        title = _text(body.get("title"))
        if not title:
            raise HTTPException(status_code=400, detail="系列案件名称不能为空")
        initial_preview_token = _text(body.get("initial_import_preview_id"))
        initial_preview = import_previews.get(initial_preview_token) if initial_preview_token else None
        if initial_preview_token and not initial_preview:
            raise HTTPException(status_code=404, detail="首批子案导入预览已失效，请重新上传")
        if initial_preview and int(initial_preview.get("series_id") or 0) != 0:
            raise HTTPException(status_code=400, detail="首批子案导入预览无效")
        initial_parsed = (initial_preview or {}).get("parsed") or {}
        if initial_preview and initial_parsed.get("errors"):
            raise HTTPException(status_code=400, detail="首批子案表格仍有错误")
        stamp = _now()
        with get_db() as conn:
            existing = conn.execute(
                "SELECT case_kind,source_status FROM case_manager_cases WHERE title=? AND source_status<>'deleted' LIMIT 1",
                (title,),
            ).fetchone()
            if existing:
                location = "归档案件" if _text(existing["source_status"]) == "archived" else "在办案件"
                kind = "系列案件" if _text(existing["case_kind"]) == "series" else "普通案件"
                raise HTTPException(status_code=409, detail=f"{location}中已有同名{kind}，请打开原记录或修改系列案件名称")
            deleted_matches = conn.execute(
                "SELECT id FROM case_manager_cases WHERE title=? AND source_status='deleted' ORDER BY id",
                (title,),
            ).fetchall()
            for deleted in deleted_matches:
                deleted_id = int(deleted["id"])
                tombstone_title = f"{title} · 已删除 #{deleted_id}"
                conn.execute(
                    "UPDATE case_manager_cases SET title=?,updated_at=? WHERE id=?",
                    (tombstone_title, stamp, deleted_id),
                )
                audit(conn, deleted_id, "renamed", "为重新创建同名案件释放已删除记录的名称")
            cur = conn.execute(
                """INSERT INTO case_manager_cases
                   (title,case_type,current_stage,court_contact,party_contact,memo,source_status,deleted_at,archived_at,pinned,
                    created_at,updated_at,case_kind,parent_series_id,series_code,series_order,series_display_name)
                   VALUES (?,?,?,?,?,?,'active','','',0,?,?,'series',NULL,'',0,'')""",
                (title, _text(body.get("case_type")), _text(body.get("current_stage")), _text(body.get("court_contact")),
                 _text(body.get("party_contact")), _text(body.get("memo")), stamp, stamp),
            )
            series_id = int(cur.lastrowid)
            profile = body.get("profile") or {}
            initial_columns = initial_parsed.get("columns") if initial_parsed.get("generic") else None
            conn.execute(
                """INSERT INTO case_manager_series_profiles
                   (series_case_id,common_facts,common_contract_terms,series_memo,custom_fields_json,public_field_order_json,public_field_settings_json,table_columns_json,table_settings_json,created_at,updated_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                (series_id, _text(profile.get("common_facts")), _text(profile.get("common_contract_terms")),
                 _text(profile.get("series_memo")), _json(_profile_custom_fields({"custom_fields_json": _json(profile.get("custom_fields") or [])})), _json(_profile_public_field_order({"public_field_order_json": _json(profile.get("public_field_order") or [])})), _json(_profile_public_field_settings({"public_field_settings_json": _json(profile.get("public_field_settings") or [])})), _json(_sanitize_columns(initial_columns or profile.get("table_columns") or DEFAULT_SERIES_COLUMNS)),
                 _json(DEFAULT_TABLE_SETTINGS), stamp, stamp),
            )
            replace_parties(conn, series_id, body.get("parties") or [])
            replace_claims(conn, series_id, body.get("common_claims") or [], parent_kind=True)
            series = require_case(conn, series_id, "series")
            for index, child in enumerate(body.get("children") or []):
                create_child_record(conn, series, child, index)
            if initial_preview and initial_parsed.get("generic"):
                imported_columns = _sanitize_columns(initial_parsed.get("columns") or [])
                start_order = len(body.get("children") or [])
                for offset, row in enumerate(initial_parsed.get("generic_rows") or []):
                    child_body = generic_row_to_child(row, imported_columns, start_order + offset)
                    existing = conn.execute(
                        """SELECT id FROM case_manager_cases
                           WHERE parent_series_id=? AND case_kind='series_child' AND series_code=? AND source_status<>'deleted'""",
                        (series_id, child_body["series_code"]),
                    ).fetchone()
                    if existing:
                        update_child_record(conn, int(existing["id"]), child_body)
                    else:
                        create_child_record(conn, series, child_body, start_order + offset)
            audit(conn, series_id, "created", "创建系列案件")
            conn.commit()
            if initial_preview_token:
                import_previews.pop(initial_preview_token, None)
            return series_detail(conn, series_id)

    @app.get("/api/case-manager/series/{series_id}")
    def get_series(series_id: int):
        with get_db() as conn:
            return series_detail(conn, series_id)

    @app.put("/api/case-manager/series/{series_id}")
    async def update_series(series_id: int, request: Request):
        body = await request.json()
        with get_db() as conn:
            current = require_case(conn, series_id, "series")
            title = _text(body.get("title")) or _text(current.get("title"))
            existing = conn.execute(
                """SELECT case_kind,source_status FROM case_manager_cases
                   WHERE title=? AND id<>? AND source_status<>'deleted' LIMIT 1""",
                (title, series_id),
            ).fetchone()
            if existing:
                location = "归档案件" if _text(existing["source_status"]) == "archived" else "在办案件"
                kind = "系列案件" if _text(existing["case_kind"]) == "series" else "普通案件"
                raise HTTPException(status_code=409, detail=f"{location}中已有同名{kind}，请打开原记录或修改系列案件名称")
            stamp = _now()
            conn.execute(
                """UPDATE case_manager_cases SET title=?,case_type=?,current_stage=?,court_contact=?,party_contact=?,memo=?,updated_at=? WHERE id=?""",
                (title, _text(body.get("case_type")), _text(body.get("current_stage")), _text(body.get("court_contact")),
                 _text(body.get("party_contact")), _text(body.get("memo")), stamp, series_id),
            )
            if title != current.get("title"):
                for table in ("case_manager_todos", "case_manager_hearings", "case_manager_preservations", "case_manager_appeals"):
                    conn.execute(f"UPDATE {table} SET title=?,updated_at=? WHERE case_id=?", (title, stamp, series_id))
            if title != current.get("title"):
                renamed_series = {**current, "id": series_id, "title": title}
                children = conn.execute(
                    "SELECT id,series_code,series_display_name FROM case_manager_cases WHERE parent_series_id=? AND case_kind='series_child'",
                    (series_id,),
                ).fetchall()
                for child in children:
                    sync_child_title(conn, renamed_series, int(child["id"]), _text(child["series_display_name"]), _text(child["series_code"]), stamp)
            profile = body.get("profile") or {}
            current_profile = _row(conn.execute("SELECT * FROM case_manager_series_profiles WHERE series_case_id=?", (series_id,)).fetchone())
            next_columns = profile.get("table_columns") if "table_columns" in profile else _profile_columns(current_profile)
            next_settings = profile.get("table_settings") if "table_settings" in profile else _profile_settings(current_profile)
            next_custom_fields = profile.get("custom_fields") if "custom_fields" in profile else _profile_custom_fields(current_profile)
            next_public_field_order = profile.get("public_field_order") if "public_field_order" in profile else _profile_public_field_order(current_profile)
            next_public_field_settings = profile.get("public_field_settings") if "public_field_settings" in profile else _profile_public_field_settings(current_profile)
            conn.execute(
                """UPDATE case_manager_series_profiles SET common_facts=?,common_contract_terms=?,series_memo=?,custom_fields_json=?,public_field_order_json=?,public_field_settings_json=?,table_columns_json=?,table_settings_json=?,updated_at=? WHERE series_case_id=?""",
                (_text(profile.get("common_facts")), _text(profile.get("common_contract_terms")), _text(profile.get("series_memo")),
                 _json(_profile_custom_fields({"custom_fields_json": _json(next_custom_fields)})), _json(_profile_public_field_order({"public_field_order_json": _json(next_public_field_order)})), _json(_profile_public_field_settings({"public_field_settings_json": _json(next_public_field_settings)})), _json(_sanitize_columns(next_columns)), _json(_profile_settings({"table_settings_json": _json(next_settings)})), stamp, series_id),
            )
            if "parties" in body:
                replace_parties(conn, series_id, body.get("parties") or [])
            if "common_claims" in body:
                replace_claims(conn, series_id, body.get("common_claims") or [], parent_kind=True)
            audit(conn, series_id, "updated", "修改系列共性信息")
            conn.commit()
            return series_detail(conn, series_id)

    @app.get("/api/case-manager/series/{series_id}/children")
    def list_children(series_id: int, q: str = "", stage: str = "", status: str = "all", risk: str = "", sort: str = "series_order", direction: str = "asc", page: int = 1, page_size: int = 100):
        with get_db() as conn:
            require_case(conn, series_id, "series")
            where = ["c.parent_series_id=?", "c.case_kind='series_child'"]
            args: List[Any] = [series_id]
            if status in {"active", "archived", "deleted"}:
                where.append("c.source_status=?")
                args.append(status)
            if stage:
                where.append("c.current_stage=?")
                args.append(stage)
            if risk == "overdue":
                where.append("""(
                    EXISTS (SELECT 1 FROM case_manager_todos t WHERE t.case_id=c.id AND COALESCE(t.status,'pending')<>'completed' AND t.remind_at<>'' AND date(t.remind_at)<date('now','localtime'))
                    OR EXISTS (SELECT 1 FROM case_manager_hearings h WHERE h.case_id=c.id AND h.hearing_at<>'' AND date(h.hearing_at)<date('now','localtime'))
                    OR EXISTS (SELECT 1 FROM case_manager_preservations p WHERE p.case_id=c.id AND p.deadline_at<>'' AND date(p.deadline_at)<date('now','localtime'))
                    OR EXISTS (SELECT 1 FROM case_manager_appeals a WHERE a.case_id=c.id AND a.deadline_at<>'' AND date(a.deadline_at)<date('now','localtime'))
                )""")
            if q:
                pattern = f"%{q}%"
                where.append("""(c.series_code LIKE ? OR c.series_display_name LIKE ? OR c.case_focus LIKE ? OR c.other_differences LIKE ?
                    OR EXISTS (SELECT 1 FROM case_manager_parties p WHERE p.case_id=c.id AND (p.name LIKE ? OR p.role LIKE ?))
                    OR EXISTS (SELECT 1 FROM case_manager_case_numbers n WHERE n.case_id=c.id AND (n.case_no LIKE ? OR n.court LIKE ?))
                    OR EXISTS (SELECT 1 FROM case_manager_claims cl WHERE cl.case_id=c.id AND cl.claim_text LIKE ?)
                    OR EXISTS (SELECT 1 FROM case_manager_milestones m WHERE m.case_id=c.id AND (m.title LIKE ? OR m.notes LIKE ?))
                    OR EXISTS (SELECT 1 FROM case_manager_series_cell_values cv WHERE cv.child_case_id=c.id AND cv.value_text LIKE ?))""")
                args.extend([pattern] * 12)
            sql_where = " AND ".join(where)
            total = int(conn.execute(f"SELECT count(*) FROM case_manager_cases c WHERE {sql_where}", args).fetchone()[0])
            size = max(1, min(int(page_size or 100), 200))
            offset = (max(1, int(page or 1)) - 1) * size
            order_column = {
                "series_order": "c.series_order", "code": "c.series_code", "name": "c.series_display_name",
                "stage": "c.current_stage", "status": "c.source_status", "updated": "c.updated_at",
            }.get(sort, "c.series_order")
            order_direction = "DESC" if direction.lower() == "desc" else "ASC"
            rows = [_row(r) for r in conn.execute(
                f"SELECT c.* FROM case_manager_cases c WHERE {sql_where} ORDER BY {order_column} {order_direction},c.id {order_direction} LIMIT ? OFFSET ?",
                [*args, size, offset],
            ).fetchall()]
            items = [child_summary(conn, item) for item in rows]
            return {"items": items, "total": total, "page": max(1, page), "page_size": size}

    @app.post("/api/case-manager/series/{series_id}/children")
    async def create_child(series_id: int, request: Request):
        body = await request.json()
        with get_db() as conn:
            series = require_case(conn, series_id, "series")
            requested_order = body.get("series_order")
            child_id = create_child_record(conn, series, body, int(requested_order) if requested_order is not None else None)
            audit(conn, series_id, "updated", f"新增子案 {_text(body.get('display_name'))}")
            conn.commit()
            return child_detail(conn, child_id)

    @app.post("/api/case-manager/series/{series_id}/children/insert")
    async def insert_children(series_id: int, request: Request):
        body = await request.json()
        count = max(1, min(int(body.get("count") or 1), 100))
        position = max(0, int(body.get("position") or 0))
        with get_db() as conn:
            series = require_case(conn, series_id, "series")
            visible_count = int(conn.execute(
                """SELECT count(*) FROM case_manager_cases
                   WHERE parent_series_id=? AND case_kind='series_child' AND source_status<>'deleted'""",
                (series_id,),
            ).fetchone()[0])
            position = min(position, visible_count)
            created_ids: List[int] = []
            for offset in range(count):
                sequence = int(conn.execute(
                    "SELECT count(*)+1 FROM case_manager_cases WHERE parent_series_id=?",
                    (series_id,),
                ).fetchone()[0])
                child_id = create_child_record(
                    conn,
                    series,
                    {"display_name": f"新子案 {sequence}"},
                    position + offset,
                )
                created_ids.append(child_id)
            batch_id = uuid.uuid4().hex
            audit(conn, series_id, "updated", f"在主表第 {position + 1} 行前插入 {count} 个子案", batch_id=batch_id)
            conn.commit()
            return {
                "ok": True,
                "batch_id": batch_id,
                "items": [child_detail(conn, child_id) for child_id in created_ids],
            }

    @app.put("/api/case-manager/series/{series_id}/children/order")
    async def reorder_children(series_id: int, request: Request):
        body = await request.json()
        requested = [int(value) for value in body.get("child_ids") or [] if int(value or 0) > 0]
        if len(requested) != len(set(requested)):
            raise HTTPException(status_code=400, detail="子案排序中存在重复项目")
        with get_db() as conn:
            require_case(conn, series_id, "series")
            rows = conn.execute(
                """SELECT id FROM case_manager_cases
                   WHERE parent_series_id=? AND case_kind='series_child' AND source_status<>'deleted'
                   ORDER BY series_order,id""",
                (series_id,),
            ).fetchall()
            existing = [int(row[0]) for row in rows]
            if set(requested) != set(existing) or len(requested) != len(existing):
                raise HTTPException(status_code=400, detail="请在清除筛选后对全部可见子案排序")
            stamp = _now()
            for order, child_id in enumerate(requested):
                conn.execute(
                    "UPDATE case_manager_cases SET series_order=?,updated_at=? WHERE id=?",
                    (order, stamp, child_id),
                )
            audit(conn, series_id, "updated", "调整系列子案行顺序")
            conn.commit()
        return {"ok": True, "child_ids": requested}

    @app.get("/api/case-manager/series/{series_id}/children/{child_id}")
    def get_child(series_id: int, child_id: int):
        with get_db() as conn:
            child = child_detail(conn, child_id)
            if int(child["series_id"]) != series_id:
                raise HTTPException(status_code=404, detail="子案件不属于该系列")
            return child

    @app.put("/api/case-manager/series/{series_id}/children/{child_id}")
    async def update_child(series_id: int, child_id: int, request: Request):
        body = await request.json()
        with get_db() as conn:
            child = require_case(conn, child_id, "series_child")
            if int(child.get("parent_series_id") or 0) != series_id:
                raise HTTPException(status_code=404, detail="子案件不属于该系列")
            update_child_record(conn, child_id, body)
            conn.commit()
            return child_detail(conn, child_id)

    @app.patch("/api/case-manager/series/{series_id}/children/{child_id}/cell")
    async def update_child_cell(series_id: int, child_id: int, request: Request):
        body = await request.json()
        field = _text(body.get("field"))
        if not field:
            raise HTTPException(status_code=400, detail="缺少要修改的列")
        with get_db() as conn:
            update_inline_value(conn, series_id, child_id, field, body.get("value"))
            conn.execute("DELETE FROM case_manager_series_cell_formats WHERE child_case_id=? AND column_key=?", (child_id, field))
            conn.commit()
            return child_summary(conn, require_case(conn, child_id, "series_child"))

    @app.patch("/api/case-manager/series/{series_id}/children/{child_id}/cell-format")
    async def update_child_cell_format(series_id: int, child_id: int, request: Request):
        body = await request.json()
        field = _text(body.get("field"))
        if not field:
            raise HTTPException(status_code=400, detail="缺少要设置格式的列")
        with get_db() as conn:
            child = require_case(conn, child_id, "series_child")
            if int(child.get("parent_series_id") or 0) != series_id:
                raise HTTPException(status_code=404, detail="子案件不属于该系列")
            profile = _row(conn.execute("SELECT * FROM case_manager_series_profiles WHERE series_case_id=?", (series_id,)).fetchone())
            if field not in {item["key"] for item in _profile_columns(profile)}:
                raise HTTPException(status_code=400, detail="该列不存在或已被删除")
            normalized = replace_cell_format(conn, series_id, child_id, field, body.get("format"))
            audit(conn, child_id, "updated", f"表格单元格格式：{field}")
            conn.commit()
            return {"field": field, "format": normalized}

    @app.delete("/api/case-manager/series/{series_id}/children/{child_id}")
    def delete_child(series_id: int, child_id: int):
        stamp = _now()
        with get_db() as conn:
            child = require_case(conn, child_id, "series_child")
            if int(child.get("parent_series_id") or 0) != series_id:
                raise HTTPException(status_code=404, detail="子案件不属于该系列")
            if child.get("source_status") != "deleted":
                conn.execute(
                    "UPDATE case_manager_cases SET pre_delete_status=?,source_status='deleted',deleted_at=?,updated_at=? WHERE id=?",
                    (_text(child.get("source_status")) or "active", stamp, stamp, child_id),
                )
            audit(conn, child_id, "deleted", "从系列主表删除子案并移入垃圾站")
            audit(conn, series_id, "updated", f"删除子案 {_text(child.get('series_display_name'))}")
            conn.commit()
        return {"ok": True}

    @app.put("/api/case-manager/series/{series_id}/columns")
    async def update_series_columns(series_id: int, request: Request):
        body = await request.json()
        columns = _sanitize_columns(body.get("columns") or [])
        deleted_keys = {_text(value) for value in body.get("deleted_keys") or [] if _text(value)}
        settings = _profile_settings({"table_settings_json": _json(body.get("settings") or DEFAULT_TABLE_SETTINGS)})
        with get_db() as conn:
            require_case(conn, series_id, "series")
            conn.execute(
                "UPDATE case_manager_series_profiles SET table_columns_json=?,table_settings_json=?,updated_at=? WHERE series_case_id=?",
                (_json(columns), _json(settings), _now(), series_id),
            )
            if deleted_keys:
                marks = ",".join("?" for _ in deleted_keys)
                conn.execute(
                    f"DELETE FROM case_manager_series_cell_values WHERE series_case_id=? AND column_key IN ({marks})",
                    [series_id, *sorted(deleted_keys)],
                )
                conn.execute(
                    f"DELETE FROM case_manager_series_cell_formats WHERE series_case_id=? AND column_key IN ({marks})",
                    [series_id, *sorted(deleted_keys)],
                )
            audit(conn, series_id, "updated", "调整系列子案主表列", changes=[{"columns": columns, "deleted_keys": sorted(deleted_keys)}])
            conn.commit()
        return {"items": columns, "settings": settings}

    @app.put("/api/case-manager/series/{series_id}/detail-field-mappings")
    async def update_detail_field_mappings(series_id: int, request: Request):
        body = await request.json()
        mappings = _profile_detail_field_mappings({"detail_field_mappings_json": _json(body.get("mappings") or [])})
        with get_db() as conn:
            require_case(conn, series_id, "series")
            profile = _row(conn.execute(
                "SELECT * FROM case_manager_series_profiles WHERE series_case_id=?", (series_id,)
            ).fetchone())
            column_keys = {item["key"] for item in _profile_columns(profile)} | BUILTIN_COLUMN_KEYS
            mappings = [item for item in mappings if item["source_key"] in column_keys]
            conn.execute(
                "UPDATE case_manager_series_profiles SET detail_field_mappings_json=?,updated_at=? WHERE series_case_id=?",
                (_json(mappings), _now(), series_id),
            )
            audit(conn, series_id, "updated", "调整子案详情字段映射", changes=[{"detail_field_mappings": mappings}])
            conn.commit()
        return {"mappings": mappings}

    def selected_child_ids(conn: Any, series_id: int, body: Dict[str, Any]) -> List[int]:
        ids = [int(value) for value in body.get("child_ids") or [] if int(value or 0) > 0]
        if body.get("scope") == "filter":
            filters = body.get("filters") or {}
            where = ["parent_series_id=?", "case_kind='series_child'"]
            args: List[Any] = [series_id]
            if _text(filters.get("stage")):
                where.append("current_stage=?")
                args.append(_text(filters.get("stage")))
            if _text(filters.get("status")) in {"active", "archived", "deleted"}:
                where.append("source_status=?")
                args.append(_text(filters.get("status")))
            ids = [int(r[0]) for r in conn.execute(f"SELECT id FROM case_manager_cases WHERE {' AND '.join(where)}", args).fetchall()]
        valid = []
        for child_id in ids:
            row = conn.execute("SELECT 1 FROM case_manager_cases WHERE id=? AND parent_series_id=? AND case_kind='series_child'", (child_id, series_id)).fetchone()
            if row:
                valid.append(child_id)
        if not valid:
            raise HTTPException(status_code=400, detail="请选择系列子案")
        return valid

    @app.post("/api/case-manager/series/{series_id}/children/bulk")
    async def bulk_children(series_id: int, request: Request):
        body = await request.json()
        operation = _text(body.get("operation"))
        payload = body.get("payload") or {}
        batch_id = uuid.uuid4().hex
        stamp = _now()
        with get_db() as conn:
            require_case(conn, series_id, "series")
            ids = selected_child_ids(conn, series_id, body)
            for child_id in ids:
                before = require_case(conn, child_id, "series_child")
                if operation == "set_fields":
                    allowed = {"current_stage", "court_contact", "party_contact", "memo", "facts_mode", "facts_difference", "contract_mode", "contract_difference", "other_differences", "case_focus"}
                    values = {key: payload.get(key) for key in allowed if key in payload}
                    if values:
                        assignments = ",".join(f"{key}=?" for key in values)
                        conn.execute(f"UPDATE case_manager_cases SET {assignments},updated_at=? WHERE id=?", [*values.values(), stamp, child_id])
                    if "parties" in payload:
                        mode = _text(payload.get("party_operation")) or "replace"
                        existing = [_row(r) for r in conn.execute("SELECT party_group,role,name,credit_code FROM case_manager_parties WHERE case_id=? ORDER BY sort_order,id", (child_id,)).fetchall()]
                        incoming = payload.get("parties") or []
                        if mode == "remove":
                            keys = {(_text(x.get("party_group")), _text(x.get("role")), _text(x.get("name"))) for x in incoming}
                            next_items = [x for x in existing if (_text(x.get("party_group")), _text(x.get("role")), _text(x.get("name"))) not in keys]
                        else:
                            next_items = incoming if mode == "replace" else [*existing, *incoming]
                        replace_parties(conn, child_id, next_items)
                    if "case_numbers" in payload:
                        mode = _text(payload.get("number_operation")) or "replace"
                        existing = [_row(r) for r in conn.execute("SELECT * FROM case_manager_case_numbers WHERE case_id=? ORDER BY sort_order,id", (child_id,)).fetchall()]
                        incoming = payload.get("case_numbers") or []
                        if mode == "remove":
                            keys = {(_text(x.get("number_type")), _text(x.get("case_no")), _text(x.get("court"))) for x in incoming}
                            next_items = [x for x in existing if (_text(x.get("number_type")), _text(x.get("case_no")), _text(x.get("court"))) not in keys]
                        else:
                            next_items = incoming if mode == "replace" else [*existing, *incoming]
                        replace_case_numbers(conn, child_id, next_items)
                    if "claims" in payload:
                        mode = _text(payload.get("claim_operation")) or "replace"
                        existing = [_row(r) for r in conn.execute("SELECT * FROM case_manager_claims WHERE case_id=? ORDER BY sort_order,id", (child_id,)).fetchall()]
                        incoming = payload.get("claims") or []
                        if mode == "remove":
                            keys = {(_text(x.get("claim_code")), _text(x.get("claim_text"))) for x in incoming}
                            next_items = [x for x in existing if (_text(x.get("claim_code")), _text(x.get("claim_text"))) not in keys]
                        else:
                            next_items = incoming if mode == "replace" else [*existing, *incoming]
                        replace_claims(conn, child_id, next_items)
                elif operation == "add_milestone":
                    item = payload.get("milestone") or {}
                    existing = [_row(r) for r in conn.execute("SELECT * FROM case_manager_milestones WHERE case_id=? ORDER BY sort_order,id", (child_id,)).fetchall()]
                    replace_milestones(conn, child_id, [*existing, item])
                    if payload.get("sync_stage") and _text(payload.get("current_stage")):
                        conn.execute("UPDATE case_manager_cases SET current_stage=?,updated_at=? WHERE id=?", (_text(payload.get("current_stage")), stamp, child_id))
                elif operation in {"add_todo", "add_hearing"}:
                    child = require_case(conn, child_id, "series_child")
                    if operation == "add_todo":
                        conn.execute("""INSERT INTO case_manager_todos(case_id,title,todo_title,remind_at,notes,status,created_at,updated_at,source_key)
                            VALUES (?,?,?,?,?,'pending',?,?,?)""", (child_id, child["title"], _text(payload.get("todo_title")), _text(payload.get("remind_at")), _text(payload.get("notes")), stamp, stamp, uuid.uuid4().hex))
                    else:
                        conn.execute("""INSERT INTO case_manager_hearings(case_id,title,reminder_type,hearing_at,location,case_no,court,court_contact,party_contact,notes,created_at,updated_at,source_key)
                            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""", (child_id, child["title"], _text(payload.get("reminder_type")) or "开庭", _text(payload.get("hearing_at")), _text(payload.get("location")), _text(payload.get("case_no")), _text(payload.get("court")), _text(payload.get("court_contact")), _text(payload.get("party_contact")), _text(payload.get("notes")), stamp, stamp, uuid.uuid4().hex))
                elif operation in {"archive", "restore", "delete"}:
                    next_status = {"archive": "archived", "restore": "active", "delete": "deleted"}[operation]
                    pre = _text(before.get("source_status")) or "active"
                    if pre != next_status:
                        deleted_at = stamp if next_status == "deleted" else ""
                        archived_at = stamp if next_status == "archived" else ""
                        conn.execute("UPDATE case_manager_cases SET source_status=?,pre_delete_status=?,deleted_at=?,archived_at=?,updated_at=? WHERE id=?", (next_status, pre if next_status == "deleted" else _text(before.get("pre_delete_status")), deleted_at, archived_at, stamp, child_id))
                    if operation == "restore":
                        conn.execute("UPDATE case_manager_cases SET source_status='active',archived_at='',deleted_at='',updated_at=? WHERE id=? AND case_kind='series'", (stamp, series_id))
                else:
                    raise HTTPException(status_code=400, detail="无效的批量操作")
                audit(conn, child_id, "updated", f"批量操作：{operation}", batch_id=batch_id)
            conn.execute("INSERT INTO case_manager_series_batches(batch_id,series_case_id,operation,summary_json,created_at) VALUES (?,?,?,?,?)", (batch_id, series_id, operation, _json({"count": len(ids), "payload": payload}), stamp))
            audit(conn, series_id, "updated", f"对子案执行批量操作：{operation}（{len(ids)}案）", batch_id=batch_id)
            conn.commit()
            return {"ok": True, "batch_id": batch_id, "count": len(ids)}

    @app.post("/api/case-manager/series/{series_id}/archive")
    async def archive_series(series_id: int, request: Request):
        body = await request.json()
        force = bool(body.get("force"))
        stamp = _now()
        with get_db() as conn:
            require_case(conn, series_id, "series")
            active = int(conn.execute("SELECT count(*) FROM case_manager_cases WHERE parent_series_id=? AND case_kind='series_child' AND source_status='active'", (series_id,)).fetchone()[0])
            if active and not force:
                raise HTTPException(status_code=409, detail=f"仍有 {active} 个在办子案；请逐案结案或确认强制归档")
            batch_id = uuid.uuid4().hex
            if force:
                conn.execute("UPDATE case_manager_cases SET source_status='archived',archived_at=?,deleted_at='',updated_at=? WHERE parent_series_id=? AND case_kind='series_child' AND source_status='active'", (stamp, stamp, series_id))
            conn.execute("UPDATE case_manager_cases SET source_status='archived',archived_at=?,deleted_at='',pinned=0,updated_at=? WHERE id=?", (stamp, stamp, series_id))
            audit(conn, series_id, "archived", "强制全部结案并归档" if force else "归档系列案件", batch_id=batch_id)
            conn.commit()
        return {"ok": True, "batch_id": batch_id}

    @app.post("/api/case-manager/series/{series_id}/restore")
    def restore_series(series_id: int):
        with get_db() as conn:
            require_case(conn, series_id, "series")
            conn.execute("UPDATE case_manager_cases SET source_status='active',archived_at='',deleted_at='',updated_at=? WHERE id=?", (_now(), series_id))
            audit(conn, series_id, "restored", "恢复系列案件；子案状态保持不变")
            conn.commit()
        return {"ok": True}

    @app.delete("/api/case-manager/series/{series_id}")
    def delete_series(series_id: int):
        stamp = _now()
        with get_db() as conn:
            require_case(conn, series_id, "series")
            rows = conn.execute("SELECT id,source_status FROM case_manager_cases WHERE id=? OR parent_series_id=?", (series_id, series_id)).fetchall()
            for row in rows:
                conn.execute("UPDATE case_manager_cases SET pre_delete_status=?,source_status='deleted',deleted_at=?,archived_at='',updated_at=? WHERE id=?", (_text(row["source_status"]) or "active", stamp, stamp, row["id"]))
            audit(conn, series_id, "deleted", "系列案件及全部子案移入垃圾站")
            conn.commit()
        return {"ok": True}

    @app.get("/api/case-manager/series/{series_id}/daily-reminders")
    def daily_reminders(series_id: int, date: str = "", item_type: str = ""):
        with get_db() as conn:
            require_case(conn, series_id, "series")
            rows = conn.execute(
                """SELECT 'todo' item_type,t.id,t.case_id,t.todo_title label,t.remind_at due_at,t.notes detail,c.series_code,CASE WHEN c.case_kind='series' THEN '系列事项' ELSE c.series_display_name END series_display_name
                   FROM case_manager_todos t JOIN case_manager_cases c ON c.id=t.case_id WHERE (c.parent_series_id=? OR c.id=?) AND COALESCE(t.status,'pending')<>'completed'
                   UNION ALL SELECT 'hearing',h.id,h.case_id,h.reminder_type,h.hearing_at,h.location,c.series_code,CASE WHEN c.case_kind='series' THEN '系列事项' ELSE c.series_display_name END
                   FROM case_manager_hearings h JOIN case_manager_cases c ON c.id=h.case_id WHERE (c.parent_series_id=? OR c.id=?)
                   UNION ALL SELECT 'preservation',p.id,p.case_id,p.preservation_type,p.deadline_at,p.asset_content,c.series_code,CASE WHEN c.case_kind='series' THEN '系列事项' ELSE c.series_display_name END
                   FROM case_manager_preservations p JOIN case_manager_cases c ON c.id=p.case_id WHERE (c.parent_series_id=? OR c.id=?)
                   UNION ALL SELECT 'appeal',a.id,a.case_id,a.appeal_type,a.deadline_at,a.case_no,c.series_code,CASE WHEN c.case_kind='series' THEN '系列事项' ELSE c.series_display_name END
                   FROM case_manager_appeals a JOIN case_manager_cases c ON c.id=a.case_id WHERE (c.parent_series_id=? OR c.id=?)""",
                (series_id, series_id, series_id, series_id, series_id, series_id, series_id, series_id),
            ).fetchall()
            items = [_row(r) for r in rows if _text(r["due_at"])]
            if date:
                items = [item for item in items if _text(item.get("due_at"))[:10] == date[:10]]
            if item_type:
                items = [item for item in items if item.get("item_type") == item_type]
            groups: Dict[Tuple[str, str], List[Dict[str, Any]]] = {}
            for item in items:
                groups.setdefault((_text(item["due_at"])[:10], item["item_type"]), []).append(item)
            return {"groups": [{"date": key[0], "item_type": key[1], "count": len(values), "items": values} for key, values in sorted(groups.items())]}

    def workbook_payload(workbook: Any) -> Dict[str, Any]:
        result: Dict[str, Any] = {"series": {}, "children": [], "parties": [], "case_numbers": [], "claims": [], "milestones": []}
        errors: List[Dict[str, Any]] = []
        sheet_map = {"系列信息": "series", "子案件": "children", "当事人": "parties", "案号": "case_numbers", "请求项": "claims", "进度节点": "milestones"}
        headers = {
            "children": ["子案编号", "子案名称", "案件类型", "当前阶段", "案情模式", "案情差异", "合同模式", "合同差异", "其他差异", "侧重点", "备忘"],
            "parties": ["子案编号", "类别", "诉讼地位", "名称", "统一社会信用代码"],
            "case_numbers": ["子案编号", "性质", "案号", "受理机构", "承办人联系方式", "一审立案时间", "一审判决送达时间", "二审立案时间", "二审判决生效时间"],
            "claims": ["子案编号", "共同请求编号", "处理方式", "请求编号", "请求内容", "金额", "币种", "计算口径", "依据", "备注"],
            "milestones": ["子案编号", "节点类型", "标题", "状态", "时间", "备注"],
        }
        if "系列信息" in workbook.sheetnames:
            ws = workbook["系列信息"]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0]:
                    result["series"] = {"title": _text(row[0]), "case_type": _text(row[1] if len(row) > 1 else ""), "common_facts": _text(row[2] if len(row) > 2 else ""), "common_contract_terms": _text(row[3] if len(row) > 3 else ""), "series_memo": _text(row[4] if len(row) > 4 else "")}
                    break
        for sheet_name, key in sheet_map.items():
            if key == "series" or sheet_name not in workbook.sheetnames:
                continue
            ws = workbook[sheet_name]
            actual = [_text(cell.value) for cell in ws[1]]
            missing = [name for name in headers[key] if name not in actual]
            if missing:
                errors.append({"sheet": sheet_name, "row": 1, "message": f"缺少列：{'、'.join(missing)}"})
                continue
            positions = {name: actual.index(name) for name in headers[key]}
            for row_number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not any(value not in (None, "") for value in values):
                    continue
                record = {name: values[index] if index < len(values) else "" for name, index in positions.items()}
                result[key].append(record)
                if key == "children" and not (_text(record["子案编号"]) and _text(record["子案名称"])):
                    errors.append({"sheet": sheet_name, "row": row_number, "message": "子案编号和子案名称不能为空"})
        codes = [_text(item["子案编号"]) for item in result["children"]]
        duplicates = sorted({code for code in codes if code and codes.count(code) > 1})
        if duplicates:
            errors.append({"sheet": "子案件", "row": 0, "message": f"子案编号重复：{'、'.join(duplicates)}"})
        result["errors"] = errors
        return result

    def generic_workbook_payload(workbook: Any) -> Dict[str, Any]:
        preferred = next(
            (name for name in ("系列案件总表", "系列子案主表") if name in workbook.sheetnames),
            workbook.sheetnames[0],
        )
        worksheet = workbook[preferred]
        best_row = 1
        best_score = -1
        if preferred == "系列案件总表" and worksheet.auto_filter.ref:
            _, best_row, _, _ = range_boundaries(worksheet.auto_filter.ref)
        else:
            for row_number in range(1, min(15, worksheet.max_row) + 1):
                values = [_text(cell.value) for cell in worksheet[row_number]]
                nonempty = sum(1 for value in values if value)
                semantic = sum(2 for value in values if _header_builtin(value))
                score = nonempty + semantic
                if nonempty >= 2 and score > best_score:
                    best_score, best_row = score, row_number
        headers = [_text(cell.value) for cell in worksheet[best_row]]
        columns: List[Dict[str, Any]] = []
        seen: set[str] = set()
        default_by_key = {item["key"]: item for item in DEFAULT_SERIES_COLUMNS}
        for index, label in enumerate(headers):
            if not label:
                continue
            builtin = _header_builtin(label)
            if not builtin:
                compact = label.replace(" ", "")
                if not any(item.get("key") == "display_name" for item in columns) and any(token in compact for token in ("姓名", "名称", "房号", "铺号")):
                    builtin = "display_name"
            key = builtin or _custom_key(label)
            if key in seen:
                key = _custom_key(f"{label}-{index}")
                builtin = ""
            seen.add(key)
            base = dict(default_by_key.get(builtin) or {})
            party_like = any(token in label for token in ("原告", "被告", "申请人", "被申请人", "第三人", "执行人", "业主", "购房人", "承租人", "出租人"))
            source = base.get("source") or ("parties" if party_like else "custom")
            semantic_type = "amount" if _is_amount_header(label) and builtin != "claim_totals" else ""
            field_type = "amount" if semantic_type else (base.get("field_type") or ("lookup" if party_like else "text"))
            columns.append({
                "key": key, "label": label, "field_type": field_type, "source": source,
                "width": base.get("width", max(110, min(240, len(label) * 22))), "excel_index": index,
                "recognized_as": builtin or ("金额列（自动汇总）" if semantic_type else "自定义列"),
                "semantic_type": semantic_type, "currency": "CNY",
            })
        if not columns:
            return {"generic": True, "sheet": preferred, "header_row": best_row, "columns": [], "generic_rows": [], "errors": [{"sheet": preferred, "row": best_row, "message": "没有识别到表头"}]}
        keys = {item["key"] for item in columns}
        has_stable_key = "series_code" in keys
        if "series_code" not in keys:
            columns.insert(0, {**DEFAULT_SERIES_COLUMNS[0], "excel_index": -1, "recognized_as": "自动生成"})
        if "display_name" not in keys:
            columns.insert(1, {**DEFAULT_SERIES_COLUMNS[1], "excel_index": -1, "recognized_as": "自动生成"})
        rows: List[Dict[str, Any]] = []
        missing_stable_key_rows: List[int] = []
        for row_number, values in enumerate(worksheet.iter_rows(min_row=best_row + 1, values_only=True), start=best_row + 1):
            if not any(value not in (None, "") for value in values):
                continue
            item: Dict[str, Any] = {"_row": row_number}
            for column in columns:
                index = int(column.get("excel_index", -1))
                if index >= 0:
                    item[column["key"]] = _text(values[index] if index < len(values) else "")
            if not _text(item.get("series_code")):
                if has_stable_key:
                    missing_stable_key_rows.append(row_number)
                item["series_code"] = f"{len(rows) + 1:03d}"
            if not _text(item.get("display_name")):
                candidate = next((_text(item.get(column["key"])) for column in columns if int(column.get("excel_index", -1)) >= 0 and _text(item.get(column["key"]))), "")
                item["display_name"] = candidate or f"第{row_number}行子案"
            rows.append(item)
        errors: List[Dict[str, Any]] = []
        codes = [item["series_code"] for item in rows]
        duplicates = sorted({code for code in codes if codes.count(code) > 1})
        if duplicates:
            errors.append({"sheet": preferred, "row": 0, "message": f"识别到重复编号：{'、'.join(duplicates)}"})
        return {
            "generic": True, "sheet": preferred, "header_row": best_row, "columns": columns,
            "generic_rows": rows, "errors": errors, "has_stable_key": has_stable_key,
            "missing_stable_key_rows": missing_stable_key_rows,
        }

    def generic_row_to_child(row: Dict[str, Any], columns: List[Dict[str, Any]], index: int) -> Dict[str, Any]:
        child: Dict[str, Any] = {}
        for key in EDITABLE_CASE_FIELDS:
            if key in row:
                child[key] = row.get(key)
        child["series_code"] = _text(row.get("series_code")) or f"{index + 1:03d}"
        child["display_name"] = _text(row.get("display_name")) or f"第{index + 1}个子案"
        if "clients" in row or "opponents" in row:
            parties: List[Dict[str, Any]] = []
            for field, group, role in (("clients", "client", "委托人"), ("opponents", "other", "对方当事人")):
                names = _text(row.get(field)).replace("，", "、").replace(",", "、").split("、")
                parties.extend({"party_group": group, "role": role, "name": name.strip()} for name in names if name.strip())
            child["parties"] = parties
        if "case_no" in row or "court" in row:
            child["case_numbers"] = [{"case_no": _text(row.get("case_no")), "court": _text(row.get("court"))}]
        child["custom_values"] = {
            column["key"]: row.get(column["key"], "")
            for column in columns if column["key"] not in BUILTIN_COLUMN_KEYS
        }
        return child

    def generic_import_diff(conn: Any, series_id: int, parsed: Dict[str, Any]) -> Dict[str, Any]:
        """Compare an uploaded main sheet with the current non-deleted child rows."""
        columns = parsed.get("columns") or []
        rows = parsed.get("generic_rows") or []
        comparable_columns = [
            column for column in columns
            if column.get("field_type") != "readonly" and column.get("source") != "derived"
        ]
        current_children = [_row(row) for row in conn.execute(
            """SELECT * FROM case_manager_cases
               WHERE parent_series_id=? AND case_kind='series_child' AND source_status<>'deleted'
               ORDER BY series_order,id""",
            (series_id,),
        ).fetchall()]
        current_by_code: Dict[str, Dict[str, Any]] = {}
        for child in current_children:
            current_by_code[_text(child.get("series_code"))] = child_summary(conn, child)

        def value_for(item: Dict[str, Any], column: Dict[str, Any]) -> str:
            key = _text(column.get("key"))
            if key in item:
                return _text(item.get(key))
            return _text((item.get("custom_values") or {}).get(key))

        def field_changes(before: Dict[str, Any], after: Dict[str, Any]) -> List[Dict[str, str]]:
            changes: List[Dict[str, str]] = []
            for column in comparable_columns:
                old_value = value_for(before, column)
                new_value = value_for(after, column)
                if old_value != new_value:
                    changes.append({
                        "key": _text(column.get("key")), "label": _text(column.get("label")),
                        "before": old_value, "after": new_value,
                    })
            return changes

        incoming_by_code: Dict[str, Dict[str, Any]] = {}
        changes: List[Dict[str, Any]] = []
        for row in rows:
            code = _text(row.get("series_code"))
            incoming_by_code[code] = row
            current = current_by_code.get(code)
            if current is None:
                changes.append({
                    "type": "added", "code": code, "name": _text(row.get("display_name")),
                    "fields": [
                        {"key": _text(column.get("key")), "label": _text(column.get("label")),
                         "before": "", "after": value_for(row, column)}
                        for column in comparable_columns if value_for(row, column)
                    ],
                })
                continue
            fields = field_changes(current, row)
            if fields:
                changes.append({"type": "updated", "code": code, "name": _text(row.get("display_name")) or _text(current.get("display_name")), "fields": fields})

        deleted = []
        if parsed.get("has_stable_key"):
            for code, current in current_by_code.items():
                if code and code not in incoming_by_code:
                    deleted.append({
                        "type": "deleted", "code": code,
                        "name": _text(current.get("display_name")), "fields": [],
                    })
            changes.extend(deleted)
        summary = {
            "added": sum(1 for item in changes if item["type"] == "added"),
            "updated": sum(1 for item in changes if item["type"] == "updated"),
            "deleted": sum(1 for item in changes if item["type"] == "deleted"),
            "unchanged": max(0, len(current_by_code) - sum(1 for item in changes if item["type"] in {"updated", "deleted"})),
            "before": len(current_by_code),
            "after": len(rows),
        }
        return {
            "summary": summary,
            "changes": changes[:300],
            "truncated": len(changes) > 300,
            "deletions_enabled": bool(parsed.get("has_stable_key")),
        }

    def smart_workbook_payload(workbook: Any) -> Dict[str, Any]:
        if "系列案件总表" in workbook.sheetnames or "系列子案主表" in workbook.sheetnames:
            return generic_workbook_payload(workbook)
        if "子案件" in workbook.sheetnames:
            expected = {"子案编号", "子案名称", "案件类型", "当前阶段", "案情模式", "案情差异", "合同模式", "合同差异", "其他差异", "侧重点", "备忘"}
            actual = {_text(cell.value) for cell in workbook["子案件"][1]}
            if expected.issubset(actual):
                return workbook_payload(workbook)
        return generic_workbook_payload(workbook)

    @app.post("/api/case-manager/series-create/import-preview")
    async def preview_initial_series_import(file: UploadFile = File(...)):
        try:
            data = await file.read()
            workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            parsed = generic_workbook_payload(workbook)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"无法读取Excel：{exc}") from exc
        preview_id = uuid.uuid4().hex
        import_previews[preview_id] = {"series_id": 0, "parsed": parsed, "created_at": _now()}
        return {
            "preview_id": preview_id,
            "mode": "smart_table",
            "sheet": parsed.get("sheet"),
            "header_row": parsed.get("header_row"),
            "counts": {"children": len(parsed.get("generic_rows") or []), "columns": len(parsed.get("columns") or [])},
            "columns": [{key: value for key, value in column.items() if key != "excel_index"} for column in parsed.get("columns") or []],
            "samples": (parsed.get("generic_rows") or [])[:8],
            "errors": parsed["errors"],
        }

    @app.post("/api/case-manager/series/{series_id}/import-preview")
    async def preview_series_import(series_id: int, file: UploadFile = File(...)):
        with get_db() as conn:
            require_case(conn, series_id, "series")
            existing_child_count = int(conn.execute(
                """SELECT count(*) FROM case_manager_cases
                   WHERE parent_series_id=? AND case_kind='series_child' AND source_status<>'deleted'""",
                (series_id,),
            ).fetchone()[0])
        try:
            data = await file.read()
            workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            parsed = smart_workbook_payload(workbook)
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"无法读取Excel：{exc}") from exc
        diff = None
        if parsed.get("generic"):
            if existing_child_count and not parsed.get("has_stable_key"):
                parsed.setdefault("errors", []).append({
                    "sheet": parsed.get("sheet", ""), "row": parsed.get("header_row", 1),
                    "message": "当前系列已有子案，必须保留“编号/序号/子案编号”列，系统才能安全识别修改、删除和回退对象",
                })
            elif existing_child_count and parsed.get("missing_stable_key_rows"):
                missing_rows = "、".join(str(row) for row in parsed["missing_stable_key_rows"][:20])
                suffix = "等" if len(parsed["missing_stable_key_rows"]) > 20 else ""
                parsed.setdefault("errors", []).append({
                    "sheet": parsed.get("sheet", ""), "row": 0,
                    "message": f"第 {missing_rows} 行缺少编号，无法安全判断是新增还是修改{suffix}；请补齐编号后再导入",
                })
            with get_db() as conn:
                diff = generic_import_diff(conn, series_id, parsed)
        preview_id = uuid.uuid4().hex
        import_previews[preview_id] = {"series_id": series_id, "parsed": parsed, "created_at": _now()}
        if parsed.get("generic"):
            return {
                "preview_id": preview_id,
                "mode": "smart_table",
                "sheet": parsed.get("sheet"),
                "header_row": parsed.get("header_row"),
                "counts": {"children": len(parsed.get("generic_rows") or []), "columns": len(parsed.get("columns") or [])},
                "columns": [{key: value for key, value in column.items() if key != "excel_index"} for column in parsed.get("columns") or []],
                "samples": (parsed.get("generic_rows") or [])[:5],
                "errors": parsed["errors"],
                "diff": diff,
                "has_stable_key": bool(parsed.get("has_stable_key")),
                "missing_stable_key_rows": parsed.get("missing_stable_key_rows") or [],
            }
        return {"preview_id": preview_id, "mode": "standard_workbook", "counts": {key: len(parsed[key]) for key in ("children", "parties", "case_numbers", "claims", "milestones")}, "errors": parsed["errors"]}

    @app.post("/api/case-manager/series/{series_id}/import-commit")
    async def commit_series_import(series_id: int, request: Request):
        body = await request.json()
        preview_token = _text(body.get("preview_id"))
        preview = import_previews.get(preview_token)
        if not preview or int(preview["series_id"]) != series_id:
            raise HTTPException(status_code=404, detail="导入预览已失效，请重新上传")
        parsed = preview["parsed"]
        if parsed["errors"]:
            raise HTTPException(status_code=400, detail="导入文件仍有错误，未写入任何数据")
        if parsed.get("generic"):
            batch_id = uuid.uuid4().hex
            columns = _sanitize_columns(parsed.get("columns") or [])
            with get_db() as conn:
                series = require_case(conn, series_id, "series")
                diff = generic_import_diff(conn, series_id, parsed)
                history_id = create_series_history(
                    conn, series_id, "excel_import",
                    {**(diff.get("summary") or {}), "batch_id": batch_id},
                )
                conn.execute(
                    "UPDATE case_manager_series_profiles SET table_columns_json=?,updated_at=? WHERE series_case_id=?",
                    (_json(columns), _now(), series_id),
                )
                incoming_codes = {
                    _text(row.get("series_code")) for row in parsed.get("generic_rows") or []
                    if _text(row.get("series_code"))
                }
                existing_children = [_row(row) for row in conn.execute(
                    """SELECT id,series_code,source_status FROM case_manager_cases
                       WHERE parent_series_id=? AND case_kind='series_child' AND source_status<>'deleted'""",
                    (series_id,),
                ).fetchall()]
                if parsed.get("has_stable_key"):
                    for child in existing_children:
                        if _text(child.get("series_code")) in incoming_codes:
                            continue
                        stamp = _now()
                        conn.execute(
                            """UPDATE case_manager_cases SET pre_delete_status=COALESCE(NULLIF(pre_delete_status,''),source_status),
                               source_status='deleted',deleted_at=?,updated_at=? WHERE id=?""",
                            (stamp, stamp, child["id"]),
                        )
                        audit(conn, int(child["id"]), "deleted", f"Excel导入中因行缺失移入垃圾站：编号 {_text(child.get('series_code'))}", batch_id=batch_id)
                for index, row in enumerate(parsed.get("generic_rows") or []):
                    child_body = generic_row_to_child(row, columns, index)
                    existing = conn.execute(
                        """SELECT id,source_status FROM case_manager_cases
                           WHERE parent_series_id=? AND case_kind='series_child' AND series_code=?
                           ORDER BY CASE WHEN source_status='deleted' THEN 1 ELSE 0 END,id DESC LIMIT 1""",
                        (series_id, child_body["series_code"]),
                    ).fetchone()
                    if existing:
                        update_child_record(conn, int(existing["id"]), child_body)
                        conn.execute(
                            "UPDATE case_manager_cases SET source_status='active',deleted_at='',pre_delete_status='',series_order=? WHERE id=?",
                            (index, existing["id"]),
                        )
                    else:
                        create_child_record(conn, series, child_body, index)
                audit(conn, series_id, "imported", f"智能表格导入 {len(parsed.get('generic_rows') or [])} 个子案、{len(columns)} 列；新增 {diff['summary']['added']}、修改 {diff['summary']['updated']}、移入垃圾站 {diff['summary']['deleted']}", batch_id=batch_id)
                conn.execute(
                    "INSERT INTO case_manager_series_batches(batch_id,series_case_id,operation,summary_json,created_at) VALUES (?,?,?,?,?)",
                    (batch_id, series_id, "smart_excel_import", _json({"children": len(parsed.get("generic_rows") or []), "columns": len(columns), **diff["summary"], "history_id": history_id}), _now()),
                )
                conn.commit()
            import_previews.pop(preview_token, None)
            return {"ok": True, "batch_id": batch_id, "history_id": history_id, "children": len(parsed.get("generic_rows") or []), "columns": len(columns), "diff": diff}
        by_code: Dict[str, Dict[str, Any]] = {}
        for item in parsed["children"]:
            code = _text(item["子案编号"])
            by_code[code] = {
                "series_code": code, "display_name": _text(item["子案名称"]), "case_type": _text(item["案件类型"]),
                "current_stage": _text(item["当前阶段"]), "facts_mode": _text(item["案情模式"]), "facts_difference": _text(item["案情差异"]),
                "contract_mode": _text(item["合同模式"]), "contract_difference": _text(item["合同差异"]), "other_differences": _text(item["其他差异"]),
                "case_focus": _text(item["侧重点"]), "memo": _text(item["备忘"]), "parties": [], "case_numbers": [], "claims": [], "milestones": [],
            }
        for item in parsed["parties"]:
            target = by_code.get(_text(item["子案编号"]))
            if target is not None:
                target["parties"].append({"party_group": "client" if _text(item["类别"]) in {"委托人", "client"} else "other", "role": _text(item["诉讼地位"]), "name": _text(item["名称"]), "credit_code": _text(item["统一社会信用代码"])})
        for item in parsed["case_numbers"]:
            target = by_code.get(_text(item["子案编号"]))
            if target is not None:
                target["case_numbers"].append({"number_type": _text(item["性质"]), "case_no": _text(item["案号"]), "court": _text(item["受理机构"]), "handler_contact": _text(item["承办人联系方式"]), "first_filing_at": _text(item["一审立案时间"]), "first_judgment_served_at": _text(item["一审判决送达时间"]), "second_filing_at": _text(item["二审立案时间"]), "second_judgment_effective_at": _text(item["二审判决生效时间"])})
        for item in parsed["claims"]:
            target = by_code.get(_text(item["子案编号"]))
            if target is not None:
                target["claims"].append({"parent_claim_ref": _text(item["共同请求编号"]), "claim_mode": _text(item["处理方式"]), "claim_code": _text(item["请求编号"]), "claim_text": _text(item["请求内容"]), "amount_text": _text(item["金额"]), "currency": _text(item["币种"]), "calculation_basis": _text(item["计算口径"]), "legal_basis": _text(item["依据"]), "notes": _text(item["备注"])})
        for item in parsed["milestones"]:
            target = by_code.get(_text(item["子案编号"]))
            if target is not None:
                target["milestones"].append({"milestone_type": _text(item["节点类型"]), "title": _text(item["标题"]), "milestone_status": _text(item["状态"]), "happened_at": _text(item["时间"]), "notes": _text(item["备注"])})
        batch_id = uuid.uuid4().hex
        with get_db() as conn:
            series = require_case(conn, series_id, "series")
            current_children = [_row(row) for row in conn.execute(
                """SELECT id,series_code,source_status FROM case_manager_cases
                   WHERE parent_series_id=? AND case_kind='series_child' AND source_status<>'deleted'""",
                (series_id,),
            ).fetchall()]
            current_codes = {_text(item.get("series_code")) for item in current_children}
            incoming_codes = set(by_code)
            import_summary = {
                "added": len(incoming_codes - current_codes),
                "updated": len(incoming_codes & current_codes),
                "deleted": len(current_codes - incoming_codes),
                "unchanged": 0,
                "before": len(current_codes),
                "after": len(incoming_codes),
                "batch_id": batch_id,
            }
            history_id = create_series_history(conn, series_id, "excel_import", import_summary)
            imported_series = parsed.get("series") or {}
            if imported_series:
                stamp = _now()
                conn.execute("UPDATE case_manager_cases SET case_type=COALESCE(NULLIF(?,''),case_type),updated_at=? WHERE id=?", (_text(imported_series.get("case_type")), stamp, series_id))
                conn.execute(
                    """UPDATE case_manager_series_profiles SET common_facts=?,common_contract_terms=?,series_memo=?,updated_at=?
                       WHERE series_case_id=?""",
                    (_text(imported_series.get("common_facts")), _text(imported_series.get("common_contract_terms")),
                     _text(imported_series.get("series_memo")), stamp, series_id),
                )
            common_claims = []
            for item in parsed["claims"]:
                if _text(item["子案编号"]):
                    continue
                common_claims.append({"claim_code": _text(item["请求编号"]), "claim_text": _text(item["请求内容"]),
                    "amount_text": _text(item["金额"]), "currency": _text(item["币种"]),
                    "calculation_basis": _text(item["计算口径"]), "legal_basis": _text(item["依据"]), "notes": _text(item["备注"])})
            if common_claims:
                replace_claims(conn, series_id, common_claims, parent_kind=True)
            common_rows = [_row(r) for r in conn.execute(
                "SELECT id,claim_code FROM case_manager_claims WHERE case_id=? ORDER BY sort_order,id", (series_id,)
            ).fetchall()]
            common_by_code = {_text(item.get("claim_code")): int(item["id"]) for item in common_rows}
            common_ids = {int(item["id"]) for item in common_rows}
            for child in by_code.values():
                for claim in child["claims"]:
                    ref = _text(claim.pop("parent_claim_ref", ""))
                    if not ref:
                        claim["parent_claim_id"] = None
                    elif ref in common_by_code:
                        claim["parent_claim_id"] = common_by_code[ref]
                    elif ref.isdigit() and int(ref) in common_ids:
                        claim["parent_claim_id"] = int(ref)
                    else:
                        raise HTTPException(status_code=400, detail=f"共同请求编号“{ref}”不存在")
            for current in current_children:
                if _text(current.get("series_code")) in incoming_codes:
                    continue
                stamp = _now()
                conn.execute(
                    """UPDATE case_manager_cases SET pre_delete_status=COALESCE(NULLIF(pre_delete_status,''),source_status),
                       source_status='deleted',deleted_at=?,updated_at=? WHERE id=?""",
                    (stamp, stamp, current["id"]),
                )
                audit(conn, int(current["id"]), "deleted", f"Excel导入中因行缺失移入垃圾站：编号 {_text(current.get('series_code'))}", batch_id=batch_id)
            for index, child in enumerate(by_code.values()):
                existing = conn.execute(
                    """SELECT id,source_status FROM case_manager_cases
                       WHERE parent_series_id=? AND case_kind='series_child' AND series_code=?
                       ORDER BY CASE WHEN source_status='deleted' THEN 1 ELSE 0 END,id DESC LIMIT 1""",
                    (series_id, child["series_code"]),
                ).fetchone()
                if existing:
                    update_child_record(conn, int(existing["id"]), child)
                    conn.execute(
                        "UPDATE case_manager_cases SET source_status='active',deleted_at='',pre_delete_status='',series_order=? WHERE id=?",
                        (index, existing["id"]),
                    )
                else:
                    create_child_record(conn, series, child, index)
            audit(conn, series_id, "imported", f"Excel导入 {len(by_code)} 个子案；新增 {import_summary['added']}、修改 {import_summary['updated']}、移入垃圾站 {import_summary['deleted']}", batch_id=batch_id)
            conn.execute("INSERT INTO case_manager_series_batches(batch_id,series_case_id,operation,summary_json,created_at) VALUES (?,?,?,?,?)", (batch_id, series_id, "excel_import", _json({"children": len(by_code), **import_summary, "history_id": history_id}), _now()))
            conn.commit()
        import_previews.pop(preview_token, None)
        return {"ok": True, "batch_id": batch_id, "history_id": history_id, "children": len(by_code), "diff": {"summary": import_summary}}

    @app.get("/api/case-manager/series/{series_id}/history")
    def list_series_history(series_id: int):
        with get_db() as conn:
            require_case(conn, series_id, "series")
            rows = conn.execute(
                """SELECT history_id,operation,summary_json,created_at
                   FROM case_manager_series_history
                   WHERE series_case_id=? ORDER BY datetime(created_at) DESC,history_id DESC LIMIT ?""",
                (series_id, SERIES_HISTORY_LIMIT),
            ).fetchall()
        items = []
        for row in rows:
            try:
                summary = json.loads(row[2] or "{}")
            except (TypeError, ValueError, json.JSONDecodeError):
                summary = {}
            items.append({
                "history_id": int(row[0]), "operation": _text(row[1]),
                "summary": summary if isinstance(summary, dict) else {}, "created_at": row[3],
            })
        return {"items": items, "limit": SERIES_HISTORY_LIMIT}

    @app.post("/api/case-manager/series/{series_id}/history/{history_id}/restore")
    def restore_series_history(series_id: int, history_id: int):
        with get_db() as conn:
            require_case(conn, series_id, "series")
            row = conn.execute(
                """SELECT history_id,summary_json,snapshot_json FROM case_manager_series_history
                   WHERE history_id=? AND series_case_id=?""",
                (history_id, series_id),
            ).fetchone()
            if not row:
                raise HTTPException(status_code=404, detail="历史版本不存在或已超过保留期限")
            try:
                snapshot = json.loads(row[2] or "{}")
            except (TypeError, ValueError, json.JSONDecodeError) as exc:
                raise HTTPException(status_code=409, detail="历史版本数据已损坏，无法回退") from exc
            current_children = int(conn.execute(
                """SELECT count(*) FROM case_manager_cases
                   WHERE parent_series_id=? AND case_kind='series_child' AND source_status<>'deleted'""",
                (series_id,),
            ).fetchone()[0])
            restore_history_id = create_series_history(
                conn, series_id, "history_restore",
                {"restored_history_id": int(history_id), "before_children": current_children},
            )
            restore_series_snapshot(conn, series_id, snapshot)
            batch_id = uuid.uuid4().hex
            audit(conn, series_id, "restored", f"回退到导入历史版本 #{history_id}", batch_id=batch_id)
            conn.commit()
            return {"ok": True, "history_id": restore_history_id, "restored_history_id": int(history_id)}

    @app.get("/api/case-manager/series/{series_id}/export")
    def export_series(series_id: int):
        with get_db() as conn:
            data = series_detail(conn, series_id, include_children=False)
            children = [child_detail(conn, int(r[0])) for r in conn.execute("SELECT id FROM case_manager_cases WHERE parent_series_id=? AND case_kind='series_child' ORDER BY series_order,id", (series_id,)).fetchall()]
        workbook = openpyxl.Workbook()
        table_columns = _profile_columns(data.get("profile") or {})
        visible_children = [child for child in children if child.get("source_status") != "deleted"]
        main_sheet = workbook.active
        main_sheet.title = "系列案件总表"
        main_sheet.sheet_properties.tabColor = "A63D35"
        main_sheet.sheet_view.showGridLines = False
        main_sheet.sheet_view.zoomScale = 90

        title_fill = PatternFill("solid", fgColor="A63D35")
        section_fill = PatternFill("solid", fgColor="F3E7E5")
        label_fill = PatternFill("solid", fgColor="F7F3F2")
        header_fill = PatternFill("solid", fgColor="273449")
        white_font = Font(name="微软雅黑", color="FFFFFF", bold=True)
        title_font = Font(name="微软雅黑", color="FFFFFF", bold=True, size=16)
        label_font = Font(name="微软雅黑", color="6B3A35", bold=True, size=10)
        body_font = Font(name="微软雅黑", color="172033", size=10)
        light_side = Side(style="thin", color="D9DEE7")
        medium_side = Side(style="medium", color="B8C0CC")
        cell_border = Border(left=light_side, right=light_side, top=light_side, bottom=light_side)
        section_border = Border(left=medium_side, right=medium_side, top=medium_side, bottom=medium_side)
        max_columns = max(len(table_columns), 8)
        last_summary_col = get_column_letter(max_columns)

        def style_merged_row(row: int, fill: PatternFill, border: Border) -> None:
            for column_index in range(1, max_columns + 1):
                cell = main_sheet.cell(row=row, column=column_index)
                cell.fill = fill
                cell.border = border

        main_sheet.merge_cells(f"A1:{last_summary_col}1")
        main_sheet["A1"] = f"{data['title']}｜系列案件总表"
        main_sheet["A1"].fill = title_fill
        main_sheet["A1"].font = title_font
        main_sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
        style_merged_row(1, title_fill, Border())
        main_sheet.row_dimensions[1].height = 34

        stats = data.get("stats") or {}
        total_texts = []
        for currency, raw_amount in (stats.get("claim_totals") or {}).items():
            try:
                amount_text = format(Decimal(_text(raw_amount)), ",.2f")
            except InvalidOperation:
                amount_text = _text(raw_amount)
            symbol = {"CNY": "¥", "USD": "$", "HKD": "HK$", "EUR": "€"}.get(currency, "")
            total_texts.append(f"{currency} {symbol}{amount_text}".strip())
        request_total = "；".join(total_texts) or "—"
        summary_pairs = [
            ("系列名称", data.get("title", "")),
            ("案由", data.get("case_type", "") or "—"),
            ("子案总数", stats.get("total", 0)),
            ("在办", stats.get("active", 0)),
            ("已归档", stats.get("archived", 0)),
            ("垃圾站", stats.get("deleted", 0)),
            ("逾期事项", stats.get("overdue", 0)),
            ("未来7日", stats.get("next_7_days", 0)),
        ]
        for index, (label, value) in enumerate(summary_pairs):
            row = 3 + index // 4
            column = 1 + (index % 4) * 2
            label_cell = main_sheet.cell(row=row, column=column, value=label)
            value_cell = main_sheet.cell(row=row, column=column + 1, value=value)
            label_cell.fill = label_fill
            label_cell.font = label_font
            value_cell.font = Font(name="微软雅黑", color="172033", bold=True, size=10)
            for cell in (label_cell, value_cell):
                cell.border = cell_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        main_sheet.row_dimensions[3].height = 26
        main_sheet.row_dimensions[4].height = 26

        amount_note = "、".join(stats.get("amount_columns") or [])
        amount_note = f"按主表金额列（{amount_note}）汇总" if amount_note else "按有效请求项汇总"
        main_sheet["A6"] = "请求金额"
        main_sheet["A6"].fill = section_fill
        main_sheet["A6"].font = label_font
        main_sheet["A6"].border = section_border
        main_sheet.merge_cells(f"B6:{last_summary_col}6")
        main_sheet["B6"] = f"{request_total}　·　{amount_note}"
        main_sheet["B6"].font = Font(name="微软雅黑", color="A63D35", bold=True, size=11)
        main_sheet["B6"].alignment = Alignment(vertical="center", wrap_text=True)
        for column_index in range(2, max_columns + 1):
            main_sheet.cell(row=6, column=column_index).border = section_border
        main_sheet.row_dimensions[6].height = 30

        narrative_rows = (
            (8, "共同基本案情", data["profile"].get("common_facts", "")),
            (10, "共同合同条款", data["profile"].get("common_contract_terms", "")),
            (12, "系列备忘", data["profile"].get("series_memo", "")),
        )
        for row, label, value in narrative_rows:
            main_sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_columns)
            main_sheet.cell(row=row, column=1, value=label).fill = section_fill
            main_sheet.cell(row=row, column=1).font = label_font
            main_sheet.cell(row=row, column=1).alignment = Alignment(vertical="center")
            style_merged_row(row, section_fill, section_border)
            main_sheet.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=max_columns)
            value_cell = main_sheet.cell(row=row + 1, column=1, value=_text(value) or "暂无记录")
            value_cell.font = body_font
            value_cell.alignment = Alignment(vertical="top", wrap_text=True)
            style_merged_row(row + 1, PatternFill(fill_type=None), section_border)
            main_sheet.row_dimensions[row].height = 22
            main_sheet.row_dimensions[row + 1].height = min(72, max(28, 18 * ((len(_text(value)) // 90) + 1)))

        table_header_row = 15
        for column_index, column in enumerate(table_columns, start=1):
            cell = main_sheet.cell(row=table_header_row, column=column_index, value=column["label"])
            cell.fill = header_fill
            cell.font = white_font
            cell.border = cell_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        main_sheet.row_dimensions[table_header_row].height = 30

        for child in visible_children:
            parties = child.get("parties") or []
            numbers = child.get("case_numbers") or []
            milestones = child.get("milestones") or []
            complete = [item for item in milestones if item.get("milestone_status") == "completed"]
            planned = [item for item in milestones if item.get("milestone_status") == "planned"]
            values = {
                **child,
                "display_name": child.get("display_name") or child.get("series_display_name") or "",
                "clients": "、".join(_text(item.get("name")) for item in parties if item.get("party_group") == "client" and _text(item.get("name"))),
                "opponents": "、".join(_text(item.get("name")) for item in parties if item.get("party_group") != "client" and _text(item.get("name"))),
                "case_no": "\n".join(_text(item.get("case_no")) for item in numbers if _text(item.get("case_no"))),
                "court": "\n".join(_text(item.get("court")) for item in numbers if _text(item.get("court"))),
                "latest_milestone": _text((complete[-1] if complete else {}).get("title")),
                "next_milestone": _text((planned[0] if planned else {}).get("title")),
                "claim_totals": "；".join(f"{currency} {amount}" for currency, amount in (child.get("claim_totals") or {}).items()),
            }
            values.update(child.get("custom_values") or {})
            output_row: List[Any] = []
            for column in table_columns:
                raw_value = values.get(column["key"], "")
                if column.get("semantic_type") == "amount":
                    _, parsed_amount = parse_table_amount(raw_value, _text(column.get("currency")) or "CNY")
                    raw_value = float(parsed_amount) if parsed_amount is not None else raw_value
                output_row.append(raw_value)
            main_sheet.append(output_row)

        last_data_row = max(table_header_row, main_sheet.max_row)
        for row in main_sheet.iter_rows(min_row=table_header_row + 1, max_row=last_data_row, max_col=len(table_columns)):
            wrapped_lines = 1
            for index, cell in enumerate(row):
                column = table_columns[index]
                cell.font = body_font
                cell.border = cell_border
                cell.alignment = Alignment(
                    horizontal="right" if column.get("semantic_type") == "amount" else "left",
                    vertical="top",
                    wrap_text=True,
                )
                if column.get("semantic_type") == "amount" and isinstance(cell.value, (int, float, Decimal)):
                    cell.number_format = '#,##0.00'
                text_length = len(_text(cell.value))
                wrapped_lines = max(wrapped_lines, min(10, (text_length // max(10, int(column.get("width", 150)) // 8)) + 1))
            main_sheet.row_dimensions[row[0].row].height = min(180, max(24, wrapped_lines * 18))

        for column_index, column in enumerate(table_columns, start=1):
            width = max(11, min(48, int(column.get("width") or 150) / 7.2))
            main_sheet.column_dimensions[get_column_letter(column_index)].width = width
        main_sheet.auto_filter.ref = f"A{table_header_row}:{get_column_letter(len(table_columns))}{last_data_row}"
        settings = data.get("profile", {}).get("table_settings") or DEFAULT_TABLE_SETTINGS
        freeze_columns = min(max(int(settings.get("freeze_columns", 2)), 0), len(table_columns))
        freeze_rows = min(max(int(settings.get("freeze_rows", 0)), 0), len(visible_children))
        main_sheet.freeze_panes = main_sheet.cell(row=table_header_row + 1 + freeze_rows, column=freeze_columns + 1)
        main_sheet.print_title_rows = f"{table_header_row}:{table_header_row}"
        main_sheet.print_area = f"A1:{get_column_letter(max(len(table_columns), 8))}{last_data_row}"
        main_sheet.page_setup.orientation = "landscape"
        main_sheet.page_setup.fitToWidth = 1
        main_sheet.page_setup.fitToHeight = 0
        main_sheet.sheet_properties.pageSetUpPr.fitToPage = True
        main_sheet.sheet_properties.outlinePr.summaryBelow = True

        first = workbook.create_sheet("系列信息")
        first.append(["系列名称", "案件类型", "共同基本案情", "共同合同条款", "系列备忘"])
        first.append([data["title"], data.get("case_type", ""), data["profile"].get("common_facts", ""), data["profile"].get("common_contract_terms", ""), data["profile"].get("series_memo", "")])
        specs = {
            "子案件": ["子案编号", "子案名称", "案件类型", "当前阶段", "案情模式", "案情差异", "合同模式", "合同差异", "其他差异", "侧重点", "备忘"],
            "当事人": ["子案编号", "类别", "诉讼地位", "名称", "统一社会信用代码"],
            "案号": ["子案编号", "性质", "案号", "受理机构", "承办人联系方式", "一审立案时间", "一审判决送达时间", "二审立案时间", "二审判决生效时间"],
            "请求项": ["子案编号", "共同请求编号", "处理方式", "请求编号", "请求内容", "金额", "币种", "计算口径", "依据", "备注"],
            "进度节点": ["子案编号", "节点类型", "标题", "状态", "时间", "备注"],
        }
        sheets = {}
        for name, headers in specs.items():
            sheets[name] = workbook.create_sheet(name)
            sheets[name].append(headers)
        common_code_by_id = {int(claim["id"]): claim.get("claim_code", "") for claim in data.get("common_claims") or []}
        for claim in data.get("common_claims") or []:
            sheets["请求项"].append(["", claim.get("claim_code", ""), "common", claim.get("claim_code", ""), claim.get("claim_text", ""), claim.get("amount_text", ""), claim.get("currency", ""), claim.get("calculation_basis", ""), claim.get("legal_basis", ""), claim.get("notes", "")])
        for child in children:
            code = child.get("series_code", "")
            sheets["子案件"].append([code, child.get("display_name", ""), child.get("case_type", ""), child.get("current_stage", ""), child.get("facts_mode", ""), child.get("facts_difference", ""), child.get("contract_mode", ""), child.get("contract_difference", ""), child.get("other_differences", ""), child.get("case_focus", ""), child.get("memo", "")])
            for party in child["parties"]:
                sheets["当事人"].append([code, "委托人" if party.get("party_group") == "client" else "其他当事人", party.get("role", ""), party.get("name", ""), party.get("credit_code", "")])
            for item in child["case_numbers"]:
                sheets["案号"].append([code, item.get("number_type", ""), item.get("case_no", ""), item.get("court", ""), item.get("handler_contact", ""), item.get("first_filing_at", ""), item.get("first_judgment_served_at", ""), item.get("second_filing_at", ""), item.get("second_judgment_effective_at", "")])
            for claim in child["claims"]:
                sheets["请求项"].append([code, common_code_by_id.get(int(claim.get("parent_claim_id") or 0), ""), claim.get("claim_mode", ""), claim.get("claim_code", ""), claim.get("claim_text", ""), claim.get("amount_text", ""), claim.get("currency", ""), claim.get("calculation_basis", ""), claim.get("legal_basis", ""), claim.get("notes", "")])
            for item in child["milestones"]:
                sheets["进度节点"].append([code, item.get("milestone_type", ""), item.get("title", ""), item.get("milestone_status", ""), item.get("happened_at", ""), item.get("notes", "")])
        for sheet in workbook.worksheets[1:]:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_state = "veryHidden"
        workbook.active = 0
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        filename = f"系列案件-{data['title']}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
        )
