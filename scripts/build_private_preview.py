#!/usr/bin/env python3
"""Build a data-free Apple Silicon Stella preview from a Git source allowlist."""
from __future__ import annotations
import argparse
import ast
import base64
import hashlib
import json
import os
import plistlib
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

ROOT=Path(__file__).resolve().parent.parent
TEXT_SUFFIXES={'.py','.swift','.html','.js','.css','.md'}
BRAND_ASSETS={'frontend/stella-brand.png','frontend/distribution-brand.png','frontend/user-manual.md','frontend/static/site.webmanifest'}
FORBIDDEN_PARTS={'.git','uploads','article_images','data','outputs','browser-profile','browser-profiles','__pycache__'}
SECRET_PATTERNS=[re.compile(r'(?<![A-Za-z0-9])sk-[A-Za-z0-9_-]{20,}'),re.compile(r'(?:ghp|github_pat)_[A-Za-z0-9_]{20,}'),re.compile(r'-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----')]


def allowed(path, profile='macos'):
    p=Path(path)
    if profile == 'macos' and path in {'backend/personal_center.py', 'backend/feishu_family_finance_sync.py', 'frontend/personal-center.html'}: return False
    if profile == 'windows' and (p.name.startswith('feishu_') or p.name.startswith('recording') or p.suffix == '.swift'): return False
    if set(p.parts)&FORBIDDEN_PARTS or p.name.startswith(('local_', '.env')): return False
    return (p.parts[0] in ('backend','frontend') and p.suffix in TEXT_SUFFIXES
            or path in BRAND_ASSETS
            or path.startswith('frontend/static/') and p.name.startswith(('app-icon-','apple-touch-icon')) and p.suffix=='.png')


def public_replacements(root):
    """Derive private defaults from committed source, never from live user data."""
    replacements = {}
    targets = {
        'backend/automated_filing.py': {'DEFAULT_SERVICE_CONTACT', 'FIRM_LAWYER_PHONES', 'FIRM_AGENT_PROFILES'},
        'backend/case_automation.py': {'DEFAULT_FIRM'},
        'backend/main.py': {'default_clients', 'default_parties'},
    }
    def strings(value):
        if isinstance(value, str):
            yield value
        elif isinstance(value, dict):
            for key, item in value.items():
                yield from strings(key)
                yield from strings(item)
        elif isinstance(value, (tuple, list)):
            for item in value:
                yield from strings(item)
    for path, names in targets.items():
        result = subprocess.run(['git', 'show', 'HEAD:' + path], cwd=root, capture_output=True, text=True)
        if result.returncode:
            continue
        for node in ast.walk(ast.parse(result.stdout)):
            if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id in names for t in node.targets):
                for value in strings(ast.literal_eval(node.value)):
                    if len(value) > 1 and value not in {'mobile', 'address', 'document_number', 'practice_number', 'name', 'credit_code', 'bank', 'account'}:
                        replacements[value] = '' if any(c.isdigit() for c in value) else '示例资料（虚构）'
    # Distinct fictional names preserve keyed library behavior without real identities.
    names = [v for v in replacements if 1 < len(v) <= 3 and not any(c.isdigit() for c in v)]
    for index, value in enumerate(names):
        replacements[value] = '示例律师' + '甲乙丙丁戊己庚辛壬癸'[index]
    replacements['刘律师'] = '新用户'
    return replacements


def sanitize_public_text(text, replacements):
    for original in sorted(replacements, key=len, reverse=True):
        text = text.replace(original, replacements[original])
    text = re.sub(r'/Users/[A-Za-z0-9_.-]+', '/Users/example', text)
    return text


def public_source_text(path, text, replacements, profile='macos'):
    text = sanitize_public_text(text, replacements)
    if profile == 'macos' and path == 'backend/main.py':
        for module, function, args in (
            ('personal_center', 'register_personal_center_routes', 'app, get_db, FRONTEND_DIR'),
            ('feishu_family_finance_sync', 'register_feishu_family_finance_sync', 'app, get_db'),
        ):
            text = text.replace(f'from .{module} import {function}\n{function}({args})', '')
            text = text.replace(f'    from .{module} import {function}\n    {function}({args})', '    pass')
    if profile == 'macos' and path in {'frontend/app.js', 'frontend/app-bootstrap.js'}:
        # Remove the catalog entry before its order/customization adapters consume it.
        text = re.sub(r"^  \{ id: 'personal_center', title: '个人中心'.*?\n", '', text, flags=re.M)
        # Compact bootstrap is regenerated from this sanitized canonical app.js below.
        text = text.replace('  renderPersonalCenterSecurity();', '')
    if profile == 'macos' and path == 'frontend/index.html':
        text = re.sub(r'<section[^>]*id="settings-section-personal-center".*?</section>', '', text, flags=re.S)
        text = text.replace('管理应用登录与个人中心的独立访问保护。', '管理应用登录和访问保护。')
    if path == 'frontend/user-manual.md':
        # Original screenshots are runtime/user material, never copied into a public bundle.
        text = re.sub(r'!\[([^]]*)\]\(/static/manual-images/[^)]+\)', r'（界面说明：\1）', text)
    return text


def isolated_source(root,stage,profile='macos'):
    manifest=[]
    replacements = public_replacements(root)
    entries=subprocess.check_output(['git','ls-tree','-rz','HEAD'],cwd=root).split(b'\0')
    for entry in entries:
        if not entry: continue
        metadata,path_bytes=entry.split(b'\t',1); mode,_,blob=metadata.split(); path=path_bytes.decode()
        if not allowed(path, profile): continue
        if mode not in (b'100644',b'100755'): raise RuntimeError('Refusing non-regular source entry: '+path)
        content=subprocess.check_output(['git','cat-file','blob',blob.decode()],cwd=root)
        if Path(path).suffix in TEXT_SUFFIXES:
            text=public_source_text(path, content.decode('utf-8'), replacements, profile)
            if any(pattern.search(text) for pattern in SECRET_PATTERNS): raise RuntimeError('Credential pattern in source: '+path)
            # In the frozen app, the recipient's editable settings must live outside the signed bundle.
            for base in ('Path(__file__).resolve().parent.parent','Path(ai_config.__file__).resolve().parent.parent','Path(ai_config_module.__file__).resolve().parent.parent'):
                for filename in ('local_ai_secrets.py','local_ai_prefs.json'):
                    text=text.replace(f'{base} / "{filename}"',f'Path(os.environ.get("LEGAL_KB_CONFIG_DIR", str({base}))) / "{filename}"')
            if path == 'frontend/index.html':
                brand = subprocess.run(['git', 'show', 'HEAD:frontend/distribution-brand.png'], cwd=root, capture_output=True)
                if brand.returncode == 0:
                    logo = 'data:image/png;base64,' + base64.b64encode(brand.stdout).decode()
                    seed = '<script>try{if(!localStorage.getItem("legal_kb_brand_logo")){localStorage.setItem("legal_kb_brand_logo",' + json.dumps(logo) + ')}document.documentElement.classList.add("has-custom-brand-logo")}catch(e){}</script>'
                    text = text.replace('<head>', '<head>' + seed, 1)
            content=text.encode('utf-8')
        target=stage/path; target.parent.mkdir(parents=True,exist_ok=True); target.write_bytes(content)
        manifest.append({'path':path,'sha256':hashlib.sha256(content).hexdigest()})
    builder_dir = stage/'scripts'
    builder_dir.mkdir()
    shutil.copyfile(root/'scripts/build_frontend.cjs', builder_dir/'build_frontend.cjs') if (root/'scripts/build_frontend.cjs').is_file() else None
    if (builder_dir/'build_frontend.cjs').is_file():
        (stage/'docs/architecture').mkdir(parents=True)
        node_env = {**os.environ, 'NODE_PATH': str(root/'node_modules')}
        subprocess.run(['node', str(builder_dir/'build_frontend.cjs')], cwd=stage, env=node_env, check=True, stdout=subprocess.DEVNULL)
        for item in manifest:
            item['sha256'] = hashlib.sha256((stage/item['path']).read_bytes()).hexdigest()
    entry=root/'scripts/distribution/preview_entry.py'
    shutil.copyfile(entry,stage/'preview_entry.py')
    # The public build uses exactly the canonical frontend styles, without preview overrides.
    # User-authored binary document templates are never copied. Generate plain trial templates anew.
    from docx import Document
    templates=stage/'backend/archive_templates'; templates.mkdir()
    for category in ('litigation','non_litigation'):
        for kind in ('summary','catalog'):
            doc=Document(); doc.core_properties.author='Stella Preview'
            doc.add_heading('案卷归档示例模板（体验版）',0)
            doc.add_paragraph('本文件为空白体验模板，不包含任何真实案件材料。')
            if kind=='summary':
                table=doc.add_table(rows=8,cols=2)
                for row,label in zip(table.rows,['案件名称','委托人','对方当事人','案号','承办律师','收案日期','结案日期','案件摘要']):
                    row.cells[0].text=label
            else:
                table=doc.add_table(rows=1,cols=4)
                for cell,label in zip(table.rows[0].cells,['序号','材料名称','日期','备注']): cell.text=label
            doc.save(templates/f'{category}_{kind}.docx')
    (stage/'SOURCE_MANIFEST.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    return manifest


def public_bundle_info(version, build_number, key, metrics_url):
    from urllib.parse import urlparse
    if not re.fullmatch(r'[0-9]+\.[0-9]+\.[0-9]+', version):
        raise ValueError('version must be major.minor.patch')
    if not re.fullmatch(r'[1-9][0-9]{0,17}', str(build_number)):
        raise ValueError('build number must be a positive integer')
    if len(base64.b64decode(key, validate=True)) != 32:
        raise ValueError('invalid update public key')
    url = urlparse(metrics_url)
    if url.scheme != 'https' or not url.hostname or url.path != '/v1/events' or url.query or url.fragment or url.username:
        raise ValueError('metrics URL must be an HTTPS /v1/events endpoint')
    return dict(CFBundleIdentifier='com.stella.public.native', CFBundleDisplayName='Stella', CFBundleName='Stella',
        CFBundleExecutable='StellaPublic', LegalKBFrontendPath='', CFBundleShortVersionString=version,
        CFBundleVersion=str(build_number), LSMinimumSystemVersion='12.0',
        SUFeedURL='https://liu-zongjin.github.io/stella-downloads/appcast.xml', SUPublicEDKey=key,
        SUEnableAutomaticChecks=True, SUAutomaticallyUpdate=False, SUAllowsAutomaticUpdates=False, SUEnableSystemProfiling=False,
        SUVerifyUpdateBeforeExtraction=True, StellaMetricsURL=metrics_url)


def build(output, keep_stage=False, *, version, build_number, metrics_url):
    if os.uname().machine!='arm64': raise RuntimeError('Build on an Apple Silicon Mac')
    from distribution.sparkle_tools import tools, public_key, swift_annotation_flags
    changed = subprocess.check_output(['git', 'status', '--porcelain', '--untracked-files=normal'], cwd=ROOT, text=True)
    if changed.strip():
        raise RuntimeError('Commit reviewed sources before building a public release')
    sdk = tools()
    bundle_info = public_bundle_info(version, build_number, public_key(sdk), metrics_url)
    output.mkdir(parents=True,exist_ok=True)
    workspace=Path(tempfile.mkdtemp(prefix='stella-private-preview-'))
    stage=workspace/'source'; stage.mkdir()
    manifest=isolated_source(ROOT,stage)
    app=output/'Stella.app'
    if app.exists(): raise RuntimeError('Output already exists; use a new output directory')
    macos=app/'Contents/MacOS'; resources=app/'Contents/Resources'; macos.mkdir(parents=True); resources.mkdir()
    build_env=os.environ.copy()
    build_env.pop('PYTHONPATH',None)
    for key in list(build_env):
        if key.endswith(('_API_KEY','_AUTH_CODE','_AUTHORIZATION','_WEB_COOKIE','_ACCESS_TOKEN')): build_env.pop(key,None)
    command=[shutil.which('python3'),'-m','PyInstaller','--noconfirm','--clean','--onedir','--name','stella-preview-server',
        '--distpath',str(workspace/'frozen'),'--workpath',str(workspace/'work'),'--specpath',str(workspace),
        '--paths',str(stage),'--add-data',f'{stage}/frontend:frontend', '--add-data',f'{stage}/backend:backend',
        '--collect-submodules','backend','--collect-all','jieba','--collect-all','ddddocr','--collect-all','onnxruntime', '--hidden-import','uvicorn.logging',
        '--hidden-import','uvicorn.loops.auto','--hidden-import','uvicorn.protocols.http.auto','--hidden-import','uvicorn.lifespan.on',
        '--exclude-module','torch','--exclude-module','tensorflow','--exclude-module','matplotlib',
        '--exclude-module','scipy','--exclude-module','tkinter',
        str(stage/'preview_entry.py')]
    subprocess.run(command,cwd=stage,env=build_env,check=True)
    frozen=workspace/'frozen/stella-preview-server'
    # PyInstaller freezes modules separately; do not also distribute raw backend sources.
    for source in (frozen/'_internal/backend').rglob('*.py'):
        source.unlink()
    shutil.copytree(frozen,resources/'Server')
    swift=ROOT/'native-shell/Sources/main.swift'
    native_source=workspace/'main.swift'
    native_source.write_text(swift.read_text().replace(':8000',':18865'))
    subprocess.run(['xcrun','swiftc','-O','-whole-module-optimization','-target','arm64-apple-macosx12.0',
                    '-framework','Cocoa','-framework','WebKit','-framework','LocalAuthentication',
                    '-framework','EventKit','-framework','UserNotifications', '-D', 'STELLA_PUBLIC',
                    *swift_annotation_flags(), '-F', str(sdk), '-framework', 'Sparkle', '-Xlinker', '-rpath',
                    '-Xlinker', '@executable_path/../Frameworks', str(native_source),
                    str(ROOT/'native-shell/Sources/PublicDistribution.swift'), '-o', str(macos/'StellaPublic')],check=True)
    subprocess.run(['xcrun','swiftc','-O','-target','arm64-apple-macosx12.0','-framework','EventKit',
                    str(ROOT/'backend/calendar_bridge.swift'),'-o',str(macos/'APEXCalendarBridge')],check=True)
    info=plistlib.loads((ROOT/'native-shell/Info.plist').read_bytes())
    info.update(bundle_info)
    frameworks = app/'Contents/Frameworks'
    frameworks.mkdir()
    shutil.copytree(sdk/'Sparkle.framework', frameworks/'Sparkle.framework', symlinks=True)
    (app/'Contents/Info.plist').write_bytes(plistlib.dumps(info))
    shutil.copyfile(ROOT/'native-shell/StellaIcon.png',resources/'StellaBrand.png')
    icon=ROOT/'dist/Stella.app/Contents/Resources/AppIcon.icns'
    if icon.exists(): shutil.copyfile(icon,resources/'AppIcon.icns')
    # The native executable owns its backend process, including update/quit shutdown.
    for filename in ('LICENSE', 'THIRD_PARTY_NOTICES.md'):
        shutil.copyfile(ROOT/filename, resources/filename)
    shutil.copytree(ROOT/'third_party', resources/'third_party', dirs_exist_ok=True)
    subprocess.run(['codesign','--force','--deep','--sign','-',str(app)],check=True)
    subprocess.run([str(resources/'Server/stella-preview-server'),'--runtime',str(workspace/'validation-runtime'),'--self-test'],
                   cwd=workspace,env=build_env,check=True)
    audit_bundle(app)
    readme=output/'使用说明.txt'
    readme.write_text('''Stella 公开体验版（Apple 芯片 Mac，macOS 12 或更新）

1. 解压，将 Stella.app 拖到“应用程序”或您自己的 Applications 文件夹。
2. 打开应用即可使用，无需安装 Python，无需 GitHub 账号。
3. 软件没有开发者公证。若 macOS 提示无法验证开发者，请按系统设置 → 隐私与安全性中的官方“仍要打开”流程确认来源；不要关闭系统安全保护。
4. 第一次启动自动生成两个虚构案件、三篇全新编写的示例文章、少量联系人、待办和演示日程。个人财务为空。
5. AI API、MCP、Skills 管理页保留；所有密钥和账号配置为空，请填写您自己的信息。默认团队身份、联系方式、证件号码及真实客户资料均已替换为虚构占位或清空。
6. 体验数据独立保存在 ~/Library/Application Support/StellaPublic，退出应用会关闭随包后端。
7. 本包不含开发者的数据库、案件附件、上传文件、个人财务、AI聊天/长期记忆、API/MCP密钥、邮箱授权码、浏览器会话或钥匙串导出。
8. 自动读取司法短信已默认关闭。法院登录、邮箱、AI、MCP等联网服务须自行配置；法院服务包含与本机相同的浏览器驱动及本地验证码识别依赖；Chrome、Office/PDF 转换等外部程序由使用者按需要安装。
9. 仅分发应用，不公开源代码；使用许可及第三方声明随包提供。
''',encoding='utf-8')
    for filename in ('LICENSE', 'THIRD_PARTY_NOTICES.md'):
        shutil.copyfile(ROOT/filename, output/filename)
    shutil.copytree(ROOT/'third_party', output/'third_party', dirs_exist_ok=True)
    report={'source_commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(),
            'version':version, 'build':str(build_number), 'minimum_macos':'12.0', 'sparkle':'2.10.0', 'source_files':len(manifest),'architecture':'arm64','runtime_data_included':False,
            'team_defaults_sanitized':True,'raw_backend_source_included':False,'credentials_included':False,'fictional_data_created_on_first_run':True}
    (output/'包装检查.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    archive=output/f'Stella-macOS-AppleSilicon-v{version}-{build_number}.zip'
    subprocess.run(['ditto','-c','-k','--norsrc','--noextattr','--keepParent',str(app),str(archive)],check=True)
    # Keep the update archive app-only to preserve Sparkle/translocation behavior.
    # Ship license notices inside the bundle, not beside it in the update archive.
    (output/'SHA256.txt').write_text(hashlib.sha256(archive.read_bytes()).hexdigest()+'  '+archive.name+'\n')
    print('PACKAGE_READY',archive)
    if not keep_stage: shutil.rmtree(workspace)
    return archive


def audit_bundle(app):
    for path in app.rglob('*'):
        if not path.is_file(): continue
        relative=path.relative_to(app)
        public_ca_bundle = path.name == 'cacert.pem' and path.parent.name == 'certifi'
        if (path.suffix in ('.db','.sqlite','.sqlite3','.pem','.key','.p12') and not public_ca_bundle) or path.name.startswith(('local_', '.env')):
            raise RuntimeError('Runtime/private file in bundle: '+str(relative))
        own_source = 'backend' in relative.parts or 'frontend' in relative.parts
        if own_source and set(relative.parts)&FORBIDDEN_PARTS:
            raise RuntimeError('User data directory in bundle: '+str(relative))
        if path.suffix in TEXT_SUFFIXES and any(p.search(path.read_text(errors='ignore')) for p in SECRET_PATTERNS):
            raise RuntimeError('Credential pattern in bundle: '+str(relative))


if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--output',type=Path,required=True); parser.add_argument('--keep-stage',action='store_true')
    parser.add_argument('--version', required=True)
    parser.add_argument('--build', required=True, dest='build_number')
    parser.add_argument('--metrics-url', required=True)
    args=parser.parse_args()
    build(args.output.resolve(), args.keep_stage, version=args.version,
          build_number=args.build_number, metrics_url=args.metrics_url)
